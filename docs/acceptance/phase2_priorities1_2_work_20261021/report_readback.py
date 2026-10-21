"""Independent readback of generated synthetic reports and canonical SQLite."""
import csv,json,sqlite3,sys,re
from pathlib import Path
from decimal import Decimal
from openpyxl import load_workbook
from docx import Document
import fitz
root=Path(sys.argv[1]);results=[]
for kind in ('solid_slab','joist_block','joist_foam'):
 p=root/kind;db=sqlite3.connect(p/'db/projects.db')
 project=json.loads(db.execute('select payload from projects where id=?',('P',)).fetchone()[0]);db.close()
 expected={r['source_id']:(Decimal(str(r['quantities'][0]['amount'])),r['quantities'][0]['unit']) for r in project['takeoffs']}
 assert len(expected)==len(project['takeoffs'])
 for b in project['boq']:
  assert b['unit_price'] is None and b['total'] is None
  assert Decimal(str(b['quantity']))==sum(expected[s][0] for s in b['source_ids'])
 for ext in ('csv','xlsx','docx','pdf'):
  path=p/('roof.'+ext)
  if ext=='csv':
   assert path.read_bytes().startswith(b'\xef\xbb\xbf')
   rows=list(csv.reader(path.open(encoding='utf-8-sig',newline='')))
   actual={r[4]:Decimal(r[2]) for r in rows if len(r)>=7 and r[4] in expected}
   assert actual=={k:v[0] for k,v in expected.items()}
   assert all(r[3]==expected[r[4]][1] for r in rows if len(r)>=7 and r[4] in expected)
  elif ext=='xlsx':
   w=load_workbook(path);assert all(s.sheet_view.rightToLeft for s in w)
   rows=list(w['ریز متره'].values)
   actual={r[4]:Decimal(str(r[2])) for r in rows if len(r)>=7 and r[4] in expected}
   assert actual=={k:v[0] for k,v in expected.items()}
  elif ext=='docx':
   d=Document(path);assert 'bidi' in d._element.xml
   rows=[[c.text for c in r.cells] for t in d.tables for r in t.rows]
   actual={r[4]:Decimal(r[2]) for r in rows if len(r)>=7 and r[4] in expected}
   assert actual=={k:v[0] for k,v in expected.items()}
  else:
   with fitz.open(path) as pdf:
    text=''.join(page.get_text() for page in pdf);compact=re.sub(r'\s+','',text)
    assert all(s in compact for s in expected)
    assert all(str(v[0]).rstrip('0').rstrip('.') in text for v in expected.values())
    assert all(page.get_text().strip() for page in pdf)
   assert any('\u0600'<=c<='\u06ff' or '\ufb50'<=c<='\ufeff' for c in text)
  results.append({'family':kind,'format':ext,'source_rows':len(expected),'status':'PASS'})
print(json.dumps(results,ensure_ascii=False,indent=2))
