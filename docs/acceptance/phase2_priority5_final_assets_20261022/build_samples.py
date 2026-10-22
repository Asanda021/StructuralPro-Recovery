from pathlib import Path
import sys,json
sys.path[:0]=['.','tests']
from test_phase2_priority5_closure import crosswalk_case,opening_model,test_ifc_column_explicit_detail_uses_existing_specialist_and_boq,test_ifc_revision_amendment_uses_existing_journal_without_history_rewrite
from core.takeoff import manual_formula_register as formula
from core.drawings.ifc_pipeline import quantity_formula_data
from core.drawings.bim_quantities import compare_volume_qto
from PySide6.QtWidgets import QApplication,QLineEdit,QTreeWidget
from PySide6.QtCore import Qt,QCoreApplication,QEvent
from app.graphical_takeoff import GraphicalTakeoffDialog
root=Path('/workspace/scratch/61dce18379e6/p5-evidence/samples');root.mkdir(exist_ok=True)
service,p,drawing,session,proposal,kw=crosswalk_case(root)
link=service.save_ifc_crosswalk('P',**kw)
data=quantity_formula_data(service,'P',proposal['id'],'Length',author='author')
formula.save(service,'P','IFC-Q',data,expected_revision=0,reason='Synthetic explicit QTO, not material inference')
formula.review(service,'P','IFC-Q',expected_revision=1,reviewer='reviewer',note='Independent registered role; synthetic fixture')
formula.commit(service,'P','IFC-Q',expected_revision=1)
qt=QApplication.instance() or QApplication([])
w=GraphicalTakeoffDialog(app_service=service,project_id='P');w.show();w.open_drawing(str(p));w.ifc_revision.setText('R1');w.open_ifc_inspector()
dialog=w._ifc_inspector;dialog.findChild(QLineEdit,'ifc_search').setText(proposal['global_id']);tree=dialog.findChild(QTreeWidget,'ifc_tree');tree.setCurrentItem(tree.topLevelItem(0).child(0));qt.processEvents()
dialog.grab().save(str(root/'ifc-inspector.png'));dialog.close()
other=w.open_ifc_crosswalk(link['id']);w.resize(1400,850);other.resize(1400,850);qt.processEvents()
w.grab().save(str(root/'ifc-model.png'));other.grab().save(str(root/'cad-linked.png'))
for fmt in ('xlsx','pdf','docx','csv'):service.export_manual_formula_report('P',root/('ifc-qto.'+fmt),format=fmt)
project=service.open_project('P');(root/'project-evidence.json').write_text(json.dumps(project,ensure_ascii=False,indent=2))
from openpyxl import load_workbook
from docx import Document
import csv,fitz
book=load_workbook(root/'ifc-qto.xlsx',data_only=False);assert all(s.sheet_view.rightToLeft for s in book.worksheets);assert str(book['ریز متره'].cell(2,6).value)=='4'
word=Document(root/'ifc-qto.docx');texts=' '.join(c.text for t in word.tables for r in t.rows for c in r.cells);assert 'IFC-Q' in texts and 'مقدار' in texts and '4' in texts
with open(root/'ifc-qto.csv',encoding='utf-8-sig') as f:rows=list(csv.DictReader(f))
assert any('4' in str(list(r.values())) for r in rows)
with fitz.open(root/'ifc-qto.pdf') as pdf:
 pdf[0].get_pixmap().save(root/'report-page1.png');assert len(pdf)>0 and any('IFC' in pg.get_text() for pg in pdf)
result={'generic_quantity':'4','unit':'m','boq_rows':len(project['boq']),'takeoff_rows':len(project['takeoffs']),'crosswalks':len(project['ifc_2d_crosswalks']),'sqlite_integrity':service.store.integrity_check(),'exports_read':['xlsx','pdf','docx','csv'],'synthetic':True,'office_visual_verified':False}
w.close();QCoreApplication.sendPostedEvents(None,QEvent.DeferredDelete);qt.processEvents();service.store.close()
void=root/'real-void.ifc';gid,_=opening_model(void);result['volume_oracle']=compare_volume_qto(void,revision='R1',global_id=gid)
column=root/'column';column.mkdir();test_ifc_column_explicit_detail_uses_existing_specialist_and_boq(column)
revision=root/'revision';revision.mkdir();test_ifc_revision_amendment_uses_existing_journal_without_history_rewrite(revision,True)
result['specialist_oracle']={'concrete_m3':'0.6','formwork_m2':'2.7','reinforcement_kg':'21.488'}
result['revision_oracle']={'original_m':'4','effective_m':'5','delta_m':'1','new_global_id':True,'original_takeoffs_preserved':True}
(root/'sample-result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2));print(json.dumps(result,ensure_ascii=False))
