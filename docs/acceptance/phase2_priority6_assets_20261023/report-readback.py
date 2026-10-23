from pathlib import Path
from decimal import Decimal
import sys,json,csv,hashlib,re
from openpyxl import load_workbook
from docx import Document
from pypdf import PdfReader
from core.platform.application import StructuralProApp
p=Path(sys.argv[1]);out=Path(sys.argv[2]);service=StructuralProApp(p/'app')
boq=sum(Decimal(str(r['quantity'])) for r in service.open_project('P')['boq']);assert boq==6
w=load_workbook(p/'effective.xlsx',data_only=True)
rows=list(w['انتقال مالکیت IFC'].values)[1:]
current=[r for r in rows if r[1]=='جاری؛ مؤثر']
assert sorted((r[2],Decimal(str(r[4])),r[5]) for r in current)==[('B2',Decimal(2),'m'),('B3',Decimal(4),'m')]
assert sum(Decimal(str(r[4])) for r in current)==boq
assert all(s.sheet_view.rightToLeft for s in w)
doc=Document(p/'effective.docx');text='\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
assert 'B2' in text and 'B3' in text and 'جاری' in text
pdf=PdfReader(p/'effective.pdf');pdftext='\n'.join(page.extract_text() for page in pdf.pages)
assert 'B2' in pdftext and 'B3' in pdftext and re.search('[\u0600-\u06ff]',pdftext)
with (p/'effective.csv').open(encoding='utf-8-sig') as f: csvrows=list(csv.reader(f))
flat='\n'.join(','.join(r) for r in csvrows);assert 'B2' in flat and 'B3' in flat and 'جاری' in flat
result={'status':'PASS','sqlite_effective_m':str(boq),'current_members':{'B2':'2','B3':'4'},'pdf_pages':len(pdf.pages),'docx_tables':len(doc.tables),'xlsx_rtl':True,'csv_rows':len(csvrows),'farsi_text_extracted':True,'files':{f:hashlib.sha256((p/('effective.'+f)).read_bytes()).hexdigest() for f in ('pdf','xlsx','docx','csv')},'windows_office_tested':False}
service.store.close();out.write_text(json.dumps(result,ensure_ascii=False,indent=2));print(json.dumps(result,ensure_ascii=False))
