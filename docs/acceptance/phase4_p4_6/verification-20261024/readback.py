"""Read all forty scenario exports, without rewriting any output."""
import csv,hashlib,json,re,sys,zipfile
from pathlib import Path
import fitz,openpyxl
from docx import Document
from docx.oxml.ns import qn
root=Path(sys.argv[1]);results=[]
for kind in ['lookahead','resources','materials','purchases','daily','obstacles','meetings','correspondence','package','dashboard']:
    for fmt in ['pdf','xlsx','docx','csv']:
        p=root/(kind+'.'+fmt);row={'name':p.name,'size':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()}
        assert row['size']>100
        if fmt=='pdf':
            with fitz.open(p) as d:
                text='\n'.join(page.get_text() for page in d)
                row.update(pages=len(d),text_characters=len(text),fonts=sorted({f[3] for page in d for f in page.get_fonts()}),images=sum(len(page.get_images()) for page in d))
                assert text.strip() and '\ufffd' not in text
                if kind in ('dashboard','package'):assert len(d)<=15
        elif fmt=='xlsx':
            w=openpyxl.load_workbook(p)
            text='\n'.join(str(c.value) for s in w for r in s for c in r if c.value is not None)
            row.update(sheets=[{'name':s.title,'rows':s.max_row,'cols':s.max_column,'rtl':s.sheet_view.rightToLeft,'images':len(s._images),'charts':len(s._charts)} for s in w])
            assert all(s.sheet_view.rightToLeft for s in w)
            assert not any(c.data_type=='e' for s in w for r in s for c in r)
            w.close()
        elif fmt=='docx':
            d=Document(p);text='\n'.join([x.text for x in d.paragraphs]+[c.text for t in d.tables for r in t.rows for c in r.cells])
            row.update(editable_tables=len(d.tables),table_rtl=[t.table_direction is not None and bool(t._tbl.tblPr.find(qn('w:bidiVisual')) is not None) for t in d.tables],images=len(d.inline_shapes),bidi_elements=len(d._element.findall('.//'+qn('w:bidi'))))
            assert d.tables and all(row['table_rtl']) and row['bidi_elements']>0
            with zipfile.ZipFile(p) as z:assert z.testzip() is None
        else:
            text=p.read_text(encoding='utf-8-sig');records=list(csv.reader(text.splitlines()));row.update(rows=len(records),utf8_bom=p.read_bytes().startswith(b'\xef\xbb\xbf'))
            assert len(records)>1 and row['utf8_bom']
        row['persian_characters']=len(re.findall('[\u0600-\u06ff]',text))
        row['jalali_dates']=sorted(set(re.findall('[۱۲12][۰-۹0-9]{3}/[۰-۹0-9]{2}/[۰-۹0-9]{2}',text)))[:8]
        assert row['persian_characters']>0
        results.append(row)
out={'files':results,'count':len(results),'scope':'structural readback only; visual inspection and Microsoft Office are separate gates'}
Path(sys.argv[2]).write_text(json.dumps(out,ensure_ascii=False,indent=2))
print(json.dumps({'count':len(results),'pdf_pages':{x['name']:x['pages'] for x in results if 'pages'in x}},ensure_ascii=False))
