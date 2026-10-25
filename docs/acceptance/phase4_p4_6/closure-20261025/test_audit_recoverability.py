"""Recover the full source hashes and long decimal tokens from every Phase4 report."""
from pathlib import Path
import re,json

def test_all_phase4_reports_retain_csv_exact_audit_tokens():
 import fitz,openpyxl
 from docx import Document
 root=Path('/tmp/sp-closure-scenario-final');checks=[];missing=[]
 for source in sorted(root.rglob('*.csv')):
  relative=source.relative_to(root).as_posix()
  if relative.startswith('prior/engineering/'):continue # estimate/input contracts are separate
  text=source.read_text(encoding='utf-8-sig')
  hashes=set(re.findall(r'(?<![0-9a-fA-F])[0-9a-fA-F]{64}(?![0-9a-fA-F])',text))
  decimals=set(re.findall(r'(?<![\w.])[+-]?\d+\.\d{4,}(?![\w.])',text))
  for fmt in ('pdf','docx','xlsx'):
   p=source.with_suffix('.'+fmt)
   if fmt=='pdf':
    with fitz.open(p) as d:printed='\n'.join(page.get_text() for page in d)
   elif fmt=='docx':
    d=Document(p);printed='\n'.join([x.text for x in d.paragraphs]+[c.text for t in d.tables for r in t.rows for c in r.cells])
   else:
    d=openpyxl.load_workbook(p);printed='\n'.join(str(c.value) for s in d for r in s for c in r if c.value is not None);d.close()
   printed=''.join(printed.split());absent=[x for x in hashes|decimals if x not in printed]
   if absent:missing.append({'report':str(p.relative_to(root)),'tokens':absent})
   checks.append({'report':str(p.relative_to(root)),'full_hashes':len(hashes),'long_decimal_tokens':len(decimals),'missing':absent})
 result={'reports_checked':len(checks),'checks':checks,'missing':missing}
 Path('docs/acceptance/phase4_p4_6/closure-20261025/audit-recoverability.json').write_text(json.dumps(result,indent=2))
 assert len(checks)==45 and not missing,missing
