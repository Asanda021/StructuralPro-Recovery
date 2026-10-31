"""P8/P2: independent end-to-end five-floor PDF-to-price/statement evidence.

Real PyMuPDF source pages, reviewed calibrated SQLite drawing sessions,
actual BOQ, user-supplied CSV with explicit independent review, Decimal
estimate, period draft, and native report readback. This is a synthetic
engineering golden, NOT physical Windows/Office UAT or approval to pay.
"""
from __future__ import annotations

from copy import deepcopy
from decimal import Decimal
from hashlib import sha256
import json

import pytest

from core.platform.application import StructuralProApp
from core.reports.aec_office_package import audit_boq
from test_phase7_p2_multifloor_bundle import building, export


RATES = {
    "AREA_1": ("QA_RATE_1", Decimal("10")),
    "AREA_2": ("QA_RATE_2", Decimal("20")),
    "AREA_3": ("QA_RATE_3", Decimal("30")),
    "AREA_4": ("QA_RATE_4", Decimal("40")),
    "AREA_5": ("QA_RATE_5", Decimal("50")),
}
EXPECTED_QUANTITIES = {
    "AREA_1": Decimal("20"),
    "AREA_2": Decimal("20"),
    "AREA_3": Decimal("19.75"),  # explicit 10x10 px opening on floor 3
    "AREA_4": Decimal("20"),
    "AREA_5": Decimal("20"),
}
EXPECTED_ESTIMATE_IRR = Decimal("2992.5")


def _rates_file(tmp_path):
    path = tmp_path / "independent-reviewed-rates-1405.csv"
    lines = ["year,group,chapter,code,description,unit,unit_price"]
    for code, (rate_code, unit_rate) in RATES.items():
        lines.append(
            f"1405,ابنیه,01,{rate_code},Synthetic {code},m2,{unit_rate}"
        )
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def _independent_boq_assertions(project):
    assert {row["item_code"] for row in project["boq"]} == set(RATES)
    assert len(project["boq"]) == len(RATES)
    assert len(project["takeoffs"]) == 5
    for row in project["boq"]:
        code = row["item_code"]
        assert Decimal(str(row["quantity"])) == EXPECTED_QUANTITIES[code]
        assert row["unit"] == "m2"
        assert len(row.get("source_ids") or []) == 1
    assert sum((Decimal(str(row["quantity"])) for row in project["boq"]), Decimal(0)) == Decimal("99.75")
    assert audit_boq(project)["reconciled"] is True


def _price_and_estimate(app, tmp_path):
    source = _rates_file(tmp_path)
    review = app.import_project_pricebook(
        "P5", source, year=1405,
        approved_by="independent-test-price-reviewer",
        currency="IRR", source_version="QA-1405-1",
    )
    assert review["sha256"] == sha256(source.read_bytes()).hexdigest()
    original = app.open_project("P5")
    _independent_boq_assertions(original)
    mapping = {
        index: RATES[row["item_code"]][0]
        for index, row in enumerate(original["boq"])
    }
    updated = app.review_boq_price_mapping(
        "P5", mapping,
        reviewer="independent-test-boq-reviewer",
        expected_digest=app.store.current_digest("P5"),
    )
    assert len(updated) == 5
    priced = app.open_project("P5")
    assert audit_boq(priced, require_pricebook=True)["reconciled"] is True
    for row in priced["boq"]:
        rate_code, rate = RATES[row["item_code"]]
        assert row["price_code"] == rate_code
        assert Decimal(str(row["unit_price"])) == rate
    estimate = app.recalculate_estimate("P5")
    assert estimate["finalizable"] is True
    assert Decimal(str(estimate["cost"]["grand_total"])) == EXPECTED_ESTIMATE_IRR
    assert estimate["financial_policy_status"] == "LEGACY_PREVIEW_NOT_FINANCIALLY_APPROVED"
    return source, estimate


def test_two_real_pdf_sources_five_floors_to_boq_price_estimate_statement_and_reports(building, tmp_path):
    """One continuous SQLite project, source-backed and with independent arithmetic."""
    import pymupdf
    from openpyxl import load_workbook
    app, project, selections = building
    _independent_boq_assertions(project)
    assert {str(item["system"]) for item in project["takeoffs"]} == {
        f"L{n}" for n in range(1, 6)
    }
    source_digests = {
        project["drawing_sessions"][x["session_id"]]["source_sha256"] for x in selections
    }
    assert len(source_digests) == 2

    before = deepcopy(project)
    initial_digest = app.store.current_digest("P5")
    with pytest.raises(ValueError):
        app.recalculate_estimate("P5")
    assert app.store.current_digest("P5") == initial_digest
    assert app.open_project("P5") == before

    # Export the source-verified drawing bundle before changing its project revision.
    destination = tmp_path / "reviewed-2d-five-floor"
    bundle = export(building, destination)
    manifest = json.loads((destination / "manifest.json").read_text(encoding="utf-8"))
    assert bundle["manifest"] == manifest
    assert len(manifest["pages"]) == 5
    assert len(manifest["files"]) == 14
    assert all(
        sha256((destination / row["name"]).read_bytes()).hexdigest() == row["sha256"]
        for row in manifest["files"]
    )
    assert app.store.current_digest("P5") == initial_digest

    price_source, estimate = _price_and_estimate(app, tmp_path)
    assert Decimal(str(estimate["cost"]["base"])) == EXPECTED_ESTIMATE_IRR

    current_quantities = {
        RATES[row["item_code"]][0]: str(EXPECTED_QUANTITIES[row["item_code"]])
        for row in app.open_project("P5")["boq"]
    }
    period = app.save_statement_period(
        "P5", period_no=1, current_quantities=current_quantities, persist=True
    )
    assert Decimal(str(period["gross_current"])) == EXPECTED_ESTIMATE_IRR
    assert period["financial_policy_status"] == "LEGACY_RECORD_NOT_FINANCIALLY_APPROVED"
    assert app.open_project("P5")["statement_periods"][0]["number"] == 1

    # A period draft is not an approved payment certificate.
    with pytest.raises(ValueError):
        app.office_aec_report(
            "P5", "statement", "pdf", tmp_path / "unapproved-statement.pdf",
            statement=period,
        )
    assert not (tmp_path / "unapproved-statement.pdf").exists()

    for fmt in ("pdf", "xlsx"):
        path = tmp_path / f"reviewed-estimate.{fmt}"
        app.office_aec_report("P5", "estimate", fmt, path)
        assert path.is_file() and path.stat().st_size > 200
        if fmt == "pdf":
            with pymupdf.open(path) as document:
                assert len(document) >= 1
                assert document[0].get_pixmap().width > 0
        else:
            workbook = load_workbook(path, read_only=True)
            assert len(workbook.sheetnames) >= 1
            workbook.close()

    final = deepcopy(app.open_project("P5"))
    final_digest = app.store.current_digest("P5")
    assert audit_boq(final, require_pricebook=True)["reconciled"]
    assert sum(
        (Decimal(str(x["quantity"])) for x in final["boq"]), Decimal(0)
    ) == Decimal("99.75")
    # Reopen a second real SQLite handle before the original fixture closes.
    second = StructuralProApp(tmp_path / "store")
    try:
        assert second.open_project("P5") == final
        assert second.store.current_digest("P5") == final_digest
        assert second.store.db.execute("pragma integrity_check").fetchone()[0] == "ok"
        assert not second.store.db.execute("pragma foreign_key_check").fetchall()
    finally:
        second.store.close()
    assert sha256(price_source.read_bytes()).hexdigest() == final["user_pricebook_review"]["sha256"]


def test_modified_user_pricebook_blocks_both_estimate_and_export_without_sqlite_write(building, tmp_path):
    app, original, _ = building
    source, estimate = _price_and_estimate(app, tmp_path)
    committed = deepcopy(app.open_project("P5"))
    digest = app.store.current_digest("P5")
    source.write_text(
        source.read_text(encoding="utf-8") + "1405,ابنیه,01,UNREVIEWED,Unreviewed,m2,1\n",
        encoding="utf-8",
    )
    with pytest.raises(ValueError):
        app.recalculate_estimate("P5")
    with pytest.raises(ValueError):
        app.office_aec_report("P5", "estimate", "pdf", tmp_path / "blocked.pdf")
    assert not (tmp_path / "blocked.pdf").exists()
    assert app.open_project("P5") == committed
    assert app.store.current_digest("P5") == digest


def test_unreviewed_financial_mapping_never_turns_boq_into_assumed_estimate(building, tmp_path):
    app, original, _ = building
    _independent_boq_assertions(original)
    digest = app.store.current_digest("P5")
    rates = _rates_file(tmp_path)
    # No approved user import: even a plausible CSV on disk cannot authorize price.
    with pytest.raises(ValueError):
        app.review_boq_price_mapping(
            "P5", {0: "QA_RATE_1"},
            reviewer="independent-test-boq-reviewer",
            expected_digest=digest,
        )
    assert app.store.current_digest("P5") == digest
    assert app.open_project("P5") == original
    assert not app.open_project("P5").get("user_pricebook_review")
