"""Independent five-floor scenario, never a substitute for historical ownership."""
from pathlib import Path
import pytest

def test_five_floor_acceptance(tmp_path):
 from scripts.verify_stair_stage3 import verify
 result=verify(None,tmp_path/'scenario')
 assert result['explicitly_synthetic'] and result['boq_rows']==35
 assert result['effective_versions']=={'L1':2,'L2':2,'L3':3,'L4':2,'L5':2}
 assert result['steel_by_floor_kg']=={'L1':'98.5680','L2':'98.5680','L3':'103.8960','L4':'98.5680','L5':'98.5680'}
 assert result['total_steel_kg']=='498.1680'
 assert result['original_takeoffs_preserved'] and result['other_floors_unchanged']
 # Reopen all four effective-version reports; amended and original values coexist.
 report=tmp_path/'scenario/reports'
 import csv
 from openpyxl import load_workbook
 from docx import Document
 import fitz
 from decimal import Decimal
 csvfile=report/'stair-L3.csv';assert csvfile.read_bytes().startswith(b'\xef\xbb\xbf');text=csvfile.read_text(encoding='utf-8-sig');assert list(csv.reader(text.splitlines()))
 assert all(x in text for x in ('اصلاحیه اعمال‌شده','19.7136','25.0416','5.328','103.896'))
 w=load_workbook(report/'stair-L3.xlsx');assert all(sheet.sheet_view.rightToLeft for sheet in w)
 cells=[c for sheet in w for row in sheet for c in row if c.value is not None]
 assert any(c.data_type=='n' and Decimal(str(c.value))==Decimal('25.0416') for c in cells)
 d=Document(report/'stair-L3.docx');word='\n'.join(c.text for table in d.tables for row in table.rows for c in row.cells);assert 'اصلاحیه اعمال‌شده' in word and '25.0416' in word
 pdf=fitz.open(report/'stair-L3.pdf');assert len(pdf)>1;pdftext=''.join(page.get_text() for page in pdf);assert '25.0416' in pdftext and '103.896' in pdftext


@pytest.mark.parametrize('bad',['concrete','forms','bars','overlap'])
def test_cross_floor_duplicate_blocked(tmp_path,bad):
 from core.platform.application import StructuralProApp
 from core.takeoff.stair_detail import project_gate
 from core.takeoff.reviewed_aec import _id
 from test_stair_work_stage2 import sample,commit,refresh
 import json
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');d=sample(tmp_path);commit(a,d);folder=tmp_path/'other';folder.mkdir();other=sample(folder,sid='ST2',floor='L2');old=json.loads(Path(d['stair_detail']['ownership_document']['path']).read_text());o=json.loads(Path(other['stair_detail']['ownership_document']['path']).read_text())
 if bad!='overlap':
  op={'concrete':'waist_concrete','forms':'formwork','bars':'rebar_flight_main'}[bad]
  x=next(x for x in o['allocations'] if x['operation']==op);original=next(x for x in old['allocations'] if x['operation']==op)
  if bad=='bars':x['bar_pieces'][0]=original['bar_pieces'][0]
  else:x['scope_id']=original['scope_id']
 refresh(other,folder,o);a.save_aec_system_draft('P',other)
 with pytest.raises(ValueError):project_gate(a.open_project('P'),_id('stair','L2','ST2'))
 a.store.close()
