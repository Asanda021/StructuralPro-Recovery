"""Explicit synthetic rectangular schedules; never a design or real approval."""
from copy import deepcopy
from decimal import Decimal
from hashlib import sha256
import pytest
from core.platform.application import StructuralProApp
from core.takeoff.reviewed_aec import preview_system
from core.takeoff.engine import TakeoffEngine

KINDS=('solid_slab','joist_block','joist_foam')

def sample(kind='solid_slab'):
 from core.takeoff.reviewed_aec import empty_system
 d=empty_system(kind,'L1','S1');d.update(measurement_policy='roof_specialist_v1',building_id='B1',source={'mode':'manual','source_id':'SYNTHETIC-'+kind,'plan_ref':'پلان مصنوعی','detail_ref':'دیتایل مصنوعی'})
 na=lambda:{'status':'not_applicable','approved_by':'کنترل مصنوعی','source_ref':'دیتایل مصنوعی','reason':'عدم کاربرد صرفاً در نمونه مصنوعی'}
 box=lambda i,c,x,y,z,l,w,h:{'id':i,'component':c,'x':str(x),'y':str(y),'z':str(z),'length':str(l),'width':str(w),'height':str(h),'source_ref':'دیتایل مصنوعی','approved_by':'کنترل مصنوعی'}
 g={'unit':'m','shape':'rectangular','length':'4','width':'4','depth':'.25' if kind!='solid_slab' else '.2','count':'1','repeat_scope':'identical_roofs','action':'two_way','source_ref':'پلان مصنوعی','approved_by':'کنترل مصنوعی'}
 regions=[box('C1','concrete',0,0,0,4,4,.2)]
 pieces=[];layout=[]
 if kind!='solid_slab':
  g.update(joist_supply_scope='purchased_joists',topping_thickness='.05',joist_axis='x',joist_count='2',joist_spacing='2')
  regions=[box('T','topping',0,0,.2,4,4,.05)]
  for i,y in enumerate([0,2]):
   j=box('J'+str(i),'',0,y,0,4,.2,.2);j.pop('component');layout.append(j)
   regions += [box('R'+str(i),'ribs',0,y,.05,4,.2,.15),box('P'+str(i),'precast',0,y,0,4,.2,.05),box('V'+str(i),'void',0,y+.2,0,4,1.8,.2)]
   pieces.append({'id':'I'+str(i),'region_id':'V'+str(i),'count':'2','length':'2','width':'1.8','height':'.2','nx':'2','ny':'1','nz':'1','source_ref':'چیدمان مصنوعی','spec_ref':'قطعه مصنوعی'})
 bars={k:na() for k in ('main_x','main_y','distribution','thermal','rib','additional','opening','solid','tie')}
 for k in (('main_x','main_y') if kind=='solid_slab' else ('rib','thermal')):
  bars[k]={'status':'confirmed','approved_by':'کنترل مصنوعی','source_ref':'دیتایل مصنوعی','series':[{'id':'B-'+k,'physical_key':'physical-'+k,'diameter_mm':'12','supply_scope':'site','count':'10','cut_length_m':'4','unit_weight_kg_m':'.888','unit_weight_source':'جدول مصنوعی صریح','shape_ref':'برش مصنوعی','source_ref':'دیتایل مصنوعی','cut_parts':[{'id':'straight','length_m':'4','source_ref':'طول مصنوعی'}]}]}
 d['roof_detail']={'geometry':g,'regions':regions,'joist_layout':layout,'deductions':[],'deductions_ref':'عدم بازشو مصنوعی','rebars':bars,'pieces':pieces,'pieces_ref':'چیدمان مصنوعی','formwork':{'status':'confirmed','approved_by':'کنترل مصنوعی','source_ref':'قالب مصنوعی','faces':[{'id':'F1','physical_key':'underside','length_m':'4','width_m':'4','deduction_m2':'0','source_ref':'قالب مصنوعی'}]},'shoring':na()}
 return d

@pytest.mark.parametrize('kind',KINDS)
def test_exact_component_reference(kind):
 r=preview_system(sample(kind),TakeoffEngine());assert r['ready'],r['blockers']
 totals={}
 for row in r['rows']:totals[row['item_code']]=totals.get(row['item_code'],Decimal(0))+Decimal(str(row['quantity']))
 assert totals['concrete']==Decimal('3.2' if kind=='solid_slab' else '1.04')
 assert totals['rebar']==Decimal('71.040') and totals['formwork']==16
 if kind!='solid_slab':assert totals['precast_concrete']==Decimal('.08') and totals['joists']==2 and totals['blocks' if kind=='joist_block' else 'foam']==4

@pytest.mark.parametrize('fault',['unit','missing','overlap','duplicate','bars','cut','partition','piece','joist','review'])
def test_fail_closed(fault):
 d=sample('joist_block');v=d['roof_detail']
 if fault=='unit':v['geometry']['unit']='cm'
 if fault=='missing':del v['geometry']['depth']
 if fault=='overlap':v['regions'][2]['y']='0.1'
 if fault=='duplicate':v['regions'].append(deepcopy(v['regions'][0]))
 if fault=='bars':v['rebars']['rib']['series'][0].pop('shape_ref')
 if fault=='cut':v['rebars']['rib']['series'][0]['cut_length_m']='5'
 if fault=='partition':v['regions'].pop()
 if fault=='piece':v['pieces'][0]['count']='3'
 if fault=='joist':v['geometry']['joist_count']='3'
 if fault=='review':v['rebars']['rib']['status']='unreviewed'
 assert not preview_system(d,TakeoffEngine())['ready']


def setup(tmp_path,kind):
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');d=sample(kind)
 a.save_aec_system_draft('P',d,expected_revision=0)
 a.review_roof_system('P',kind,'L1','S1',expected_revision=1,reviewer='بازبین مصنوعی',note='کنترل مستقل دیتایل')
 a.commit_aec_system('P',kind,'L1','S1',expected_revision=2)
 return a,d

@pytest.mark.parametrize('kind',KINDS)
def test_sqlite_amendment_and_boq(tmp_path,kind):
 a,d=setup(tmp_path,kind);before=deepcopy(a.open_project('P'));d=deepcopy(a.get_aec_system('P',kind,'L1','S1')['data']);d.pop('roof_review',None);d['roof_detail']['geometry']['count']='2'
 path=tmp_path/'new.txt';path.write_text('explicit synthetic revised detail')
 kw=dict(expected_revision=2,created_by='سازنده مصنوعی',reason='تغییر تکرار مستند',evidence={'evidence_path':str(path),'evidence_sha256':sha256(path.read_bytes()).hexdigest(),'evidence_reference':'مدرک مصنوعی'})
 am=a.create_roof_amendment('P',kind,'L1','S1',d,**kw)
 with pytest.raises(ValueError):a.create_roof_amendment('P',kind,'L1','S1',d,**kw)
 with pytest.raises(ValueError):a.review_beam_amendment('P',am['id'],expected_revision=1,reviewer=kw['created_by'],note='ممنوع')
 a.review_beam_amendment('P',am['id'],expected_revision=1,reviewer='مستقل دیگر',note='کنترل مصنوعی')
 a.apply_beam_amendment('P',am['id'],expected_revision=1)
 p=a.open_project('P');assert p['reviewed_aec_systems']==before['reviewed_aec_systems'] and p['takeoffs']==before['takeoffs']
 for old,new in zip(before['boq'],p['boq']):assert Decimal(str(new['quantity']))==2*Decimal(str(old['quantity']))
 with pytest.raises(ValueError):a.apply_beam_amendment('P',am['id'],expected_revision=1)
 assert a.get_aec_system('P',kind,'L1','S1')['revision']==3
 a.store.close();a=StructuralProApp(tmp_path/'db');assert a.get_aec_system('P',kind,'L1','S1')['data']['roof_detail']['geometry']['count']=='2'
 assert a.store.db.execute('pragma integrity_check').fetchone()[0]=='ok';a.store.close()

@pytest.mark.parametrize('case',['concurrent','source','commit_review','financial','immutable','evidence','identity','stale_amendment','policy','duplicate_bar'])
def test_workflow_guards(tmp_path,case):
 a,d=setup(tmp_path,'solid_slab')
 if case=='immutable':
  p=a.open_project('P');next(iter(p['reviewed_aec_systems'].values()))['data']['roof_detail']['geometry']['count']='3'
  with pytest.raises(ValueError):a.store.save('P',p)
 elif case in ('concurrent','source','commit_review','policy'):
  b=sample();b['system_id']='S2';b['source']['source_id']='new'
  a.save_aec_system_draft('P',b,expected_revision=0)
  if case=='concurrent':
   b['building_id']='B2';a.save_aec_system_draft('P',b,expected_revision=1)
   with pytest.raises(RuntimeError):a.save_aec_system_draft('P',b,expected_revision=1)
  elif case=='source':
   b['system_id']='S3'
   with pytest.raises(ValueError):a.save_aec_system_draft('P',b,expected_revision=0)
  elif case=='policy':
   b.pop('measurement_policy')
   with pytest.raises(ValueError):a.save_aec_system_draft('P',b,expected_revision=1)
  else:
   with pytest.raises(ValueError):a.commit_aec_system('P','solid_slab','L1','S2',expected_revision=1)
 elif case=='duplicate_bar':
  b=sample();b['roof_detail']['rebars']['main_y']['series'][0]['physical_key']='physical-main_x'
  assert not preview_system(b,a.takeoff)['ready']
 else:
  b=deepcopy(a.get_aec_system('P','solid_slab','L1','S1')['data']);b.pop('roof_review');b['roof_detail']['geometry']['count']='2'
  path=tmp_path/'rev.txt';path.write_text('synthetic revision');e={'evidence_path':str(path),'evidence_sha256':sha256(path.read_bytes()).hexdigest(),'evidence_reference':'مدرک جدید'}
  kw=dict(expected_revision=2,created_by='مصنوعی',reason='مصنوعی',evidence=e)
  if case=='identity':
   b['building_id']='B2'
   with pytest.raises(ValueError):a.create_roof_amendment('P','solid_slab','L1','S1',b,**kw)
  else:
   am=a.create_roof_amendment('P','solid_slab','L1','S1',b,**kw)
   if case=='evidence':
    path.write_text('changed')
    with pytest.raises(ValueError):a.review_beam_amendment('P',am['id'],expected_revision=1,reviewer='دیگری',note='کنترل')
   else:
    a.review_beam_amendment('P',am['id'],expected_revision=1,reviewer='دیگری',note='کنترل')
    if case=='financial':
     p=a.open_project('P');p['contract']={'id':'locked'};a.store.save('P',p)
     with pytest.raises(ValueError):a.apply_beam_amendment('P',am['id'],expected_revision=1)
    else:
     a.apply_beam_amendment('P',am['id'],expected_revision=1)
     with pytest.raises(RuntimeError):a.create_roof_amendment('P','solid_slab','L1','S1',b,**kw)
 a.store.close()

@pytest.mark.parametrize('kind',KINDS)
def test_copy_delete_and_reports(tmp_path,kind):
 a,d=setup(tmp_path,kind)
 c=a.aec_workflow().copy_roof('P',kind,'L1','S1',new_system_id='S2',new_source_id='COPY',reason='کپی مصنوعی')
 assert not a.preview_aec_system('P',kind,'L1','S2')['ready']
 a.aec_workflow().delete_roof('P',kind,'L1','S2',expected_revision=c['revision'],reason='حذف پیش‌نویس مصنوعی')
 with pytest.raises(ValueError):a.aec_workflow().delete_roof('P',kind,'L1','S1',expected_revision=2,reason='ممنوع')
 for fmt in ('pdf','xlsx','docx','csv'):
  path=tmp_path/('roof.'+fmt);a.export_roof_detail_report('P',kind,'L1','S1',path,format=fmt);assert path.stat().st_size>100
  if fmt=='pdf':
   import pymupdf
   with pymupdf.open(path) as r:assert all(p.get_text().strip() for p in r)
  elif fmt=='xlsx':
   from openpyxl import load_workbook
   w=load_workbook(path);assert all(s.sheet_view.rightToLeft for s in w)
  elif fmt=='docx':
   from docx import Document
   assert Document(path).tables
  else:assert 'CSV' in path.read_text(encoding='utf-8-sig')
 a.store.close()

@pytest.mark.parametrize('case',['unknown_geometry','unknown_detail','hidden_bars','physical_roof','drawing_scale','duplicate_store_amendment'])
def test_specialist_boundary_regressions(tmp_path,case):
 d=sample()
 if case=='unknown_geometry':
  d['roof_detail']['geometry']['slope']='.1';assert not preview_system(d,TakeoffEngine())['ready']
 elif case=='unknown_detail':
  d['roof_detail']['extra_deductions']=[{'volume':'1'}];assert not preview_system(d,TakeoffEngine())['ready']
 elif case=='hidden_bars':
  d['roof_detail']['rebars']['distribution']['series']=deepcopy(d['roof_detail']['rebars']['main_x']['series']);assert not preview_system(d,TakeoffEngine())['ready']
 elif case=='physical_roof':
  a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');a.save_aec_system_draft('P',d,expected_revision=0)
  other=sample('joist_block')
  with pytest.raises(ValueError):a.save_aec_system_draft('P',other,expected_revision=0)
  a.store.close()
 elif case=='drawing_scale':
  from core.takeoff.roof_detail import validate_roof_source
  d['source'].update(mode='drawing')
  with pytest.raises(ValueError,match='مقیاس'):validate_roof_source(d['source'],'S1')
 else:
  a,d=setup(tmp_path,'solid_slab');d.pop('roof_review',None);d['roof_detail']['geometry']['count']='2';p=tmp_path/'new.txt';p.write_text('synthetic')
  am=a.create_roof_amendment('P','solid_slab','L1','S1',d,expected_revision=2,created_by='مصنوعی',reason='تغییر',evidence={'evidence_path':str(p),'evidence_sha256':sha256(p.read_bytes()).hexdigest(),'evidence_reference':'جدید'})
  project=a.open_project('P');clone=deepcopy(am);clone['id']='forged-duplicate';project['beam_amendments'][clone['id']]=clone
  with pytest.raises(ValueError):a.store.save('P',project,manual_formula_command={'action':'beam_create','id':clone['id']})
  a.store.close()


def test_spatial_opening_across_regions_and_actual_two_connections(tmp_path):
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');d=sample();v=d['roof_detail'];r=v['regions'][0];r['length']='2';r2=deepcopy(r);r2.update(id='C2',x='2');v['regions'].append(r2)
 opening={'id':'OPEN-1','kind':'opening','x':'1.5','y':'1','z':'0','length':'1','width':'1','height':'.2','source_ref':'بازشوی مصنوعی','approved_by':'مصنوعی'};v['deductions']=[opening]
 group=deepcopy(v['rebars']['main_x']);group['series'][0].update(id='B-opening',physical_key='opening-bars');v['rebars']['opening']=group;v['formwork']['faces'][0]['deduction_m2']='1'
 result=preview_system(d,a.takeoff);assert result['ready'],result['blockers'];concrete=[r for r in result['rows'] if r['item_code']=='concrete'];assert sum(Decimal(str(r['quantity'])) for r in concrete)==Decimal('3.0');assert [r['params']['deducted_m3'] for r in concrete]==[.1,.1]
 a.save_aec_system_draft('P',d,expected_revision=0);b=StructuralProApp(tmp_path/'db');assert b.get_aec_system('P','solid_slab','L1','S1')['revision']==1
 a.review_roof_system('P','solid_slab','L1','S1',expected_revision=1,reviewer='مستقل',note='کنترل')
 with pytest.raises(RuntimeError):b.save_aec_system_draft('P',d,expected_revision=1)
 b.store.close();a.store.close()


def test_report_records_cut_sources_na_and_summary(tmp_path):
 a,d=setup(tmp_path,'solid_slab');path=tmp_path/'detail.xlsx';a.export_roof_detail_report('P','solid_slab','L1','S1',path,format='xlsx')
 from openpyxl import load_workbook
 w=load_workbook(path);text=' '.join(str(c.value) for s in w for row in s for c in row)
 assert 'طول مصنوعی' in text and 'عدم کاربرد صرفاً در نمونه مصنوعی' in text
 assert w['خلاصه'].max_row>3
 a.store.close()


def test_roof_cannot_use_beam_migration_path(tmp_path):
 a,d=setup(tmp_path,'solid_slab')
 with pytest.raises(ValueError,match='مهاجرت'):a.create_roof_amendment('P','solid_slab','L1','S1',d,expected_revision=2,created_by='مصنوعی',reason='ممنوع',migration={})
 a.store.close()


def test_purchased_joist_constituents_not_paid_twice(tmp_path):
 a,d=setup(tmp_path,'joist_block');p=a.open_project('P')
 assert not any(r['item_code']=='precast_concrete' for r in p['boq'])
 a.store.close()


def test_transverse_tie_partition_and_factory_rebar_role(tmp_path):
 d=sample('joist_block');v=d['roof_detail'];new=[]
 for r in v['regions']:
  if r['component']=='ribs':
   left=deepcopy(r);left['length']='1.5';right=deepcopy(r);right.update(id=r['id']+'-right',x='2',length='2');new.extend([left,right])
  elif r['component']=='void':
   for suffix,x,length,z,height in [('left','0','1.5','0','.2'),('under','1.5','.5','0','.05'),('right','2','2','0','.2')]:
    part=deepcopy(r);part.update(id=r['id']+'-'+suffix,x=x,length=length,z=z,height=height);new.append(part)
  else:new.append(r)
 new.append({'id':'TIE','component':'ties','x':'1.5','y':'0','z':'.05','length':'.5','width':'4','height':'.15','source_ref':'کلاف مصنوعی','approved_by':'مصنوعی'})
 v['regions']=new;v['pieces']=[]
 for r in new:
  if r['component']=='void':v['pieces'].append({'id':'piece-'+r['id'],'region_id':r['id'],'count':'1','nx':'1','ny':'1','nz':'1','length':r['length'],'width':r['width'],'height':r['height'],'source_ref':'چیدمان و برش مصنوعی','spec_ref':'قطعه برش‌خورده مستند مصنوعی'})
 group=deepcopy(v['rebars']['rib']);group['series'][0].update(id='tie-bars',physical_key='tie-bars');v['rebars']['tie']=group
 v['rebars']['rib']['series'][0]['supply_scope']='factory'
 result=preview_system(d,TakeoffEngine());assert result['ready'],result['blockers']
 assert sum(Decimal(str(r['quantity'])) for r in result['rows'] if r['item_code']=='concrete')==Decimal('1.31')
 factory=next(r for r in result['rows'] if r['component']=='rib');assert factory['role']=='material'
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');a.save_aec_system_draft('P',d,expected_revision=0);a.review_roof_system('P','joist_block','L1','S1',expected_revision=1,reviewer='مستقل',note='کنترل');a.commit_aec_system('P','joist_block','L1','S1',expected_revision=2)
 assert not any(factory['source_id'] in r['source_ids'] for r in a.open_project('P')['boq']);a.store.close()

@pytest.mark.parametrize('case',['fractional_count','negative_dimension','bad_weight','missing_group','overlap_deductions','opening_rebar','beam_owner'])
def test_extra_invalid_boundaries(case):
 d=sample();v=d['roof_detail']
 if case=='fractional_count':v['geometry']['count']='1.5'
 elif case=='negative_dimension':v['regions'][0]['height']='-.2'
 elif case=='bad_weight':v['rebars']['main_x']['series'][0]['unit_weight_kg_m']='NaN'
 elif case=='missing_group':del v['rebars']['distribution']
 else:
  opening={'id':'O1','kind':'opening' if case!='beam_owner' else 'beam_joint','x':'1','y':'1','z':'0','length':'1','width':'1','height':'.2','source_ref':'مصنوعی','approved_by':'مصنوعی'};v['deductions']=[opening]
  if case=='overlap_deductions':other=deepcopy(opening);other['id']='O2';v['deductions'].append(other)
 assert not preview_system(d,TakeoffEngine())['ready']


def test_amendment_report_numeric_negative_delta_and_actual_reviewer(tmp_path):
 a,d=setup(tmp_path,'solid_slab');d['roof_detail']['rebars']['main_x']['series'][0]['count']='9';path=tmp_path/'doc.txt';path.write_text('explicit synthetic cut schedule')
 am=a.create_roof_amendment('P','solid_slab','L1','S1',d,expected_revision=2,created_by='مصنوعی',reason='تغییر صریح دیتایل',evidence={'evidence_path':str(path),'evidence_sha256':sha256(path.read_bytes()).hexdigest(),'evidence_reference':'مدرک جدید'})
 a.review_beam_amendment('P',am['id'],expected_revision=1,reviewer='بازبین واقعیِ نمونه اصلاحیه',note='کنترل مصنوعی');a.apply_beam_amendment('P',am['id'],expected_revision=1)
 p=tmp_path/'changed.xlsx';a.export_roof_detail_report('P','solid_slab','L1','S1',p,format='xlsx')
 from openpyxl import load_workbook
 w=load_workbook(p);sheet=w['ریز متره'];header=next(row for row in sheet if any(c.value=='شرح' for c in row));labels=[c.value for c in header]
 assert 'عنوان' not in labels and 'مقدار و مرجع' not in labels
 assert any(c.data_type=='n' and c.value==-3.552 for row in sheet for c in row)
 assert 'بازبین واقعیِ نمونه اصلاحیه' in ' '.join(str(c.value) for row in w['اطلاعات'] for c in row)
 a.store.close()
