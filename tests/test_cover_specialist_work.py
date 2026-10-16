"""Explicit synthetic covering layouts, not design/product acceptance."""
from copy import deepcopy
from decimal import Decimal
import pytest
from core.takeoff.reviewed_aec import empty_system,preview_system
from core.takeoff.engine import TakeoffEngine
from test_architecture_specialist_work import obj
KINDS=('ceramic_finish','stone_floor','stone_cladding')
def sample(kind='ceramic_finish'):
 d=empty_system(kind,'L1','C1');d.update(measurement_policy='architecture_specialist_v1',building_id='B1',axis='فضای مصنوعی',source=dict(mode='manual',source_id='SYNTHETIC-COVER',plan_ref='پلان مصنوعی',detail_ref='چیدمان مصنوعی'))
 g=dict(physical_surface_id='SURFACE-1',host_id='HOST-1',space_id='ROOM-1',face='floor',plane='horizontal',shape='rectangle',unit='m',length_m='3',width_m='2',count='1',repeat_ref='یک سطح مستقل مصنوعی',source_ref='هندسه مصنوعی',approved_by='کنترل مصنوعی')
 s=obj('LAYER',layer_key='FINISH-1',operation='ceramic_surface' if kind=='ceramic_finish' else 'stone_surface',x='0',y='0',width_m='3',height_m='2',material_ref='مشخصات قطعه مصنوعی',payment_unit='m2',payment_ref='پرداخت مستقل سطح نصب؛ قطعات و خرید اطلاعاتی',opening_ids='H1',deduction_ref='کسر بازشو',coverage_scope='complete',scope_ref='تمام سطح',layout_status='documented',layout_ref='قطعه‌بندی کامل و بند صفر مستند مصنوعی')
 a=dict(contract='interior_cover_v1',geometry=g,openings=[obj('H1',x='2',y='1',width_m='1',height_m='1')],surfaces=[s],stocks=[obj('ST1',length_m='2',width_m='2',material_ref=s['material_ref']),obj('ST2',length_m='2',width_m='1',material_ref=s['material_ref'])],pieces=[obj('P1',surface_id='LAYER',stock_id='ST1',piece_type='full',x='0',y='0',width_m='2',height_m='2',stock_x='0',stock_y='0',cut_ref='قطعه کامل مصنوعی'),obj('P2',surface_id='LAYER',stock_id='ST2',piece_type='cut',x='2',y='0',width_m='1',height_m='1',stock_x='0',stock_y='0',cut_ref='نقشه برش مصنوعی')],joints=[],consumptions=[])
 d['architecture_detail']=a
 if kind=='stone_cladding':g.update(face='inside',plane='vertical')
 return d

def result(d):return preview_system(d,TakeoffEngine())
@pytest.mark.parametrize('kind',KINDS)
def test_independent_area_layout_purchase(kind):
 r=result(sample(kind));assert r['ready'],r['blockers'];q={}
 for v in r['rows']:q[v['item_code']]=q.get(v['item_code'],Decimal(0))+Decimal(str(v['quantity']))
 assert q=={'ceramic_surface' if kind=='ceramic_finish' else 'stone_surface':Decimal(5),'installed_pieces':Decimal(2),'cut_pieces':Decimal(1),'purchase_stock':Decimal(2),'purchase_area':Decimal(6),'offcut_area':Decimal(1)};assert sum(v['role']=='work' for v in r['rows'])==1

@pytest.mark.parametrize('fault',['piece_overlap','stock_overlap','missing_cut','outside','hole','duplicate_hole','missing_stock','missing_source','unknown_waste','wrong_unit','missing_piece','no_layout_with_pieces','negative','plane','face','duplicate_id','duplicate_physical','stock_material','full_piece_cut','contradictory_coverage'])
def test_invalid_cover_fails_closed(fault):
 d=sample();a=d['architecture_detail'];p=a['pieces'][1]
 if fault=='piece_overlap':p['x']='1'
 if fault=='stock_overlap':p.update(stock_id='ST1',stock_x='0',stock_y='0')
 if fault=='missing_cut':p['cut_ref']=''
 if fault=='outside':p['x']='3'
 if fault=='hole':p['y']='1'
 if fault=='duplicate_hole':a['openings'].append({**a['openings'][0],'id':'H2','physical_key':'H2'})
 if fault=='missing_stock':p['stock_id']='NO'
 if fault=='missing_source':a['surfaces'][0]['material_ref']=''
 if fault=='unknown_waste':a['surfaces'][0]['waste']='.1'
 if fault=='wrong_unit':a['surfaces'][0]['payment_unit']='عدد'
 if fault=='missing_piece':a['pieces'].pop()
 if fault=='no_layout_with_pieces':a['surfaces'][0]['layout_status']='area_only'
 if fault=='negative':p['width_m']='-1'
 if fault=='plane':a['geometry']['plane']='curved'
 if fault=='face':a['geometry']['face']='inside'
 if fault=='duplicate_id':p['id']='P1'
 if fault=='duplicate_physical':p['physical_key']=a['pieces'][0]['physical_key']
 if fault=='stock_material':a['stocks'][1]['material_ref']='متفاوت'
 if fault=='full_piece_cut':p['piece_type']='full'
 if fault=='contradictory_coverage':a['surfaces'][0]['width_m']='2'
 assert not result(d)['ready']

@pytest.mark.parametrize('kind',KINDS)
def test_area_only_does_not_invent_counts(kind):
 d=sample(kind);a=d['architecture_detail'];a['pieces']=[];a['stocks']=[];a['surfaces'][0]['layout_status']='area_only';r=result(d);assert r['ready'],r['blockers'];assert len(r['rows'])==1 and r['rows'][0]['quantity']==5

def test_vertical_face_is_independent():
 d=sample();d['architecture_detail']['geometry'].update(face='inside',plane='vertical');assert result(d)['ready']

@pytest.mark.parametrize('kind',KINDS)
@pytest.mark.parametrize('case',['valid','cas','duplicate_source','duplicate_surface','legacy','same_reviewer','financial','evidence','unit','identity','copy'])
def test_workflow_store_history_boq(tmp_path,kind,case):
 from core.platform.application import StructuralProApp
 from hashlib import sha256
 a=StructuralProApp(tmp_path/'db');a.create_project('آزمون مصنوعی پوشش','P');d=sample(kind)
 if case=='legacy':
  old=empty_system(kind,'L1','C1');old['source']={'mode':'manual','source_id':'OLD'};a.save_aec_system_draft('P',old,expected_revision=0)
  with pytest.raises(ValueError):a.save_aec_system_draft('P',d,expected_revision=1)
  a.store.close();return
 a.save_aec_system_draft('P',d,expected_revision=0)
 if case=='cas':
  b=StructuralProApp(tmp_path/'db');stale=b.get_aec_system('P',kind,'L1','C1');a.save_aec_system_draft('P',d,expected_revision=1)
  with pytest.raises(RuntimeError):b.save_aec_system_draft('P',stale['data'],expected_revision=1)
  b.store.close();a.store.close();return
 if case in ('duplicate_source','duplicate_surface'):
  other=deepcopy(d);other['system_id']='OTHER'
  if case=='duplicate_source':other['architecture_detail']['geometry'].update(host_id='OTHER',physical_surface_id='OTHER')
  else:other['source']['source_id']='OTHER'
  for table in ('openings','surfaces','pieces','stocks'):
   for p in other['architecture_detail'][table]:p['physical_key']='OTHER-'+p['physical_key']
  with pytest.raises(ValueError):a.save_aec_system_draft('P',other,expected_revision=0)
  a.store.close();return
 if case=='copy':
  v=a.aec_workflow().copy_roof('P',kind,'L1','C1',new_system_id='COPY',new_source_id='COPY',reason='کپی مصنوعی')['data'];assert v['architecture_detail']['geometry']['physical_surface_id']!='SURFACE-1';assert v['architecture_detail']['pieces'][0]['stock_id']==v['architecture_detail']['stocks'][0]['id'];assert not result(v)['ready'];a.aec_workflow().delete_roof('P',kind,'L1','COPY',expected_revision=1,reason='حذف مصنوعی');a.store.close();return
 a.review_roof_system('P',kind,'L1','C1',expected_revision=1,reviewer='بازبین',note='کنترل مصنوعی');a.commit_aec_system('P',kind,'L1','C1',expected_revision=2);before=deepcopy(a.open_project('P'));assert len(before['boq'])==1 and before['boq'][0]['quantity']==5
 illegal=deepcopy(before);next(iter(illegal['reviewed_aec_systems'].values()))['data']['axis']='illegal'
 with pytest.raises(ValueError):a.store.save('P',illegal)
 d=deepcopy(a.get_aec_system('P',kind,'L1','C1')['data']);d.pop('roof_review');ar=d['architecture_detail'];ar['surfaces'][0].update(layout_status='area_only',width_m='2',opening_ids='',coverage_scope='partial',scope_ref='حذف مستند ناحیه انتهایی');ar['pieces']=[];ar['stocks']=[]
 if case=='unit':ar['surfaces'][0]['operation']='changed'
 if case=='identity':ar['geometry']['physical_surface_id']='CHANGED'
 e=tmp_path/'new.txt';e.write_text('EXPLICIT SYNTHETIC cover reduced from net 5 m2 to net 4 m2; last strip removed; procurement no longer certified.');kw=dict(expected_revision=2,created_by='تهیه‌کننده',reason='کاهش مستند سطح مصنوعی',evidence=dict(evidence_path=str(e),evidence_sha256=sha256(e.read_bytes()).hexdigest(),evidence_reference='دیتایل مصنوعی'))
 if case in ('unit','identity'):
  with pytest.raises(ValueError):a.create_roof_amendment('P',kind,'L1','C1',d,**kw)
 else:
  am=a.create_roof_amendment('P',kind,'L1','C1',d,**kw);assert Decimal(str(am['delta']['aec:'+kind+':L1:C1:LAYER']))==-1
  if case=='same_reviewer':
   with pytest.raises(ValueError):a.review_beam_amendment('P',am['id'],expected_revision=1,reviewer='تهیه‌کننده',note='نامجاز')
  elif case=='evidence':
   e.write_text('tampered')
   with pytest.raises(ValueError):a.review_beam_amendment('P',am['id'],expected_revision=1,reviewer='دیگری',note='کنترل')
  else:
   a.review_beam_amendment('P',am['id'],expected_revision=1,reviewer='دیگری',note='مستقل')
   if case=='financial':
    p=a.open_project('P');p['contract']={'id':'locked'};a.store.save('P',p);locked=deepcopy(a.open_project('P'))
    with pytest.raises(ValueError):a.apply_beam_amendment('P',am['id'],expected_revision=1)
    assert a.open_project('P')==locked
   else:
    a.apply_beam_amendment('P',am['id'],expected_revision=1);assert a.open_project('P')['takeoffs']==before['takeoffs'];assert a.open_project('P')['boq'][0]['quantity']==4;a.store.close();a=StructuralProApp(tmp_path/'db');assert a.get_aec_system('P',kind,'L1','C1')['revision']==3 and a.store.db.execute('pragma integrity_check').fetchone()[0]=='ok'
 if case=='valid':
  from core.reports.architecture_detail_report import export
  export(a,'P',kind,'L1','C1',tmp_path/'amended.csv',format='csv');assert 'وضعیت جدید: فاقد ردیف اطلاعاتی فعال' in (tmp_path/'amended.csv').read_text(encoding='utf-8-sig')
 a.store.close()

@pytest.mark.parametrize('kind',KINDS)
def test_four_shared_exports(tmp_path,kind):
 from core.platform.application import StructuralProApp
 from core.reports.architecture_detail_report import export
 from openpyxl import load_workbook
 from docx import Document
 import pymupdf
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');d=sample(kind);a.save_aec_system_draft('P',d,expected_revision=0);a.review_roof_system('P',kind,'L1','C1',expected_revision=1,reviewer='مستقل',note='کنترل');a.commit_aec_system('P',kind,'L1','C1',expected_revision=2)
 for fmt in ('pdf','xlsx','docx','csv'):export(a,'P',kind,'L1','C1',tmp_path/('cover.'+fmt),format=fmt)
 b=load_workbook(tmp_path/'cover.xlsx');assert all(s.sheet_view.rightToLeft for s in b);b.close();doc=Document(tmp_path/'cover.docx');assert all(t._tbl.tblPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bidiVisual') is not None for t in doc.tables);pdf=pymupdf.open(tmp_path/'cover.pdf');assert len(pdf)>0;pdf.close();t=(tmp_path/'cover.csv').read_text(encoding='utf-8-sig');assert 'نقشه برش مصنوعی' in t and 'ST2' in t and 'اطلاعاتی' in t;a.store.close()


def test_explicit_joint_geometry_and_grout_volume():
 d=sample();a=d['architecture_detail'];a['pieces'][0]['width_m']='1.9';a['stocks'][0]['length_m']='1.9';a['joints']=[obj('J1',surface_id='LAYER',x='1.9',y='0',width_m='.1',height_m='2')];a['consumptions']=[obj('GR',surface_id='LAYER',operation='grout',basis='constant_thickness',thickness_m='.01',rate='',rate_ref='ضخامت بند مصنوعی',material_ref='بند مصنوعی',quantity_unit='m3')];r=result(d);assert r['ready'],r['blockers'];assert Decimal(str(next(v['quantity'] for v in r['rows'] if v['item_code']=='grout')))==Decimal('.002')

@pytest.mark.parametrize('fault',['missing_rate_ref','missing_rate','wrong_unit','duplicate','no_joints','rate_and_thickness'])
def test_consumption_fail_closed(fault):
 d=sample();a=d['architecture_detail'];p=obj('A1',surface_id='LAYER',operation='adhesive',basis='area_rate',thickness_m='',rate='1.25',rate_ref='گواهی کاملاً مصنوعی مصرف',material_ref='چسب مصنوعی',quantity_unit='kg');a['consumptions']=[p]
 if fault=='missing_rate_ref':p['rate_ref']=''
 if fault=='missing_rate':p['rate']=''
 if fault=='wrong_unit':p['quantity_unit']='m3'
 if fault=='duplicate':a['consumptions'].append({**p,'id':'A2','physical_key':'A2'})
 if fault=='no_joints':p.update(operation='grout',basis='constant_thickness',rate='',thickness_m='.01',quantity_unit='m3')
 if fault=='rate_and_thickness':p['thickness_m']='.01'
 assert not result(d)['ready']


def test_explicit_material_consumption_is_not_a_second_payment(tmp_path):
 from core.platform.application import StructuralProApp
 d=sample();d['architecture_detail']['consumptions']=[obj('A1',surface_id='LAYER',operation='adhesive',basis='area_rate',thickness_m='',rate='1.25',rate_ref='گواهی کاملاً مصنوعی مصرف',material_ref='چسب مصنوعی',quantity_unit='kg')];r=result(d);assert r['ready'];assert next(v['quantity'] for v in r['rows'] if v['item_code']=='adhesive')==Decimal('6.25')
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');a.save_aec_system_draft('P',d,expected_revision=0);a.review_roof_system('P','ceramic_finish','L1','C1',expected_revision=1,reviewer='بازبین',note='کنترل');a.commit_aec_system('P','ceramic_finish','L1','C1',expected_revision=2);assert len(a.open_project('P')['boq'])==1;a.store.close()

@pytest.mark.parametrize('fault',['different_layer_key','two_faces_same_surface','wall_dimensions','wall_face_identity'])
def test_physical_surface_guard_cannot_be_bypassed(tmp_path,fault):
 from core.platform.application import StructuralProApp
 from test_architecture_specialist_work import sample as wall_sample
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');d=sample()
 if fault.startswith('wall_'):
  wall=wall_sample();a.save_aec_system_draft('P',wall,expected_revision=0);g=wall['architecture_detail']['geometry'];surface=wall['architecture_detail']['surfaces'][0];d['architecture_detail']['geometry'].update(host_id=g['physical_wall_id'],physical_surface_id=surface['surface_id'],face='inside',plane='vertical')
  if fault=='wall_face_identity':d['architecture_detail']['geometry']['physical_surface_id']='OTHER'
 else:
  a.save_aec_system_draft('P',d,expected_revision=0);d=deepcopy(d);d['system_id']='C2';d['source']['source_id']='OTHER'
  for name in ('openings','surfaces','stocks','pieces'):
   for p in d['architecture_detail'][name]:p['physical_key']='OTHER-'+p['physical_key']
  if fault=='different_layer_key':d['architecture_detail']['surfaces'][0]['layer_key']='OTHER'
  else:d['architecture_detail']['geometry'].update(face='inside',plane='vertical')
 with pytest.raises(ValueError):a.save_aec_system_draft('P',d,expected_revision=0)
 a.store.close()


def test_stored_five_floor_scenario_extends_all_previous_layers(tmp_path):
 from scripts.verify_architecture_step6 import create_sample
 from scripts.verify_floor_step6_package2 import extend_sample as floors
 from scripts.verify_insulation_step6_package3 import extend_sample as insulation
 from scripts.verify_interior_cover_step6_package2 import extend_sample
 create_sample(tmp_path/'p1');floors(tmp_path/'p1',tmp_path/'p2');insulation(tmp_path/'p2',tmp_path/'p3');r=extend_sample(tmp_path/'p3',tmp_path/'p4');assert r['previous_records_preserved'] and r['previous_boq_preserved'] and r['other_four_floor_quantities_unchanged'];assert r['boq_rows']==58 and Decimal(r['cover_floor_area_m2'])==94 and Decimal(r['cover_wall_area_m2'])==Decimal('54.93') and Decimal(r['skirting_length_m'])==20


def test_zero_cuts_are_known_from_full_layout_not_an_unknown_default():
 d=sample();d['architecture_detail']['stocks'][1]['length_m']='1';d['architecture_detail']['pieces'][1]['piece_type']='full';r=result(d);assert r['ready'];assert next(v for v in r['rows'] if v['item_code']=='installed_pieces')['params']['cut_count']==0


def test_unregistered_litre_unit_is_blocked_at_preview_not_only_at_commit():
 d=sample();d['architecture_detail']['consumptions']=[obj('LIQUID',surface_id='LAYER',operation='adhesive',basis='area_rate',thickness_m='',rate='1.25',rate_ref='مصرف فرض آزمون نرم‌افزار؛ نه تأیید ماده',material_ref='چسب مایع آزمون',quantity_unit='litre')];assert not result(d)['ready']
