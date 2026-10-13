from decimal import Decimal
from pathlib import Path
import pytest
from core.platform.application import StructuralProApp
from test_manual_formula_persistence import data

def test_reports_exact_text_live_safe_and_injection(tmp_path):
 s=StructuralProApp(tmp_path/'db');s.create_project('مصنوعی','P')
 s.save_manual_formula('P','A',data(expression='n*2',variables={'n':'10'},unit='m2',title='=HYPERLINK("x")'),expected_revision=0,reason='اول')
 s.save_manual_formula('P','B',data(source_id='B',member='B',expression='n',variables={'n':'12345678901234567890.123456789'}),expected_revision=0,reason='اول')
 assert callable(getattr(s,'export_manual_formula_report',None)), 'گزارش فرمول و تاریخچه وجود ندارد'
 path=tmp_path/'report.xlsx';s.export_manual_formula_report('P',path,format='xlsx')
 from openpyxl import load_workbook
 w=load_workbook(path);assert len(w.sheetnames)>=10;assert w['محاسبه امن']['C2'].data_type=='f';assert w['محاسبه امن']['C3'].data_type=='s';assert 'دقت' in w['محاسبه امن']['D3'].value
 assert any(c.value=='=HYPERLINK("x")' and c.data_type=='s' for row in w['ریز متره'] for c in row)
 assert any(c.value=='12345678901234567890.123456789' and c.data_type=='s' for row in w['ریز متره'] for c in row)
 for sheet in w:
  assert sheet.sheet_view.rightToLeft
  for row in sheet:
   for c in row:
    if c.data_type=='f':assert 'HYPERLINK' not in c.value
 pdf=tmp_path/'report.pdf';s.export_manual_formula_report('P',pdf,format='pdf')
 import fitz
 d=fitz.open(pdf);assert len(d)>0
 for page in d:assert page.get_pixmap().width>0
 s.store.close()
