"""Versioned layouts and atomic batches over real SQLite/project sources."""
import importlib.util
from dataclasses import asdict
from hashlib import sha256
from pathlib import Path
import json
import csv
import pytest
from test_phase7_p2_multifloor_bundle import building
from core.reports.designer import ReportLayout

def api():
    assert importlib.util.find_spec("core.reports.custom_batch")
    return __import__("core.reports.custom_batch",fromlist=["TemplateStore"])

def template(**changes):
    return dict(schema_version=2,template_id="measures",version=1,report_type="boq",
        layout=asdict(ReportLayout(title="گزارش سفارشی متره",columns=["item_no","description","quantity","unit","source_id"],group_by="floor",header="دفتر فنی",footer="بازبینی منابع الزامی است")),
        filters={},page_format="A4",orientation="landscape",show_source=True,include_totals=True,logo=None)|changes

def test_template_roundtrip_version_audit_and_explicit_migration(building):
    app,p,_=building;m=api();store=m.TemplateStore(app,"P5");before=app.store.current_digest("P5")
    first=store.save(template(),actor="مهندس آزمون",reason="ثبت قالب")
    again=m.TemplateStore(app,"P5");assert again.load("measures",1)==first
    second=store.save(template(version=2,filters={"floor":["L1","L3"]}),actor="مهندس آزمون",reason="فیلتر طبقات")
    assert second["previous_sha256"]==first["sha256"]
    assert again.load("measures",2)["template"]["filters"]=={"floor":["L1","L3"]}
    with pytest.raises(ValueError):store.save(template(version=2),actor="a",reason="duplicate")
    legacy={"schema_version":1,"template_id":"old","version":1,"layout":asdict(ReportLayout())}
    migrated=m.migrate_template_v1(legacy);assert migrated["schema_version"]==2 and migrated["layout"]["rtl"] is True
    assert app.store.current_digest("P5")==before

@pytest.mark.parametrize("patch",[
    {"schema_version":99},{"template_id":"../escape"},{"layout":{}},
    {"filters":{"unknown":"x"}},{"orientation":"automatic"},{"page_format":"A0"},
    {"logo":"../../outside.png"},{"version":True},{"report_type":"ai-audit"},
])
def test_invalid_template_rejected_before_storage(building,patch):
    m=api();app,_,_=building;store=m.TemplateStore(app,"P5")
    with pytest.raises(ValueError):store.save(template(**patch),actor="مهندس",reason="ورودی نامعتبر")
    assert store.list_versions()==[]

def test_batch_two_selected_floors_four_formats_manifest_and_no_project_write(building,tmp_path):
    import pymupdf
    from docx import Document
    from openpyxl import load_workbook
    m=api();app,p,_=building;store=m.TemplateStore(app,"P5");store.save(template(),actor="مهندس",reason="مرجع")
    digest=app.store.current_digest("P5")
    result=m.export_custom_report_batch(app,project_id="P5",expected_project_version=p["_meta"]["version"],
        requests=({"template_id":"measures","version":1,"filters":{"floor":["L1"]}},
                  {"template_id":"measures","version":1,"filters":{"floor":["L3"]}}),destination=tmp_path/"batch")
    out=Path(result["destination"]);manifest=json.loads((out/"manifest.json").read_text(encoding="utf-8"))
    assert len(manifest["reports"])==2 and len(manifest["files"])==8
    assert manifest["project_version"]==p["_meta"]["version"] and manifest["created_at"]
    for f in manifest["files"]:
        path=out/f["name"];assert sha256(path.read_bytes()).hexdigest()==f["sha256"];assert path.stat().st_size==f["size"]
    for n,expected in ((1,"20"),(2,"19.75")):
        stem=f"report_{n:03}"
        with (out/(stem+".csv")).open(encoding="utf-8-sig",newline="") as stream:rows=list(csv.reader(stream))
        assert any(expected in row for row in rows)
        doc=Document(out/(stem+".docx"));assert "w:bidi" in doc._element.xml
        assert "دفتر فنی" in doc.sections[0].header.paragraphs[0].text
        assert "بازبینی منابع" in doc.sections[0].footer.paragraphs[0].text
        book=load_workbook(out/(stem+".xlsx"));assert all(s.sheet_view.rightToLeft and s.print_area for s in book)
        with pymupdf.open(out/(stem+".pdf")) as pdf:assert pdf.page_count>0 and pdf[0].rect.width>pdf[0].rect.height
    assert app.open_project("P5")==p and app.store.current_digest("P5")==digest

@pytest.mark.parametrize("case",["empty","missing_floor","stale","tampered_template","changed_source"])
def test_invalid_or_stale_batch_never_publishes(building,tmp_path,case):
    m=api();app,p,selections=building;store=m.TemplateStore(app,"P5");store.save(template(),actor="مهندس",reason="مرجع")
    requests=({"template_id":"measures","version":1,"filters":{}},);version=p["_meta"]["version"]
    if case=="empty":requests=()
    if case=="missing_floor":requests=({"template_id":"measures","version":1,"filters":{"floor":["missing"]}},)
    if case=="stale":version+=1
    if case=="tampered_template":
        path=store.root/"measures"/"1.json";path.write_text(path.read_text(encoding="utf-8").replace("گزارش سفارشی","دستکاری شده"),encoding="utf-8")
    if case=="changed_source":
        path=Path(p["drawing_sessions"][selections[0]["session_id"]]["drawing_source"]);path.write_bytes(path.read_bytes()+b"changed")
    with pytest.raises((ValueError,RuntimeError)):
        m.export_custom_report_batch(app,project_id="P5",expected_project_version=version,requests=requests,destination=tmp_path/"bad")
    assert not (tmp_path/"bad").exists()

def test_batch_exporter_failure_removes_staging(building,tmp_path,monkeypatch):
    from core.reports.project_report import ProjectReport
    m=api();app,p,_=building;store=m.TemplateStore(app,"P5");store.save(template(),actor="مهندس",reason="مرجع")
    old=ProjectReport.export
    def fail(self,path,fmt):
        if fmt=="docx":raise RuntimeError("injected failure")
        return old(self,path,fmt)
    monkeypatch.setattr(ProjectReport,"export",fail)
    with pytest.raises(RuntimeError):
        m.export_custom_report_batch(app,project_id="P5",expected_project_version=p["_meta"]["version"],requests=({"template_id":"measures","version":1,"filters":{}},),destination=tmp_path/"bad")
    assert not (tmp_path/"bad").exists() and not list(tmp_path.glob("structuralpro_p7_batch_*"))

@pytest.mark.parametrize("orientation,wide",[("portrait",False),("landscape",True)])
def test_page_format_and_logo_are_real_in_three_formats(building,tmp_path,orientation,wide):
    from PIL import Image
    import pymupdf
    from docx import Document
    from openpyxl import load_workbook
    m=api();app,p,_=building;store=m.TemplateStore(app,"P5")
    logo=tmp_path/"logo.png";Image.new("RGB",(120,40),(20,80,150)).save(logo);name=store.import_logo(logo)
    store.save(template(page_format="A3",orientation=orientation,logo=name),actor="مهندس",reason="صفحه و لوگو")
    m.export_custom_report_batch(app,project_id="P5",expected_project_version=p["_meta"]["version"],requests=({"template_id":"measures","version":1,"filters":{}},),destination=tmp_path/"with_logo")
    base=tmp_path/"with_logo"
    with pymupdf.open(base/"report_001.pdf") as pdf:
        assert (pdf[0].rect.width>pdf[0].rect.height)==wide
        assert max(pdf[0].rect.width,pdf[0].rect.height)>1190 and pdf[0].get_images()
    section=Document(base/"report_001.docx").sections[0]
    assert (section.page_width>section.page_height)==wide and "w:drawing" in section.header._element.xml
    book=load_workbook(base/"report_001.xlsx")
    assert all(s.page_setup.orientation==orientation and s.page_setup.paperSize==8 for s in book)
    assert sum(len(s._images) for s in book)==1
    (store.assets/name).write_bytes(b"changed")
    with pytest.raises(ValueError,match="لوگو"):store.load("measures",1)

def test_explicit_empty_filter_clears_saved_defaults(building):
    m=api();app,p,_=building
    result=m.preview_custom_report(app,project_id="P5",expected_project_version=p["_meta"]["version"],template=template(filters={"floor":["L1"]}),filters={})
    assert {r.get("floor") for r in result["rows"] if r.get("item_no") is not None}=={f"L{i}" for i in range(1,6)}

def test_changed_project_during_render_cannot_publish(building,tmp_path,monkeypatch):
    from core.reports.project_report import ProjectReport
    m=api();app,p,_=building;store=m.TemplateStore(app,"P5");store.save(template(),actor="مهندس",reason="مرجع")
    old=ProjectReport.export
    def change(self,path,fmt):
        result=old(self,path,fmt)
        if fmt=="csv":monkeypatch.setattr(app.store,"current_digest",lambda _:"changed")
        return result
    monkeypatch.setattr(ProjectReport,"export",change)
    with pytest.raises(RuntimeError,match="پروژه"):
        m.export_custom_report_batch(app,project_id="P5",expected_project_version=p["_meta"]["version"],requests=({"template_id":"measures","version":1,"filters":{}},),destination=tmp_path/"changed")
    assert not (tmp_path/"changed").exists()


def test_persisted_quantity_and_boq_cannot_disagree_with_source_measurement(building):
    from copy import deepcopy
    m=api();app,p,_=building;forged=deepcopy(p)
    forged["takeoffs"][0]["quantities"][0]["amount"]=40
    forged["boq"][0]["quantity"]=40
    with pytest.raises(ValueError,match="منبع|کمیت|متره"):
        m._source_rows(forged,"boq")

def test_source_change_after_last_preview_is_rejected(building,tmp_path,monkeypatch):
    m=api();app,p,selections=building;store=m.TemplateStore(app,"P5");store.save(template(),actor="مهندس",reason="مرجع")
    old=m.preview_custom_report;calls=[]
    def changing(*args,**kwargs):
        result=old(*args,**kwargs);calls.append(1)
        if len(calls)==2:
            source=Path(p["drawing_sessions"][selections[0]["session_id"]]["drawing_source"])
            source.write_bytes(source.read_bytes()+b"changed after final preview")
        return result
    monkeypatch.setattr(m,"preview_custom_report",changing)
    with pytest.raises(ValueError,match="منبع"):
        m.export_custom_report_batch(app,project_id="P5",expected_project_version=p["_meta"]["version"],requests=({"template_id":"measures","version":1,"filters":{}},),destination=tmp_path/"stale_source")
    assert not (tmp_path/"stale_source").exists()

def test_grand_total_without_grouping_is_in_actual_documents(building,tmp_path):
    from docx import Document
    from openpyxl import load_workbook
    m=api();app,p,_=building;store=m.TemplateStore(app,"P5")
    t=template();t["layout"]["group_by"]="";store.save(t,actor="مهندس",reason="جمع کل")
    m.export_custom_report_batch(app,project_id="P5",expected_project_version=p["_meta"]["version"],requests=({"template_id":"measures","version":1,"filters":{}},),destination=tmp_path/"total")
    doc=Document(tmp_path/"total/report_001.docx")
    text=" ".join(c.text for tab in doc.tables for row in tab.rows for c in row.cells)
    assert "جمع مقدار" in text and "99.75" in text
    book=load_workbook(tmp_path/"total/report_001.xlsx")
    assert any(str(c.value)=="99.75" for row in book["خلاصه"] for c in row)
