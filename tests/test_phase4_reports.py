from decimal import Decimal
from pathlib import Path
import pytest
from core.platform.application import StructuralProApp
from test_phase4_progress import weights,work,project

@pytest.mark.parametrize('kind',['wbs','activities','schedule','progress','curve'])
@pytest.mark.parametrize('fmt',['xlsx','pdf','docx','csv'])
def test_real_reports_readback(project,tmp_path,kind,fmt):
    a,p=project;assert hasattr(a,'project_controls_report'),'project reports adapter absent'
    weights(a,p);work(a,p,'A','5','2026-01-06','A');work(a,p,'B','5','2026-01-06','B')
    from test_phase4_structure import cmd
    cmd(a,'progress_report',date='2026-01-06',reference='R')
    path=tmp_path/(kind+'.'+fmt);a.project_controls_report('P',kind,fmt,path,as_of='2026-01-06')
    assert path.stat().st_size>100
    if fmt=='xlsx':
        from openpyxl import load_workbook
        b=load_workbook(path);assert all(s.sheet_view.rightToLeft for s in b)
        if kind in {'schedule','curve'}:assert any(s._images for s in b)
    elif fmt=='docx':
        from docx import Document
        d=Document(path);assert d.tables
        if kind in {'schedule','curve'}:assert d.inline_shapes
    elif fmt=='pdf':
        import pymupdf
        with pymupdf.open(path) as pdf:
            assert len(pdf)>0 and pdf[0].get_text()
            if kind in {'schedule','curve'}:assert any(page.get_images() for page in pdf)
    else:assert 'آزمایشی' in path.read_text(encoding='utf-8-sig') or path.read_text(encoding='utf-8-sig')

def test_figure_persian_has_single_bidi_pass():
    from PIL import Image,ImageDraw,ImageFont,ImageChops,features
    from core.reports.project_control_figures import canvas
    if not features.check_feature('raqm'):
        pytest.skip('Native complex text oracle unavailable')
    title='منحنی پیشرفت برنامه و کارکرد'
    produced,_,_=canvas(title,100)
    oracle=Image.new('RGB',(1300,100),'white');font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',18)
    ImageDraw.Draw(oracle).text((700,18),title,font=font,fill='#203044',direction='rtl')
    assert ImageChops.difference(produced,oracle).getbbox() is None,'diagram text was reordered twice'
