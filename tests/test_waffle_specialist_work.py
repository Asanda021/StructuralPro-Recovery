"""Rectangular, explicit synthetic waffle schedules, independently summed references."""
from copy import deepcopy
from decimal import Decimal
import pytest
from test_roof_specialist_work import sample as base_sample
from core.takeoff.reviewed_aec import empty_system,preview_system
from core.takeoff.engine import TakeoffEngine

KINDS=('waffle_one_way','waffle_two_way')

def sample(kind='waffle_one_way'):
 d=base_sample();d.update(kind=kind,components=empty_system(kind,'L1','S1')['components']);d['source']['source_id']='SYNTHETIC-'+kind
 rd=d['roof_detail'];g=rd['geometry'];g.update(depth='.25',topping_thickness='.05',rib_axis='x');g['action']='one_way' if kind==KINDS[0] else 'two_way'
 def box(i,c,x,y,l,w,z='0',h='.2'):
  return dict(id=i,component=c,x=str(x),y=str(y),z=z,length=str(l),width=str(w),height=h,source_ref='دیتایل مصنوعی وافل',approved_by='کنترل مصنوعی')
 rd['regions']=[box('TOP','topping',0,0,4,4,'.2','.05')];rd['rib_layout']=[];rd['pieces']=[]
 for i,y in enumerate(('0','2')):
  b=box('RX'+str(i),'',0,y,4,'.2');b.pop('component');b['axis']='x';rd['rib_layout'].append(b)
 if kind==KINDS[1]:
  for i,x in enumerate(('0','2')):
   b=box('RY'+str(i),'',x,0,'.2',4);b.pop('component');b['axis']='y';rd['rib_layout'].append(b)
 xs=list(map(Decimal,('0','.2','2','2.2','4'))) if kind==KINDS[1] else list(map(Decimal,('0','4')))
 ys=list(map(Decimal,('0','.2','2','2.2','4')))
 for i,(x,X) in enumerate(zip(xs,xs[1:])):
  for j,(y,Y) in enumerate(zip(ys,ys[1:])):
   a=j%2==0;b=kind==KINDS[1] and i%2==0
   c='junctions' if a and b else 'ribs_x' if a else 'ribs_y' if b else 'void'
   ident=f'CELL-{i}-{j}';rd['regions'].append(box(ident,c,x,y,X-x,Y-y))
   if c=='void':rd['pieces'].append(dict(id='M-'+ident,physical_key='mold-'+ident,region_id=ident,count='1',length=str(X-x),width=str(Y-y),height='.2',nx='1',ny='1',nz='1',shape='rectangular_prism',source_ref='جانمایی مصنوعی قالب',spec_ref='قالب منشوری مصنوعی',approved_by='کنترل مصنوعی'))
 rd['rebars']['thermal']=deepcopy(rd['rebars']['main_x']);s=rd['rebars']['thermal']['series'][0];s.update(id='B-thermal',physical_key='physical-thermal')
 if kind==KINDS[0]:rd['rebars']['main_y']=dict(status='not_applicable',approved_by='کنترل مصنوعی',source_ref='دیتایل مصنوعی',reason='نمونه مصنوعی یک‌طرفه')
 return d

@pytest.mark.parametrize('kind',KINDS)
def test_independent_reference(kind):
 r=preview_system(sample(kind),TakeoffEngine());assert r['ready'],r['blockers']
 totals={}
 for row in r['rows']:totals[row['item_code']]=totals.get(row['item_code'],Decimal(0))+Decimal(str(row['quantity']))
 # 16*.05 + 2*4*.2*.2 [+2*4*.2*.2 -4*.2*.2*.2].
 assert totals['concrete']==Decimal('1.12' if kind==KINDS[0] else '1.408')
 assert totals['waffle_molds']==(2 if kind==KINDS[0] else 4)
 assert totals['rebar']==Decimal('71.04' if kind==KINDS[0] else '106.56')
 assert totals['formwork']==16

@pytest.mark.parametrize('kind',KINDS)
@pytest.mark.parametrize('operation',['sqlite_amendment_and_boq','copy_delete_and_reports'])
def test_shared_workflow(tmp_path,kind,operation,monkeypatch):
 import test_roof_specialist_work as base
 monkeypatch.setattr(base,'sample',sample)
 getattr(base,'test_'+operation)(tmp_path,kind)

@pytest.mark.parametrize('kind',KINDS)
@pytest.mark.parametrize('fault',['same_axis','missing_ribs','rib_source','partition','junction','mold_count','mold_source','mold_shape','mold_duplicate','bar_identity','bar_detail','missing_direction','unit','deduction','unreviewed','unknown','factory','joist'])
def test_invalid_waffle(kind,fault):
 d=sample(kind);rd=d['roof_detail']
 if fault=='same_axis':
  b=deepcopy(rd['rib_layout'][0]);b['id']='DUP';rd['rib_layout'].append(b)
 elif fault=='missing_ribs':rd.pop('rib_layout')
 elif fault=='rib_source':rd['rib_layout'][0]['source_ref']=''
 elif fault=='partition':rd['regions'].pop()
 elif fault=='junction':next(p for p in rd['regions'] if p['component'].startswith('ribs') or p['component']=='junctions')['component']='solid_zones'
 elif fault=='mold_count':rd['pieces'][0]['count']='2'
 elif fault=='mold_source':rd['pieces'][0]['spec_ref']=''
 elif fault=='mold_shape':rd['pieces'][0]['shape']='sloped'
 elif fault=='mold_duplicate':
  p=deepcopy(rd['pieces'][0]);p['id']='DUP';rd['pieces'].append(p)
 elif fault=='bar_identity':rd['rebars']['thermal']['series'][0]['physical_key']='physical-main_x'
 elif fault=='bar_detail':rd['rebars']['main_x']['series'][0]['cut_parts']=[]
 elif fault=='missing_direction':rd['geometry']['rib_axis']=None
 elif fault=='unit':rd['geometry']['unit']='cm'
 elif fault=='deduction':rd['deductions']=[dict(id='O',kind='opening',x='0',y='0',z='0',length='1',width='1',height='.25',source_ref='مصنوعی',approved_by='مصنوعی')]
 elif fault=='unreviewed':rd['rebars']['thermal']['status']='unreviewed'
 elif fault=='unknown':rd['geometry']['guess']='1'
 elif fault=='factory':rd['rebars']['thermal']['series'][0]['supply_scope']='factory'
 elif fault=='joist':rd['joist_layout']=[{}]
 assert not preview_system(d,TakeoffEngine())['ready']

@pytest.mark.parametrize('kind',KINDS)
def test_solid_zone_opening_and_joint_reference(kind):
 d=sample(kind);rd=d['roof_detail'];p=next(p for p in rd['regions'] if p['component']=='void');p['component']='solid_zones';p['height']='.25';rd['pieces']=[m for m in rd['pieces'] if m['region_id']!=p['id']]
 # Replace the continuous topping with measured cells around a full-height solid zone.
 rd['regions']=[v for v in rd['regions'] if v['component']!='topping']
 tops=[]
 for cell in rd['regions']:
  if cell is p:continue
  top=deepcopy(cell);top.update(id='TOP-'+cell['id'],component='topping',z='.2',height='.05');tops.append(top)
 rd['regions'].extend(tops)
 rd['rebars']['solid']=deepcopy(rd['rebars']['main_x']);rd['rebars']['solid']['series'][0].update(id='solid-bars',physical_key='solid-bars')
 rd['rebars']['opening']=deepcopy(rd['rebars']['main_x']);rd['rebars']['opening']['series'][0].update(id='opening-bars',physical_key='opening-bars')
 for label,x,y in [('O',p['x'],p['y']),('J',str(Decimal(p['x'])+Decimal('.4')),p['y'])]:
  rd['deductions'].append(dict(id=label,kind='opening' if label=='O' else 'beam_joint',owner_source_id='BEAM' if label=='J' else '',x=x,y=y,z='0',length='.2',width='.2',height='.25',source_ref='مصنوعی',approved_by='مصنوعی'))
 r=preview_system(d,TakeoffEngine());assert r['ready'],r['blockers']
 expected=Decimal('1.12' if kind==KINDS[0] else '1.408')+Decimal(p['length'])*Decimal(p['width'])*Decimal('.2')-Decimal('.02')
 assert sum(Decimal(str(v['quantity'])) for v in r['rows'] if v['item_code']=='concrete')==expected
 rd['deductions'].append(deepcopy(rd['deductions'][0]));rd['deductions'][-1]['id']='DUP'
 assert not preview_system(d,TakeoffEngine())['ready']

@pytest.mark.parametrize('kind',KINDS)
@pytest.mark.parametrize('case',['policy','financial','concurrent','source','immutable','mold_payable'])
def test_persistence_guards(tmp_path,kind,case):
 from core.platform.application import StructuralProApp
 from hashlib import sha256
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');d=sample(kind)
 if case=='policy':
  d.pop('measurement_policy');assert not preview_system(d,a.takeoff)['ready'];a.store.close();return
 a.save_aec_system_draft('P',d,expected_revision=0)
 if case=='source':
  d['system_id']='DUP'
  with pytest.raises(ValueError):a.save_aec_system_draft('P',d,expected_revision=0)
 elif case=='concurrent':
  b=StructuralProApp(tmp_path/'db');a.review_roof_system('P',kind,'L1','S1',expected_revision=1,reviewer='مستقل',note='کنترل')
  with pytest.raises(RuntimeError):b.save_aec_system_draft('P',d,expected_revision=1)
  b.store.close()
 else:
  a.review_roof_system('P',kind,'L1','S1',expected_revision=1,reviewer='مستقل',note='کنترل');a.commit_aec_system('P',kind,'L1','S1',expected_revision=2)
  if case=='mold_payable':assert not any(r['item_code']=='waffle_molds' for r in a.open_project('P')['boq'])
  elif case=='immutable':
   p=a.open_project('P');next(iter(p['reviewed_aec_systems'].values()))['data']['roof_detail']['geometry']['count']='5'
   with pytest.raises(ValueError):a.store.save('P',p)
  else:
   d=deepcopy(a.get_aec_system('P',kind,'L1','S1')['data']);d.pop('roof_review');d['roof_detail']['geometry']['count']='2';path=tmp_path/'rev.txt';path.write_text('explicit synthetic revised document')
   am=a.create_roof_amendment('P',kind,'L1','S1',d,expected_revision=2,created_by='سازنده',reason='تغییر مستند مصنوعی',evidence={'evidence_path':str(path),'evidence_sha256':sha256(path.read_bytes()).hexdigest(),'evidence_reference':'مرجع مصنوعی'})
   a.review_beam_amendment('P',am['id'],expected_revision=1,reviewer='بازبین دیگر',note='کنترل');p=a.open_project('P');p['contract']={'id':'locked'};a.store.save('P',p)
   with pytest.raises(ValueError):a.apply_beam_amendment('P',am['id'],expected_revision=1)
 a.store.close()


def test_one_way_y_axis_reference():
 d=sample();rd=d['roof_detail'];rd['geometry']['rib_axis']='y'
 for p in rd['regions']+rd['rib_layout']:
  p['x'],p['y']=p['y'],p['x'];p['length'],p['width']=p['width'],p['length']
  if p.get('component')=='ribs_x':p['component']='ribs_y'
  if 'axis' in p:p['axis']='y'
 for p in rd['pieces']:
  p['length'],p['width']=p['width'],p['length'];p['nx'],p['ny']=p['ny'],p['nx']
 rd['rebars']['main_y'],rd['rebars']['main_x']=rd['rebars']['main_x'],rd['rebars']['main_y']
 r=preview_system(d,TakeoffEngine());assert r['ready'],r['blockers']
 assert sum(Decimal(str(p['quantity'])) for p in r['rows'] if p['item_code']=='concrete')==Decimal('1.12')

@pytest.mark.parametrize('kind',KINDS)
def test_editable_report_quantities_and_sources(tmp_path,kind):
 from core.platform.application import StructuralProApp
 from openpyxl import load_workbook
 from docx import Document
 from docx.oxml.ns import qn
 a=StructuralProApp(tmp_path/'db');a.create_project('نمونه مصنوعی','P');d=sample(kind);a.save_aec_system_draft('P',d,expected_revision=0)
 path=tmp_path/'detail.xlsx';a.export_roof_detail_report('P',kind,'L1','S1',path,format='xlsx');w=load_workbook(path);sheet=w['ریز متره'];headers=next(row for row in sheet if [c.value for c in row][:3]==['ردیف','شرح','مقدار']);start=headers[0].row+1
 text=' '.join(str(c.value) for row in sheet for c in row)
 for p in d['roof_detail']['rib_layout']:assert p['id'] in text and p['source_ref'] in text
 nums=[row[2] for row in sheet.iter_rows(min_row=start) if row[2].value is not None]
 assert nums and all(c.data_type=='n' for c in nums)
 summary={row[0].value:row[1].value for row in w['خلاصه'].iter_rows(min_row=4)}
 assert summary['گروه: بتن درجاریز (متر مکعب)']==(1.12 if kind==KINDS[0] else 1.408)
 path=tmp_path/'detail.docx';a.export_roof_detail_report('P',kind,'L1','S1',path,format='docx');doc=Document(path)
 assert all(t._tbl.tblPr.find(qn('w:bidiVisual')) is not None for t in doc.tables)
 a.store.close()


def test_inventory_preserves_85_manual_families_and_waffle_variants():
 from core.takeoff.manual_formula_register import catalog
 from core.product.aec_item_inventory_audit import assert_priority_inventory
 assert len(catalog())==85
 assert assert_priority_inventory()['valid']

@pytest.mark.parametrize('fault',['missing','unknown','duplicate','wrong_parent','top_level','not_list'])
def test_inventory_variants_fail_closed(tmp_path,fault):
 import json
 from pathlib import Path
 from core.product.aec_item_inventory_audit import verify_priority_inventory
 p=Path(__file__).resolve().parents[1]/'docs/aec_priority_1_9_master_inventory_20261013.json';d=json.loads(p.read_text());parent=next(r for r in d['items'] if r['code']=='waffle');variants=parent.get('review_contract_variants',[])
 if fault=='missing':parent['review_contract_variants']=[]
 elif fault=='unknown':parent['review_contract_variants']=variants+['unknown_waffle']
 elif fault=='duplicate':parent['review_contract_variants']=variants+['waffle_one_way']
 elif fault=='wrong_parent':
  parent['review_contract_variants']=[];next(r for r in d['items'] if r['code']=='solid_slab')['review_contract_variants']=variants
 elif fault=='top_level':
  child=deepcopy(parent);child.update(code='waffle_one_way',review_contract_variants=[]);d['items'].append(child)
  d['priority_family_count']=len(d['items']);d['component_mentions_count']=sum(len(r['components']) for r in d['items'])
 elif fault=='not_list':parent['review_contract_variants']='waffle_one_way'
 target=tmp_path/'invalid.json';target.write_text(json.dumps(d,ensure_ascii=False))
 assert not verify_priority_inventory(target)['valid']
