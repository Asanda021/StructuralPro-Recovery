"""Explicit synthetic horizontal insulation, not product/design acceptance."""
from copy import deepcopy
from decimal import Decimal
from hashlib import sha256
import pytest
from test_floor_specialist_work import sample as floor_sample
from core.takeoff.reviewed_aec import empty_system,preview_system
from core.takeoff.engine import TakeoffEngine
KINDS=('waterproofing','thermal_insulation','acoustic_insulation')
def sample(kind='thermal_insulation'):
 d=floor_sample();d['kind']=kind;d['components']=empty_system(kind,'L1','F1')['components'];d['source']['source_id']='SYNTHETIC-INSULATION-'+kind
 g=d['architecture_detail']['geometry'];g.update(installation_plane='horizontal',installation_ref='سطح افقی و تراز مستند مصنوعی')
 codes={'waterproofing':'membrane_layer','thermal_insulation':'thermal_layer','acoustic_insulation':'acoustic_layer'}
 for p in d['architecture_detail']['layers']:
  p['operation']=codes[kind];p.update(purchase_quantity='',purchase_unit='',purchase_ref='')
  if kind=='waterproofing':p['payment_unit']='m2'
 return d

@pytest.mark.parametrize('kind',KINDS)
def test_independent_net_area_and_volume(kind):
 r=preview_system(sample(kind),TakeoffEngine());assert r['ready'],r['blockers'];assert [Decimal(str(v['quantity'])) for v in r['rows']]==([Decimal(19)]*3 if kind=='waterproofing' else [Decimal('1.9'),Decimal('.76'),Decimal(19)])
 assert [Decimal(str(v['params']['net_volume_m3'])) for v in r['rows']]==[Decimal('1.9'),Decimal('.76'),Decimal('.19')]
 assert all(v['source_id'].startswith('aec:'+kind+':') and v['role']=='work' for v in r['rows'])

@pytest.mark.parametrize('kind',KINDS)
@pytest.mark.parametrize('fault',['plane','plane_ref','zero','negative','unit','overlap','duplicate_hole','unknown_rolls','purchase','operation','shape','holes','waste','source','liquid','variable'])
def test_fail_closed(kind,fault):
 d=sample(kind);a=d['architecture_detail'];g=a['geometry'];p=a['layers'][1]
 if fault=='plane':g['installation_plane']='sloped'
 if fault=='plane_ref':g['installation_ref']=''
 if fault=='zero':p['thickness_m']='0'
 if fault=='negative':p['thickness_m']='-.1'
 if fault=='unit':p['payment_unit']='litre'
 if fault=='overlap':p['bottom_m']='.09'
 if fault=='duplicate_hole':a['openings'].append({**a['openings'][0],'id':'OTHER','physical_key':'OTHER'})
 if fault=='unknown_rolls':p['roll_count']='2'
 if fault=='purchase':p.update(purchase_quantity='20',purchase_unit='m2',purchase_ref='فاقد جانمایی')
 if fault=='operation':p['operation']='floor_screed'
 if fault=='shape':g['shape']='irregular'
 if fault=='holes':p['opening_ids']=''
 if fault=='waste':p['waste']='.1'
 if fault=='source':p['material_ref']=''
 if fault=='liquid':p['operation']='liquid_membrane'
 if fault=='variable':p['method']='variable_thickness'
 assert not preview_system(d,TakeoffEngine())['ready']

@pytest.mark.parametrize('kind',KINDS)
@pytest.mark.parametrize('case',['valid','cas','duplicate','cross_floor','legacy','same_reviewer','financial','evidence','unit','identity','copy'])
def test_existing_store_boq_amendment(tmp_path,kind,case):
 from core.platform.application import StructuralProApp
 a=StructuralProApp(tmp_path/'db');a.create_project('آزمون مصنوعی عایق','P');d=sample(kind)
 if case=='legacy':
  old=empty_system(kind,'L1','F1');old['source']={'mode':'manual','source_id':'OLD'};a.save_aec_system_draft('P',old,expected_revision=0)
  with pytest.raises(ValueError):a.save_aec_system_draft('P',d,expected_revision=1)
  a.store.close();return
 a.save_aec_system_draft('P',d,expected_revision=0)
 if case=='cas':
  b=StructuralProApp(tmp_path/'db');stale=b.get_aec_system('P',kind,'L1','F1');a.save_aec_system_draft('P',d,expected_revision=1)
  with pytest.raises(RuntimeError):b.save_aec_system_draft('P',stale['data'],expected_revision=1)
  b.store.close();a.store.close();return
 if case in ('duplicate','cross_floor'):
  other=deepcopy(d) if case=='duplicate' else floor_sample();other['system_id']='OTHER';other['source']['source_id']='OTHER'
  if case=='cross_floor':
   for table in ('openings','layers'):
    for p in other['architecture_detail'][table]:
     p['physical_key']='OTHER-'+p['physical_key'];p['id']='OTHER-'+p['id']
     if 'layer_key' in p:p['layer_key']='OTHER-'+p['layer_key']
     if p.get('opening_ids'):p['opening_ids']='OTHER-H1'
  with pytest.raises(ValueError):a.save_aec_system_draft('P',other,expected_revision=0)
  a.store.close();return
 if case=='copy':
  v=a.aec_workflow().copy_roof('P',kind,'L1','F1',new_system_id='COPY',new_source_id='COPY',reason='آزمون مصنوعی')['data'];assert v['architecture_detail']['geometry']['physical_region_id']!='FLOOR-1';assert not preview_system(v,a.takeoff)['ready'];a.aec_workflow().delete_roof('P',kind,'L1','COPY',expected_revision=1,reason='حذف مصنوعی');a.store.close();return
 a.review_roof_system('P',kind,'L1','F1',expected_revision=1,reviewer='بازبین',note='کنترل مصنوعی');a.commit_aec_system('P',kind,'L1','F1',expected_revision=2);before=deepcopy(a.open_project('P'));assert sum(len(r['source_ids']) for r in before['boq'])==3;assert sum(Decimal(str(r['quantity'])) for r in before['boq'])==(Decimal('57') if kind=='waterproofing' else Decimal('21.66'))
 illegal=deepcopy(before);next(iter(illegal['reviewed_aec_systems'].values()))['data']['axis']='illegal'
 with pytest.raises(ValueError):a.store.save('P',illegal)
 d=deepcopy(a.get_aec_system('P',kind,'L1','F1')['data']);d.pop('roof_review');d['architecture_detail']['layers'][1]['thickness_m']='.05';d['architecture_detail']['layers'][2]['bottom_m']='.15'
 if case=='unit':d['architecture_detail']['layers'][1]['payment_unit']='kg';d['architecture_detail']['layers'][1].update(density_kg_m3='100',density_ref='مصنوعی')
 if case=='identity':d['architecture_detail']['geometry']['physical_region_id']='CHANGED'
 e=tmp_path/'new.txt';e.write_text('EXPLICIT SYNTHETIC new insulation thickness .04 -> .05; third layer moves to .15.');kw=dict(expected_revision=2,created_by='تهیه‌کننده',reason='دیتایل اصلاحی مصنوعی',evidence=dict(evidence_path=str(e),evidence_sha256=sha256(e.read_bytes()).hexdigest(),evidence_reference='مصنوعی'))
 if case in ('unit','identity'):
  with pytest.raises(ValueError):a.create_roof_amendment('P',kind,'L1','F1',d,**kw)
 else:
  am=a.create_roof_amendment('P',kind,'L1','F1',d,**kw)
  if case=='same_reviewer':
   with pytest.raises(ValueError):a.review_beam_amendment('P',am['id'],expected_revision=1,reviewer='تهیه‌کننده',note='نامجاز')
  elif case=='evidence':
   e.write_text('tampered')
   with pytest.raises(ValueError):a.review_beam_amendment('P',am['id'],expected_revision=1,reviewer='دیگری',note='کنترل')
  else:
   a.review_beam_amendment('P',am['id'],expected_revision=1,reviewer='دیگری',note='مستقل')
   if case=='financial':
    p=a.open_project('P');p['contract']={'id':'locked'};a.store.save('P',p);locked=deepcopy(a.open_project('P'))
    with pytest.raises(ValueError):a.apply_beam_amendment('P',am['id'],expected_revision=1)
    assert a.open_project('P')==locked
   else:
    a.apply_beam_amendment('P',am['id'],expected_revision=1);assert a.open_project('P')['takeoffs']==before['takeoffs'];a.store.close();a=StructuralProApp(tmp_path/'db');assert a.get_aec_system('P',kind,'L1','F1')['revision']==3;assert a.store.db.execute('pragma integrity_check').fetchone()[0]=='ok'
 a.store.close()

@pytest.mark.parametrize('kind',KINDS)
def test_four_reports_reopen_exact_decimals(tmp_path,kind):
 from core.platform.application import StructuralProApp
 from core.reports.architecture_detail_report import export
 from openpyxl import load_workbook
 from docx import Document
 import pymupdf,csv
 a=StructuralProApp(tmp_path/'db');a.create_project('نمونه کاملاً مصنوعی عایق','P');d=sample(kind);d['architecture_detail']['layers']=d['architecture_detail']['layers'][1:2]
 if kind!='waterproofing':d['architecture_detail']['layers'][0]['thickness_m']='.040000000000000000001'
 a.save_aec_system_draft('P',d,expected_revision=0);a.review_roof_system('P',kind,'L1','F1',expected_revision=1,reviewer='مستقل',note='کنترل مصنوعی');a.commit_aec_system('P',kind,'L1','F1',expected_revision=2)
 for fmt in ('pdf','xlsx','docx','csv'):export(a,'P',kind,'L1','F1',tmp_path/('insulation.'+fmt),format=fmt)
 b=load_workbook(tmp_path/'insulation.xlsx');assert all(s.sheet_view.rightToLeft for s in b)
 if kind!='waterproofing':assert any(c.value=='0.760000000000000000019' for s in b for r in s for c in r)
 b.close();doc=Document(tmp_path/'insulation.docx');assert all(t._tbl.tblPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bidiVisual') is not None for t in doc.tables)
 pdf=pymupdf.open(tmp_path/'insulation.pdf');assert len(pdf)>0;assert all(p.rect.width>0 for p in pdf);pdf.close()
 text=(tmp_path/'insulation.csv').read_text(encoding='utf-8-sig');assert 'سطح نصب' in text and 'خرید، رول' in text
 if kind!='waterproofing':assert '0.760000000000000000019' in text
 a.store.close()

@pytest.mark.parametrize('kind',KINDS)
def test_explicit_density_no_material_defaults(kind):
 d=sample(kind);d['architecture_detail']['layers']=d['architecture_detail']['layers'][:1];p=d['architecture_detail']['layers'][0];p.update(payment_unit='kg',density_kg_m3='25',density_ref='گواهی چگالی کاملاً مصنوعی')
 r=preview_system(d,TakeoffEngine())
 if kind=='waterproofing':assert not r['ready']
 else:assert r['ready'] and Decimal(str(r['rows'][0]['quantity']))==Decimal('47.5')


def test_duplicate_source_with_disjoint_geometry_rejected(tmp_path):
 from core.platform.application import StructuralProApp
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');d=sample();a.save_aec_system_draft('P',d,expected_revision=0);other=deepcopy(d);other['system_id']='OTHER';other['architecture_detail']['geometry']['physical_region_id']='OTHER'
 for table in ('openings','layers'):
  for p in other['architecture_detail'][table]:p['physical_key']='OTHER-'+p['physical_key']
 with pytest.raises(ValueError):a.save_aec_system_draft('P',other,expected_revision=0)
 a.store.close()


def test_extends_saved_five_floor_data(tmp_path):
 from scripts.verify_architecture_step6 import create_sample
 from scripts.verify_floor_step6_package2 import extend_sample as floors
 from scripts.verify_insulation_step6_package3 import extend_sample
 create_sample(tmp_path/'p1');floors(tmp_path/'p1',tmp_path/'p2');r=extend_sample(tmp_path/'p2',tmp_path/'p3');assert r['previous_records_preserved'] and r['previous_boq_preserved'] and r['other_four_floor_quantities_unchanged'];assert r['boq_rows']==43 and r['insulation_effective_totals']['thermal_layer|m3']=='5.51'


def test_insulation_evidence_requires_all_six_limited_gates_and_no_final_acceptance():
 from core.product.aec_evidence_matrix import MATRIX_FILE,assert_family_evidence
 import json
 rows={r['code']:r for r in json.loads(MATRIX_FILE.read_text())['rows']}
 for kind in KINDS:
  r=rows[kind];assert r['acceptance_status']=='limited_documented';assert not r['fully_accepted'];assert all(r['gates'][g]=='not_verified' for g in ('drawing_takeoff','real_engineering_case','windows_office_acceptance'));assert all(r['gates'][g]=='documented_limited' for g in ('engineering_calculation','manual_ui','sqlite_history','review_amendment','boq_dedup','report_exports'))
 assert assert_family_evidence()['family_count']==85


@pytest.mark.parametrize('kind,code',[(k,c) for k,codes in {'waterproofing':('membrane_layer','bituminous_layer','insulation_protection','insulation_separator'),'thermal_insulation':('thermal_layer','insulation_protection','insulation_separator'),'acoustic_insulation':('acoustic_layer','insulation_protection','insulation_separator')}.items() for c in codes])
def test_each_explicit_operation(kind,code):
 d=sample(kind);d['architecture_detail']['layers']=d['architecture_detail']['layers'][:1];d['architecture_detail']['layers'][0]['operation']=code;r=preview_system(d,TakeoffEngine());assert r['ready'],r['blockers'];assert r['rows'][0]['item_code']==code and Decimal(str(r['rows'][0]['quantity']))==(Decimal('19') if kind=='waterproofing' else Decimal('1.9'))
