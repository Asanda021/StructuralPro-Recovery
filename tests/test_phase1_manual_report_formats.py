"""Use the existing manual report dataset and shared exporters losslessly."""
from copy import deepcopy
from pathlib import Path
import csv
import pytest
from core.platform.application import StructuralProApp
from test_manual_formula_persistence import data

@pytest.mark.parametrize('fmt',['docx','csv'])
def test_manual_report_editable_formats_preserve_exact_data_and_sqlite(tmp_path,fmt):
 app=StructuralProApp(tmp_path/'db');app.create_project('نمونه گزارش مستند','P')
 try:
  number='12345678901234567890.123456789'
  app.save_manual_formula('P','A',data(expression='n',variables={'n':number},variable_units={'n':'m3'},title='=HYPERLINK("unsafe")'),expected_revision=0,reason='ورود مستند مصنوعی')
  before=deepcopy(app.open_project('P'));revision=app.store.db.execute('select * from revisions order by project_id,version').fetchall()
  out=tmp_path/('manual.'+fmt);app.export_manual_formula_report('P',out,format=fmt)
  if fmt=='docx':
   from docx import Document
   from docx.oxml.ns import qn
   doc=Document(out);assert doc.tables
   assert all(t._tbl.tblPr.find(qn('w:bidiVisual')) is not None for t in doc.tables)
   text='\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
   assert '=HYPERLINK("unsafe")' in text
  else:
   assert out.read_bytes().startswith(b'\xef\xbb\xbf')
   rows=list(csv.reader(out.open(encoding='utf-8-sig',newline='')))
   text='\n'.join(v for r in rows for v in r)
   assert not any(v.lstrip().startswith('=HYPERLINK') for r in rows for v in r)
  assert number in text and 'بررسی‌نشده' in text and 'متر مکعب' in text
  assert app.open_project('P')==before and not before['boq']
  assert app.store.db.execute('select * from revisions order by project_id,version').fetchall()==revision
 finally:app.store.close()
