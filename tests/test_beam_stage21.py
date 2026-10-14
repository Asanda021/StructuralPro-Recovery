from PySide6.QtWidgets import QApplication
from copy import deepcopy
from decimal import Decimal
import pytest
from core.takeoff.beam_detail import preview_beam,BeamDetailWorkflow
from core.takeoff.engine import TakeoffEngine
from core.platform.application import StructuralProApp
from test_stirrup_zones import complete_beam,valid


def specialist(count=1):
 d=complete_beam();d['measurement_policy']='beam_specialist_v1';d['source']['plan_ref']='پلان مصنوعی B1';d['source']['detail_ref']='دیتایل مصنوعی D1'
 d['geometry'].update(count=count,variable_units={'length':'m','width':'m','depth':'m','count':'عدد'},deduction_scope='total_group',deduction_scope_ref='کسورات کل گروه صریح',cover_m='.03',cover_ref='پوشش صریح در دیتایل')
 d['formwork'].update(area_scope='per_member',area_scope_ref='سطح یک تیر هم‌تیپ')
 d['bars']['stirrups'].update(count_scope='per_member',count_scope_ref='جانمایی یک تیر هم‌تیپ')
 for i,z in enumerate(d['bars']['stirrups']['segments']):z.update(id='zone'+str(i),nominal_diameter_mm='8',shape_ref='خاموت مستطیلی D1',shape_kind='rectangular',section_width_m='.2',section_depth_m='.35',section_ref='ابعاد بیرونی صریح D1',cut_length_ref='طول محور برش شامل خم/قلاب D1')
 return d


def row(report,suffix):return next(x for x in report['rows'] if x['source_id'].endswith(suffix))

def test_repeated_beam_ambiguous_counts_and_formwork_block():
 d=complete_beam();d['geometry']['count']=2
 r=preview_beam(d,TakeoffEngine());assert not r['ready'],'دامنه تعداد و قالب تیرهای هم‌تیپ نامشخص پذیرفته شد'

def test_specialist_repeats_exact_independent_reference():
 d=specialist(2);r=preview_beam(d,TakeoffEngine());assert r['ready'],r['blockers']
 assert Decimal(str(row(r,':concrete')['quantity']))==Decimal('1.5')
 assert Decimal(str(row(r,':formwork')['quantity']))==Decimal('12')
 assert Decimal(str(row(r,':stirrups')['quantity']))==Decimal('38.5125')
 assert row(r,':stirrups')['params']['count_scope']=='per_member'

def test_source_backed_variable_sections_and_deductions():
 d=specialist();d['geometry']['sections']=[{'id':'s1','start_m':'0','end_m':'2','width_m':'.3','depth_m':'.5','source_ref':'D1','approved_by':'مهندس'},{'id':'s2','start_m':'2','end_m':'5','width_m':'.4','depth_m':'.6','source_ref':'D2','approved_by':'مهندس'}]
 d['geometry'].pop('deductions_m3');d['geometry']['intersections']=[{'id':'joint-column','volume_m3':'.02','source_ref':'اتصال مصنوعی J1','approved_by':'مهندس'}]
 r=preview_beam(d,TakeoffEngine());assert r['ready'],r['blockers'];assert Decimal(str(row(r,':concrete')['quantity']))==Decimal('1.0')
 assert len(row(r,':concrete')['params']['sections'])==2

@pytest.mark.parametrize('mutate',[
 lambda d:d['geometry'].pop('cover_m'),
 lambda d:d['geometry']['variable_units'].update(width='m2'),
 lambda d:d['bars']['stirrups']['segments'][0].update(count=99),
 lambda d:d['bars']['stirrups']['segments'][0].update(unit_weight_kg_m='NaN'),
 lambda d:d['bars']['stirrups']['segments'][0].pop('nominal_diameter_mm'),
 lambda d:d['source'].pop('detail_ref'),
 lambda d:d['geometry'].update(intersections=[{'id':'x','volume_m3':'.01','source_ref':'J','approved_by':'مهندس'}]*2),
])
def test_missing_or_invalid_specialist_detail_blocks(mutate):
 d=specialist();mutate(d);assert not preview_beam(d,TakeoffEngine())['ready']


def test_duplicate_independent_bar_identity_blocks():
 d=specialist();bar={'status':'confirmed','approved_by':'مهندس','source_ref':'D1','count_scope':'per_member','count_scope_ref':'یک تیر','segments':[{'id':'same-physical','position':'top','count':2,'nominal_diameter_mm':16,'cut_length_m':5.4,'cut_length_ref':'محور شامل خم و مهار','unit_weight_kg_m':1.58,'unit_weight_source':'کاتالوگ صریح','source_ref':'D1'}]}
 d['bars']['longitudinal']=deepcopy(bar);d['bars']['additional']=deepcopy(bar)
 assert not preview_beam(d,TakeoffEngine())['ready'],'یک میلگرد فیزیکی در دو گروه دوباره شمرده شد'


def test_specialist_sqlite_boq_and_safe_reopen(tmp_path):
 s=StructuralProApp(tmp_path);s.create_project('مصنوعی','P');d=specialist(2);s.save_beam_detail_draft('P',d,expected_revision=0)
 r=s.preview_beam_detail('P','L1','B1');assert r['ready'],r['blockers'];s.commit_beam_detail('P','L1','B1',expected_revision=1);p=s.open_project('P');assert len(p['takeoffs'])==3;assert len({x['source_id'] for x in p['takeoffs']})==3
 with pytest.raises(ValueError):s.commit_beam_detail('P','L1','B1',expected_revision=1)
 s.store.close();s=StructuralProApp(tmp_path);assert s.get_beam_detail('P','L1','B1')['data']==d;assert len(s.open_project('P')['takeoffs'])==3;s.store.close()

def test_fractional_large_repeat_count_is_not_rounded_to_integer():
 d=specialist();d['geometry']['count']='9007199254740992.5';assert not preview_beam(d,TakeoffEngine())['ready']

def test_specialist_stations_cannot_pass_on_float_tolerance():
 d=specialist();z=d['bars']['stirrups']['segments'][0];z['last_bar_m']='1.150000001'
 assert not preview_beam(d,TakeoffEngine())['ready'],'جانمایی ناسازگار با تساهل float پذیرفته شد'

def test_stirrup_section_dimensions_are_mandatory():
 d=specialist();d['bars']['stirrups']['segments'][0].pop('section_width_m',None)
 assert not preview_beam(d,TakeoffEngine())['ready']

def test_beam_specialist_report_exact_units_and_pdf(tmp_path):
 from openpyxl import load_workbook
 import fitz
 s=StructuralProApp(tmp_path/'db');s.create_project('تیر مصنوعی','P');s.save_beam_detail_draft('P',specialist(2),expected_revision=0)
 assert callable(getattr(s,'export_beam_detail_report',None)),'گزارش تفصیلی تیر متصل نیست'
 s.export_beam_detail_report('P','L1','B1',tmp_path/'beam.xlsx',format='xlsx');w=load_workbook(tmp_path/'beam.xlsx');text=' '.join(str(c.value) for sh in w for rr in sh for c in rr)
 assert '38.5125' in text and 'کیلوگرم' in text and 'zone0' in text
 s.export_beam_detail_report('P','L1','B1',tmp_path/'beam.pdf',format='pdf')
 with fitz.open(tmp_path/'beam.pdf') as doc:assert len(doc)>0 and all(p.get_pixmap().width>0 for p in doc)
 s.store.close()

def test_native_beam_combined_scope_review_store_report_and_boq(tmp_path,monkeypatch,nonblocking_error_dialogs):
 from PySide6.QtWidgets import QApplication,QMessageBox
 from PySide6.QtCore import Qt
 from PySide6.QtTest import QTest
 from app.beam_detail_panel import BeamDetailPanel
 global _qt
 _qt=QApplication.instance() or QApplication([]);s=StructuralProApp(tmp_path);s.create_project('مصنوعی','P');d=BeamDetailPanel(s,lambda:'P',lambda:None);d.show();_qt.processEvents()
 assert hasattr(d,'specialist_policy'),'پنل با قرارداد تخصصی یکپارچه نیست'
 d.floor.setText('L1');d.beam.setText('B1');d.reviewer.setText('مهندس مصنوعی');d.plan_ref.setText('پلان مصنوعی P1');d.detail_ref.setText('دیتایل مصنوعی D1')
 for key,value in {'length':'5','width':'.3','depth':'.5','count':'2','deductions_m3':'0','cover_m':'.03','source_ref':'D1','cover_ref':'D1','deduction_scope_ref':'کسورات کل گروه'}.items():d.geometry[key].setText(value)
 d.deduction_scope.setCurrentIndex(d.deduction_scope.findData('total_group'));d.formwork['net_area_m2'].setText('6.5');d.formwork['source_ref'].setText('سطوح واقعی F1');d.formwork['area_scope_ref'].setText('یک تیر');d.formwork_scope.setCurrentIndex(d.formwork_scope.findData('per_member'))
 for kind,w in d.bars.items():w['status'].setCurrentIndex(w['status'].findData('not_applicable'));w['source_ref'].setText('D1');w['reason'].setText('عدم کاربرد صریح برای نمونه مصنوعی')
 w=d.bars['longitudinal'];w['status'].setCurrentIndex(w['status'].findData('confirmed'))
 for key,value in {'id':'top-1','count':'4','nominal_diameter_mm':'16','cut_length_m':'5.5','cut_length_ref':'محور شامل مهار/خم D1','unit_weight_kg_m':'1.58','unit_weight_source':'کاتالوگ مصنوعی','count_scope_ref':'یک تیر'}.items():w[key].setText(value)
 d.bar_scopes['longitudinal'].setCurrentIndex(d.bar_scopes['longitudinal'].findData('per_member'));d.bar_positions['longitudinal'].setCurrentIndex(d.bar_positions['longitudinal'].findData('top'))
 QTest.mouseClick(d.save_button,Qt.LeftButton);_qt.processEvents();assert s.get_beam_detail('P','L1','B1') is not None,nonblocking_error_dialogs;assert s.get_beam_detail('P','L1','B1')['revision']==1
 QTest.mouseClick(d.load_button,Qt.LeftButton);_qt.processEvents();assert d.formwork_scope.currentData()=='per_member'
 QTest.mouseClick(d.preview_button,Qt.LeftButton);_qt.processEvents();assert '69.52' in d.message.text() and '13' in d.message.text()
 # Explicit synthetic owner's confirmation only; assertion checks real SQLite/BOQ.
 from PySide6.QtCore import QTimer
 confirmed=[]
 def confirm():
  box=_qt.activeModalWidget();assert isinstance(box,QMessageBox);QTest.mouseClick(box.button(QMessageBox.Yes),Qt.LeftButton);confirmed.append(True)
 QTimer.singleShot(100,confirm)
 QTest.mouseClick(d.commit_button,Qt.LeftButton);_qt.processEvents();assert len(s.open_project('P')['takeoffs'])==3 and confirmed
 d.grab().save(str(tmp_path/'beam-native.png'));d.close();d.deleteLater();_qt.processEvents();s.store.close()

def test_beam_report_does_not_execute_formula_injection(tmp_path):
 from openpyxl import load_workbook
 s=StructuralProApp(tmp_path);s.create_project('مصنوعی','P');d=specialist();d['source']['detail_ref']='=HYPERLINK("bad")';s.save_beam_detail_draft('P',d,expected_revision=0);s.export_beam_detail_report('P','L1','B1',tmp_path/'safe.xlsx',format='xlsx')
 w=load_workbook(tmp_path/'safe.xlsx');cells=[c for sheet in w for rr in sheet for c in rr]
 assert any(c.value=='=HYPERLINK("bad")' and c.data_type=='s' for c in cells);assert not any(c.data_type=='f' for c in cells);s.store.close()

def test_native_advanced_geometry_is_not_silently_overwritten(tmp_path,nonblocking_error_dialogs):
 from PySide6.QtCore import Qt
 from PySide6.QtTest import QTest
 from app.beam_detail_panel import BeamDetailPanel
 global _qt
 _qt=QApplication.instance() or QApplication([]);s=StructuralProApp(tmp_path);s.create_project('مصنوعی','P');draft=specialist();draft['geometry']['intersections']=[];s.save_beam_detail_draft('P',draft,expected_revision=0);before=deepcopy(s.get_beam_detail('P','L1','B1'))
 d=BeamDetailPanel(s,lambda:'P',lambda:None);d.floor.setText('L1');d.beam.setText('B1');d.show();_qt.processEvents();QTest.mouseClick(d.load_button,Qt.LeftButton);_qt.processEvents()
 d.geometry['length'].setText('6');QTest.mouseClick(d.save_button,Qt.LeftButton);_qt.processEvents();assert s.get_beam_detail('P','L1','B1')==before and nonblocking_error_dialogs
 d.close();d.deleteLater();_qt.processEvents();s.store.close()

def test_native_stirrup_editor_preserves_specialist_detail(tmp_path,nonblocking_error_dialogs):
 from PySide6.QtCore import Qt
 from PySide6.QtTest import QTest
 from app.beam_stirrup_zone_editor import StirrupZoneEditor
 global _qt
 _qt=QApplication.instance() or QApplication([]);zones=specialist()['bars']['stirrups']['segments'];d=StirrupZoneEditor(beam_length_m=5,zones=zones,specialist=True);d.show();_qt.processEvents()
 QTest.mouseClick(d.save_button,Qt.LeftButton);_qt.processEvents();assert d.zones is not None and d.zones[0]['id']=='zone0' and d.zones[0]['shape_kind']=='rectangular' and d.zones[0]['section_width_m']=='.2'
 assert d.zones[0]['cut_length_ref']==zones[0]['cut_length_ref'];d.close();d.deleteLater();_qt.processEvents()

def test_specialist_cannot_be_downgraded_without_new_review(tmp_path):
 s=StructuralProApp(tmp_path);s.create_project('مصنوعی','P');d=specialist();s.save_beam_detail_draft('P',d,expected_revision=0);d.pop('measurement_policy')
 with pytest.raises(ValueError):s.save_beam_detail_draft('P',d,expected_revision=1)
 s.store.close()

def _commit_worker(path,event,queue):
 s=StructuralProApp(path);event.wait()
 try:s.commit_beam_detail('P','L1','B1',expected_revision=1);queue.put('COMMITTED')
 except (ValueError,RuntimeError):queue.put('BLOCKED')
 finally:s.store.close()

def test_beam_real_two_process_commit_and_sqlite_integrity(tmp_path):
 import multiprocessing as mp,sqlite3
 s=StructuralProApp(tmp_path);s.create_project('مصنوعی','P');s.save_beam_detail_draft('P',specialist(2),expected_revision=0);s.store.close();ctx=mp.get_context('spawn');event=ctx.Event();queue=ctx.Queue();ps=[ctx.Process(target=_commit_worker,args=(str(tmp_path),event,queue)) for _ in range(2)]
 for p in ps:p.start()
 event.set();results=[queue.get(timeout=15) for p in ps]
 for p in ps:p.join(15);assert p.exitcode==0
 assert sorted(results)==['BLOCKED','COMMITTED'];s=StructuralProApp(tmp_path);assert len(s.open_project('P')['takeoffs'])==3;s.store.close()
 for path in tmp_path.rglob('*.db'):
  with sqlite3.connect(path) as c:assert c.execute('PRAGMA integrity_check').fetchone()[0]=='ok'

@pytest.mark.parametrize('mode',['fixed','zones','joint'])
def test_three_synthetic_scenarios_reference_export_reopen_boq(tmp_path,mode):
 s=StructuralProApp(tmp_path/'db');s.create_project('مصنوعی؛ طراحی تأییدشده نیست','P');d=specialist()
 if mode=='fixed':d['bars']['stirrups']={'status':'not_applicable','approved_by':'مهندس مصنوعی','source_ref':'دیتایل صریح نمونه بدون خاموت؛ طراحی نیست','reason':'آزمون مستقل بتن و قالب'}
 if mode=='joint':
  d['geometry'].pop('deductions_m3');d['geometry']['intersections']=[{'id':'joint:test','volume_m3':'.05','source_ref':'هندسه مصنوعی اتصال','approved_by':'مهندس مصنوعی'}]
  d['bars']['additional']={'status':'confirmed','approved_by':'مهندس مصنوعی','source_ref':'تقویتی D2','count_scope':'per_member','count_scope_ref':'یک تیر','segments':[{'id':'support-top','position':'support','count':2,'nominal_diameter_mm':16,'cut_length_m':'1.5','cut_length_ref':'محور برش شامل مهار صریح','unit_weight_kg_m':'1.58','unit_weight_source':'مرجع مصنوعی وزن','source_ref':'D2'}]}
 s.save_beam_detail_draft('P',d,expected_revision=0);report=s.preview_beam_detail('P','L1','B1');assert report['ready'],report['blockers']
 expected=Decimal('.70') if mode=='joint' else Decimal('.75');assert Decimal(str(row(report,':concrete')['quantity']))==expected
 s.export_beam_detail_report('P','L1','B1',tmp_path/(mode+'.xlsx'),format='xlsx');s.export_beam_detail_report('P','L1','B1',tmp_path/(mode+'.pdf'),format='pdf');s.commit_beam_detail('P','L1','B1',expected_revision=1);s.store.close()
 s=StructuralProApp(tmp_path/'db');assert s.get_beam_detail('P','L1','B1')['status']=='committed' and len(s.open_project('P')['takeoffs'])==len(report['rows']);s.store.close()
