"""Explicit synthetic schedules; no real engineering approval is implied."""
from copy import deepcopy
from hashlib import sha256
import pytest
from core.platform.application import StructuralProApp
from test_foundation_tie_stage23_chat import sample

KINDS=['tie_beam','isolated_footing','strip_footing','mat_footing','pile_cap']

def setup_case(tmp_path,kind):
 app=StructuralProApp(tmp_path/'db');app.create_project('نمونه مصنوعی','P')
 d=sample(kind);app.save_structural_member_draft('P',d,expected_revision=0)
 app.commit_structural_member('P',kind,'L1','T1',expected_revision=1)
 evidence=tmp_path/'detail.txt';evidence.write_text('Synthetic revised detail, explicitly not a real plan')
 kw=dict(expected_revision=1,created_by='سازنده مصنوعی',reason='تغییر دیتایل مصنوعی',evidence={'evidence_path':str(evidence),'evidence_sha256':sha256(evidence.read_bytes()).hexdigest(),'evidence_reference':'دیتایل جدید مصنوعی'})
 return app,d,kw

@pytest.mark.parametrize('kind',KINDS)
def test_effective_amendment_preserves_original_and_rebuilds_boq(tmp_path,kind):
 app,d,kw=setup_case(tmp_path,kind);original=deepcopy(app.open_project('P'))
 d['geometry']['count']='3'
 a=app.create_foundation_amendment('P',kind,'L1','T1',d,**kw)
 assert any(v!='0' for v in a['delta'].values())
 app.review_beam_amendment('P',a['id'],expected_revision=1,reviewer='بازبین مستقل مصنوعی',note='کنترل مصنوعی')
 app.apply_beam_amendment('P',a['id'],expected_revision=1)
 p=app.open_project('P');assert p['takeoffs']==original['takeoffs'];assert p['reviewed_structural_members']==original['reviewed_structural_members'];assert p['boq']!=original['boq']
 assert app.get_structural_member('P',kind,'L1','T1')['revision']==2
 with pytest.raises(ValueError):app.apply_beam_amendment('P',a['id'],expected_revision=1)
 app.store.close();app=StructuralProApp(tmp_path/'db');assert app.get_structural_member('P',kind,'L1','T1')['data']['geometry']['count']=='3'

@pytest.mark.parametrize('invalid',['self_review','evidence','identity','stale','financial'])
def test_invalid_foundation_amendments(tmp_path,invalid):
 app,d,kw=setup_case(tmp_path,'tie_beam');d['geometry']['count']='3'
 if invalid=='identity':
  d['member_id']='different'
  with pytest.raises(ValueError):app.create_foundation_amendment('P','tie_beam','L1','T1',d,**kw)
  return
 if invalid=='stale':
  kw['expected_revision']=2
  with pytest.raises(RuntimeError):app.create_foundation_amendment('P','tie_beam','L1','T1',d,**kw)
  return
 a=app.create_foundation_amendment('P','tie_beam','L1','T1',d,**kw)
 if invalid=='evidence':
  from pathlib import Path
  Path(kw['evidence']['evidence_path']).write_text('changed')
  with pytest.raises(ValueError):app.review_beam_amendment('P',a['id'],expected_revision=1,reviewer='مستقل',note='کنترل')
 elif invalid=='self_review':
  with pytest.raises(ValueError):app.review_beam_amendment('P',a['id'],expected_revision=1,reviewer=kw['created_by'],note='کنترل')
 else:
  app.review_beam_amendment('P',a['id'],expected_revision=1,reviewer='مستقل',note='کنترل')
  p=app.open_project('P');p['contract']={'id':'locked'};app.store.save('P',p)
  with pytest.raises(ValueError):app.apply_beam_amendment('P',a['id'],expected_revision=1)

@pytest.mark.parametrize('fmt',['pdf','xlsx'])
def test_foundation_report_real_file(tmp_path,fmt):
 app,d,kw=setup_case(tmp_path,'tie_beam')
 path=tmp_path/('foundation.'+fmt)
 app.export_foundation_detail_report('P','tie_beam','L1','T1',path,format=fmt)
 if fmt=='pdf':
  import fitz
  with fitz.open(path) as doc:assert len(doc)>0 and any(p.get_text().strip() for p in doc)
 else:
  from openpyxl import load_workbook
  book=load_workbook(path);assert book.active.max_row>10
 app.store.close()


def test_native_foundation_load_edit_save(tmp_path,nonblocking_error_dialogs):
 from PySide6.QtWidgets import QApplication
 from PySide6.QtCore import Qt,QTimer
 from PySide6.QtTest import QTest
 from app.reviewed_structural_manual_editor import open_manual_structural_editor
 global _qt
 _qt=QApplication.instance() or QApplication([])
 app=StructuralProApp(tmp_path);app.create_project('مصنوعی','P');app.save_structural_member_draft('P',sample(),expected_revision=0)
 errors=[]
 def run():
  dialog=_qt.activeModalWidget()
  try:
   c=dialog.column_controls;c['kind'].setCurrentIndex(c['kind'].findData('tie_beam'));c['id'].setText('T1');QTest.mouseClick(c['load'],Qt.LeftButton)
   assert c['specialist'].isChecked()
   QTest.mouseClick(c['review'],Qt.LeftButton);QTest.mouseClick(c['save'],Qt.LeftButton)
   saved=app.get_structural_member('P','tie_beam','L1','T1');assert saved['revision']==2,nonblocking_error_dialogs
   assert saved['data']['measurement_policy']=='foundation_tie_specialist_v1'
   assert saved['data']['rebars']==sample()['rebars']
  except Exception as e:errors.append(repr(e))
  finally:dialog.reject()
 QTimer.singleShot(100,run);open_manual_structural_editor(app,project_id='P',floor_id='L1');assert not errors,errors
 app.store.close()


def test_copy_new_identity_and_delete_draft_only(tmp_path):
 app,d,kw=setup_case(tmp_path,'tie_beam');original=deepcopy(app.open_project('P'))
 copied=app.copy_structural_member('P','tie_beam','L1','T1',new_member_id='T2',new_source_id='synthetic:new-source',expected_revision=1,reason='نمونه مصنوعی مستقل')
 assert copied['revision']==1 and copied['data']['member_id']=='T2'
 assert copied['data']['source']['source_id']=='synthetic:new-source'
 assert copied['data']['geometry'].get('approved_by') is None
 assert not app.preview_structural_member('P','tie_beam','L1','T2')['ready']
 with pytest.raises(ValueError):app.delete_structural_member_draft('P','tie_beam','L1','T1',expected_revision=1,reason='ممنوع')
 with pytest.raises(RuntimeError):app.delete_structural_member_draft('P','tie_beam','L1','T2',expected_revision=2,reason='تعارض')
 app.delete_structural_member_draft('P','tie_beam','L1','T2',expected_revision=1,reason='حذف پیش‌نویس مصنوعی')
 assert app.get_structural_member('P','tie_beam','L1','T2') is None
 p=app.open_project('P');assert p['takeoffs']==original['takeoffs'];assert p['structural_member_history'][-1]['reason']=='حذف پیش‌نویس مصنوعی'
 app.store.close()


def test_duplicate_open_amendment_and_direct_mutation(tmp_path):
 app,d,kw=setup_case(tmp_path,'tie_beam');d['geometry']['count']='3'
 app.create_foundation_amendment('P','tie_beam','L1','T1',d,**kw)
 with pytest.raises(ValueError):app.create_foundation_amendment('P','tie_beam','L1','T1',d,**kw)
 p=app.open_project('P');next(iter(p['reviewed_structural_members'].values()))['data']['geometry']['count']='99'
 with pytest.raises(ValueError):app.store.save('P',p)
 app.store.close()


def test_native_foundation_amendment_clicks(tmp_path,nonblocking_error_dialogs):
 from PySide6.QtWidgets import QApplication
 from PySide6.QtCore import Qt,QTimer
 from PySide6.QtTest import QTest
 from app.reviewed_structural_manual_editor import open_manual_structural_editor
 global _qt
 _qt=QApplication.instance() or QApplication([]);app,d,kw=setup_case(tmp_path,'tie_beam');original=deepcopy(app.open_project('P'));errors=[]
 def run():
  dialog=_qt.activeModalWidget()
  try:
   c=dialog.column_controls;c['kind'].setCurrentIndex(c['kind'].findData('tie_beam'));c['id'].setText('T1');QTest.mouseClick(c['load'],Qt.LeftButton)
   c['numeric']['count'].setText('3');c['creator'].setText(kw['created_by']);c['reason'].setText(kw['reason']);c['reviewer'].setText('بازبین مستقل');c['note'].setText('کنترل مصنوعی');c['evidence_path'].setText(kw['evidence']['evidence_path']);c['evidence_ref'].setText(kw['evidence']['evidence_reference'])
   for action in ('create','review','apply'):QTest.mouseClick(c['amend'][action],Qt.LeftButton)
   assert app.get_structural_member('P','tie_beam','L1','T1')['revision']==2
   assert app.open_project('P')['takeoffs']==original['takeoffs']
   assert c['copy_member'].isEnabled() and c['delete_member'].isEnabled()
   c['new_member'].setText('T2');c['new_source'].setText('synthetic:T2');QTest.mouseClick(c['copy_member'],Qt.LeftButton)
   assert app.get_structural_member('P','tie_beam','L1','T2')['revision']==1
   QTest.mouseClick(c['delete_member'],Qt.LeftButton);assert app.get_structural_member('P','tie_beam','L1','T2') is None
  except Exception as e:errors.append(repr(e))
  finally:dialog.reject()
 QTimer.singleShot(100,run);open_manual_structural_editor(app,project_id='P',floor_id='L1');assert not errors,errors
 app.store.close()


def test_foundation_report_labels_are_persian(tmp_path):
 from openpyxl import load_workbook
 app,d,kw=setup_case(tmp_path,'tie_beam');path=tmp_path/'persian.xlsx'
 app.export_foundation_detail_report('P','tie_beam','L1','T1',path,format='xlsx')
 book=load_workbook(path);labels=[str(c.value) for row in book.active for c in row if c.value is not None and not str(c.value).startswith('structural:')]
 assert not any('approved_by' in v or 'intersection_scope' in v or 'longitudinal' in v for v in labels)
 app.store.close()


def test_deleted_identity_cannot_be_reused(tmp_path):
 app=StructuralProApp(tmp_path);app.create_project('مصنوعی','P');d=sample();app.save_structural_member_draft('P',d,expected_revision=0)
 app.delete_structural_member_draft('P','tie_beam','L1','T1',expected_revision=1,reason='حذف صریح')
 with pytest.raises(ValueError):app.save_structural_member_draft('P',d,expected_revision=0)
 app.store.close()


def test_copy_rechecks_source_in_write_transaction(tmp_path,monkeypatch):
 from core.takeoff.reviewed_structural import ReviewedStructuralWorkflow
 app,d,kw=setup_case(tmp_path,'tie_beam');original_save=ReviewedStructuralWorkflow.save_draft
 def concurrent_change(flow,*args,**kwargs):
  changed=deepcopy(d);changed['geometry']['count']='3'
  a=app.create_foundation_amendment('P','tie_beam','L1','T1',changed,**kw)
  app.review_beam_amendment('P',a['id'],expected_revision=1,reviewer='مستقل',note='کنترل')
  app.apply_beam_amendment('P',a['id'],expected_revision=1)
  return original_save(flow,*args,**kwargs)
 monkeypatch.setattr(ReviewedStructuralWorkflow,'save_draft',concurrent_change)
 with pytest.raises(RuntimeError):app.copy_structural_member('P','tie_beam','L1','T1',new_member_id='T2',new_source_id='synthetic:T2',expected_revision=1,reason='کپی')
 assert app.get_structural_member('P','tie_beam','L1','T2') is None
 app.store.close()
