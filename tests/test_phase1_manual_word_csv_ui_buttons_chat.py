"""Real offscreen Qt clicks for both manual-takeoff Word/CSV report paths.

UI actions must use the already-existing exporters and StructuralProApp API;
reports must not commit a preview, change a revision, or modify SQLite/BOQ.
These tests intentionally require PySide6 and a real application store.
"""
from copy import deepcopy
import csv

import pytest

from PySide6.QtCore import Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox

from core.platform.application import StructuralProApp
from app.manual_takeoff import ManualTakeoffDialog
from test_manual_formula_persistence import data


@pytest.fixture
def qt():
    app = QApplication.instance() or QApplication([])
    yield app
    app.processEvents()


@pytest.mark.parametrize("fmt,button_name", (
    ("docx", "formula_docx_button"),
    ("csv", "formula_csv_button"),
))
def test_full_manual_formula_report_export_button_is_real_and_readonly(
    tmp_path, monkeypatch, qt, fmt, button_name
):
    svc = StructuralProApp(tmp_path / "db")
    svc.create_project("آزمون خروجی Word و CSV", "P")
    svc.save_manual_formula(
        "P", "A", data(), expected_revision=0, reason="مبنای آزمایشی مستند"
    )
    before = deepcopy(svc.open_project("P"))
    revision_count = svc.store.db.execute(
        "SELECT COUNT(*) FROM revisions WHERE project_id='P'"
    ).fetchone()[0]
    out = tmp_path / ("formula." + fmt)
    dialogs = []
    monkeypatch.setattr(
        QFileDialog, "getSaveFileName",
        lambda *a, **k: (str(out), ""),
    )
    monkeypatch.setattr(
        QMessageBox, "warning",
        lambda *a, **k: dialogs.append(str(a[-1])),
    )
    dlg = ManualTakeoffDialog(svc, "P")
    try:
        dlg.show()
        dlg.manual_tabs.setCurrentIndex(1)
        qt.processEvents()
        button = getattr(dlg, button_name)
        assert button.isEnabled() and button.isVisible()
        QTest.mouseClick(button, Qt.LeftButton)
        qt.processEvents()
        assert not dialogs, dialogs
        assert out.is_file() and out.stat().st_size > 100
        if fmt == "csv":
            assert out.read_bytes().startswith(b"\xef\xbb\xbf")
            with out.open(encoding="utf-8-sig", newline="") as handle:
                records = list(csv.reader(handle))
            assert records and any("ستون" in v for row in records for v in row)
        else:
            from docx import Document
            report = Document(out)
            assert report.tables
        assert "گزارش مستند ذخیره شد" in dlg.status.text()
        assert svc.open_project("P") == before
        assert svc.store.db.execute(
            "SELECT COUNT(*) FROM revisions WHERE project_id='P'"
        ).fetchone()[0] == revision_count
        assert not svc.open_project("P")["boq"]
    finally:
        dlg.close()
        dlg.deleteLater()
        qt.processEvents()
        svc.store.close()


@pytest.mark.parametrize("fmt,button_name", (
    ("docx", "export_docx_button"),
    ("csv", "export_csv_button"),
))
def test_visible_manual_preview_export_button_does_not_commit_boq(
    tmp_path, monkeypatch, qt, fmt, button_name
):
    svc = StructuralProApp(tmp_path / "db")
    svc.create_project("آزمون پیش‌نمایش ریزمتره", "P")
    out = tmp_path / ("preview." + fmt)
    errors = []
    monkeypatch.setattr(
        QFileDialog, "getSaveFileName",
        lambda *a, **k: (str(out), ""),
    )
    monkeypatch.setattr(
        QMessageBox, "warning",
        lambda *a, **k: errors.append(str(a[-1])),
    )
    dlg = ManualTakeoffDialog(svc, "P")
    try:
        dlg.show()
        dlg.manual_tabs.setCurrentIndex(0)
        dlg.floor_id.setText("L1")
        dlg.entry_text.setPlainText(
            "ستون تعداد=1 عرض=0.4 عمق=0.4 ارتفاع=3"
        )
        qt.processEvents()
        QTest.mouseClick(dlg.add_button, Qt.LeftButton)
        qt.processEvents()
        assert dlg.workbench and len(dlg.workbench.records) == 1, errors
        # The preview itself appends an audit event; reporting must not change
        # the already-previewed state, revisions, takeoffs or BOQ.
        baseline = deepcopy(svc.open_project("P"))
        button = getattr(dlg, button_name)
        assert button.isEnabled() and button.isVisible()
        QTest.mouseClick(button, Qt.LeftButton)
        qt.processEvents()
        assert not errors, errors
        assert out.is_file() and out.stat().st_size > 100
        if fmt == "csv":
            assert out.read_bytes().startswith(b"\xef\xbb\xbf")
            with out.open(encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.reader(handle))
            assert rows and any("پیش‌نمایش بررسی‌نشده" in val for row in rows for val in row)
        else:
            from docx import Document
            assert Document(out).tables
        assert "گزارش ردیف‌های قابل مشاهده ذخیره شد" in dlg.status.text()
        assert svc.open_project("P") == baseline
    finally:
        dlg.close()
        dlg.deleteLater()
        qt.processEvents()
        svc.store.close()
