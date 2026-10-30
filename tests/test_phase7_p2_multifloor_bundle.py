"""Synthetic five-floor/two-document Golden, never a commissioned project."""
from hashlib import sha256
import csv
import json
import os
from pathlib import Path
import shutil

import pytest

from core.drawings.graphical_takeoff import Point
from core.drawings.takeoff_session import DrawingTakeoffSession
from core.platform.application import StructuralProApp
import core.reports.source_backed_2d_office_bundle as bundle


@pytest.fixture
def building(tmp_path):
    import pymupdf
    app = StructuralProApp(tmp_path / "store")
    app.create_project("داده مصنوعی آزمون پنج طبقه", "P5")
    selections = []
    for source_no, count in ((0, 3), (1, 2)):
        path = tmp_path / f"source_{source_no}.pdf"
        with pymupdf.open() as doc:
            for _ in range(count):
                doc.new_page(width=220, height=220).draw_rect(pymupdf.Rect(10, 10, 200, 200))
            doc.save(path)
        session = DrawingTakeoffSession(str(path))
        items = []
        for page in range(1, count + 1):
            n = page + 3 * source_no
            app.add_project_floor("P5", f"طبقه {n}", level=f"L{n}")
            app.add_project_member("P5", f"W{n}", f"L{n}", "سقف")
            session.calibrate(page, 100, 5)
            item = session.add_polygon(
                [Point(x,y) for x,y in ((20,20),(120,20),(120,100),(20,100))],
                holes=[[Point(x,y) for x,y in ((30,30),(40,30),(40,40),(30,40))]] if n == 3 else (),
                opening_refs=["O3"] if n == 3 else (),
                page=page, floor_id=f"L{n}", member_id=f"W{n}",
                quantity_owner="سند مصنوعی آزمون", source=f"test-floor-{n}", takeoff_code=f"AREA_{n}",
            )
            items.append(item)
        receipt = app.save_drawing_takeoff_session("P5", session)
        app.commit_drawing_takeoff_to_boq("P5", receipt["session_id"], [i.id for i in items], expected_session_revision=receipt["revision"])
        current = app.open_project("P5")["drawing_sessions"][receipt["session_id"]]
        for page,item in enumerate(items,1):
            selections.append(dict(session_id=receipt["session_id"], expected_session_revision=current["revision"], page=page, selected_takeoff_ids=(item.id,)))
    snapshot = app.open_project("P5")
    yield app, snapshot, tuple(selections)
    app.store.close()


def export(building, target, selections=None):
    app, snapshot, original = building
    fn = getattr(bundle, "export_verified_2d_multifloor_bundle", None)
    assert callable(fn), "single atomic multi-page export is missing"
    return fn(app, project_id="P5", expected_project_version=snapshot["_meta"]["version"],
              revision="R1", selections=selections if selections is not None else original, destination=target)


def test_one_five_floor_two_source_bundle_readback_and_immutable_store(building, tmp_path):
    import pymupdf
    from docx import Document
    from openpyxl import load_workbook
    app,snapshot,_ = building
    digest = app.store.current_digest("P5")
    target = tmp_path / "one_bundle"
    result = export(building, target)
    manifest = json.loads((target / "manifest.json").read_text(encoding="utf-8"))
    assert result["manifest"] == manifest
    assert len(manifest["pages"]) == 5
    assert len({p["source_sha256"] for p in manifest["pages"]}) == 2
    assert {l["floor_id"] for p in manifest["pages"] for l in p["links"]} == {f"L{i}" for i in range(1,6)}
    assert manifest["pages"][2]["links"][0]["opening_refs"] == ["O3"]
    for item in manifest["files"]:
        path = target / item["name"]
        assert item["sha256"] == sha256(path.read_bytes()).hexdigest()
        assert item["size"] == path.stat().st_size > 0
    assert len(manifest["files"]) == 14
    with pymupdf.open(target / "report.pdf") as pdf:
        assert sum(bool(p.get_images()) for p in pdf) == 5
        text = "\n".join(p.get_text() for p in pdf)
        assert all(f"W{i}" in text for i in range(1,6))
        assert all(p.get_pixmap().width > 0 for p in pdf)
        pdf[-1].get_pixmap().save(target / "qa-last-page.png")
    doc = Document(target / "report.docx")
    assert len(doc.inline_shapes) == 5
    assert "w:bidi" in doc._element.xml
    book = load_workbook(target / "report.xlsx")
    assert sum(len(s._images) for s in book) == 5
    index = book["فهرست نقشه‌ها"]
    assert index.max_row == 6
    assert all(index.cell(n,1).hyperlink for n in range(2,7))
    assert all(index.cell(n,7).hyperlink.target.startswith("#'ریز متره'!") for n in range(2,7))
    assert all(s.sheet_view.rightToLeft and s.print_area for s in book)
    with (target / "report.csv").open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream))
    assert all(any(f"W{i}" in str(row) for row in rows) for i in range(1,6))
    assert app.open_project("P5") == snapshot
    assert app.store.current_digest("P5") == digest
    if os.environ.get("STRUCTURALPRO_P7_GOLDEN_DIR"):
        shutil.copytree(target, Path(os.environ["STRUCTURALPRO_P7_GOLDEN_DIR"]), dirs_exist_ok=True)


@pytest.mark.parametrize("kind", ["empty", "duplicate", "stale", "source"])
def test_invalid_multifloor_selection_never_publishes(building, tmp_path, kind):
    _,snapshot,selections = building
    if kind == "empty": selections = ()
    elif kind == "duplicate": selections += selections[:1]
    elif kind == "stale": selections = (selections[0] | {"expected_session_revision": 999},)
    else:
        path = Path(snapshot["drawing_sessions"][selections[-1]["session_id"]]["drawing_source"])
        path.write_bytes(path.read_bytes() + b"modified")
    with pytest.raises(ValueError): export(building, tmp_path / "invalid", selections)
    assert not (tmp_path / "invalid").exists()


@pytest.mark.parametrize("kind", ["exporter", "digest", "late_source"])
def test_multifloor_publication_is_atomic_on_mid_export_failure(building, tmp_path, monkeypatch, kind):
    from core.reports.project_report import ProjectReport
    app,snapshot,selections = building
    original = ProjectReport.export
    calls = []
    def intervene(self, path, fmt):
        calls.append(fmt)
        result = original(self,path,fmt)
        if fmt == "csv":
            if kind == "exporter": raise RuntimeError("injected exporter failure")
            if kind == "digest": monkeypatch.setattr(app.store,"current_digest",lambda _: "changed")
            if kind == "late_source":
                source = Path(snapshot["drawing_sessions"][selections[-1]["session_id"]]["drawing_source"])
                source.write_bytes(source.read_bytes() + b"changed")
        return result
    monkeypatch.setattr(ProjectReport,"export",intervene)
    with pytest.raises((RuntimeError,ValueError)): export(building,tmp_path / "failed")
    assert calls
    assert not (tmp_path / "failed").exists()
    assert not list(tmp_path.glob("structuralpro_p7_multi_*"))


def test_first_source_changed_during_last_crosswalk_check_is_rejected(building, tmp_path, monkeypatch):
    _, snapshot, selections = building
    real = bundle.link_source_backed_2d_overlay_to_boq
    calls = []
    def changing(*args, **kwargs):
        result = real(*args, **kwargs)
        calls.append(kwargs)
        if len(calls) == len(selections)*2:
            source = Path(snapshot["drawing_sessions"][selections[0]["session_id"]]["drawing_source"])
            source.write_bytes(source.read_bytes() + b"changed during final check")
        return result
    monkeypatch.setattr(bundle, "link_source_backed_2d_overlay_to_boq", changing)
    with pytest.raises(ValueError, match="منبع"): export(building, tmp_path / "stale")
    assert not (tmp_path / "stale").exists()
