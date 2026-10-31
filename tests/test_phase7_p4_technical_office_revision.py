"""P7 P4: report ONLY from authenticated persisted SQLite revisions.

Synthetic vector source and real ProjectStore fixture, not customer data.
"""
from copy import deepcopy
from decimal import Decimal
from hashlib import sha256
import csv
import json

import pytest

from test_phase7_p2_persisted_boq_crosswalk import _persisted
from core.reports.technical_office_revision import (
    preview_revision_quantity_changes,
    export_revision_quantity_package,
)


def _history(tmp_path):
    app, project, opts, item = _persisted(tmp_path)
    records = app.store.revisions("P")
    assert records and records[-1]["version"] == project["_meta"]["version"]
    initial = next(r["version"] for r in records if not r["payload"].get("takeoffs"))
    latest = project["_meta"]["version"]
    args = dict(project_id="P", before_version=initial, after_version=latest)
    return app, project, args, item


def test_real_revision_added_takeoff_has_exact_delta_and_unknown_financial_effect(tmp_path):
    app, project, args, item = _history(tmp_path)
    before = deepcopy(project)
    digest = app.store.current_digest("P")
    try:
        report = preview_revision_quantity_changes(app, **args)
        assert report["readonly"] is True
        assert report["financial_impact"] is None
        assert report["financial_impact_status"] == "NOT_DETERMINABLE_WITHOUT_APPROVED_PRICES"
        assert report["before_version"] < report["after_version"]
        assert len(report["rows"]) == 1
        row = report["rows"][0]
        assert row["مقدار قبلی"] == "0"
        assert Decimal(row["مقدار جدید"]) == Decimal(row["quantity"])
        assert Decimal(row["اختلاف مقدار"]) == Decimal(row["quantity"])
        assert row["وضعیت تغییر"] == "افزوده"
        assert row["source_id"] == project["takeoffs"][0]["source_id"]
        assert row["member"] == "P1"
        assert Decimal(report["delta_by_unit"][row["unit"]]) == Decimal(row["quantity"])
        assert project == before == app.open_project("P")
        assert app.store.current_digest("P") == digest
    finally:
        app.store.close()


def test_real_revision_four_format_bundle_and_sha_integrity(tmp_path):
    import pymupdf
    from docx import Document
    from openpyxl import load_workbook

    app, project, args, item = _history(tmp_path)
    before = deepcopy(project)
    digest = app.store.current_digest("P")
    dest = tmp_path / "revision_bundle"
    try:
        result = export_revision_quantity_package(
            app, **args, expected_current_version=project["_meta"]["version"],
            destination=dest,
        )
        names = {p.name for p in dest.iterdir()}
        assert names == {
            "revision_changes.pdf", "revision_changes.docx",
            "revision_changes.xlsx", "revision_changes.csv", "manifest.json",
        }
        manifest = json.loads((dest / "manifest.json").read_text(encoding="utf-8"))
        assert manifest == result["manifest"]
        assert manifest["readonly"] is True
        assert manifest["financial_impact"] is None
        assert manifest["source_project_digest"] == digest
        assert manifest["before_version"] == args["before_version"]
        assert manifest["after_version"] == args["after_version"]
        for entry in manifest["files"]:
            file = dest / entry["name"]
            assert file.stat().st_size == entry["size"] > 0
            assert sha256(file.read_bytes()).hexdigest() == entry["sha256"]
        with pymupdf.open(dest / "revision_changes.pdf") as pdf:
            assert len(pdf) and all(p.rect.width > 0 for p in pdf)
        doc = Document(dest / "revision_changes.docx")
        assert "w:bidi" in doc._element.xml
        assert "اختلاف مقدار" in doc._element.xml
        book = load_workbook(dest / "revision_changes.xlsx")
        assert all(sheet.sheet_view.rightToLeft for sheet in book)
        assert any("اختلاف مقدار" == cell.value
                   for sheet in book for row in sheet for cell in row)
        with (dest / "revision_changes.csv").open(encoding="utf-8-sig", newline="") as stream:
            csv_rows = list(csv.reader(stream, strict=True))
        assert any("اختلاف مقدار" in row for row in csv_rows)
        assert any(project["takeoffs"][0]["source_id"] in row for row in csv_rows)
        assert project == before == app.open_project("P")
        assert app.store.current_digest("P") == digest
    finally:
        app.store.close()


@pytest.mark.parametrize("version_args", [
    {"before_version": 0, "after_version": 2},
    {"before_version": True, "after_version": 2},
    {"before_version": 2, "after_version": 2},
    {"before_version": 9999, "after_version": 10000},
])
def test_invalid_or_missing_historical_versions_are_rejected(tmp_path, version_args):
    app, project, args, _ = _history(tmp_path)
    try:
        with pytest.raises(ValueError):
            preview_revision_quantity_changes(app, **(args | version_args))
    finally:
        app.store.close()


def test_source_hash_changed_rejected_and_no_package(tmp_path):
    app, project, args, _ = _history(tmp_path)
    target = tmp_path / "rejected"
    original = (tmp_path / "plan.pdf")
    try:
        original.write_bytes(original.read_bytes() + b"corrupted-after-revision")
        with pytest.raises(ValueError):
            export_revision_quantity_package(
                app, **args, expected_current_version=project["_meta"]["version"],
                destination=target,
            )
        assert not target.exists()
    finally:
        app.store.close()


def test_tampered_revision_payload_fails_reused_boq_audit(tmp_path, monkeypatch):
    app, project, args, _ = _history(tmp_path)
    original = app.store.revisions
    def tampered(pid):
        results = original(pid)
        record = next(r for r in results if r["version"] == args["after_version"])
        record["payload"]["boq"][0]["quantity"] = "9999"
        return results
    monkeypatch.setattr(app.store, "revisions", tampered)
    try:
        with pytest.raises(ValueError):
            preview_revision_quantity_changes(app, **args)
        assert app.open_project("P") == project
    finally:
        app.store.close()


def test_unrelated_project_or_current_version_fails_closed(tmp_path):
    app, project, args, _ = _history(tmp_path)
    try:
        with pytest.raises(KeyError):
            preview_revision_quantity_changes(app, project_id="OTHER",
                                             before_version=1, after_version=2)
        with pytest.raises(ValueError, match="نسخه"):
            export_revision_quantity_package(
                app, **args, expected_current_version=project["_meta"]["version"] - 1,
                destination=tmp_path / "stale",
            )
        assert not (tmp_path / "stale").exists()
    finally:
        app.store.close()


def test_existing_output_is_never_overwritten(tmp_path):
    app, project, args, _ = _history(tmp_path)
    out = tmp_path / "existing"
    out.mkdir()
    try:
        with pytest.raises(FileExistsError):
            export_revision_quantity_package(
                app, **args, expected_current_version=project["_meta"]["version"],
                destination=out,
            )
    finally:
        app.store.close()
