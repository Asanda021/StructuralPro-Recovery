"""Synthetic source-backed integration, independent arithmetic and existing ledgers."""
from copy import deepcopy
from decimal import Decimal
from hashlib import sha256
from pathlib import Path
import pytest
from core.takeoff.reviewed_aec import preview_system
from core.takeoff.engine import TakeoffEngine
from test_floor_specialist_work import sample as floor_sample

def sample(tmp_path):
 d=floor_sample();d['system_id']='S1';d['source']['source_id']='SYNTHETIC-SLOPE'
 g=d['architecture_detail']['geometry'];doc=tmp_path/'slope-source.txt';doc.write_text('EXPLICIT SYNTHETIC SOFTWARE TEST ONLY; horizontal datum -1.2; rectangle 5x4; corner thickness .2/.3/.4; opening (1,1) 1x1; no material recipe or price.')
 g.update(bottom_m='0',coordinate_frame_ref='Shared measured local axes: synthetic',material_ref='مصالح صرفاً مصنوعی',source_path=str(doc),source_sha256=sha256(doc.read_bytes()).hexdigest(),layer_key='SLOPE-LAYER',payment_ref='فقط حجم اجرای شیب‌بندی؛ بدون خرید یا پوشش',interface_ref='مرجع تطبیق محور و تراز و بازشوهای میزبان؛ مصنوعی')
 z=dict(id='Z1',physical_key='SLOPE-Z1',source_ref='ZONE-SOURCE',approved_by='مهندس مصنوعی',datum_ref='DATUM',x_m='0',y_m='0',length_m='5',width_m='4',base_elevation_m='-1.2',thickness_00_m='.2',thickness_10_m='.3',thickness_01_m='.4')
 h=dict(id='H1',physical_key='SLOPE-H1',source_ref='HOLE-SOURCE',approved_by='مهندس مصنوعی',zone_id='Z1',x_m='1',y_m='1',length_m='1',width_m='1',penetration='full_depth')
 d['architecture_detail']=dict(contract='planar_slope_aec_v1',geometry=g,zones=[z],openings=[h]);return d

def result(d):return preview_system(d,TakeoffEngine())

def commit(a,d):
 a.save_aec_system_draft('P',d,expected_revision=0)
 a.review_roof_system('P','floor_finish','L1','S1',expected_revision=1,reviewer='بازبین مستقل',note='کنترل مصنوعی')
 a.commit_aec_system('P','floor_finish','L1','S1',expected_revision=2)

def test_reference_exact_volume_and_verified_document(tmp_path):
 r=result(sample(tmp_path));assert r['ready'],r['blockers'];p=r['rows'][0]['params']
 # Integral: 20*(.2+.3+.4+.5)/4 - 1*(.2+.1*1.5/5+.2*1.5/4)
 assert Decimal(str(r['rows'][0]['quantity']))==Decimal('6.695')
 assert p['gross_volume_m3']=='7.0' and Decimal(p['deducted_volume_m3'])==Decimal('.305')
 assert p['source_sha256_verified'] is True and r['rows'][0]['role']=='work'

@pytest.mark.parametrize('fault',['unit','float','missing','digest','document','unknown','disconnected','overlap','outside','base','duplicate','hole_owner'])
def test_fail_closed(tmp_path,fault):
 d=sample(tmp_path);a=d['architecture_detail'];g=a['geometry'];z=a['zones'][0]
 if fault=='unit':g['unit']='cm'
 if fault=='float':z['thickness_00_m']=.2
 if fault=='missing':z['datum_ref']=''
 if fault=='digest':g['source_sha256']='0'*64
 if fault=='document':g['source_path']=str(tmp_path/'absent')
 if fault=='unknown':g['waste']='.05'
 if fault in ('disconnected','overlap'):a['zones'].append({**z,'id':'Z2','physical_key':'Z2-P','source_ref':'Z2-S','x_m':'6' if fault=='disconnected' else '1'})
 if fault=='outside':z['x_m']='1'
 if fault=='base':z['base_elevation_m']='0'
 if fault=='duplicate':a['openings'].append(deepcopy(a['openings'][0]))
 if fault=='hole_owner':a['openings'][0]['zone_id']='ABSENT'
 assert not result(d)['ready']

@pytest.mark.parametrize('case',['valid','copy','cas','two_connections','self_review','direct','financial','duplicate','identity','repeat','legacy','region','floor_volume','insulation_volume','host_hole','physical'])
def test_shared_workflow_and_ownership(tmp_path,case):
 from core.platform.application import StructuralProApp
 a=StructuralProApp(tmp_path/'db');a.create_project('پروژه مصنوعی','P');d=sample(tmp_path)
 if case in ('legacy','floor_volume','insulation_volume','host_hole'):
  old=floor_sample()
  if case=='legacy':old['system_id']='S1'
  if case=='insulation_volume':
   from test_insulation_specialist_work import sample as ins_sample
   old=ins_sample('thermal_insulation')
  a.save_aec_system_draft('P',old,expected_revision=0)
  with pytest.raises(ValueError):a.save_aec_system_draft('P',d if case!='host_hole' else {**d,'architecture_detail':{**d['architecture_detail'],'openings':[]}},expected_revision=1 if case=='legacy' else 0)
  a.store.close();return
 if case=='cas':
  a.save_aec_system_draft('P',d,expected_revision=0)
  with pytest.raises(RuntimeError):a.save_aec_system_draft('P',d,expected_revision=0)
  a.store.close();return
 if case=='self_review':
  a.save_aec_system_draft('P',d,expected_revision=0)
  with pytest.raises(ValueError):a.review_roof_system('P','floor_finish','L1','S1',expected_revision=1,reviewer='مهندس مصنوعی',note='خودبازبینی')
  a.store.close();return
 commit(a,d);before=deepcopy(a.open_project('P'));assert len(before['boq'])==1
 if case=='copy':
  r=a.aec_workflow().copy_roof('P','floor_finish','L1','S1',new_system_id='S2',new_source_id='COPY',reason='کپی مصنوعی');v=r['data']['architecture_detail'];assert v['geometry']['physical_region_id']!='FLOOR-1';assert v['openings'][0]['zone_id']==v['zones'][0]['id'];assert not result(r['data'])['ready'];a.aec_workflow().delete_roof('P','floor_finish','L1','S2',expected_revision=1,reason='حذف مصنوعی');a.store.close();return
 if case in ('region','physical'):
  other=deepcopy(d);other['system_id']='S2';other['source']['source_id']='OTHER'
  if case=='physical':other['architecture_detail']['geometry']['physical_region_id']='OTHER'
  with pytest.raises(ValueError):a.save_aec_system_draft('P',other,expected_revision=0)
  assert a.open_project('P')==before;a.store.close();return
 if case=='direct':
  illegal=deepcopy(before);next(iter(illegal['reviewed_aec_systems'].values()))['data']['axis']='changed'
  with pytest.raises(ValueError):a.store.save('P',illegal)
  a.store.close();return
 if case=='two_connections':
  b=StructuralProApp(tmp_path/'db');p,token=b.aec_workflow()._get_project('P');a.create_project('یک پروژه مستقل','OTHER');current=a.open_project('P');current['name']='نام جدید';a.store.save('P',current)
  with pytest.raises(RuntimeError):b.store.save('P',p,expected_digest=token)
  b.store.close();a.store.close();return
 d=deepcopy(a.get_aec_system('P','floor_finish','L1','S1')['data']);d.pop('roof_review');z=d['architecture_detail']['zones'][0]
 for k in ('thickness_00_m','thickness_10_m','thickness_01_m'):z[k]=str(Decimal(z[k])+Decimal('.1'))
 if case=='identity':d['architecture_detail']['geometry']['layer_key']='CHANGED'
 e=tmp_path/'amend.txt';e.write_text('EXPLICIT SYNTHETIC +.1m all three corners; delta volume 19*.1=1.9 m3')
 kw=dict(expected_revision=2,created_by='سازنده اصلاحیه',reason='اصلاح مصنوعی',evidence=dict(evidence_path=str(e),evidence_sha256=sha256(e.read_bytes()).hexdigest(),evidence_reference='مدرک جدید مصنوعی'))
 if case=='identity':
  with pytest.raises(ValueError):a.create_roof_amendment('P','floor_finish','L1','S1',d,**kw)
  a.store.close();return
 am=a.create_roof_amendment('P','floor_finish','L1','S1',d,**kw);assert sum(Decimal(str(v)) for v in am['delta'].values())==Decimal('1.9')
 if case=='duplicate':
  with pytest.raises(ValueError):a.review_beam_amendment('P',am['id'],expected_revision=1,reviewer='سازنده اصلاحیه',note='خودبازبینی')
 a.review_beam_amendment('P',am['id'],expected_revision=1,reviewer='بازبین اصلاحیه',note='کنترل مستقل')
 if case=='financial':
  p=a.open_project('P');p['statement_periods']=[{'id':'LOCK'}];a.store.save('P',p)
  with pytest.raises(ValueError):a.apply_beam_amendment('P',am['id'],expected_revision=1)
 else:
  a.apply_beam_amendment('P',am['id'],expected_revision=1)
  if case=='repeat':
   with pytest.raises((ValueError,RuntimeError)):a.apply_beam_amendment('P',am['id'],expected_revision=1)
  p=a.open_project('P');assert p['takeoffs']==before['takeoffs'];assert Decimal(str(p['boq'][0]['quantity']))==Decimal('8.595')
 a.store.close();a=StructuralProApp(tmp_path/'db');assert a.store.db.execute('pragma integrity_check').fetchone()[0]=='ok';a.store.close()

@pytest.mark.parametrize('missing_hole',[False,True])
def test_slope_above_registered_floor_and_host_hole(tmp_path,missing_hole):
 from core.platform.application import StructuralProApp
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');old=floor_sample();a.save_aec_system_draft('P',old,expected_revision=0)
 d=sample(tmp_path);ar=d['architecture_detail'];ar['geometry']['bottom_m']='.15';ar['zones'][0]['base_elevation_m']='-1.05'
 if missing_hole:
  ar['openings']=[]
  with pytest.raises(ValueError,match='بازشو'):a.save_aec_system_draft('P',d,expected_revision=0)
 else:commit(a,d);assert len(a.open_project('P')['boq'])==1
 a.store.close()

def test_four_reports_reopen(tmp_path):
 from core.platform.application import StructuralProApp
 from core.reports.architecture_detail_report import export
 from openpyxl import load_workbook
 from docx import Document
 import pymupdf,csv
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');commit(a,sample(tmp_path))
 for fmt in ('pdf','xlsx','docx','csv'):export(a,'P','floor_finish','L1','S1',tmp_path/('slope.'+fmt),format=fmt)
 with pymupdf.open(tmp_path/'slope.pdf') as p:assert len(p)>0
 w=load_workbook(tmp_path/'slope.xlsx');assert all(s.sheet_view.rightToLeft for s in w);assert any(str(c.value)=='6.695' for s in w for row in s for c in row);w.close()
 assert Document(tmp_path/'slope.docx').tables
 text=(tmp_path/'slope.csv').read_text(encoding='utf-8-sig');assert '6.695' in text and 'ناخالص' in text and 'هش مدرک' in text and 'ضخامت گوشه' in text;a.store.close()

def test_two_adjacent_regions_exactly_retain_reference_volume(tmp_path):
 d=sample(tmp_path);a=d['architecture_detail'];z=a['zones'][0];z.update(length_m='2.5',thickness_10_m='.25')
 a['zones'].append({**z,'id':'Z2','physical_key':'SL-Z2','source_ref':'SECOND-ZONE','x_m':'2.5','thickness_00_m':'.25','thickness_10_m':'.3','thickness_01_m':'.45'})
 r=result(d);assert r['ready'],r['blockers'];assert sum(Decimal(v['quantity']) for v in r['rows'])==Decimal('6.695')

def test_source_canonical_alias_and_tampering_are_blocked(tmp_path):
 d=sample(tmp_path);d['architecture_detail']['geometry']['source_ref']='MASTER-source';d['architecture_detail']['zones'][0]['source_ref']='MASTER-SOURCE';assert not result(d)['ready']
 d=sample(tmp_path);Path(d['architecture_detail']['geometry']['source_path']).write_text('TAMPERED');assert not result(d)['ready']

def test_incompatible_coordinate_frames_on_one_host_are_blocked(tmp_path):
 from core.platform.application import StructuralProApp
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');d=sample(tmp_path);commit(a,d);other=sample(tmp_path);other['system_id']='S2';other['source']['source_id']='OTHER';g=other['architecture_detail']['geometry'];g.update(bottom_m='1',layer_key='SECOND-LAYER',coordinate_frame_ref='DIFFERENT ORIGIN');other['architecture_detail']['zones'][0].update(id='Z2',physical_key='Z2-P',base_elevation_m='-.2');other['architecture_detail']['openings'][0].update(id='H2',physical_key='H2-P',zone_id='Z2')
 with pytest.raises(ValueError,match='مختصات'):a.save_aec_system_draft('P',other,expected_revision=0)
 a.store.close()

def test_long_decimal_survives_sqlite_boq_and_excel(tmp_path):
 from core.platform.application import StructuralProApp
 from core.reports.architecture_detail_report import export
 from openpyxl import load_workbook
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');d=sample(tmp_path);g=d['architecture_detail']['geometry'];d['architecture_detail']['zones'][0]['thickness_10_m']='.300000000000000000001';f=Path(g['source_path']);f.write_text(f.read_text()+' EXPLICIT SYNTHETIC precision case t10=.300000000000000000001 m.');g['source_sha256']=sha256(f.read_bytes()).hexdigest();commit(a,d);a.store.close();a=StructuralProApp(tmp_path/'db')
 expected=Decimal('6.6950000000000000000097') # 6.695 + (20/2 - 1*1.5/5)*1e-21
 assert Decimal(str(a.open_project('P')['boq'][0]['quantity']))==expected
 export(a,'P','floor_finish','L1','S1',tmp_path/'precise.xlsx',format='xlsx');w=load_workbook(tmp_path/'precise.xlsx');assert any(c.value==str(expected) for s in w for row in s for c in row);w.close();a.store.close()
