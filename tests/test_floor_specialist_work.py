"""Explicit synthetic rectangular floor strata; no guessed material consumption."""
from copy import deepcopy
from decimal import Decimal
import pytest
from core.takeoff.reviewed_aec import empty_system,preview_system
from core.takeoff.engine import TakeoffEngine
from test_architecture_specialist_work import obj

def sample():
 d=empty_system('floor_finish','L1','F1');d.update(measurement_policy='architecture_specialist_v1',building_id='B1',axis='اتاق مصنوعی',source=dict(mode='manual',source_id='SYNTHETIC-FLOOR',plan_ref='پلان مصنوعی',detail_ref='دیتایل مصنوعی'))
 g=dict(physical_region_id='FLOOR-1',space_id='ROOM-1',unit='m',shape='rectangle',length_m='5',width_m='4',base_level_m='-1.2',level_ref='تراز مبنای مصنوعی',count='1',repeat_ref='یک ناحیه مستقل مصنوعی',source_ref='پلان مصنوعی',approved_by='کنترل مصنوعی')
 def layer(i,code,bottom,t,unit):return obj(i,layer_key=i,operation=code,method='constant_rectangular',x='0',y='0',width_m='5',height_m='4',bottom_m=bottom,thickness_m=t,material_ref='مشخصات کاملاً مصنوعی '+code,payment_unit=unit,payment_ref='مبنای پرداخت مستقل همین لایه',opening_ids='H1',deduction_ref='کسر فقط بازشوی تمام لایه‌ها',coverage_scope='complete',scope_ref='کل سطح خالص',density_kg_m3='',density_ref='',purchase_quantity='',purchase_unit='',purchase_ref='')
 d['architecture_detail']=dict(geometry=g,openings=[obj('H1',x='1',y='1',width_m='1',height_m='1',penetration='all_layers')],layers=[layer('FILL','floor_fill','0','.1','m3'),layer('SCREED','floor_screed','.1','.04','m3'),layer('PROTECT','floor_protection','.14','.01','m2')])
 return d

def totals(d):
 r=preview_system(d,TakeoffEngine());assert r['ready'],r['blockers'];out={}
 for v in r['rows']:out[v['item_code']]=out.get(v['item_code'],Decimal(0))+Decimal(str(v['quantity']))
 return out,r

def test_independent_multilayer_reference():
 t,r=totals(sample());assert t=={'floor_fill':Decimal('1.9'),'floor_screed':Decimal('.76'),'floor_protection':Decimal('19')};assert all(v['params']['gross_area_m2']==20 and v['params']['net_area_m2']==19 for v in r['rows'])
 assert Decimal(str(r['rows'][1]['params']['top_level_m']))==Decimal('-1.06')

@pytest.mark.parametrize('fault',['zero','negative','variable','unit','overlap','duplicate_opening','outside','incomplete','missing_source','opening_link','duplicate_layer','unknown_waste','density','purchase_partial','wrong_policy','partial_masquerade','layer_definition'])
def test_floor_fail_closed(fault):
 d=sample();a=d['architecture_detail'];p=a['layers'][1]
 if fault=='zero':p['thickness_m']='0'
 if fault=='negative':p['thickness_m']='-.1'
 if fault=='variable':p['method']='variable_thickness'
 if fault=='unit':p['payment_unit']='litre'
 if fault=='overlap':p['bottom_m']='.09'
 if fault=='duplicate_opening':a['openings'].append({**a['openings'][0],'id':'H2','physical_key':'other'})
 if fault=='outside':p['width_m']='6'
 if fault=='incomplete':p['thickness_m']=''
 if fault=='missing_source':p['material_ref']=''
 if fault=='opening_link':p['opening_ids']=''
 if fault=='duplicate_layer':a['layers'].append({**p,'id':'SECOND','physical_key':'second'})
 if fault=='unknown_waste':p['waste']='.1'
 if fault=='density':p['payment_unit']='kg'
 if fault=='purchase_partial':p['purchase_quantity']='1'
 if fault=='wrong_policy':d.pop('measurement_policy')
 if fault=='partial_masquerade':p['width_m']='3'
 if fault=='layer_definition':a['layers'].append({**p,'id':'SECOND','physical_key':'second','x':'0','width_m':'1','bottom_m':'.3'})
 assert not preview_system(d,TakeoffEngine())['ready']

@pytest.mark.parametrize('code',['floor_fill','floor_screed','floor_concrete','floor_subbase','floor_separator','floor_protection'])
def test_independent_layer_types(code):
 d=sample();d['architecture_detail']['layers']=d['architecture_detail']['layers'][1:2];d['architecture_detail']['layers'][0]['operation']=code;assert totals(d)[0][code]==Decimal('.76')


def test_explicit_mass_purchase_and_repeat_are_distinct():
 d=sample();d['architecture_detail']['geometry'].update(count='2',repeat_ref='دو ناحیه هم‌تیپ مشخص مصنوعی');p=d['architecture_detail']['layers'][1];p.update(payment_unit='kg',density_kg_m3='1500',density_ref='گواهی مصنوعی چگالی',purchase_quantity='1.1',purchase_unit='m3',purchase_ref='حجم خرید مستند برای هر ناحیه مستقل')
 t,r=totals(d);assert t['floor_screed']==2280 and t['floor_purchase']==Decimal('2.2');assert next(v for v in r['rows'] if v['item_code']=='floor_purchase')['role']=='material'

@pytest.mark.parametrize('case',['valid','same_reviewer','duplicate','financial','evidence','unit','identity','cas','legacy','direct_legacy','physical','duplicate_source','copy'])
def test_floor_existing_sqlite_and_amendment(tmp_path,case):
 from core.platform.application import StructuralProApp
 from hashlib import sha256
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی کف','P');d=sample()
 if case in ('legacy','direct_legacy'):
  old=empty_system('floor_finish','L1','F1');old['source']={'mode':'manual','source_id':'LEGACY'};a.save_aec_system_draft('P',old,expected_revision=0)
  with pytest.raises(ValueError):
   if case=='legacy':a.save_aec_system_draft('P',d,expected_revision=1)
   else:
    p=a.open_project('P');next(iter(p['reviewed_aec_systems'].values()))['data']=d;a.store.save('P',p)
  a.store.close();return
 a.save_aec_system_draft('P',d,expected_revision=0)
 if case=='cas':
  with pytest.raises(RuntimeError):a.save_aec_system_draft('P',d,expected_revision=0)
  a.store.close();return
 if case in ('physical','duplicate_source'):
  other=deepcopy(d);other['system_id']='F2'
  if case=='physical':other['source']['source_id']='OTHER'
  with pytest.raises(ValueError):a.save_aec_system_draft('P',other,expected_revision=0)
  a.store.close();return
 if case=='copy':
  r=a.aec_workflow().copy_roof('P','floor_finish','L1','F1',new_system_id='F2',new_source_id='COPY',reason='کپی مصنوعی');v=r['data']['architecture_detail'];assert v['geometry']['physical_region_id']!='FLOOR-1';assert v['layers'][0]['opening_ids']==v['openings'][0]['id'];assert not preview_system(r['data'],a.takeoff)['ready'];a.aec_workflow().delete_roof('P','floor_finish','L1','F2',expected_revision=1,reason='حذف کنترل‌شده مصنوعی');a.store.close();return
 with pytest.raises(ValueError):a.commit_aec_system('P','floor_finish','L1','F1',expected_revision=1)
 a.review_roof_system('P','floor_finish','L1','F1',expected_revision=1,reviewer='مستقل',note='کنترل لایه‌ها');a.commit_aec_system('P','floor_finish','L1','F1',expected_revision=2);before=deepcopy(a.open_project('P'))
 illegal=deepcopy(before);next(iter(illegal['reviewed_aec_systems'].values()))['data']['axis']='changed'
 with pytest.raises(ValueError):a.store.save('P',illegal)
 d=deepcopy(a.get_aec_system('P','floor_finish','L1','F1')['data']);d.pop('roof_review');d['architecture_detail']['layers'][1]['thickness_m']='.05';d['architecture_detail']['layers'][2]['bottom_m']='.15'
 e=tmp_path/'new.txt';e.write_text('EXPLICIT SYNTHETIC revised floor thickness .04 -> .05 m; protective layer moves to .15 m.');kw=dict(expected_revision=2,created_by='سازنده',reason='اصلاح مستند مصنوعی',evidence=dict(evidence_path=str(e),evidence_sha256=sha256(e.read_bytes()).hexdigest(),evidence_reference='مدرک مصنوعی'))
 if case in ('unit','identity'):
  if case=='unit':d['architecture_detail']['layers'][1]['payment_unit']='m2'
  else:d['architecture_detail']['geometry']['physical_region_id']='OTHER'
  with pytest.raises(ValueError):a.create_roof_amendment('P','floor_finish','L1','F1',d,**kw)
 else:
  am=a.create_roof_amendment('P','floor_finish','L1','F1',d,**kw);assert Decimal(str(am['delta']['aec:floor_finish:L1:F1:SCREED']))==Decimal('.19')
  if case=='duplicate':
   with pytest.raises(ValueError):a.create_roof_amendment('P','floor_finish','L1','F1',d,**kw)
  elif case=='same_reviewer':
   with pytest.raises(ValueError):a.review_beam_amendment('P',am['id'],expected_revision=1,reviewer='سازنده',note='نامجاز')
  elif case=='evidence':
   e.write_text('tampered')
   with pytest.raises(ValueError):a.review_beam_amendment('P',am['id'],expected_revision=1,reviewer='دیگری',note='کنترل')
  else:
   a.review_beam_amendment('P',am['id'],expected_revision=1,reviewer='دیگری',note='کنترل مستقل')
   if case=='financial':
    p=a.open_project('P');p['contract']={'id':'locked'};a.store.save('P',p);locked=deepcopy(a.open_project('P'))
    with pytest.raises(ValueError):a.apply_beam_amendment('P',am['id'],expected_revision=1)
    assert a.open_project('P')==locked
   else:
    a.apply_beam_amendment('P',am['id'],expected_revision=1);assert a.open_project('P')['takeoffs']==before['takeoffs'];assert a.open_project('P')['boq']!=before['boq'];a.store.close();a=StructuralProApp(tmp_path/'db');assert a.get_aec_system('P','floor_finish','L1','F1')['revision']==3;assert a.store.db.execute('pragma integrity_check').fetchone()[0]=='ok'
 a.store.close()


def test_purchase_not_paid_twice_and_exact_report_reopen(tmp_path):
 from core.platform.application import StructuralProApp
 from core.reports.architecture_detail_report import export
 from openpyxl import load_workbook
 from docx import Document
 import pymupdf,csv
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');d=sample();d['architecture_detail']['layers']=d['architecture_detail']['layers'][1:2];p=d['architecture_detail']['layers'][0];p.update(thickness_m='.040000000000000000001',purchase_quantity='1.1',purchase_unit='m3',purchase_ref='خرید مستند مصنوعی برای هر ناحیه')
 a.save_aec_system_draft('P',d,expected_revision=0);a.review_roof_system('P','floor_finish','L1','F1',expected_revision=1,reviewer='مستقل',note='کنترل');a.commit_aec_system('P','floor_finish','L1','F1',expected_revision=2)
 assert len(a.open_project('P')['boq'])==1;assert Decimal(str(a.open_project('P')['boq'][0]['quantity']))==Decimal('.760000000000000000019')
 for fmt in ('pdf','xlsx','docx','csv'):export(a,'P','floor_finish','L1','F1',tmp_path/('floor.'+fmt),format=fmt)
 b=load_workbook(tmp_path/'floor.xlsx');assert all(s.sheet_view.rightToLeft for s in b);assert any(c.value=='0.760000000000000000019' for s in b for r in s for c in r);b.close()
 doc=Document(tmp_path/'floor.docx');assert doc.tables;assert all(t._tbl.tblPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bidiVisual') is not None for t in doc.tables)
 pdf=pymupdf.open(tmp_path/'floor.pdf');assert len(pdf)>0;pdf.close()
 text=(tmp_path/'floor.csv').read_text(encoding='utf-8-sig');assert '0.760000000000000000019' in text and 'تراز' in text and 'خرید' in text;a.store.close()


def test_extends_saved_five_floor_scenario_without_recreating_walls(tmp_path):
 from scripts.verify_architecture_step6 import create_sample
 from scripts.verify_floor_step6_package2 import extend_sample
 reference=tmp_path/'package1';create_sample(reference);r=extend_sample(reference,tmp_path/'extended');assert r['previous_wall_records_preserved'] and r['other_four_floor_quantities_unchanged'] and r['sqlite_integrity']=='ok';assert r['floor_effective_totals']['floor_screed']=='3.99'


def test_same_physical_layer_cannot_move_to_another_elevation_in_another_record(tmp_path):
 from core.platform.application import StructuralProApp
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');d=sample();a.save_aec_system_draft('P',d,expected_revision=0);other=deepcopy(d);other['system_id']='F2';other['source']['source_id']='OTHER';other['architecture_detail']['openings'][0].update(id='HO',physical_key='HO');other['architecture_detail']['layers']=other['architecture_detail']['layers'][1:2];other['architecture_detail']['layers'][0].update(id='OTHER',physical_key='OTHER',bottom_m='.4',opening_ids='HO')
 with pytest.raises(ValueError):a.save_aec_system_draft('P',other,expected_revision=0)
 a.store.close()


def test_floor_two_sqlite_connections_reject_stale_version(tmp_path):
 from core.platform.application import StructuralProApp
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');d=sample();a.save_aec_system_draft('P',d,expected_revision=0);b=StructuralProApp(tmp_path/'db');stale=b.get_aec_system('P','floor_finish','L1','F1');updated=deepcopy(d);updated['axis']='فضای مستقل جدید';a.save_aec_system_draft('P',updated,expected_revision=1)
 with pytest.raises(RuntimeError):b.save_aec_system_draft('P',stale['data'],expected_revision=1)
 assert b.get_aec_system('P','floor_finish','L1','F1')['data']['axis']=='فضای مستقل جدید';a.store.close();b.store.close()


def test_applied_amendment_cannot_reuse_another_regions_physical_part(tmp_path):
 from core.platform.application import StructuralProApp
 from hashlib import sha256
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');d=sample();a.save_aec_system_draft('P',d,expected_revision=0);a.review_roof_system('P','floor_finish','L1','F1',expected_revision=1,reviewer='مستقل',note='کنترل');a.commit_aec_system('P','floor_finish','L1','F1',expected_revision=2)
 other=sample();other['system_id']='F2';other['source']['source_id']='OTHER';g=other['architecture_detail'];g['geometry']['physical_region_id']='FLOOR-2'
 for p in g['openings']+g['layers']:
  p['id']='SECOND-'+p['id'];p['physical_key']='SECOND-'+p['physical_key']
  if p.get('opening_ids'):p['opening_ids']='SECOND-'+p['opening_ids']
 a.save_aec_system_draft('P',other,expected_revision=0);proposal=deepcopy(d);proposal['architecture_detail']['layers'][0]['physical_key']=g['layers'][0]['physical_key'];e=tmp_path/'e.txt';e.write_text('EXPLICIT SYNTHETIC')
 am=a.create_roof_amendment('P','floor_finish','L1','F1',proposal,expected_revision=2,created_by='سازنده',reason='اصلاح مصنوعی',evidence=dict(evidence_path=str(e),evidence_sha256=sha256(e.read_bytes()).hexdigest(),evidence_reference='مدرک مصنوعی'));a.review_beam_amendment('P',am['id'],expected_revision=1,reviewer='دیگری',note='کنترل');before=deepcopy(a.open_project('P'))
 with pytest.raises(ValueError):a.apply_beam_amendment('P',am['id'],expected_revision=1)
 assert a.open_project('P')==before;a.store.close()
