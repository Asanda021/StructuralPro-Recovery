"""Synthetic acceptance; native Linux Qt is not Windows/engineering approval."""
from copy import deepcopy
from decimal import Decimal
from pathlib import Path
import pytest
from PySide6.QtCore import Qt,QTimer
from PySide6.QtWidgets import QApplication,QDialog,QTableWidgetItem
from PySide6.QtTest import QTest
from app.beam_component_editor import BeamComponentEditor
from test_chat_circular_column import with_spiral

_qt=None
def qt():
 global _qt
 _qt=QApplication.instance() or QApplication([])
 return _qt

def test_native_copy_regenerates_nested_identity_and_clears_review():
 qt();row=with_spiral()['rebars']['ties']['segments'][0]
 row.update(reviewed_by='old-reviewer',approved_by='old-reviewer',review_status='confirmed',review_digest='old')
 row['end_parts_m']=[{'id':'EP1','kind':'hook','length_m':'.1','source_ref':'D','approved_by':'old','review_digest':'old'}]
 child=BeamComponentEditor('series',[row]);child.show()
 try:
  QTest.mouseClick(child.copy_button,Qt.LeftButton);assert child.table.rowCount()==2
  copy=child.collect(1);assert copy['id']!=row['id'];assert copy['end_parts_m'][0]['id']!='EP1'
  assert copy['review_status']=='needs_review'
  for target in (copy,copy['end_parts_m'][0]):
   assert not any(k in target for k in ('approved_by','review_digest','reviewed_by'))
 finally:child.close()

def test_mixed_pdf_column_allocates_full_numeric_width(tmp_path):
 from core.reports.exporters import _pdf_column_plan,_pdf_font,_pdf_text
 from reportlab.pdfbase import pdfmetrics
 value='0.7539822368615503772310344119870806922073206558500240'
 rows=[{'field':'بتن','value':value},{'field':'منبع','value':'D-synthetic'}]
 font=_pdf_font()
 plan=_pdf_column_plan([('field','عنوان'),('value','مقدار و مرجع')],rows,font=font,usable_width=535)
 width=next(w for panel in plan for k,l,w in panel if k=='value')
 assert width>=pdfmetrics.stringWidth(_pdf_text(value),font,8)+18

def test_mixed_pdf_overlong_number_fails_closed(tmp_path):
 from core.reports.exporters import export_pdf
 with pytest.raises(ValueError,match='بیش از حد طولانی'):
  export_pdf([{'value':'1'*300},{'value':'مدرک'}],tmp_path/'too-long.pdf',columns=[('value','مقدار')])

def test_native_spiral_add_edit_delete_copy_review_and_end_parts(nonblocking_error_dialogs):
 from app.beam_component_editor import shown
 q=qt();row=with_spiral()['rebars']['ties']['segments'][0]
 extra=[(k,k) for k in ('shape_kind','start_m','end_m','pitch_m','turns','axis_diameter_m','cover_m','end_condition','end_condition_ref')]
 child=BeamComponentEditor('series',[row],extra_fields=extra,spiral_mode=True,spiral_column_diameter='.4');child.show();errors=[]
 try:
  assert 'cut_length_m' not in dict(child.fields)
  QTest.mouseClick(child.add_button,Qt.LeftButton);assert child.table.rowCount()==2
  added=deepcopy(row);added['id']='SP-added'
  for j,(key,label) in enumerate(child.fields):child.table.item(1,j).setText(str(shown(key,added.get(key,''))))
  # Explicitly confirm the empty end-detail table; do not infer missing parts.
  def confirm_no_extensions():
   part=q.activeModalWidget();QTest.mouseClick(part.save_button,Qt.LeftButton)
  QTimer.singleShot(50,confirm_no_extensions);QTest.mouseClick(child.nested_button,Qt.LeftButton)
  child.reviewer.setText('بازبین مصنوعی');QTest.mouseClick(child.review_button,Qt.LeftButton)
  assert child.collect(1)['approved_by']=='بازبین مصنوعی',nonblocking_error_dialogs
  QTest.mouseClick(child.copy_button,Qt.LeftButton);assert child.table.rowCount()==3
  assert child.collect(2)['review_status']=='needs_review'
  QTest.mouseClick(child.delete_button,Qt.LeftButton);assert child.table.rowCount()==2
  QTest.mouseClick(child.delete_button,Qt.LeftButton);assert child.table.rowCount()==1
  child.table.setCurrentCell(0,0)
  child.table.item(0,next(j for j,f in enumerate(child.fields) if f[0]=='pitch_m')).setText('.125')
  child.table.item(0,next(j for j,f in enumerate(child.fields) if f[0]=='turns')).setText('8')
  assert child.collect(0)['review_status']=='needs_review'
  def nested():
   part=q.activeModalWidget()
   try:
    assert isinstance(part,BeamComponentEditor) and part.kind=='parts'
    QTest.mouseClick(part.add_button,Qt.LeftButton)
    values={'kind':'hook','length_m':'.1','source_ref':'D-synthetic-hook'}
    for j,(key,label) in enumerate(part.fields):
     if key in values:part.table.item(0,j).setText(str(shown(key,values[key])))
    part.reviewer.setText('بازبین جزء مصنوعی');QTest.mouseClick(part.review_button,Qt.LeftButton)
    assert part.collect(0)['review_status']=='confirmed'
    QTest.mouseClick(part.save_button,Qt.LeftButton)
   except Exception as exc:errors.append(str(exc));part.reject()
  QTimer.singleShot(50,nested);QTest.mouseClick(child.nested_button,Qt.LeftButton)
  assert not errors,errors
  assert child.collect(0)['end_condition']=='detailed_extensions'
  assert len(child.collect(0)['end_parts_m'])==1
  QTest.mouseClick(child.review_button,Qt.LeftButton);assert child.collect(0)['review_status']=='confirmed',nonblocking_error_dialogs
  QTest.mouseClick(child.calculate_button,Qt.LeftButton);assert 'ناقص' not in child.total.text(),nonblocking_error_dialogs
  QTest.mouseClick(child.save_button,Qt.LeftButton);assert len(child.rows)==1
 finally:child.close()

def test_native_circular_formwork_review_and_missing_detail_block(nonblocking_error_dialogs):
 qt();d=with_spiral();surface=d['formwork']['surfaces'][0]
 extra=[(k,k) for k in ('coverage','area_method')]
 child=BeamComponentEditor('surfaces',[surface],extra_fields=extra,circular_formwork_mode=True);child.show()
 try:
  assert 'gross_area_m2' not in dict(child.fields)
  child.reviewer.setText('بازبین مصنوعی');QTest.mouseClick(child.review_button,Qt.LeftButton)
  assert child.collect(0)['review_status']=='confirmed',nonblocking_error_dialogs
  child.table.item(0,next(j for j,f in enumerate(child.fields) if f[0]=='coverage')).setText('')
  QTest.mouseClick(child.review_button,Qt.LeftButton)
  assert child.collect(0)['review_status']=='needs_review'
  assert not child.collect(0).get('approved_by')
 finally:child.close()

def test_native_circular_main_save_reopen_boq_amendment_and_reports(tmp_path,nonblocking_error_dialogs,monkeypatch):
 from core.platform.application import StructuralProApp
 from app.reviewed_structural_manual_editor import open_manual_structural_editor
 from PySide6.QtWidgets import QFileDialog
 from openpyxl import load_workbook
 import pymupdf
 q=qt();app=StructuralProApp(tmp_path/'project');app.create_project('مصنوعی','P');app.save_structural_member_draft('P',with_spiral(),expected_revision=0);errors=[]
 def run():
  dialog=q.activeModalWidget()
  try:
   c=dialog.column_controls;c['id'].setText('C-round');QTest.mouseClick(c['load'],Qt.LeftButton)
   assert c['shape'].currentData()=='circular_column'
   assert c['numeric']['diameter'].text()=='.4' or Decimal(c['numeric']['diameter'].text())==Decimal('.4')
   c['group'].setCurrentIndex(c['group'].findData('ties'))
   def edit_zone():
    child=q.activeModalWidget()
    try:
     assert isinstance(child,BeamComponentEditor) and child.spiral_mode
     child.table.setCurrentCell(0,0)
     for key,value in [('pitch_m','.125'),('turns','8')]:child.table.item(0,next(j for j,f in enumerate(child.fields) if f[0]==key)).setText(value)
     child.reviewer.setText('بازبین سری مصنوعی');QTest.mouseClick(child.review_button,Qt.LeftButton)
     edited=next(child.collect(i) for i in range(child.table.rowCount()) if child.collect(i)['id']=='SP1')
     assert edited['review_status']=='confirmed',nonblocking_error_dialogs
     QTest.mouseClick(child.save_button,Qt.LeftButton)
    except Exception as exc:errors.append(str(exc));child.reject()
   QTimer.singleShot(50,edit_zone);QTest.mouseClick(c['advanced'],Qt.LeftButton);assert not errors,errors
   assert 'component_review' not in c['state']['data']
   QTest.mouseClick(c['review'],Qt.LeftButton);QTest.mouseClick(c['save'],Qt.LeftButton)
   assert app.get_structural_member('P','column','L1','C-round')['revision']==2,nonblocking_error_dialogs
   QTest.mouseClick(c['load'],Qt.LeftButton);assert next(x for x in c['state']['data']['rebars']['ties']['segments'] if x['id']=='SP1')['pitch_m']=='.125'
   QTest.mouseClick(c['commit'],Qt.LeftButton);original=deepcopy(app.open_project('P'))
   assert len(original['boq'])==4,nonblocking_error_dialogs
   c['numeric']['count'].setText('3');c['creator'].setText('پیشنهاددهنده مصنوعی');c['reason'].setText('مدرک مصنوعی اصلاح تعداد هم‌تیپ')
   QTest.mouseClick(c['amend']['create'],Qt.LeftButton)
   c['reviewer'].setText('بازبین مستقل مصنوعی');c['note'].setText('کنترل مصنوعی قطر و مارپیچ')
   QTest.mouseClick(c['amend']['review'],Qt.LeftButton);QTest.mouseClick(c['amend']['apply'],Qt.LeftButton)
   assert app.get_structural_member('P','column','L1','C-round')['revision']==3,nonblocking_error_dialogs
   assert app.open_project('P')['reviewed_structural_members']==original['reviewed_structural_members']
   for fmt in ('xlsx','pdf'):
    out=tmp_path/('native-circular.'+fmt)
    monkeypatch.setattr(QFileDialog,'getSaveFileName',lambda *a,chosen=str(out),**k:(chosen,''))
    QTest.mouseClick(c['reports'][fmt],Qt.LeftButton);assert out.exists(),nonblocking_error_dialogs
   book=load_workbook(tmp_path/'native-circular.xlsx');assert book.active.max_row>10;book.close()
   pdf=pymupdf.open(tmp_path/'native-circular.pdf');assert len(pdf)>0
   # Exact numeric text must stay inside the actual table grid on each page.
   for page in pdf:
    lines=[item for draw in page.get_drawings() for item in draw['items'] if item[0]=='l']
    vertical=[item[1].x for item in lines if abs(item[1].x-item[2].x)<.01]
    assert vertical
    for word in page.get_text('words'):
     if len(word[4])>20 and all(c in '0123456789.-' for c in word[4]):
      assert word[0]>=min(vertical)-1 and word[2]<=max(vertical)+1
   pdf.close()
  except Exception as exc:errors.append(str(exc))
  finally:dialog.reject()
 QTimer.singleShot(50,run);open_manual_structural_editor(app,project_id='P',floor_id='L1');assert not errors,errors
 p=app.open_project('P');report=app.preview_structural_member('P','column','L1','C-round')
 for row in report['rows']:
  matches=[b for b in p['boq'] if b['source_id']==row['source_id']];assert len(matches)==1 and Decimal(str(matches[0]['quantity']))==Decimal(str(row['quantity']))
 app.store.close();app=StructuralProApp(tmp_path/'project');assert app.open_project('P')['boq']==p['boq'];app.store.close()
