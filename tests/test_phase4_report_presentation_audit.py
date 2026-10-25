"""Readable Phase4 display must retain exact, recoverable audit values."""
from copy import deepcopy
import pytest
from openpyxl import load_workbook
from core.reports.project_report import build_report
from core.reports.exporters import export_xlsx

HASH='0123456789abcdef'*4
EXACT='24.66666666666666666666666667'

def test_missing_summary_is_explicit_and_reports_real_row_count(tmp_path):
    path=export_xlsx([{'شرح':'رکورد واقعی آزمایش'}],tmp_path/'report.xlsx')
    sheet=load_workbook(path)['خلاصه']
    values=[cell.value for row in sheet for cell in row]
    assert 'تعداد ردیف‌های گزارش' in values and 1 in values
    assert 'خلاصه عددی ثبت نشده است' in values

@pytest.mark.parametrize('fmt',['xlsx','docx','pdf'])
def test_compact_phase4_has_full_hash_and_exact_percentage_in_audit(tmp_path,fmt):
    rows=[{'بخش':'پیشرفت','مقدار و وضعیت':'واقعی: '+EXACT+'٪','مرجع و نسخه':HASH}]
    before=deepcopy(rows)
    report=build_report('گزارش آزمایشی',rows,metadata={'_phase4_presentation':True})
    path=report.export(tmp_path/('report.'+fmt),fmt)
    assert rows==before and report.as_rows()==before
    if fmt=='xlsx':
        book=load_workbook(path)
        detail='\n'.join(str(c.value) for r in book.worksheets[0] for c in r if c.value is not None)
        assert EXACT not in detail and HASH not in detail
        assert '≈24.67٪' in detail and '[هش 1]' in detail
        assert 'ممیزی مقادیر' in book.sheetnames
        text='\n'.join(str(c.value) for s in book for r in s for c in r if c.value is not None)
    elif fmt=='docx':
        from docx import Document
        d=Document(path);text='\n'.join([p.text for p in d.paragraphs]+[c.text for t in d.tables for r in t.rows for c in r.cells])
        assert '≈24.67٪' in text
    else:
        import fitz
        with fitz.open(path) as d:text='\n'.join(p.get_text() for p in d)
    assert HASH in ''.join(text.split())
    assert EXACT in ''.join(text.split())
    assert '_phase4_presentation' not in text

def test_phase4_csv_remains_exact_and_logical(tmp_path):
    rows=[{'مقدار و وضعیت':EXACT+'٪','مرجع و نسخه':HASH}]
    path=build_report('آزمایش',rows,metadata={'_phase4_presentation':True}).export(tmp_path/'raw.csv','csv')
    text=path.read_text(encoding='utf-8-sig')
    assert EXACT in text and HASH in text and '_phase4_presentation' not in text

@pytest.mark.parametrize('raw,shown',[('100.000000','100٪'),('0.000000','0٪'),('-0.000001','≈0٪'),('-25.000000','-25٪'),('92.500000','92.5٪'),('0.005','≈0.01٪')])
def test_percent_display_preserves_sign_and_original_token(raw,shown):
    from core.reports.exporters import _phase4_presentation
    rows=[{'عنوان':'واقعی درصد','مقدار و مرجع':raw}]
    display,columns,audit=_phase4_presentation(rows,[('عنوان','عنوان'),('مقدار و مرجع','مقدار و مرجع')],{'_phase4_presentation':True})
    assert display[0]['مقدار و مرجع']==shown
    assert audit[0]['value']==raw and rows[0]['مقدار و مرجع']==raw

def test_identical_hash_prefixes_still_have_distinct_audit_references():
    from core.reports.exporters import _phase4_presentation
    one='0'*64;two='0'*63+'1'
    rows=[{'مرجع':one},{'مرجع':two}]
    display,_,audit=_phase4_presentation(rows,[('مرجع','مرجع')],{'_phase4_presentation':True})
    assert display[0]['مرجع']!=display[1]['مرجع']
    assert [entry['value'] for entry in audit]==[one,two]
