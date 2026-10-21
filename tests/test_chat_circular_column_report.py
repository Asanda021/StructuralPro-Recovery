"""Backend regression: current circular-column details survive a real XLSX export.

The UI/PDF requires separate runtime dependencies, so this test does not claim
Windows or Persian print acceptance.
"""
import importlib.util
from pathlib import Path
from openpyxl import load_workbook

from core.platform.application import StructuralProApp


def test_circular_spiral_xlsx_contains_traceable_geometry(tmp_path):
    source=Path(__file__).with_name('test_chat_circular_column.py')
    spec=importlib.util.spec_from_file_location('circular_fixture',source)
    module=importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    app=StructuralProApp(tmp_path)
    app.create_project('پروژه مصنوعی','P')
    data=module.with_spiral()
    app.save_structural_member_draft('P',data,expected_revision=0)
    app.commit_structural_member('P','column','L1','C-round',expected_revision=1)
    out=tmp_path/'column-round.xlsx'
    app.export_column_detail_report('P','L1','C-round',out,format='xlsx')
    assert out.is_file() and out.stat().st_size>0
    wb=load_workbook(out,read_only=True)
    try:
        values=[str(cell) for row in wb.active.values for cell in row if cell is not None]
        assert any('مارپیچ' in v for v in values)
        assert any('pi_fixed_50_decimal_v1' in v for v in values)
        assert any('axis_diameter_m' in v or 'قطر محور مارپیچ' in v for v in values)
    finally:
        wb.close()
    app.store.close()


def test_circular_report_localizes_spiral_values_without_claiming_pdf(tmp_path):
    """Inspect real XLSX cells only; no Qt/PDF/Windows acceptance is implied."""
    import importlib.util
    from core.platform.application import StructuralProApp
    source=Path(__file__).with_name('test_chat_circular_column.py')
    spec=importlib.util.spec_from_file_location('circular_fixture_localized',source)
    module=importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    app=StructuralProApp(tmp_path)
    try:
        app.create_project('پروژه مصنوعی','P')
        record=module.with_spiral()
        app.save_structural_member_draft('P',record,expected_revision=0)
        app.commit_structural_member('P','column','L1','C-round',expected_revision=1)
        output=tmp_path/'localized_circular.xlsx'
        app.export_column_detail_report('P','L1','C-round',output,format='xlsx')
        book=load_workbook(output,read_only=True)
        try:
            strings=[str(cell) for row in book.active.values for cell in row if cell is not None]
            assert 'ستون دایره‌ای' in strings
            assert 'مارپیچ' in strings
            assert 'بدون طول اضافی با تأیید صریح' in strings
            assert any('سیاست عددی پی' in cell for cell in strings)
        finally:book.close()
    finally:app.store.close()

