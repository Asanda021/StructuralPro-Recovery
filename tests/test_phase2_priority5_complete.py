"""Real IFC, canonical immutable correction journal and SQLite acceptance."""
from copy import deepcopy
from decimal import Decimal
from hashlib import sha256
import json
import pytest
from test_phase2_priority5_closure import identity_case
from test_phase2_priority5_ifc_lineage_chat import _review


def case(tmp_path, mode='split', delta=False, crosswalks=None, identity=None, names=('B1','B2','B3')):
    import ifcopenshell as ifc
    from core.drawings.ifc_pipeline import quantity_formula_data
    if identity is None:
        app,path,proof,_,_=identity_case(tmp_path)
    else:app,path,proof=identity
    existing={x['id'] for x in app.open_project('P')['members']}
    for member in names:
        if member not in existing:app.add_project_member('P',member,'F1','beam')
    model=ifc.open(str(path));first=model.by_type('IfcBeam')[0]
    if identity is not None:first.GlobalId=ifc.guid.new()
    placement=model.create_entity('IfcLocalPlacement',PlacementRelTo=first.ObjectPlacement.PlacementRelTo,RelativePlacement=model.create_entity('IfcAxis2Placement3D',Location=model.create_entity('IfcCartesianPoint',Coordinates=(7.,0.,0.))))
    second=model.create_entity('IfcBeam',GlobalId=ifc.guid.new(),Name='Second explicit beam',ObjectPlacement=placement)
    storey=first.ContainedInStructure[0].RelatingStructure
    model.create_entity('IfcRelContainedInSpatialStructure',GlobalId=ifc.guid.new(),RelatedElements=[second],RelatingStructure=storey)
    def length(obj,value):
        definitions=[x.RelatingPropertyDefinition for x in obj.IsDefinedBy if x.RelatingPropertyDefinition.is_a('IfcElementQuantity')]
        if definitions:definitions[0].Quantities[0].LengthValue=float(value)
        else:
            q=model.create_entity('IfcQuantityLength',Name='Length',LengthValue=float(value))
            qs=model.create_entity('IfcElementQuantity',GlobalId=ifc.guid.new(),Name='Qto_Beam',MethodOfMeasurement='Explicit synthetic segment schedule',Quantities=[q])
            model.create_entity('IfcRelDefinesByProperties',GlobalId=ifc.guid.new(),RelatedObjects=[obj],RelatingPropertyDefinition=qs)
        # Real swept solid: 0.2 x 0.3 section, extrusion along the member X axis.
        # Building rotation maps that axis to world Y: lengths are independently
        # observable as 6m or adjacent 2m/4m intervals [7,9] and [9,13].
        axes=model.create_entity('IfcAxis2Placement3D',Location=model.create_entity('IfcCartesianPoint',Coordinates=(0.,0.,0.)),Axis=model.create_entity('IfcDirection',DirectionRatios=(1.,0.,0.)),RefDirection=model.create_entity('IfcDirection',DirectionRatios=(0.,1.,0.)))
        profile=model.create_entity('IfcRectangleProfileDef',ProfileType='AREA',XDim=.2,YDim=.3)
        solid=model.create_entity('IfcExtrudedAreaSolid',SweptArea=profile,Position=axes,ExtrudedDirection=model.create_entity('IfcDirection',DirectionRatios=(0.,0.,1.)),Depth=float(value))
        rep=model.create_entity('IfcShapeRepresentation',ContextOfItems=model.by_type('IfcGeometricRepresentationContext')[0],RepresentationIdentifier='Body',RepresentationType='SweptSolid',Items=[solid])
        obj.Representation=model.create_entity('IfcProductDefinitionShape',Representations=[rep])
    length(first,6 if mode=='split' else 2);length(second,4)
    oldpath=tmp_path/'beams-R1.ifc';model.write(str(oldpath))
    old=[_review(app,oldpath,revision='R1',gid=first.GlobalId,member=names[0],proof=proof)]
    if mode=='merge':old.append(_review(app,oldpath,revision='R1',gid=second.GlobalId,member=names[1],proof=proof))
    for n,pid in enumerate(old):
        data=quantity_formula_data(app,'P',pid,'Length',author='author')
        app.save_manual_formula('P',f'old-{names[0]}-{n}',data,expected_revision=0,reason='Explicit original QTO')
        app.review_manual_formula('P',f'old-{names[0]}-{n}',expected_revision=1,reviewer='reviewer',note='Independent synthetic review')
        app.commit_manual_formula('P',f'old-{names[0]}-{n}',expected_revision=1)
    old_links=[]
    if crosswalks:
        for n,pid in enumerate(old):
            kw=drawing_mapping(app,tmp_path,pid,'old-'+str(n),crosswalks,([6] if mode=='split' else [2,4])[n])
            old_links.append(app.save_ifc_crosswalk('P',**kw,author='author',reviewer='reviewer',expected_digest=app.store.current_digest('P')))
    first.GlobalId=ifc.guid.new();second.GlobalId=ifc.guid.new()
    length(first,2 if mode=='split' else 6);length(second,5 if delta else 4)
    newpath=tmp_path/'beams-R2.ifc';model.write(str(newpath))
    new=[_review(app,newpath,revision='R2',gid=first.GlobalId,member=names[1] if mode=='split' else names[2],proof=proof)]
    if mode=='split':new.append(_review(app,newpath,revision='R2',gid=second.GlobalId,member=names[2],proof=proof))
    evidence={'policy':'reviewed_ifc_lineage_v1','project_id':'P','mode':mode,'previous_proposal_ids':old,'current_proposal_ids':new,'previous_source_sha256':sha256(oldpath.read_bytes()).hexdigest(),'current_source_sha256':sha256(newpath.read_bytes()).hexdigest(),'edges':[{'previous_proposal_id':a,'current_proposal_id':b,'physical_scope_reference':'Explicit segment schedule'} for a in old for b in new]}
    ep=tmp_path/'lineage.json';ep.write_text(json.dumps(evidence))
    event=app.review_ifc_identity_lineage('P',mode=mode,previous_proposal_ids=old,current_proposal_ids=new,author='author',reviewer='reviewer',evidence_path=str(ep),evidence_sha256=sha256(ep.read_bytes()).hexdigest(),note='Reviewed synthetic split/merge',expected_digest=app.store.current_digest('P'))
    if crosswalks:
        links=[drawing_mapping(app,tmp_path,pid,'new-'+str(n),crosswalks,([2,5 if delta else 4] if mode=='split' else [6])[n]) for n,pid in enumerate(new)]
        migration={'policy':'reviewed_ifc_crosswalk_group_v1','project_id':'P','lineage_id':event['id'],'previous_crosswalk_ids':[x['id'] for x in old_links],'current':links}
        (tmp_path/'crosswalk-group.json').write_text(json.dumps(migration))
    allocation={'policy':'reviewed_ifc_quantity_allocation_v1','project_id':'P','lineage_id':event['id'],'quantity_name':'Length','unit':'m','classification':'unspecified','previous':[{'proposal_id':p,'quantity':str(q),'scope_reference':'Explicit old member'} for p,q in zip(old,[6] if mode=='split' else [2,4])],'current':[{'proposal_id':p,'quantity':str(q),'scope_reference':'Explicit current segment'} for p,q in zip(new,[2,5 if delta else 4] if mode=='split' else [6])]}
    ep=tmp_path/'allocation.json';ep.write_text(json.dumps(allocation))
    args=dict(lineage_id=event['id'],evidence_path=str(ep),evidence_sha256=sha256(ep.read_bytes()).hexdigest(),author='author',reviewer='reviewer',expected_digest=app.store.current_digest('P'))
    return app,args,ep,allocation


@pytest.mark.parametrize('mode',['split','merge'])
def test_group_transfer_changes_only_effective_owners_and_reopens(tmp_path,mode):
    app,args,ep,data=case(tmp_path,mode)
    original=deepcopy(app.open_project('P'))
    assert hasattr(app,'apply_ifc_quantity_transfer'), 'Atomic canonical group ownership transfer is missing'
    result=app.apply_ifc_quantity_transfer('P',**args)
    after=app.open_project('P')
    assert result['status']=='applied' and result['delta']=='0'
    assert after['takeoffs']==original['takeoffs']
    assert after['manual_formula_register']==original['manual_formula_register']
    assert sum(Decimal(str(r['quantity'])) for r in after['boq'])==Decimal(6)
    assert {s for r in after['boq'] for s in r['source_ids']}=={r['source_id'] for r in result['rows']}
    from core.bim.ifc import DeepIFCAdapter
    source=after['ifc_identity_reviews'][-1]['document']['source_path']
    meshes={m.global_id:m for m in DeepIFCAdapter().read_display_meshes(source).meshes}
    for allocation in result['current']:
        mesh=meshes[allocation['global_id']]
        assert Decimal(str(max(v[1] for v in mesh.vertices)-min(v[1] for v in mesh.vertices)))==Decimal(allocation['quantity'])
    before=app.store.current_digest('P')
    with pytest.raises((ValueError,RuntimeError)):app.apply_ifc_quantity_transfer('P',**{**args,'expected_digest':before})
    assert app.store.current_digest('P')==before
    app.store.close()
    from core.platform.application import StructuralProApp
    reopened=StructuralProApp(tmp_path/'app')
    assert reopened.open_project('P')['boq']==after['boq']
    assert reopened.open_project('P')['takeoffs']==original['takeoffs']
    assert reopened.store.integrity_check()['ok']
    reopened.store.close()


@pytest.mark.parametrize('bad',['missing','duplicate','amount','classification','source','proof','self','stale','lock'])
def test_group_transfer_rejection_preserves_snapshot(tmp_path,bad):
    app,args,ep,data=case(tmp_path)
    if bad=='missing':data['current'].pop()
    elif bad=='duplicate':data['current'][1]=data['current'][0]
    elif bad=='amount':data['current'][0]['quantity']='3'
    elif bad=='classification':data['classification']='gross'
    elif bad=='source':
        project=app.open_project('P');p=project['ifc_identity_reviews'][-1]['document']['source_path']
        with open(p,'a') as f:f.write('\n/* Changed source */')
    elif bad=='proof':ep.write_text('{}')
    elif bad=='self':args['reviewer']='author'
    elif bad=='stale':args['expected_digest']='stale'
    elif bad=='lock':
        p=deepcopy(app.open_project('P'));p['contract']={'reference':'Locked synthetic contract'}
        app.store.save('P',p,expected_digest=app.store.current_digest('P'));args['expected_digest']=app.store.current_digest('P')
    if bad not in ('proof',):ep.write_text(json.dumps(data));args['evidence_sha256']=sha256(ep.read_bytes()).hexdigest()
    snapshot=deepcopy(app.open_project('P'));before=app.store.current_digest('P')
    with pytest.raises((ValueError,RuntimeError,PermissionError)):app.apply_ifc_quantity_transfer('P',**args)
    assert app.store.current_digest('P')==before and app.open_project('P')==snapshot
    app.store.close()


def drawing_mapping(app,tmp_path,pid,index,kind='dxf',length=6):
    from core.drawings.takeoff_session import DrawingTakeoffSession
    binding=next(x for x in app.open_project('P')['ifc_identity_reviews'] if x['proposal_id']==pid)
    if kind=='pdf':
        import fitz
        from core.drawings.graphical_takeoff import Point
        path=tmp_path/f'plan-{index}.pdf';pdf=fitz.open()
        for _ in range(3):pdf.new_page(width=100,height=100).draw_rect(fitz.Rect(0,0,length*10,2))
        pdf.save(path);pdf.close()
        session=DrawingTakeoffSession(str(path));session.current_page=2;session.calibrate(2,10,1)
        item=session.add_polygon([Point(0,0),Point(length*10,0),Point(length*10,2),Point(0,2)],floor_id=binding['floor_id'],member_id=binding['member_id'],quantity_owner='Documented synthetic scope',takeoff_code='TEST')
    else:
        import ezdxf
        from core.drawings.dwg_takeoff import DWGTakeoffEngine
        d=ezdxf.new('R2018');d.units=6;d.header['$DIMSTYLE']='Standard';entity=d.modelspace().add_line((0,0),(length,0));path=tmp_path/f'plan-{index}.dxf';d.saveas(path)
        if kind=='dwg':
            import os,subprocess
            path=path.with_suffix('.dwg')
            subprocess.run([os.environ['STRUCTURALPRO_TEST_DOTNET'],os.environ['STRUCTURALPRO_TEST_DWG_WRITER'],str(path.with_suffix('.dxf')),str(path)],check=True,capture_output=True)
        doc=DWGTakeoffEngine.with_array_support().import_file(path)
        handle=next(x.handle for x in doc.entities if x.entity_type=='LINE')
        session=DrawingTakeoffSession(str(path));item=session.add_cad_entity(doc,handle,'length',takeoff_code='TEST',floor_id=binding['floor_id'],member_id=binding['member_id'],quantity_owner='Documented scope',created_by='author',reference='Synthetic segment schedule')
        session.review_cad_entity(item.id,reviewer='reviewer',note='Independent synthetic review')
    saved=app.save_drawing_takeoff_session('P',session)
    evidence={'policy':'reviewed_reference_point_v1','project_id':'P','proposal_id':pid,'session_id':saved['session_id'],'item_id':item.id,'ifc_sha256':binding['document']['source_sha256'],'drawing_sha256':sha256(path.read_bytes()).hexdigest(),'ifc_origin_m':['1000000','9' if index.endswith('-1') else '7','0'],'drawing_anchor_xy':['0','0'],'ifc_to_drawing_si':['1','0','0','1','-1000000','-9' if index.endswith('-1') else '-7'],'physical_identity_reference':'Explicit synthetic ownership schedule','unit_evidence_reference':'Documented SI mapping'}
    ep=tmp_path/f'crosswalk-{index}.json';ep.write_text(json.dumps(evidence))
    return dict(proposal_id=pid,session_id=saved['session_id'],item_id=item.id,evidence_path=str(ep),evidence_sha256=sha256(ep.read_bytes()).hexdigest())


@pytest.mark.parametrize('mode',['split','merge'])
@pytest.mark.parametrize('kind',['dxf','pdf'])
def test_crosswalk_group_migration_is_append_only_atomic_and_reopens(tmp_path,mode,kind):
    app,args,ep,data=case(tmp_path,mode,crosswalks=kind)
    assert hasattr(app,'migrate_ifc_crosswalk_group'), 'Group crosswalk migration is missing'
    project=deepcopy(app.open_project('P'));old=project['ifc_2d_crosswalks']
    migration=json.loads((tmp_path/'crosswalk-group.json').read_text())
    proof=tmp_path/'crosswalk-group.json'
    kw=dict(lineage_id=args['lineage_id'],evidence_path=str(proof),evidence_sha256=sha256(proof.read_bytes()).hexdigest(),author='author',reviewer='reviewer',expected_digest=app.store.current_digest('P'))
    result=app.migrate_ifc_crosswalk_group('P',**kw)
    after=app.open_project('P');assert after['ifc_2d_crosswalks'][:len(old)]==old
    assert len(result)==(2 if mode=='split' else 1)
    assert {p for x in result for p in x['previous_crosswalk_ids']}=={x['id'] for x in old}
    assert all(x['lineage_id']==args['lineage_id'] and x['revision']==2 for x in result)
    assert after['boq']==project['boq'] # Crosswalk alone never pays.
    app.apply_ifc_quantity_transfer('P',**{**args,'expected_digest':app.store.current_digest('P')})
    assert sum(Decimal(str(r['quantity'])) for r in app.open_project('P')['boq'])==6
    app.store.close()
    from core.platform.application import StructuralProApp
    reopened=StructuralProApp(tmp_path/'app');assert reopened.open_project('P')['ifc_2d_crosswalks']==after['ifc_2d_crosswalks'];assert reopened.store.integrity_check()['ok'];reopened.store.close()


@pytest.mark.parametrize('mode',['split','merge'])
@pytest.mark.parametrize('kind',['dxf','pdf'])
def test_group_ownership_and_crosswalk_native_qt(tmp_path,kind,mode):
    from PySide6.QtWidgets import QApplication,QPushButton,QLineEdit,QLabel,QComboBox,QTableWidget,QScrollArea
    from PySide6.QtCore import Qt
    from PySide6.QtTest import QTest
    from app.graphical_takeoff import GraphicalTakeoffDialog
    app,args,ep,data=case(tmp_path,mode=mode,crosswalks=kind)
    source=app.open_project('P')['ifc_identity_reviews'][-1]['document']['source_path']
    qt=QApplication.instance() or QApplication([])
    w=GraphicalTakeoffDialog(app_service=app,project_id='P');w.show();w.open_drawing(source);w.ifc_revision.setText('R2');QTest.mouseClick(w.ifc_inspect_button,Qt.LeftButton)
    dialog=w._ifc_inspector
    scroll=dialog.findChild(QScrollArea)
    def click(button):
        scroll.ensureWidgetVisible(button);qt.processEvents()
        assert button.isVisible()
        QTest.mouseClick(button,Qt.LeftButton)
    try:
        button=dialog.findChild(QPushButton,'ifc_group_preview')
        assert button is not None, 'Native group ownership review controls are missing'
        dialog.findChild(QLineEdit,'ifc_author').setText('author')
        dialog.findChild(QLineEdit,'ifc_group_allocation').setText(str(ep))
        click(button)
        assert dialog.findChild(QTableWidget,'ifc_group_quantities').rowCount()==3
        assert '6' in dialog.findChild(QLabel,'ifc_group_summary').text()
        dialog.findChild(QPushButton,'ifc_group_confirm').setChecked(True)
        click(dialog.findChild(QPushButton,'ifc_group_apply'))
        assert any(a.get('kind')=='reviewed_ifc_quantity_transfer_v1' for a in app.open_project('P')['manual_formula_amendments'].values()),dialog.findChild(QLabel,'ifc_review_outcome').text()
        dialog.findChild(QLineEdit,'ifc_group_crosswalk_document').setText(str(tmp_path/'crosswalk-group.json'))
        click(dialog.findChild(QPushButton,'ifc_group_migrate'))
        links=app.open_project('P')['ifc_2d_crosswalks'];assert len(links)==3,dialog.findChild(QLabel,'ifc_review_outcome').text()
        combo=dialog.findChild(QComboBox,'ifc_group_links');combo.setCurrentIndex(combo.count()-1)
        click(dialog.findChild(QPushButton,'ifc_group_open_link'));qt.processEvents()
        other=w._ifc_crosswalk_window
        assert other is not None
        target=links[-1];w.highlight_ifc_member(target['global_id']);qt.processEvents()
        if kind in ('dxf','dwg'):
            assert other._cad_handle_items[target['drawing']['handle']].isSelected()
            other.canvas.scene.clearSelection()
            obj=other._cad_handle_items[target['drawing']['handle']]
            point=other.canvas.mapFromScene(obj.line().pointAt(.5))
            QTest.mouseClick(other.canvas.viewport(),Qt.LeftButton,pos=point);qt.processEvents()
        else:
            assert other.page==2
            other.table.clearSelection();point=other.table.visualItemRect(other.table.item(0,0)).center()
            QTest.mouseClick(other.table.viewport(),Qt.LeftButton,pos=point);qt.processEvents()
        assert w._ifc_selected_global_id==target['global_id']
        assert target['global_id'] in w._ifc_handle_items, 'The selected IFC product has no actual displayed mesh'
        assert any(item.isSelected() for item in w._ifc_handle_items[target['global_id']])
        scroll.ensureWidgetVisible(dialog.findChild(QTableWidget,'ifc_group_quantities'));qt.processEvents()
        dialog.grab().save(str(tmp_path/'group-ownership-qt.png'));other.grab().save(str(tmp_path/'group-linked-qt.png'))
    finally:w.close();qt.processEvents();app.store.close()


@pytest.mark.parametrize('bad',['missing','duplicate','old_missing','proof','source','self','stale'])
def test_crosswalk_group_failure_is_atomic(tmp_path,bad):
    app,args,_,_=case(tmp_path,crosswalks='dxf')
    ep=tmp_path/'crosswalk-group.json';data=json.loads(ep.read_text())
    if bad=='missing':data['current'].pop()
    elif bad=='duplicate':data['current'][1]=data['current'][0]
    elif bad=='old_missing':data['previous_crosswalk_ids']=['unknown']
    elif bad=='source':
        target=json.loads(open(data['current'][0]['evidence_path']).read());target['drawing_sha256']='0'*64
        p=data['current'][0]['evidence_path'];open(p,'w').write(json.dumps(target));data['current'][0]['evidence_sha256']=sha256(open(p,'rb').read()).hexdigest()
    ep.write_text(json.dumps(data))
    kw=dict(lineage_id=args['lineage_id'],evidence_path=str(ep),evidence_sha256=sha256(ep.read_bytes()).hexdigest(),author='author',reviewer='author' if bad=='self' else 'reviewer',expected_digest='stale' if bad=='stale' else app.store.current_digest('P'))
    if bad=='proof':ep.write_text('{}')
    snapshot=deepcopy(app.open_project('P'));before=app.store.current_digest('P')
    with pytest.raises((ValueError,RuntimeError,PermissionError)):app.migrate_ifc_crosswalk_group('P',**kw)
    assert app.store.current_digest('P')==before and app.open_project('P')==snapshot
    app.store.close()


def test_group_documented_delta_and_history_protection(tmp_path):
    app,args,ep,data=case(tmp_path,delta=True)
    data['quantity_change_reference']='Explicit revised segment 4 to 5 metres'
    change=tmp_path/'delta.json'
    project=app.open_project('P');event=project['ifc_identity_lineages'][0]
    change.write_text(json.dumps({'policy':'reviewed_ifc_quantity_delta_v1','project_id':'P','lineage_id':event['id'],'previous_source_sha256':event['previous_source_sha256'],'current_source_sha256':event['current_source_sha256'],'delta':'1','unit':'m','classification':'unspecified','reason':'Explicit synthetic correction to 5m segment'}))
    data['quantity_change_document']={'path':str(change),'sha256':sha256(change.read_bytes()).hexdigest()}
    ep.write_text(json.dumps(data));args['evidence_sha256']=sha256(ep.read_bytes()).hexdigest()
    result=app.apply_ifc_quantity_transfer('P',**args)
    assert result['previous_total']=='6' and result['current_total']=='7' and result['delta']=='1'
    assert sum(Decimal(str(r['quantity'])) for r in app.open_project('P')['boq'])==7
    before=app.store.current_digest('P')
    for field in ['takeoffs','manual_formula_amendments','boq']:
        project=deepcopy(app.open_project('P'))
        if field=='takeoffs':project[field][0]['quantities'][0]['amount']='99'
        elif field=='boq':project[field][0]['quantity']='99'
        else:project[field][result['id']]['delta']='99'
        with pytest.raises(ValueError):app.store.save('P',project,expected_digest=before)
        assert app.store.current_digest('P')==before
    app.store.close()


@pytest.mark.parametrize('mode',['split','merge'])
def test_group_cas_second_connection_and_consumed_owner_bypass(tmp_path,mode):
    from core.platform.application import StructuralProApp
    from core.takeoff.commit_guards import validate_commit_rows,require_source_unused
    app,args,_,_=case(tmp_path,mode,crosswalks='dxf')
    other=StructuralProApp(tmp_path/'app');stale=deepcopy(other.open_project('P'));token=other.store.current_digest('P')
    result=app.apply_ifc_quantity_transfer('P',**args)
    before=app.store.current_digest('P')
    with pytest.raises(RuntimeError):other.store.save('P',stale,expected_digest=token)
    project=deepcopy(app.open_project('P'))
    for row in result['rows']+result['consumed_rows']:
        with pytest.raises(ValueError):validate_commit_rows(project,[row],source_id=row['source_id'])
        with pytest.raises(ValueError):require_source_unused(project,row['source_id'])
        # CAD/manual cannot acquire the same reviewed physical owner under a new source.
        candidate=deepcopy(row);candidate['source_id']='different-session-and-operation';candidate['params'].pop('ifc_evidence',None)
        with pytest.raises(ValueError):validate_commit_rows(project,[candidate],source_id=candidate['source_id'])
    assert app.store.current_digest('P')==before and app.store.integrity_check()['ok']
    app.store.close();other.store.close()


@pytest.mark.parametrize('bad',['missing_delta_document','changed_delta_document','changed_lineage','unknown_owner','unit','missing_classification','changed_floor','missing_old_commit'])
def test_additional_group_allocation_fail_closed(tmp_path,bad):
    app,args,ep,data=case(tmp_path,delta=bad.startswith('missing_delta') or bad.startswith('changed_delta'))
    if bad.endswith('delta_document'):
        data['quantity_change_reference']='Explicit correction reference'
        if bad=='changed_delta_document':
            doc=tmp_path/'delta-proof.json';doc.write_text('{}')
            data['quantity_change_document']={'path':str(doc),'sha256':'0'*64}
    elif bad=='changed_lineage':
        event=app.open_project('P')['ifc_identity_lineages'][0];open(event['evidence_path'],'w').write('{}')
    elif bad=='unknown_owner':data['current'][0]['proposal_id']='missing'
    elif bad=='unit':data['unit']='m2'
    elif bad=='missing_classification':data.pop('classification')
    elif bad=='changed_floor':
        data['current'][0]['floor_id']='F2'
    else:
        # Existing nonfinancial lineage with no committed old QTO cannot be paid.
        app.store.close()
        from test_phase2_priority5_verified import quantity_lineage_case
        other=tmp_path/'nonfinancial';other.mkdir();app,args,ep,data=quantity_lineage_case(other);data['classification']='unspecified'
    ep.write_text(json.dumps(data));args['evidence_sha256']=sha256(ep.read_bytes()).hexdigest()
    snapshot=deepcopy(app.open_project('P'));before=app.store.current_digest('P')
    with pytest.raises((ValueError,RuntimeError)):app.apply_ifc_quantity_transfer('P',**args)
    assert app.store.current_digest('P')==before and app.open_project('P')==snapshot
    app.store.close()


def test_split_and_merge_in_one_project_keep_history_and_twelve_metre_boq(tmp_path):
    app,args,_,_=case(tmp_path,crosswalks='pdf')
    first_history=deepcopy(app.open_project('P')['takeoffs'])
    sub=tmp_path/'merge';sub.mkdir()
    other,merge_args,_,_=case(sub,mode='merge',crosswalks='dxf',identity=(app,tmp_path/'source.ifc',tmp_path/'mapping.txt'),names=('M1','M2','M3'))
    assert other is app
    assert sum(Decimal(str(x['quantity'])) for x in app.open_project('P')['boq'])==12
    for group,folder in [(args,tmp_path),(merge_args,sub)]:
        ep=folder/'crosswalk-group.json'
        app.migrate_ifc_crosswalk_group('P',lineage_id=group['lineage_id'],evidence_path=str(ep),evidence_sha256=sha256(ep.read_bytes()).hexdigest(),author='author',reviewer='reviewer',expected_digest=app.store.current_digest('P'))
        app.apply_ifc_quantity_transfer('P',**{**group,'expected_digest':app.store.current_digest('P')})
    project=app.open_project('P')
    assert sum(Decimal(str(x['quantity'])) for x in project['boq'])==12
    assert project['takeoffs'][:len(first_history)]==first_history
    assert len(project['manual_formula_amendments'])==2 and len(project['ifc_2d_crosswalks'])==6
    assert app.store.integrity_check()['ok'];app.store.close()


def test_crosswalk_migration_rejects_changed_old_session(tmp_path):
    from core.drawings.takeoff_session import DrawingTakeoffSession
    app,args,_,_=case(tmp_path,crosswalks='dxf')
    link=app.open_project('P')['ifc_2d_crosswalks'][0]
    saved=app.load_drawing_takeoff_session('P',link['session_id'])
    session=DrawingTakeoffSession.from_dict(saved['session']);session.edit(session.items[0].id,label='Changed measured scope; requires new review')
    app.save_drawing_takeoff_session('P',session,session_id=link['session_id'],expected_revision=saved['revision'])
    ep=tmp_path/'crosswalk-group.json';before=app.store.current_digest('P')
    with pytest.raises(ValueError):app.migrate_ifc_crosswalk_group('P',lineage_id=args['lineage_id'],evidence_path=str(ep),evidence_sha256=sha256(ep.read_bytes()).hexdigest(),author='author',reviewer='reviewer',expected_digest=before)
    assert app.store.current_digest('P')==before;app.store.close()


def test_existing_reports_read_group_journal_and_effective_boq(tmp_path):
    from openpyxl import load_workbook
    from docx import Document
    import csv
    app,args,_,_=case(tmp_path)
    event=app.apply_ifc_quantity_transfer('P',**args)
    path=tmp_path/'group-register.xlsx'
    app.export_manual_formula_report('P',path,format='xlsx')
    book=load_workbook(path,data_only=True)
    assert 'انتقال مالکیت IFC' in book.sheetnames
    values=' '.join(str(c.value) for row in book['انتقال مالکیت IFC'] for c in row)
    assert event['lineage_id'] in values and event['rows'][0]['source_id'] in values
    for fmt in ('csv','docx','pdf','xlsx'):app.export_manual_formula_report('P',tmp_path/('effective.'+fmt),format=fmt)
    # No approved price book was supplied: the priced financial report must stay closed.
    with pytest.raises(ValueError):app.report('P','pdf',tmp_path/'unapproved-financial.pdf')
    with (tmp_path/'effective.csv').open(encoding='utf-8-sig') as f:rows=list(csv.reader(f))
    assert '6' in str(rows)
    doc=Document(tmp_path/'effective.docx');assert doc.tables
    assert app.store.integrity_check()['ok'];app.store.close()
