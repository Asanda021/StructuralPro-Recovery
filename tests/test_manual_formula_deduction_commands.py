from copy import deepcopy
import pytest
from core.platform.application import StructuralProApp
from test_manual_formula_persistence import data

def test_deduction_create_edit_delete_gross_net_reopen(tmp_path):
 s=StructuralProApp(tmp_path);s.create_project('مصنوعی','P');s.save_manual_formula('P','A',data(unit='m2',expression='length*height',variables={'length':'10','height':'3'}),expected_revision=0,reason='اول')
 assert callable(getattr(s,'set_manual_formula_deduction',None)), 'فرم و گردش کسورات متصل نیست'
 item={'id':'opening:geometry:1','variable':'کسر۱','quantity':'2','unit':'m2','basis':'ورودی هندسی صریح 1×2متر','source_id':'opening:geometry:1'}
 s.set_manual_formula_deduction('P','A',item,expected_revision=1,reason='ثبت بازشو');rows=s.load_manual_formulas('P');assert rows['A']['quantity']==28
 summary=s.manual_formula_deduction_summary('P','A');assert summary=={'gross':30,'net':28,'deducted':2,'unit':'m2'}
 item['quantity']='3';s.set_manual_formula_deduction('P','A',item,expected_revision=2,reason='اصلاح هندسه صریح');assert s.load_manual_formulas('P')['A']['quantity']==27
 before=deepcopy(s.open_project('P'))
 with pytest.raises(ValueError):s.set_manual_formula_deduction('P','A',{**item,'id':'other','variable':'دیگر'},expected_revision=3,reason='هندسه تکراری')
 assert s.open_project('P')==before
 s.remove_manual_formula_deduction('P','A','opening:geometry:1',expected_revision=3,reason='حذف کنترل‌شده');assert s.load_manual_formulas('P')['A']['quantity']==30
 s.store.close();s=StructuralProApp(tmp_path);assert s.manual_formula_deduction_summary('P','A')['net']==30;s.store.close()

@pytest.mark.parametrize('bad',[{'unit':'m3'},{'quantity':'-2'},{'quantity':'NaN'},{'source_id':''}])
def test_invalid_deduction_atomic(tmp_path,bad):
 s=StructuralProApp(tmp_path);s.create_project('مصنوعی','P');s.save_manual_formula('P','A',data(unit='m2',expression='length*height',variables={'length':'10','height':'3'}),expected_revision=0,reason='اول')
 assert callable(getattr(s,'set_manual_formula_deduction',None)), 'اعتبارسنجی کسورات وجود ندارد'
 with pytest.raises(ValueError):s.set_manual_formula_deduction('P','A',{'id':'shape','variable':'کسر','quantity':'2','unit':'m2','basis':'هندسه صریح','source_id':'shape',**bad},expected_revision=1,reason='کسر')
 assert s.load_manual_formulas('P')['A']['quantity']==30;s.store.close()

def test_committed_deduction_requires_independent_amendment_and_reports(tmp_path):
 from test_manual_formula_amendments import base
 from core.takeoff.manual_formula_deductions import change_data
 s=base(tmp_path/'db');row=s.load_manual_formulas('P')['A'];item={'id':'opening:volume','source_id':'opening:volume','variable':'بازشو','quantity':'0.5','unit':'m3','basis':'حجم هندسی مصنوعی صریح'}
 with pytest.raises(ValueError):s.set_manual_formula_deduction('P','A',item,expected_revision=1,reason='تغییر مستقیم')
 a=s.create_manual_formula_amendment('P','A',change_data(row['data'],item),expected_revision=1,created_by='سازنده',reason='کسر حجم مستند');s.review_manual_formula_amendment('P',a['id'],expected_revision=1,reviewer='بازبین مستقل',note='حجم و فرمول');s.apply_manual_formula_amendment('P',a['id'],expected_revision=1)
 assert s.manual_formula_deduction_summary('P','A')=={'gross':6,'net':5.5,'deducted':0.5,'unit':'m3'}
 assert s.open_project('P')['takeoffs'][0]['quantities'][0]['amount']==6
 path=tmp_path/'deductions.xlsx';s.export_manual_formula_report('P',path,format='xlsx')
 from openpyxl import load_workbook
 w=load_workbook(path);assert 'اصلاحیه قطعی' in w.sheetnames and 'قبل و بعد کسر' in w.sheetnames;assert w['کسورات']['F2'].value=='opening:volume'
 s.export_manual_formula_report('P',tmp_path/'deductions.pdf',format='pdf');s.store.close();s=StructuralProApp(tmp_path/'db');assert s.load_manual_formulas('P')['A']['quantity']==5.5;s.store.close()
