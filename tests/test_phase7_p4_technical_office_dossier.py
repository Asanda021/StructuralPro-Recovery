"""P7 priority 4: real source-backed site/control dossier on ProjectStore."""
from copy import deepcopy
from hashlib import sha256
import csv
import json

import pytest

from test_phase4_progress import project, weights, work
from test_phase4_structure import cmd
from core.reports.technical_office_dossier import export_verified_site_controls_dossier


def _export(app, *, target, kinds=("wbs", "activities", "progress", "curve"), version=None):
    return export_verified_site_controls_dossier(
        app, project_id="P",
        expected_project_version=app.open_project("P")["_meta"]["version"] if version is None else version,
        as_of="2026-01-06", selected_kinds=kinds, destination=target,
    )


def _ready(project):
    app, proof = project
    weights(app, proof)
    work(app, proof, "A", "5", "2026-01-06", "ACT-A")
    work(app, proof, "B", "5", "2026-01-06", "ACT-B")
    cmd(app, "progress_report", date="2026-01-06", reference="PERIOD-1")
    return app


def test_real_wbs_progress_curve_dossier_four_formats_and_graph_data(project, tmp_path):
    import pymupdf
    from docx import Document
    from openpyxl import load_workbook

    app = _ready(project)
    before = deepcopy(app.open_project("P"))
    digest = app.store.current_digest("P")
    out = tmp_path / "dossier"
    actual = _export(app, target=out)
    manifest = json.loads((out / "manifest.json").read_text(encoding="utf-8"))
    assert actual["manifest"] == manifest
    assert manifest["project_version"] == before["_meta"]["version"]
    assert manifest["source_project_digest"] == digest
    assert manifest["report_count"] == 4
    assert len(manifest["reports"]) == 4
    assert not manifest["financial_mutation"]
    assert manifest["physical_office_print_status"] == "NOT_TESTED"
    assert len({row["name"] for row in manifest["files"]}) == len(manifest["files"])
    for entry in manifest["files"]:
        path = out / entry["name"]
        assert entry["sha256"] == sha256(path.read_bytes()).hexdigest()
        assert entry["size"] == path.stat().st_size > 0
    assert {p.name for p in out.iterdir()} == {f["name"] for f in manifest["files"]} | {"manifest.json"}
    assert manifest["figure_sidecars"]
    for name in manifest["figure_sidecars"]:
        assert name.endswith(".png")
    assert any(kind == "curve" for kind in [r["kind"] for r in manifest["reports"]])
    for title in ("03-progress", "04-curve"):
        with pymupdf.open(out / (title + ".pdf")) as pdf:
            assert len(pdf) and any(page.get_text().strip() for page in pdf)
        d = Document(out / (title + ".docx"))
        assert d.tables and "w:bidi" in d._element.xml
        wb = load_workbook(out / (title + ".xlsx"))
        assert all(ws.sheet_view.rightToLeft for ws in wb)
        with (out / (title + ".csv")).open(encoding="utf-8-sig", newline="") as stream:
            assert len(list(csv.reader(stream, strict=True))) > 3
    assert app.store.current_digest("P") == digest
    assert app.open_project("P") == before


def test_missing_approved_progress_yields_no_dossier(project, tmp_path):
    app, _ = project
    target = tmp_path / "empty"
    with pytest.raises(ValueError, match="پیشرفت"):
        _export(app, target=target, kinds=("progress",))
    assert not target.exists()


@pytest.mark.parametrize("kinds", [(), ("bogus",), ("wbs", "wbs"), ["wbs"]])
def test_unknown_duplicate_or_non_tuple_kinds_rejected(project, tmp_path, kinds):
    app, _ = project
    with pytest.raises(ValueError):
        _export(app, target=tmp_path / "invalid", kinds=kinds)
    assert not (tmp_path / "invalid").exists()


def test_stale_project_version_or_bad_date_blocks_export(project, tmp_path):
    app, _ = project
    out = tmp_path / "stale"
    with pytest.raises(ValueError, match="نسخه"):
        _export(app, target=out, kinds=("wbs",),
                version=app.open_project("P")["_meta"]["version"] - 1)
    with pytest.raises(ValueError, match="تاریخ"):
        export_verified_site_controls_dossier(
            app, project_id="P", expected_project_version=app.open_project("P")["_meta"]["version"],
            as_of="2026-02-30", selected_kinds=("wbs",), destination=out,
        )
    assert not out.exists()


def test_existing_destination_not_overwritten(project, tmp_path):
    app, _ = project
    path = tmp_path / "existing"
    path.mkdir()
    with pytest.raises(FileExistsError):
        _export(app, target=path, kinds=("wbs",))
