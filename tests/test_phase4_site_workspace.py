from copy import deepcopy
from pathlib import Path
import pytest
from PySide6.QtWidgets import QApplication,QPushButton
from PySide6.QtTest import QTest
from PySide6.QtCore import Qt
from test_phase4_site_resources import setup,resource

QT=QApplication.instance() or QApplication([])

def test_resource_real_buttons_dashboard_and_error(setup):
    from app.site_operations_workspace import SiteOperationsWorkspace
    a,run,proof=setup
    w=SiteOperationsWorkspace(a,'P');w.show();QT.processEvents()
    try:
        for key,value in {'actor':'maker','reason':'آزمون رابط','proof':proof['path'],'sha':proof['sha256'],'reference':proof['reference']}.items():w.fields[key].setText(value)
        w.choose('resource_create')
        for key,value in {'id':'UI-C','name':'اکیپ رابط','capacity':'1','unit':'crew','discipline_id':'S','subcontractor':'نمونه'}.items():w.operation_fields[key].setText(value)
        QTest.mouseClick(w.findChild(QPushButton,'SiteOperationSave'),Qt.LeftButton)
        assert a.open_project('P')['project_resources'][-1]['id']=='UI-C',w.status.text()
        assert 'ثبت' in w.status.text()
        version=a.open_project('P')['_meta']['version']
        QTest.mouseClick(w.findChild(QPushButton,'SiteOperationSave'),Qt.LeftButton)
        assert 'خطا' in w.status.text() and a.open_project('P')['_meta']['version']==version
        w.fields['as_of'].setText('2026-10-26')
        QTest.mouseClick(w.findChild(QPushButton,'SiteDashboardRefresh'),Qt.LeftButton)
        assert w.dashboard.rowCount()>0
        assert w.metrics.rowCount()>=10
        assert any('۱۴۰۵' in w.dashboard.item(i,1).text() for i in range(w.dashboard.rowCount()))
        assert w.layoutDirection()==Qt.RightToLeft
    finally:w.close();QT.processEvents()

@pytest.mark.parametrize('fmt',['xlsx','docx','pdf','csv'])
def test_site_reports_readback_no_mutation(setup,tmp_path,fmt):
    a,run,proof=setup;resource(run,proof);before=deepcopy(a.open_project('P'))
    path=tmp_path/('resources.'+fmt)
    a.site_report('P','resources',fmt,path,as_of='2026-10-26')
    assert path.stat().st_size>100
    if fmt=='xlsx':
        import openpyxl
        wb=openpyxl.load_workbook(path);assert all(s.sheet_view.rightToLeft for s in wb)
        assert any('اکیپ' in str(c.value) for s in wb for row in s for c in row)
    elif fmt=='docx':
        from docx import Document
        d=Document(path);assert d.tables
        assert 'اکیپ' in ''.join(c.text for t in d.tables for r in t.rows for c in r.cells)
        assert 'bidi' in d._element.xml
    elif fmt=='pdf':
        import fitz
        d=fitz.open(path);assert len(d)>0 and len(''.join(p.get_text() for p in d))>100
    else:assert 'اکیپ' in path.read_text(encoding='utf-8-sig')
    assert a.open_project('P')==before

def test_daily_issue_and_w4_selection_real_ui(setup):
    from app.site_operations_workspace import SiteOperationsWorkspace
    a,run,proof=setup;resource(run,proof)
    w=SiteOperationsWorkspace(a,'P');w.show();QT.processEvents()
    def fill(values):
        for key,value in values.items():
            f=w.operation_fields[key]
            if hasattr(f,'setPlainText'):f.setPlainText(value)
            else:f.setText(value)
    def save():QTest.mouseClick(w.findChild(QPushButton,'SiteOperationSave'),Qt.LeftButton);assert not w.status.text().startswith('خطا'),w.status.text()
    try:
        for key,value in {'actor':'maker','reason':'آزمون رابط','proof':proof['path'],'sha':proof['sha256'],'reference':'DAILY-UI'}.items():w.fields[key].setText(value)
        w.choose('daily_create');fill({'activity_id':'A','date':'2026-10-26','summary':'کارگاه آزمایشی','location':'طبقه یک','crew_ids':'C','equipment_ids':'','material_ids':'','human_count':'3','quantity':'1.25','unit':'m3','technical_notes':'کار ثبت‌شده، هنوز پیشرفت تأییدشده نیست','attachments':''});save()
        id=a.open_project('P')['project_daily_reports'][-1]['id']
        for op,actor in [('daily_review','reviewer'),('daily_approve','approver'),('daily_issue','owner')]:
            w.choose(op);w.fields['actor'].setText(actor);fill({'id':id});save()
        assert not a.open_project('P').get('progress_documents')
        w.tabs.setCurrentIndex(1);QT.processEvents();QTest.mouseClick(w.findChild(QPushButton,'SiteSourcesRefresh'),Qt.LeftButton)
        assert w.sources_table.rowCount()==1
        item=w.sources_table.item(0,0)
        QTest.mouseClick(w.sources_table.viewport(),Qt.LeftButton,pos=w.sources_table.visualItemRect(item).center())
        QTest.mouseDClick(w.sources_table.viewport(),Qt.LeftButton,pos=w.sources_table.visualItemRect(item).center());QT.processEvents()
        assert w.operation.currentData()=='package_create'
        w.fields['actor'].setText('maker');w.fields['reference'].setText('PACKAGE-UI')
        fill({'date':'2026-10-26','title':'بسته رابط','template':'site-office','template_version':'1'});save()
        assert a.open_project('P')['site_document_packages'][0]['documents'][0]['id']==id
        w.horizon.setCurrentIndex(2);w.custom_end.setText('2026-11-03');w.fields['as_of'].setText('2026-10-26');w.show_dashboard()
        assert w.dashboard.rowCount()>0
    finally:w.close();QT.processEvents()

def test_actual_main_site_menu_and_ribbon_subprocess(tmp_path):
    import subprocess,sys,os
    code='''
import sys
from pathlib import Path
root=Path(sys.argv[1]);Path.home=classmethod(lambda cls:root)
from core.platform.application import StructuralProApp
service=StructuralProApp(root/'.structuralpro');service.create_project('آزمایشی','P');service.store.close()
from PySide6 import QtWidgets
from PySide6.QtCore import QTimer,Qt,QEvent,QCoreApplication
from PySide6.QtTest import QTest
from app.site_operations_workspace import SiteOperationsWorkspace
seen=[];errors=[]
def launch(self):seen.append(self.pid);self.close();return 0
SiteOperationsWorkspace.exec=launch
Base=QtWidgets.QApplication
class Harness(Base):
 def __init__(self,*args):super().__init__(*args);QTimer.singleShot(50,self.exercise)
 def exercise(self):
  try:
   window=next(w for w in self.topLevelWidgets() if isinstance(w,QtWidgets.QMainWindow))
   pages=window.findChild(QtWidgets.QStackedWidget);pages.setCurrentIndex(1)
   pages.widget(1).findChildren(QtWidgets.QLineEdit)[0].setText('P')
   button=window.findChild(QtWidgets.QPushButton,'SiteOperationsOpen');assert button
   QTest.mouseClick(button,Qt.LeftButton);assert seen==['P']
   ribbon=next(b for b in window.findChildren(QtWidgets.QPushButton) if b.text()=='منابع و دفتر کارگاه')
   QTest.mouseClick(ribbon,Qt.LeftButton);assert seen==['P','P']
   window.close();QCoreApplication.sendPostedEvents(None,QEvent.DeferredDelete)
  except Exception as e:errors.append(repr(e))
  finally:self.quit()
QtWidgets.QApplication=Harness
from app.main import main
assert main()==0
assert not errors,errors
'''
    result=subprocess.run([sys.executable,'-c',code,str(tmp_path)],env={**os.environ,'QT_QPA_PLATFORM':'offscreen'},capture_output=True,text=True,timeout=60)
    assert result.returncode==0,result.stdout[-1000:]+result.stderr[-2000:]
