"""P7 P2: real XLSX/DOCX/PDF/CSV + SVG/PNG with SQLite-sourced BOQ evidence."""
from __future__ import annotations

import csv
from copy import deepcopy
from hashlib import sha256
import json

import pytest

from test_phase7_p2_persisted_boq_crosswalk import _persisted
from core.reports.source_backed_2d_office_bundle import export_verified_2d_office_bundle


def _export(app, project, opts, out):
    return export_verified_2d_office_bundle(
        app, project_id="P", destination=out, **opts,
    )


def test_office_bundle_uses_real_persisted_boq_and_contains_real_images(tmp_path):
    import pymupdf
    from docx import Document
    from openpyxl import load_workbook

    app, project, opts, item = _persisted(tmp_path)
    output = tmp_path / "bounded_office_bundle"
    before = deepcopy(project)
    digest = app.store.current_digest("P")
    try:
        result = _export(app, project, opts, output)
        assert output.is_dir()
        expected = {
            "report.csv", "report.xlsx", "report.docx", "report.pdf",
            "overlay.svg", "overlay.png", "manifest.json",
        }
        assert {p.name for p in output.iterdir()} == expected
        manifest = json.loads((output / "manifest.json").read_text(encoding="utf-8"))
        assert manifest == result["manifest"]
        assert manifest["schema"] == "structuralpro.phase7.source-backed-2d-office-bundle.v1"
        assert manifest["overlay"]["links"][0]["boq_row_id"] == "boq-item-no:1"
        assert manifest["overlay"]["links"][0]["takeoff_id"] == item.id
        assert manifest["snapshot_project_version"] == project["_meta"]["version"]
        assert manifest["overlay"]["source_sha256"] == project["drawing_sessions"][opts["session_id"]]["source_sha256"]
        assert manifest["quantity_authority"] is False
        assert manifest["financial_mutation"] is False
        for file in manifest["files"]:
            actual = output / file["name"]
            assert file["sha256"] == sha256(actual.read_bytes()).hexdigest()
            assert file["size"] == actual.stat().st_size > 0
        assert manifest["overlay"]["svg_sha256"] == sha256((output / "overlay.svg").read_bytes()).hexdigest()
        assert (output / "overlay.png").read_bytes().startswith(b"\x89PNG\r\n\x1a\n")

        with pymupdf.open(output / "report.pdf") as pdf:
            assert pdf.page_count > 1
            assert len(pdf[-1].get_images()) >= 1
            assert pdf[-1].rect.width > pdf[-1].rect.height
        doc = Document(output / "report.docx")
        assert len(doc.inline_shapes) >= 1
        assert "نقشه دوبعدی" in "\n".join(p.text for p in doc.paragraphs)
        book = load_workbook(output / "report.xlsx")
        assert "نقشه مستند" in book.sheetnames
        sheet = book["نقشه مستند"]
        assert sheet.sheet_view.rightToLeft
        assert len(sheet._images) == 1
        assert sheet["C4"].value == "boq-item-no:1"
        assert sheet["D4"].value == project["takeoffs"][0]["source_id"]
        assert sheet["E4"].value == str(project["boq"][0]["quantity"])
        assert all(s.sheet_view.rightToLeft for s in book)
        with (output / "report.csv").open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream))
        assert any(project["takeoffs"][0]["source_id"] in row for row in rows)
        assert any("16" in cell for row in rows for cell in row)
        assert app.store.current_digest("P") == digest
        assert project == before == app.open_project("P")
    finally:
        app.store.close()


def test_missing_commit_and_stale_versions_produce_no_files(tmp_path):
    app, project, opts, _ = _persisted(tmp_path, commit=False)
    output = tmp_path / "not_valid"
    try:
        with pytest.raises(ValueError, match="ثبت"):
            _export(app, project, opts, output)
        assert not output.exists()
        for variant in (
            {"expected_project_version": opts["expected_project_version"] + 1},
            {"expected_session_revision": opts["expected_session_revision"] + 1},
        ):
            with pytest.raises(ValueError, match="نسخه"):
                _export(app, project, opts | variant, output)
            assert not output.exists()
    finally:
        app.store.close()


def test_changed_source_bytes_and_existing_destination_fail_closed(tmp_path):
    app, project, opts, _ = _persisted(tmp_path)
    output = tmp_path / "report_set"
    try:
        output.mkdir()
        with pytest.raises(FileExistsError):
            _export(app, project, opts, output)
        output.rmdir()
        source = tmp_path / "plan.pdf"
        source.write_bytes(source.read_bytes() + b"changed after measured")
        with pytest.raises(ValueError, match="SHA-256"):
            _export(app, project, opts, output)
        assert not output.exists()
    finally:
        app.store.close()


def test_exporter_error_discards_staging_without_touching_project(tmp_path, monkeypatch):
    import core.reports.source_backed_2d_office_bundle as bundle

    app, project, opts, _ = _persisted(tmp_path)
    target = tmp_path / "discard_on_error"
    digest = app.store.current_digest("P")
    def fail(*args, **kwargs):
        raise RuntimeError("simulated render failure")
    monkeypatch.setattr(bundle, "_append_xlsx_image", fail)
    try:
        with pytest.raises(RuntimeError, match="simulated"):
            _export(app, project, opts, target)
        assert not target.exists()
        assert not list(tmp_path.glob("structuralpro_p7_p2_*"))
        assert app.store.current_digest("P") == digest
    finally:
        app.store.close()


def test_concurrent_project_digest_change_prevents_publication(tmp_path, monkeypatch):
    app, project, opts, _ = _persisted(tmp_path)
    target = tmp_path / "discard_on_change"
    real_digest = app.store.current_digest
    calls = [0]
    def changing(project_id):
        calls[0] += 1
        if calls[0] >= 3:
            return "changed_while_exporting"
        return real_digest(project_id)
    monkeypatch.setattr(app.store, "current_digest", changing)
    try:
        with pytest.raises(RuntimeError, match="پروژه"):
            _export(app, project, opts, target)
        assert calls[0] >= 3
        assert not target.exists()
    finally:
        app.store.close()
