"""Explicit synthetic wall layouts and independently measured finish faces."""
from copy import deepcopy
from decimal import Decimal
import pytest
from core.takeoff.reviewed_aec import empty_system,preview_system
from core.takeoff.engine import TakeoffEngine
KINDS=('block_wall','brick_wall','plaster_finish')
def obj(i,**kw):return dict(id=i,physical_key='physical-'+i,source_ref='دیتایل مصنوعی '+i,approved_by='کنترل مصنوعی',**kw)
def state(status='confirmed'):return dict(status=status,approved_by='کنترل مصنوعی',source_ref='مرجع مصنوعی',reason='عدم کاربرد صریح مصنوعی' if status=='not_applicable' else '')
def sample(kind='block_wall'):
 d=empty_system(kind,'L1','W1');d.update(measurement_policy='architecture_specialist_v1',building_id='B1',axis='اتاق مصنوعی',source=dict(mode='manual',source_id='SYNTHETIC-'+kind,plan_ref='پلان مصنوعی',detail_ref='دیتایل مصنوعی'))
 g=dict(physical_wall_id='WALL-1',unit='m',length_m='3.9',height_m='2.9',wall_thickness_m='.2',count='1',repeat_ref='یک دیوار مستقل مصنوعی',source_ref='پلان مصنوعی',approved_by='کنترل مصنوعی')
 layer=obj('ML1',material_kind='brick' if kind=='brick_wall' else 'cement_block',offset_m='0',thickness_m='.2',unit_length_m='.9',unit_height_m='.9',unit_depth_m='.2',head_joint_m='.1',bed_joint_m='.1',unit_spec_ref='ابعاد واقعی مصنوعی',mortar_ref='ملات در بند کامل بیرون پوش واحد بنایی؛ مصنوعی',payment_unit='m2',payment_ref='پرداخت دیوار کامل؛ واحد و ملات داخل همین عملیات')
 regions=[obj('G1',layer_id='ML1',mode='units',x='0',y='0',width_m='3.9',height_m='2.9',nx='4',ny='3')]
 def face(i,side,code,t):return obj(i,surface_id='SURFACE-'+side,face=side,layer_code=code,x='0',y='0',width_m='3.9',height_m='2.9',thickness_m=t,material_ref='مشخصات ماده مصنوعی '+code,payment_unit='m2',payment_ref='مبنای سطح اجرا',density_kg_m3='',density_ref='',opening_ids='',deduction_ref='کسر تمام تقاطع‌های بازشو با این سطح؛ مصنوعی',coverage_scope='complete',scope_ref='تمام وجه مستند')
 surfaces=[face('IN-BASE','inside','base_plaster','.02'),face('IN-FINISH','inside','white_plaster','.005'),face('OUT','outside','cement_render','.015')]
 if kind=='plaster_finish':surfaces=surfaces[:1]
 d['architecture_detail']=dict(geometry=g,openings=[],masonry=state('not_applicable' if kind=='plaster_finish' else 'confirmed'),masonry_layers=[] if kind=='plaster_finish' else [layer],unit_regions=[] if kind=='plaster_finish' else regions,finishes=state(),surfaces=surfaces,accessories=state('not_applicable'))
 return d

def totals(d):
 r=preview_system(d,TakeoffEngine());assert r['ready'],r['blockers'];t={}
 for row in r['rows']:t[row['item_code']]=t.get(row['item_code'],Decimal(0))+Decimal(str(row['quantity']))
 return t,r

@pytest.mark.parametrize('kind',KINDS)
def test_independent_reference(kind):
 t,r=totals(sample(kind))
 if kind!='plaster_finish':assert t['masonry_work']==Decimal('11.31') and t['masonry_units']==12 and t['mortar']==Decimal('.318')
 assert t['base_plaster']==Decimal('11.31')
 if kind!='plaster_finish':assert t['white_plaster']==Decimal('11.31') and t['cement_render']==Decimal('11.31')
 assert len({row['source_id'] for row in r['rows']})==len(r['rows'])

@pytest.mark.parametrize('material',['cement_block','clay_block','light_block','aac','brick'])
def test_materials_only_use_entered_unit_geometry(material):
 d=sample('brick_wall' if material=='brick' else 'block_wall');d['architecture_detail']['masonry_layers'][0]['material_kind']=material
 assert totals(d)[0]['masonry_units']==12


def with_openings():
 d=sample();a=d['architecture_detail'];a['unit_regions']=[]
 for y in range(3):
  for x in range(4):
   if (x,y) in ((0,0),(2,1)):continue
   a['unit_regions'].append(obj(f'U{x}-{y}',layer_id='ML1',mode='units',x=str(x),y=str(y),width_m='.9',height_m='.9',nx='1',ny='1'))
 for x in ('.9','1.9','2.9'):a['unit_regions'].append(obj('V'+x,layer_id='ML1',mode='mortar',x=x,y='0',width_m='.1',height_m='2.9',nx='',ny=''))
 for x in range(4):
  for y in ('.9','1.9'):a['unit_regions'].append(obj(f'H{x}-{y}',layer_id='ML1',mode='mortar',x=str(x),y=y,width_m='.9',height_m='.1',nx='',ny=''))
 a['openings']=[obj('O1',x='0',y='0',width_m='.9',height_m='.9',penetration='full_wall'),obj('O2',x='2',y='1',width_m='.9',height_m='.9',penetration='full_wall')]
 for s in a['surfaces']:s['opening_ids']='O1|O2'
 return d


def test_multiple_openings_have_no_hidden_unit_or_mortar_ratio():
 t,r=totals(with_openings());assert t['masonry_work']==Decimal('9.69') and t['masonry_units']==10 and t['mortar']==Decimal('.318')
 assert t['base_plaster']==t['white_plaster']==t['cement_render']==Decimal('9.69')
 assert all(row['role']=='material' for row in r['rows'] if row['item_code'] in ('mortar','masonry_units'))

@pytest.mark.parametrize('fault',['missing_unit','unit','zero','negative','joint','grid','unpartitioned','duplicate','overlap','thickness','unknown','missing_source','wrong_policy','accessories','partial_surface','forgot_opening','duplicate_opening','missing_density'])
def test_fail_closed(fault):
 d=with_openings() if fault in ('forgot_opening','duplicate_opening') else sample();a=d['architecture_detail'];p=a['masonry_layers'][0];s=a['surfaces'][0]
 if fault=='missing_unit':p['unit_length_m']=''
 if fault=='unit':a['geometry']['unit']='cm'
 if fault=='zero':a['geometry']['height_m']='0'
 if fault=='negative':s['thickness_m']='-1'
 if fault=='joint':p['head_joint_m']=''
 if fault=='grid':a['unit_regions'][0]['nx']='5'
 if fault=='unpartitioned':a['unit_regions'][0]['width_m']='3.8'
 if fault=='duplicate':s['physical_key']=a['unit_regions'][0]['physical_key']
 if fault=='overlap':a['surfaces'].append({**s,'id':'OTHER','physical_key':'other'})
 if fault=='thickness':p['thickness_m']='.1'
 if fault=='unknown':p['assumed_waste']='.05'
 if fault=='missing_source':p['mortar_ref']=''
 if fault=='wrong_policy':d.pop('measurement_policy')
 if fault=='accessories':a['accessories']['status']='confirmed'
 if fault=='partial_surface':s['width_m']='2'
 if fault=='forgot_opening':s['opening_ids']='O1'
 if fault=='duplicate_opening':a['openings'].append({**a['openings'][0],'id':'O3','physical_key':'other'})
 if fault=='missing_density':s['payment_unit']='kg'
 assert not preview_system(d,TakeoffEngine())['ready']


def test_two_faces_are_not_automatically_equal_and_thickness_changes_volume():
 d=sample();a=d['architecture_detail'];a['surfaces']=[a['surfaces'][0],a['surfaces'][2]];a['surfaces'][0]['payment_unit']='m3';a['surfaces'][1]['payment_unit']='m3'
 t,_=totals(d);assert t['base_plaster']==Decimal('.2262') and t['cement_render']==Decimal('.16965') and 'white_plaster' not in t

@pytest.mark.parametrize('kind',KINDS)
@pytest.mark.parametrize('case',['cas','source','physical','legacy','policy','direct_legacy'])
def test_storage_guards(tmp_path,kind,case):
 from core.platform.application import StructuralProApp
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');d=sample(kind)
 if case in ('legacy','direct_legacy'):
  old=empty_system(kind,'L1','W1');old['source']={'mode':'manual','source_id':'LEGACY'};a.save_aec_system_draft('P',old,expected_revision=0)
  with pytest.raises(ValueError):
   if case=='legacy':a.save_aec_system_draft('P',d,expected_revision=1)
   else:
    p=a.open_project('P');next(iter(p['reviewed_aec_systems'].values()))['data']=d;a.store.save('P',p)
 else:
  a.save_aec_system_draft('P',d,expected_revision=0)
  if case=='cas':
   with pytest.raises(RuntimeError):a.save_aec_system_draft('P',d,expected_revision=0)
  elif case=='policy':
   d.pop('measurement_policy')
   with pytest.raises(ValueError):a.save_aec_system_draft('P',d,expected_revision=1)
  else:
   d['system_id']='W2'
   if case=='physical':d['source']['source_id']='OTHER'
   with pytest.raises(ValueError):a.save_aec_system_draft('P',d,expected_revision=0)
 a.store.close()

@pytest.mark.parametrize('kind',KINDS)
@pytest.mark.parametrize('fault',['identity','duplicate','evidence','finance','same_reviewer','repeat_apply','valid'])
def test_amendment_and_sqlite(tmp_path,kind,fault):
 from core.platform.application import StructuralProApp
 from hashlib import sha256
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');d=sample(kind);d['architecture_detail']['surfaces'][0]['payment_unit']='m3'
 a.save_aec_system_draft('P',d,expected_revision=0)
 with pytest.raises(ValueError):a.commit_aec_system('P',kind,'L1','W1',expected_revision=1)
 a.review_roof_system('P',kind,'L1','W1',expected_revision=1,reviewer='مستقل',note='کنترل');a.commit_aec_system('P',kind,'L1','W1',expected_revision=2)
 before=deepcopy(a.open_project('P'));assert not any(r['item_code'] in ('masonry_units','mortar') for r in before['boq'])
 illegal=deepcopy(before);next(iter(illegal['reviewed_aec_systems'].values()))['data']['axis']='changed'
 with pytest.raises(ValueError):a.store.save('P',illegal)
 d=deepcopy(a.get_aec_system('P',kind,'L1','W1')['data']);d.pop('roof_review');d['architecture_detail']['surfaces'][0]['thickness_m']='.03'
 path=tmp_path/'new.txt';path.write_text('SYNTHETIC architecture detail');kw=dict(expected_revision=2,created_by='سازنده',reason='علت مصنوعی',evidence={'evidence_path':str(path),'evidence_sha256':sha256(path.read_bytes()).hexdigest(),'evidence_reference':'مرجع مصنوعی'})
 if fault=='identity':
  d['axis']='changed'
  with pytest.raises(ValueError):a.create_roof_amendment('P',kind,'L1','W1',d,**kw)
 else:
  am=a.create_roof_amendment('P',kind,'L1','W1',d,**kw)
  assert Decimal(str(next(v for k,v in am['delta'].items() if k.endswith(':IN-BASE'))))==Decimal('.1131')
  if fault=='duplicate':
   with pytest.raises(ValueError):a.create_roof_amendment('P',kind,'L1','W1',d,**kw)
  elif fault=='evidence':
   path.write_text('changed')
   with pytest.raises(ValueError):a.review_beam_amendment('P',am['id'],expected_revision=1,reviewer='دیگری',note='کنترل')
  elif fault=='same_reviewer':
   with pytest.raises(ValueError):a.review_beam_amendment('P',am['id'],expected_revision=1,reviewer='سازنده',note='کنترل')
  else:
   a.review_beam_amendment('P',am['id'],expected_revision=1,reviewer='دیگری',note='کنترل')
   if fault=='finance':
    p=a.open_project('P');p['contract']={'id':'locked'};a.store.save('P',p)
    with pytest.raises(ValueError):a.apply_beam_amendment('P',am['id'],expected_revision=1)
   else:
    a.apply_beam_amendment('P',am['id'],expected_revision=1)
    if fault=='repeat_apply':
     with pytest.raises((ValueError,RuntimeError)):a.apply_beam_amendment('P',am['id'],expected_revision=1)
    else:
     assert a.open_project('P')['takeoffs']==before['takeoffs'];assert a.open_project('P')['boq']!=before['boq']
     a.store.close();a=StructuralProApp(tmp_path/'db');assert a.get_aec_system('P',kind,'L1','W1')['revision']==3
     assert a.store.db.execute('pragma integrity_check').fetchone()[0]=='ok'
 a.store.close()


def test_suggestions_never_create_quantities_or_boq():
 from core.takeoff.architecture_detail import suggestions
 for context in ('wall','room'):
  rows=suggestions(context);assert len(rows)>=3;assert all(set(r)=={'kind','title','status'} and 'نیازمند' in r['status'] for r in rows)


def test_multiple_layers_and_long_decimal():
 d=sample();a=d['architecture_detail'];p=a['masonry_layers'][0];p['thickness_m']=p['unit_depth_m']='.1';p2=deepcopy(p);p2.update(id='ML2',physical_key='physical-ML2',offset_m='.1');a['masonry_layers'].append(p2)
 q=deepcopy(a['unit_regions'][0]);q.update(id='G2',physical_key='physical-G2',layer_id='ML2');a['unit_regions'].append(q)
 a['surfaces'][0]['payment_unit']='m3';a['surfaces'][0]['thickness_m']='.020000000000000000001'
 t,_=totals(d);assert t['masonry_units']==24 and t['mortar']==Decimal('.318') and t['base_plaster']==Decimal('.22620000000000000001131')

@pytest.mark.parametrize('kind',KINDS)
def test_four_shared_reports_reopen_with_exact_values(tmp_path,kind):
 from core.platform.application import StructuralProApp
 from core.reports.architecture_detail_report import export
 from openpyxl import load_workbook
 from docx import Document
 import fitz,csv
 a=StructuralProApp(tmp_path/'db');a.create_project('نمونه مصنوعی صریح','P');d=sample(kind);d['architecture_detail']['surfaces'][0].update(payment_unit='m3',thickness_m='.020000000000000000001')
 if kind=='plaster_finish':
  extra=deepcopy(d['architecture_detail']['surfaces'][0]);extra.update(id='WHITE',physical_key='physical-WHITE',layer_code='white_plaster',payment_unit='m2',thickness_m='.005');d['architecture_detail']['surfaces'].append(extra)
 a.save_aec_system_draft('P',d,expected_revision=0);a.review_roof_system('P',kind,'L1','W1',expected_revision=1,reviewer='مستقل',note='کنترل');a.commit_aec_system('P',kind,'L1','W1',expected_revision=2)
 exact='.22620000000000000001131'
 for fmt in ('pdf','xlsx','docx','csv'):export(a,'P',kind,'L1','W1',tmp_path/('sample.'+fmt),format=fmt)
 book=load_workbook(tmp_path/'sample.xlsx');assert all(s.sheet_view.rightToLeft for s in book)
 cells=[c for s in book for row in s for c in row];assert any(str(c.value)=='0'+exact and c.data_type=='s' for c in cells);assert any(c.data_type=='n' and c.value==11.31 for c in cells)
 book.close()
 doc=Document(tmp_path/'sample.docx');assert all(t._tbl.tblPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bidiVisual') is not None for t in doc.tables)
 text=' '.join(c.text for t in doc.tables for r in t.rows for c in r.cells);assert '0'+exact in text and 'دیوار' in text
 pdf=fitz.open(tmp_path/'sample.pdf');assert len(pdf)>0
 for page in pdf:
  assert page.get_text()
  for word in page.get_text('words'):assert word[0]>=-1 and word[1]>=-1 and word[2]<=page.rect.width+1 and word[3]<=page.rect.height+1
 pdf.close()
 with open(tmp_path/'sample.csv',encoding='utf-8-sig',newline='') as f:rows=list(csv.reader(f))
 assert any('0'+exact in cell for row in rows for cell in row)
 assert any('مرجع دیتایل' in cell for row in rows for cell in row)
 a.store.close()

@pytest.mark.parametrize('case',['overlap','separate_id','other_floor','masonry_duplicate'])
def test_cross_record_physical_or_finish_guard(tmp_path,case):
 from core.platform.application import StructuralProApp
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');d=sample();a.save_aec_system_draft('P',d,expected_revision=0)
 other=sample('plaster_finish');other['system_id']='W2';other['source']['source_id']='OTHER'
 for p in other['architecture_detail']['surfaces']:p.update(id='SECOND',physical_key='physical-second')
 if case=='separate_id':other['architecture_detail']['surfaces'][0]['surface_id']='FAKE-SECOND-SURFACE'
 if case=='other_floor':other['floor_id']='L2'
 if case=='masonry_duplicate':
  other=sample('brick_wall');other['system_id']='W2';other['source']['source_id']='OTHER'
 with pytest.raises(ValueError):a.save_aec_system_draft('P',other,expected_revision=0)
 a.store.close()


def test_partial_and_mass_are_explicit_not_assumed():
 d=sample('plaster_finish');p=d['architecture_detail']['surfaces'][0];p.update(width_m='2',coverage_scope='partial',scope_ref='فقط نوار مستقل دو متری',payment_unit='kg',density_kg_m3='1000',density_ref='گواهی کاملاً مصنوعی')
 assert totals(d)[0]['base_plaster']==116


def test_copy_remaps_opening_foreign_keys_and_resets_approval(tmp_path):
 from core.platform.application import StructuralProApp
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');d=with_openings();a.save_aec_system_draft('P',d,expected_revision=0)
 r=a.aec_workflow().copy_roof('P','block_wall','L1','W1',new_system_id='W2',new_source_id='COPY',reason='کپی مصنوعی');detail=r['data']['architecture_detail'];ids={p['id'] for p in detail['openings']}
 assert set(detail['surfaces'][0]['opening_ids'].split('|'))==ids
 assert all(p['layer_id']==detail['masonry_layers'][0]['id'] for p in detail['unit_regions'])
 assert not preview_system(r['data'],a.takeoff)['ready'];assert detail['masonry']['status']=='unreviewed'
 a.store.close()


def test_five_floor_package_one_scenario(tmp_path):
 from scripts.verify_architecture_step6 import create_sample
 r=create_sample(tmp_path/'five-floors');assert r['floors']==5 and r['database_integrity']=='ok' and r['other_four_floors_unchanged'] and r['originals_unchanged'];assert r['effective_totals']['base_plaster']=='1.2117' and set(r['blocked'])=={'paint','floor layers','ceramic cut layouts','waterproofing','suspended ceiling','facade','wall reinforcement and connectors'}

@pytest.mark.parametrize('fault',['unit_change','physical_change','code_change'])
def test_amendment_preserves_physical_identity_and_dimensional_delta(tmp_path,fault):
 from core.platform.application import StructuralProApp
 from hashlib import sha256
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');d=sample();a.save_aec_system_draft('P',d,expected_revision=0);a.review_roof_system('P','block_wall','L1','W1',expected_revision=1,reviewer='مستقل',note='کنترل');a.commit_aec_system('P','block_wall','L1','W1',expected_revision=2)
 d=deepcopy(a.get_aec_system('P','block_wall','L1','W1')['data']);d.pop('roof_review')
 if fault=='unit_change':d['architecture_detail']['surfaces'][0]['payment_unit']='m3'
 if fault=='physical_change':d['architecture_detail']['geometry']['physical_wall_id']='DIFFERENT-WALL'
 if fault=='code_change':d['architecture_detail']['surfaces'][0]['layer_code']='levelling'
 e=tmp_path/'e.txt';e.write_text('SYNTHETIC')
 with pytest.raises(ValueError):a.create_roof_amendment('P','block_wall','L1','W1',d,expected_revision=2,created_by='سازنده',reason='اصلاح مصنوعی',evidence=dict(evidence_path=str(e),evidence_sha256=sha256(e.read_bytes()).hexdigest(),evidence_reference='مصنوعی'))
 a.store.close()


def test_incomplete_draft_can_copy_without_inventing_missing_detail(tmp_path):
 from core.platform.application import StructuralProApp
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');d=sample();detail=d['architecture_detail'];detail['unit_regions']=[{'id':'UNFINISHED','physical_key':'unfinished','layer_id':''}];detail['surfaces']=[{'id':'UNFINISHED-SURFACE','physical_key':'unfinished-surface'}]
 a.save_aec_system_draft('P',d,expected_revision=0)
 r=a.aec_workflow().copy_roof('P','block_wall','L1','W1',new_system_id='W2',new_source_id='COPY',reason='کپی پیش‌نویس ناقص مصنوعی');assert not preview_system(r['data'],a.takeoff)['ready'];assert not r['data']['architecture_detail']['unit_regions'][0]['layer_id'];a.store.close()

@pytest.mark.parametrize('code,t,reference',[('base_plaster','.02','.2262'),('white_plaster','.005','.05655'),('cement_render','.015','.16965'),('levelling','.01','.1131'),('protective','.003','.03393')])
def test_independent_layer_volume_reference(code,t,reference):
 d=sample('plaster_finish');p=d['architecture_detail']['surfaces'][0];p.update(layer_code=code,thickness_m=t,payment_unit='m3',material_ref='مشخصات مستقل و کاملاً مصنوعی');assert totals(d)[0][code]==Decimal(reference)


def test_two_sqlite_connections_cannot_overwrite_a_newer_draft(tmp_path):
 from core.platform.application import StructuralProApp
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');d=sample();a.save_aec_system_draft('P',d,expected_revision=0);b=StructuralProApp(tmp_path/'db');stale=b.get_aec_system('P','block_wall','L1','W1');updated=deepcopy(d);updated['axis']='محل جدید مستند';a.save_aec_system_draft('P',updated,expected_revision=1)
 with pytest.raises(RuntimeError):b.save_aec_system_draft('P',stale['data'],expected_revision=1)
 assert b.get_aec_system('P','block_wall','L1','W1')['data']['axis']=='محل جدید مستند';a.store.close();b.store.close()
