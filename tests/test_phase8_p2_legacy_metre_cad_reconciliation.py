"""P8P2: read-only legacy metre workbook ↔ source-backed CAD gating.

All test fixtures are synthetic and contain no user-uploaded drawing/workbook.
Actual uploaded binary remains local/private and is not published in CI.
"""
from copy import deepcopy
from hashlib import sha256

import ezdxf
from openpyxl import Workbook, load_workbook
import pytest

from core.drawings.legacy_metre_cad_reconciliation import LegacyMetreCADReconciler
from core.platform.application import StructuralProApp


def legacy_book(tmp_path):
    p = tmp_path / "old metre.xlsx"
    wb = Workbook()
    a = wb.active
    a.title = "\u200cS_6_1"
    a["A1"] = "فرم متره"
    a["A2"] = "فصل ششم"
    a.append(["رديف", "شرح عمليات", "تعداد", "طول", "عرض", "ضخامت", "طول جزئی", "سطح", "حجم", "وزن", "شماره قيمت", "قیمت", "قیمت کل"])
    a["B5"] = "طول یک خط ساده"
    a["I5"] = 2
    a["L5"] = 10
    a["M5"] = 20
    a["B6"] = "قالب بندی طرف پی (مشکوک)"
    a["I6"] = 7.2
    a["L6"] = 42900
    a["M6"] = 1
    b = wb.create_sheet("S_7_1")
    b["A1"] = "فرم متره"
    b["B3"] = "شرح عمليات"
    b["L3"] = "قیمت"
    b["M3"] = "قیمت کل"
    b["B5"] = "ميلگرد"
    b["I5"] = 8
    b["L5"] = 100
    b["M5"] = 800
    wb.save(p)
    wb.close()
    return p


def summary_book(tmp_path):
    p = tmp_path / "legacy summary.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "خلاصه صورت وضعيت"
    ws["B2"] = "شماره فصل فهرست بها"
    ws["D8"] = "=[1]S_6_1!$M$9"
    ws["D9"] = 30816253.97
    ws["G9"] = 32699311.98
    wb.save(p)
    wb.close()
    return p


def real_dxf(tmp_path, *, length=2):
    path = tmp_path / "source.dxf"
    drawing = ezdxf.new("R2013")
    drawing.header["$INSUNITS"] = 6
    handle = drawing.modelspace().add_line((0, 0), (length, 0)).dxf.handle
    drawing.saveas(path)
    return path, handle


def test_legacy_workbook_rows_are_not_imported_as_pricebook_or_boq(tmp_path):
    app = StructuralProApp(tmp_path / "app")
    try:
        app.create_project("No invented quantities", "LEGACY-READONLY")
        initial = app.store.current_digest("LEGACY-READONLY")
        source = legacy_book(tmp_path)
        report = app.audit_legacy_metre_against_drawing(source)
        assert report["priced_rows"] == 3
        assert report["metre_sheets"] == 2
        assert len(report["rows"]) == 3
        assert report["cad"]["status"] == "NOT_TESTED"
        assert all(x["cad_status"].startswith("NOT_VERIFIED") for x in report["rows"])
        assert any(x["code"] == "ROW_AMOUNT_DISAGREES_WITH_PRIMARY_QUANTITY_TIMES_RATE"
                   and x["key"] == "S_6_1!6" and x["difference"] == "-308879.0"
                   for x in report["issues"])
        assert report["chapter_totals"]["6"] == "21"
        assert app.store.current_digest("LEGACY-READONLY") == initial
        project = app.open_project("LEGACY-READONLY")
        assert project["boq"] == [] and project["takeoffs"] == []
    finally:
        app.store.close()


def test_broken_summary_external_link_and_hardcoded_chapter_detected(tmp_path):
    report = LegacyMetreCADReconciler().inspect(legacy_book(tmp_path), summary_path=summary_book(tmp_path))
    codes = {x["code"] for x in report["issues"]}
    assert "EXTERNAL_LINK_NOT_INDEPENDENTLY_VERIFIED" in codes
    assert "CHAPTER_SUMMARY_NOT_NUMERIC_OR_CACHE_INVALID" in codes
    assert "CHAPTER_SUMMARY_DIFFERS_FROM_METRE" in codes
    assert len(report["summary"]["external_references"]) == 1
    assert not report["summary"]["independent_formula_evaluation"]


def test_real_dxf_explicit_handle_only_can_produce_unapproved_geometry_candidate(tmp_path):
    source = legacy_book(tmp_path)
    drawing, handle = real_dxf(tmp_path)
    digest = sha256(drawing.read_bytes()).hexdigest()
    report = LegacyMetreCADReconciler().inspect(
        source, drawing_path=drawing,
        bindings=[{"key":"S_6_1!5", "handle": handle, "metric": "length",
                   "unit": "m", "quantity_cell": "I5",
                   "source_sha256": digest, "evidence_ref": "QA source / exact line handle"}],
    )
    assert report["cad"]["matched_candidates"] == 1
    assert report["cad"]["unbound_rows"] == 2
    assert report["cad"]["verified_rows"] == 0
    assert not report["rows"][0]["cad_evidence"]["independent_human_approved"]
    assert report["rows"][0]["cad_status"] == "CANDIDATE_MATCH_AWAITING_HUMAN_APPROVAL"
    assert report["boq_rows_committed"] == 0 and not report["prices_approved"]


def test_cad_mismatched_quantity_kept_explicit_and_cannot_be_approved(tmp_path):
    drawing, handle = real_dxf(tmp_path, length=3)
    report = LegacyMetreCADReconciler().inspect(
        legacy_book(tmp_path), drawing_path=drawing,
        bindings=[{"key":"S_6_1!5", "handle":handle, "metric":"length",
                   "unit":"m", "quantity_cell":"I5", "source_sha256":sha256(drawing.read_bytes()).hexdigest(),
                   "evidence_ref":"QA different line length"}],
    )
    assert report["cad"]["mismatched_rows"] == 1
    assert report["rows"][0]["cad_status"] == "CAD_QUANTITY_MISMATCH"
    assert any(x["code"] == "CAD_QUANTITY_DIFFERS_FROM_METRE" for x in report["issues"])


@pytest.mark.parametrize("mutation", ["wrong_hash", "wrong_unit", "not_supported_volume",
                                      "missing_ref", "wrong_cell", "unknown_row", "duplicate"])
def test_unsafe_legacy_to_cad_mappings_fail_closed(tmp_path, mutation):
    drawing, handle = real_dxf(tmp_path)
    source = legacy_book(tmp_path)
    link = {"key":"S_6_1!5", "handle":handle, "metric":"length", "unit":"m",
            "quantity_cell":"I5", "source_sha256":sha256(drawing.read_bytes()).hexdigest(),
            "evidence_ref":"QA precise source / handle proof"}
    if mutation == "wrong_hash":
        link["source_sha256"] = "0"*64
    elif mutation == "wrong_unit":
        link["unit"] = "mm"
    elif mutation == "not_supported_volume":
        link["metric"], link["unit"] = "volume", "m3"
    elif mutation == "missing_ref":
        link["evidence_ref"] = ""
    elif mutation == "wrong_cell":
        link["quantity_cell"] = "I6"
    elif mutation == "unknown_row":
        link["key"] = "S_6_1!99"
    links = [link, deepcopy(link)] if mutation == "duplicate" else [link]
    with pytest.raises(ValueError):
        LegacyMetreCADReconciler().inspect(source, drawing_path=drawing, bindings=links)


def test_dwg_without_native_converter_reports_blocked_not_pass(tmp_path):
    source = legacy_book(tmp_path)
    drawing = tmp_path / "source.dwg"
    drawing.write_bytes(b"AC1018"+"not-an-executable-cad-drawing".encode("ascii"))
    report = LegacyMetreCADReconciler().inspect(source, drawing_path=drawing)
    assert report["cad"]["status"].startswith("BLOCKED_")
    assert report["cad"]["verified_rows"] == 0
    assert all(row["cad_status"].startswith("NOT_VERIFIED") for row in report["rows"])
    assert report["boq_rows_committed"] == 0


def test_wrong_legacy_headers_rejected_without_price_or_boq_changes(tmp_path):
    file = legacy_book(tmp_path)
    wb = load_workbook(file)
    wb.worksheets[0]["B3"] = "unrelated workbook"
    wb.save(file)
    wb.close()
    with pytest.raises(ValueError, match="audited source headers"):
        LegacyMetreCADReconciler().inspect(file)
