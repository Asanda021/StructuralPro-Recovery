"""Regression for observed LibreOffice row-65 split in archived ARCH-5 export."""
from docx import Document
from docx.oxml.ns import qn
from core.reports.exporters import export_docx


def test_word_keeps_each_table_record_together_without_dropping_text(tmp_path):
 rows=[{'item_no':i,'description':'شرح فارسی طولانی و قابل بازبینی '*12,'quantity':'0.00000000000001','unit':'m3','source_id':f'synthetic:explicit:{i}'} for i in range(1,90)]
 out=export_docx(rows,tmp_path/'records.docx',columns=[('item_no','ردیف'),('description','شرح'),('quantity','مقدار'),('unit','واحد'),('source_id','منبع')])
 doc=Document(out)
 for table in doc.tables:
  for row in table.rows:
   assert row._tr.get_or_add_trPr().find(qn('w:cantSplit')) is not None,'A quantity record must not be separated from its identity across pages'
 text='\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
 for row in rows:assert row['description'] in text and row['source_id'] in text
 assert rows[0]['quantity'] in text
