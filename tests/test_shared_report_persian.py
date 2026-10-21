"""Shared exporter contracts; synthetic cross-discipline values, not real prices."""
from datetime import date
from decimal import Decimal
import csv,zipfile
import pytest
from openpyxl import load_workbook
from docx import Document
from docx.oxml.ns import qn
from core.reports.exporters import export_xlsx,export_docx,export_pdf,export_csv

ROWS=[{'code':'0010','description':'بتن پی و شناژ','quantity':Decimal('2.5'),'unit':'m3','unit_price':Decimal('100'),'total':Decimal('250'),'source_id':'CAD:A-01/بتن','notes':'مدرک اصلی تغییر نکند'}, {'code':'0002','description':'میلگرد تیر و ستون','quantity':Decimal('12345678901234567890.123456789'),'unit':'kg','source_id':'SHA:'+('a'*64)}, {'code':'0003','description':'=HYPERLINK("https://example.test")','quantity':Decimal('.25'),'unit':'m2','source_id':'C:\\نقشه\\A-01.dwg'}]
COLS=[('code','کد'),('description','شرح'),('quantity','مقدار'),('unit','واحد'),('unit_price','بهای واحد'),('total','مبلغ'),('source_id','منبع'),('notes','یادداشت')]

def header(sheet):return next(i for i,row in enumerate(sheet,1) if any(c.value=='شرح' for c in row))

def test_excel_numeric_precision_formula_and_print(tmp_path):
 p=export_xlsx(ROWS,tmp_path/'report.xlsx',title='گزارش مصنوعی',columns=COLS,metadata={'project_id':'P-01'})
 w=load_workbook(p);s=w['ریز متره'];h=header(s)
 assert s.cell(h+1,3).data_type=='n' and s.cell(h+1,3).value==2.5
 assert s.cell(h+2,3).value=='12345678901234567890.123456789' and s.cell(h+2,3).data_type=='s'
 assert s.cell(h+3,2).data_type=='s';assert s.cell(h+1,1).value=='0010'
 assert s.cell(h+1,6).data_type=='f','جمع سازگار باید فرمول زنده با ورودی عددی داشته باشد'
 assert all(v.sheet_view.rightToLeft and v.print_options.horizontalCentered for v in w)
 assert all(v.oddFooter.center.text and v.oddHeader.right.text for v in w)
 assert s.cell(h+1,2).font.name and s.print_title_rows
 assert 'شمسی' in ' '.join(str(c.value) for row in w['اطلاعات'] for c in row)


def test_word_rtl_font_bidi_mixed_and_project_info(tmp_path):
 p=export_docx(ROWS,tmp_path/'report.docx',title='گزارش مصنوعی',columns=COLS,metadata={'project_id':'P-01','source_sha256':'ab'*32})
 d=Document(p);xml=d._element.xml
 assert 'P-01' in xml and 'ab'*32 in xml
 assert 'مشخصات' in xml and 'منابع' in xml
 assert 'w:bidiVisual' in xml and 'w:bidi' in xml
 from core.reports.exporters import REPORT_FONT
 assert d.styles['Normal'].font.name==REPORT_FONT
 assert '<w:rtl w:val="0"' in xml,'شناسه فنی باید اجرای مستقل چپ‌به‌راست داشته باشد'
 with zipfile.ZipFile(p) as z:
  assert any('footer' in n for n in z.namelist())
  assert any(b'PAGE' in z.read(n) for n in z.namelist() if 'footer' in n and n.endswith('.xml'))


def test_shared_date_and_default_titles():
 from core.reports.exporters import report_date,_columns
 assert report_date(date(2026,10,15))=='۱۴۰۵/۰۷/۲۳'
 assert _columns([{'description':'الف','quantity':1,'source_id':'A'}])==[('description','شرح'),('quantity','مقدار'),('source_id','منبع متره')]


def test_numeric_text_detail_context_protects_identity(tmp_path):
 rows=[{'component':'شناژ','field':'وزن میلگرد','value':'32.232'}, {'component':'شناژ','field':'شناسه منبع وزن','value':'000123'}, {'component':'اصلاحیه','field':'هش مدرک','value':'1234567890'}]
 p=export_xlsx(rows,tmp_path/'detail.xlsx',columns=[('component','جزء'),('field','عنوان'),('value','مقدار')]);s=load_workbook(p)['ریز متره'];h=header(s) if any(c.value=='شرح' for row in s for c in row) else next(i for i,row in enumerate(s,1) if any(c.value=='عنوان' for c in row))
 assert s.cell(h+1,3).data_type=='n'
 assert s.cell(h+2,3).value=='000123' and s.cell(h+3,3).data_type=='s'


def test_csv_preserves_raw_data_and_discloses_layout_limit(tmp_path):
 p=export_csv(ROWS,tmp_path/'report.csv',title='گزارش مصنوعی',columns=COLS)
 with p.open(encoding='utf-8-sig',newline='') as f:records=list(csv.reader(f))
 assert any('0010' in r for r in records)
 assert any('CSV' in str(r) and 'چیدمان' in str(r) for r in records)


def test_all_untrusted_spreadsheet_text_is_not_a_formula(tmp_path):
 p=export_xlsx([{'x':'=2+2'}],tmp_path/'injection.xlsx',title='=1+1',metadata={'=3+3':'=4+4'},columns=[('x','=5+5')])
 w=load_workbook(p)
 assert not any(c.data_type=='f' for s in w for row in s for c in row)
 assert any(c.value=='=3+3' for s in w for row in s for c in row)


def test_tiny_number_does_not_print_as_zero(tmp_path):
 p=export_xlsx([{'quantity':Decimal('0.00000000000000000001')}],tmp_path/'tiny.xlsx',columns=[('quantity','مقدار')]);w=load_workbook(p);s=w['ریز متره'];cell=next(c for row in s for c in row if c.data_type=='n' and c.value is not None)
 assert cell.value==1e-20 and 'E' in cell.number_format.upper()


def test_extra_columns_stay_distinct_and_mapped(tmp_path):
 p=export_xlsx([{'waste_percent':'1','allowance_quantity':'2'},{'later_field':'منبع دیرتر'}],tmp_path/'extra.xlsx')
 w=load_workbook(p);s=w['ریز متره'];h=next(i for i,row in enumerate(s,1) if s.cell(i,1).value=='ریز متره و برآورد')+1
 assert len({s.cell(h,c).value for c in range(1,4)})==3
 info=' '.join(str(c.value) for row in w['اطلاعات'] for c in row)
 assert all(key in info for key in ['waste_percent','allowance_quantity','later_field'])


def test_literal_print_header_and_csv_formula_guard(tmp_path):
 title='پروژه A &P <B1>'
 p=export_xlsx([],tmp_path/'header.xlsx',title=title)
 assert load_workbook(p).active.oddHeader.right.text==title.replace('&','&&')
 p=export_csv([{'description':'=HYPERLINK("bad")','quantity':-2,'source_id':'@source'}],tmp_path/'safe.csv',title='=1+1')
 records=list(csv.reader(p.open(encoding='utf-8-sig')))
 assert records[0]==["'=1+1"]
 assert any("'=HYPERLINK(\"bad\")" in r and '-2' in r and "'@source" in r for r in records)


def test_pdf_canvas_title_is_not_xml_escaped():
 from core.reports.exporters import _pdf_page_footer
 class Canvas:
  def __init__(self):self.text=[]
  def saveState(self):pass
  def restoreState(self):pass
  def setFont(self,*a):pass
  def drawRightString(self,*a):self.text.append(a[-1])
  def drawString(self,*a):self.text.append(a[-1])
 class Doc:page=1;pagesize=(800,600);rightMargin=24;leftMargin=24;width=752
 c=Canvas();_pdf_page_footer(c,Doc(),'Helvetica','A & B <C1>')
 assert 'A & B <C1>' in c.text


@pytest.mark.parametrize('path',['/tmp/مدرک.txt',r'\\server\مدرک.txt','C:/مدرک.txt','CAD:A-01/بتن'])
def test_mixed_technical_paths_are_ltr(tmp_path,path):
 from core.reports.exporters import _technical_text
 assert _technical_text(path)
 p=export_docx([{'source_id':path}],tmp_path/'path.docx')
 d=Document(p)
 cell=next(c for t in d.tables for row in t.rows for c in row.cells if c.text==path)
 assert cell.paragraphs[0]._p.get_or_add_pPr().find(qn('w:bidi')).get(qn('w:val'))=='0'


@pytest.mark.parametrize('path',['/tmp/مدرک.txt',r'\\server\مدرک.txt','CAD:A-01/بتن'])
def test_excel_mixed_paths_read_ltr(tmp_path,path):
 p=export_xlsx([{'source_id':path}],tmp_path/'path.xlsx')
 w=load_workbook(p);c=next(c for row in w.active for c in row if c.value==path)
 assert c.alignment.readingOrder==1 and c.value==path


def test_excel_boq_ordinal_is_numeric_and_codes_keep_leading_zero(tmp_path):
 rows=[dict(ROWS[0],item_no=1),dict(ROWS[1],item_no=2)]
 p=export_xlsx(rows,tmp_path/'boq.xlsx',columns=[('item_no','ردیف'),('code','کد'),('description','شرح'),('quantity','مقدار')])
 s=load_workbook(p)['ریز متره'];h=header(s)
 assert s.cell(h+1,1).data_type=='n' and s.cell(h+1,1).value==1
 assert s.cell(h+2,1).data_type=='n' and s.cell(h+2,1).value==2
 assert s.cell(h+1,2).data_type=='s' and s.cell(h+1,2).value=='0010'
 assert s.cell(h+2,4).value=='12345678901234567890.123456789'
