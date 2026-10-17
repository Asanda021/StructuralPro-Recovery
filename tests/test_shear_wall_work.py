"""Explicit synthetic specialist workflow over existing SQLite and BOQ."""
from copy import deepcopy
from hashlib import sha256
import json
from pathlib import Path
import pytest
from core.platform.application import StructuralProApp
ROOT=Path(__file__).resolve().parents[1]

def sample(tmp_path,*,pid='P',floor='L1',sid='SW1',z='0'):
 from core.takeoff.shear_wall_detail import new_detail,bind_input_digest
 d=new_detail(floor,sid);d.update(building_id='B1',axis='A-1');e=json.loads((ROOT/'examples/shear_wall_chat_reference.json').read_text());e.update(project_id=pid,floor_id=floor,wall_id=sid,physical_wall_id='PH-'+sid)
 d['axis']=e['datum']['axis_ref']
 e['source']['source_id']='SRC-'+sid
 for seg in e['segments']:seg['z_m']=z;seg['physical_part_id']=sid+'-'+seg['id']
 for bar in [b for g in e['rebars'].values() for b in g.get('series',[])]:bar['physical_bar_set_id']=sid+'-'+bar['id']
 d['source']={'mode':'manual','source_id':e['source']['source_id'],'approved_by':'تهیه‌کننده','measurement_ref':'برداشت مصنوعی','plan_ref':'پلان مصنوعی','detail_ref':'دیتایل مصنوعی'}
 detail=d['shear_wall_detail'];detail['input']=e
 p=tmp_path/(sid+'-input.json');p.write_text(json.dumps(e,ensure_ascii=False));h=sha256(p.read_bytes()).hexdigest();detail['input_document']={'path':str(p),'sha256':h,'reference':'مدرک مصنوعی هندسه و آرماتور'}
 o={'schema':'shear_wall_independent_ownership_v1','project_id':pid,'physical_wall_id':e['physical_wall_id'],'coordinate_ref':e['datum']['coordinate_ref'],'input_digest':bind_input_digest(e),'input_document_sha256':h,'prepared_by':'تنظیم‌کننده','reviewer':'بازبین مالکیت مستقل','reference':'مالکیت مستقل صریحاً مصنوعی','concrete_scopes':{s['id']:'V-'+sid+'-'+s['id'] for s in e['segments']},'formwork_scopes':{f['id']:'F-'+sid+'-'+f['id'] for f in e['formwork']['main_faces']+e['formwork']['edge_faces']},'bar_pieces':{b['physical_bar_set_id']:[sid+'-'+b['id']+'-'+str(i) for i in range(int(b['count']))] for g in e['rebars'].values() for b in g.get('series',[])}}
 p=tmp_path/(sid+'-owner.json');p.write_text(json.dumps(o,ensure_ascii=False));detail['ownership_document']={'path':str(p),'sha256':sha256(p.read_bytes()).hexdigest(),'reference':'سند مصنوعی مالکیت'};return d

def refresh(d,tmp_path):
 from core.takeoff.shear_wall_detail import bind_input_digest
 e=d['shear_wall_detail']['input'];p=tmp_path/'changed-input.json';p.write_text(json.dumps(e,ensure_ascii=False));h=sha256(p.read_bytes()).hexdigest();old=json.loads(Path(d['shear_wall_detail']['ownership_document']['path']).read_text());old.update(input_digest=bind_input_digest(e),input_document_sha256=h)
 q=tmp_path/'changed-owner.json';q.write_text(json.dumps(old,ensure_ascii=False));d['shear_wall_detail']['input_document'].update(path=str(p),sha256=h);d['shear_wall_detail']['ownership_document'].update(path=str(q),sha256=sha256(q.read_bytes()).hexdigest())

def commit(a,d):
 a.save_aec_system_draft('P',d);r=a.review_roof_system('P','shear_wall',d['floor_id'],d['system_id'],expected_revision=1,reviewer='بازبین نهایی مستقل',note='کنترل مستقل مصنوعی');return a.commit_aec_system('P','shear_wall',d['floor_id'],d['system_id'],expected_revision=r['revision'])

def test_full_workflow_and_immutable_amendment(tmp_path):
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');d=sample(tmp_path);commit(a,d);before=deepcopy(a.open_project('P'));d['shear_wall_detail']['input']['segments'][0]['length_m']='6';refresh(d,tmp_path)
 doc=tmp_path/'reason.txt';doc.write_text('EXPLICIT SYNTHETIC revision')
 evidence={'evidence_path':str(doc),'evidence_sha256':sha256(doc.read_bytes()).hexdigest(),'evidence_reference':'اصلاح مصنوعی'}
 am=a.create_roof_amendment('P','shear_wall','L1','SW1',d,expected_revision=2,created_by='سازنده اصلاحیه',reason='تغییر نقشه مصنوعی',evidence=evidence)
 a.review_beam_amendment('P',am['id'],expected_revision=1,reviewer='بازبین اصلاحیه مستقل',note='کنترل مستقل');a.apply_beam_amendment('P',am['id'],expected_revision=1)
 p=a.open_project('P');assert p['takeoffs']==before['takeoffs'];assert p['boq']!=before['boq'];assert a.get_aec_system('P','shear_wall','L1','SW1')['revision']==3
 with pytest.raises(ValueError):a.apply_beam_amendment('P',am['id'],expected_revision=1)
 a.store.close();b=StructuralProApp(tmp_path/'db');assert b.get_aec_system('P','shear_wall','L1','SW1')['revision']==3;b.store.close()

def test_preview_reference_and_missing_evidence(tmp_path):
 from core.takeoff.shear_wall_detail import preview
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');d=sample(tmp_path);r=preview(d,a.takeoff);assert r['ready'];assert [v['quantity'] for v in r['rows'] if v['item_code']=='shear_wall_concrete']==['3.5'];Path(d['shear_wall_detail']['ownership_document']['path']).unlink();assert not preview(d,a.takeoff)['ready'];a.store.close()

@pytest.mark.parametrize('bad',['missing_input','fake_hash','stale_owner','duplicate_bar','short_bar_list','self_owner_review','float','bad_unit','duplicate_opening','interface','contact','negative_orientation'])
def test_document_and_engine_fail_closed(tmp_path,bad):
 from core.takeoff.shear_wall_detail import preview
 d=sample(tmp_path);detail=d['shear_wall_detail'];e=detail['input'];ownerpath=Path(detail['ownership_document']['path']);o=json.loads(ownerpath.read_text())
 if bad=='missing_input':Path(detail['input_document']['path']).unlink()
 elif bad=='fake_hash':detail['input_document']['sha256']='f'*64
 elif bad=='stale_owner':o['input_digest']='0'*64
 elif bad=='duplicate_bar':
  ids=list(o['bar_pieces']);o['bar_pieces'][ids[1]][0]=o['bar_pieces'][ids[0]][0]
 elif bad=='short_bar_list':next(iter(o['bar_pieces'].values())).pop()
 elif bad=='self_owner_review':o['reviewer']=o['prepared_by']
 else:
  if bad=='float':e['segments'][0]['length_m']=5.0
  if bad=='bad_unit':e['source']['length_unit']='cm'
  if bad=='duplicate_opening':e['openings'].append(deepcopy(e['openings'][0]))
  if bad=='interface':e['interfaces']=[{}]
  if bad=='contact':e['formwork']['main_faces'][0]['contact_rectangles']=[{}]
  if bad=='negative_orientation':e['datum']['orientation']='negative_x'
  refresh(d,tmp_path);detail=d['shear_wall_detail'];o=None
 if o is not None:
  ownerpath.write_text(json.dumps(o,ensure_ascii=False));detail['ownership_document']['sha256']=sha256(ownerpath.read_bytes()).hexdigest()
 assert not preview(d,None)['ready']

@pytest.mark.parametrize('peer',['beam','column','foundation','tie','generic_wall','roof','legacy_takeoff'])
def test_unknown_foreign_geometry_stays_blocked(tmp_path,peer):
 from core.takeoff.shear_wall_detail import project_gate
 from core.takeoff.reviewed_aec import _id
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');d=sample(tmp_path);a.save_aec_system_draft('P',d);p=a.open_project('P')
 if peer=='beam':p['beam_detail_register']={'legacy':{'data':{}}}
 elif peer in ('column','foundation','tie'):p['reviewed_structural_members']={'legacy':{'data':{'kind':peer}}}
 elif peer in ('generic_wall','roof'):p['reviewed_aec_systems']['legacy']={'data':{'kind':'shear_wall' if peer=='generic_wall' else 'solid_slab','floor_id':'L2','system_id':'LEGACY'}}
 else:p['takeoffs']=[{'source_id':'legacy','quantities':[{'code':'concrete','amount':'1','unit':'m3'}]}]
 with pytest.raises(ValueError):project_gate(p,_id('shear_wall','L1','SW1'))
 a.store.close()

def test_cas_copy_delete_self_review_and_financial_lock(tmp_path):
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');d=sample(tmp_path);a.save_aec_system_draft('P',d);b=StructuralProApp(tmp_path/'db');b.get_aec_system('P','shear_wall','L1','SW1');a.save_aec_system_draft('P',d,expected_revision=1)
 with pytest.raises(RuntimeError):b.save_aec_system_draft('P',d,expected_revision=1)
 with pytest.raises(ValueError):a.review_roof_system('P','shear_wall','L1','SW1',expected_revision=2,reviewer='تهیه‌کننده',note='خودبازبینی ممنوع')
 cp=a.aec_workflow().copy_roof('P','shear_wall','L1','SW1',new_system_id='COPY',new_source_id='SRC-COPY',reason='کپی مصنوعی')
 assert cp['data']['shear_wall_detail']['input']['physical_wall_id']==''
 assert not a.preview_aec_system('P','shear_wall','L1','COPY')['ready']
 a.aec_workflow().delete_roof('P','shear_wall','L1','COPY',expected_revision=1,reason='حذف کپی مصنوعی')
 a.review_roof_system('P','shear_wall','L1','SW1',expected_revision=2,reviewer='بازبین مستقل',note='کنترل')
 p=a.open_project('P');p['estimate']={'locked':True};a.store.save('P',p)
 with pytest.raises(ValueError):a.commit_aec_system('P','shear_wall','L1','SW1',expected_revision=3)
 assert not a.open_project('P')['boq'];b.store.close();a.store.close()

def test_disjoint_walls_and_duplicate_bar_ownership(tmp_path):
 from core.takeoff.shear_wall_detail import project_gate
 from core.takeoff.reviewed_aec import _id
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');one=sample(tmp_path);commit(a,one)
 two=sample(tmp_path,sid='SW2',floor='L2',z='4');two['shear_wall_detail']['input']['openings']=[];refresh(two,tmp_path);a.save_aec_system_draft('P',two);assert project_gate(a.open_project('P'),_id('shear_wall','L2','SW2'))
 o_path=Path(two['shear_wall_detail']['ownership_document']['path']);o=json.loads(o_path.read_text());old=json.loads(Path(one['shear_wall_detail']['ownership_document']['path']).read_text());next(iter(o['bar_pieces'].values()))[0]=next(iter(old['bar_pieces'].values()))[0];o_path.write_text(json.dumps(o));two['shear_wall_detail']['ownership_document']['sha256']=sha256(o_path.read_bytes()).hexdigest();a.save_aec_system_draft('P',two,expected_revision=1)
 with pytest.raises(ValueError):project_gate(a.open_project('P'),_id('shear_wall','L2','SW2'))
 a.store.close()

def test_direct_store_cannot_forge_commit_or_quantities(tmp_path):
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');d=sample(tmp_path);a.save_aec_system_draft('P',d);p=a.open_project('P');next(iter(p['reviewed_aec_systems'].values()))['status']='committed'
 with pytest.raises(ValueError):a.store.save('P',p)
 assert a.get_aec_system('P','shear_wall','L1','SW1')['status']=='draft'
 assert not a.open_project('P')['boq']
 a.store.close()

@pytest.mark.parametrize('bad',['zero','negative','blank','overlap','duplicate_physical','different_datum','same_source'])
def test_dimensions_and_cross_wall_identity(tmp_path,bad):
 from core.takeoff.shear_wall_detail import preview,project_gate
 from core.takeoff.reviewed_aec import _id
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');d=sample(tmp_path)
 if bad in ('zero','negative','blank'):
  d['shear_wall_detail']['input']['segments'][0]['height_m']={'zero':'0','negative':'-1','blank':''}[bad];refresh(d,tmp_path);assert not preview(d,None)['ready']
 else:
  commit(a,d);folder=tmp_path/'other';folder.mkdir();two=sample(folder,sid='SW2',floor='L2',z='4' if bad!='overlap' else '0')
  for op in two['shear_wall_detail']['input']['openings']:op['z_m']='5' if bad!='overlap' else '1'
  if bad=='duplicate_physical':two['shear_wall_detail']['input']['physical_wall_id']=d['shear_wall_detail']['input']['physical_wall_id']
  if bad=='different_datum':two['shear_wall_detail']['input']['datum']['coordinate_ref']='DIFFERENT'
  if bad=='same_source':two['source']['source_id']=d['source']['source_id'];two['shear_wall_detail']['input']['source']['source_id']=d['source']['source_id']
  refresh(two,folder)
  if bad in ('duplicate_physical','same_source'):
   with pytest.raises(ValueError):a.save_aec_system_draft('P',two)
  else:
   a.save_aec_system_draft('P',two)
   with pytest.raises(ValueError):project_gate(a.open_project('P'),_id('shear_wall','L2','SW2'))
 a.store.close()

def test_no_silent_historical_conversion(tmp_path):
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');old=a.new_aec_system('shear_wall','L1','SW1');a.save_aec_system_draft('P',old)
 with pytest.raises(ValueError):a.save_aec_system_draft('P',sample(tmp_path),expected_revision=1)
 assert 'shear_wall_detail' not in a.get_aec_system('P','shear_wall','L1','SW1')['data'];a.store.close()

def test_lossless_report_and_reopen_four_formats(tmp_path):
 from core.takeoff.shear_wall_detail import bind_input_digest
 from decimal import Decimal,localcontext
 from openpyxl import load_workbook
 from docx import Document
 import pymupdf
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');d=sample(tmp_path);d['shear_wall_detail']['input']['segments'][0]['thickness_m']='.250000000000000001';refresh(d,tmp_path);commit(a,d)
 expected='3.500000000000000014'
 for fmt in ('pdf','xlsx','docx','csv'):a.export_shear_wall_detail_report('P','L1','SW1',tmp_path/('wall.'+fmt),format=fmt)
 w=load_workbook(tmp_path/'wall.xlsx');assert all(s.sheet_view.rightToLeft for s in w);cells=[c for s in w for row in s for c in row];assert expected in [str(c.value) for c in cells];assert any(c.data_type=='n' for c in cells);w.close()
 document=Document(tmp_path/'wall.docx');assert document.tables;assert expected in '\n'.join(c.text for t in document.tables for row in t.rows for c in row.cells);assert 'bidiVisual' in document._element.xml
 with pymupdf.open(tmp_path/'wall.pdf') as pdf:assert len(pdf)>0 and all(len(p.get_text())>50 for p in pdf)
 raw=(tmp_path/'wall.csv').read_bytes();assert raw.startswith(b'\xef\xbb\xbf') and expected in raw.decode('utf-8-sig');a.store.close()

def test_amendment_duplicate_independence_and_financial_lock(tmp_path):
 a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');d=sample(tmp_path);commit(a,d);before=deepcopy(a.open_project('P')['boq']);folder=tmp_path/'rev';folder.mkdir();d['shear_wall_detail']['input']['segments'][0]['length_m']='6';refresh(d,folder);p=folder/'reason.txt';p.write_text('EXPLICIT SYNTHETIC reason');e={'evidence_path':str(p),'evidence_sha256':sha256(p.read_bytes()).hexdigest(),'evidence_reference':'دیتایل اصلاحی'}
 def create(evidence=e):return a.create_roof_amendment('P','shear_wall','L1','SW1',d,expected_revision=2,created_by='Creator',reason='اصلاح مصنوعی',evidence=evidence)
 with pytest.raises(ValueError):create({**e,'evidence_sha256':'0'*64})
 am=create()
 with pytest.raises(ValueError):create()
 with pytest.raises(ValueError):a.apply_beam_amendment('P',am['id'],expected_revision=1)
 for actor in ('Creator','تهیه‌کننده','بازبین مصنوعی'):
  with pytest.raises(ValueError):a.review_beam_amendment('P',am['id'],expected_revision=1,reviewer=actor,note='ممنوع')
 a.review_beam_amendment('P',am['id'],expected_revision=1,reviewer='Independent',note='کنترل');p=a.open_project('P');p['estimate']={'locked':True};a.store.save('P',p)
 with pytest.raises(ValueError):a.apply_beam_amendment('P',am['id'],expected_revision=1)
 assert a.open_project('P')['boq']==before;a.store.close()

def test_pdf_summary_heading_stays_with_table(tmp_path):
 import pymupdf,unicodedata
 from core.reports.exporters import _pdf_text
 a=StructuralProApp(tmp_path/'db');a.create_project('ساختمان پنج‌طبقه؛ داده مصنوعی و آموزش نرم‌افزاری','P');d=sample(tmp_path);commit(a,d);path=tmp_path/'wall.pdf';a.export_shear_wall_detail_report('P','L1','SW1',path,format='pdf')
 def norm(v):return ''.join(c for c in unicodedata.normalize('NFKC',v) if not c.isspace() and unicodedata.category(c)!='Cf')
 needle=norm('جمع‌بندی')
 with pymupdf.open(path) as pdf:
  pages=[p for p in pdf if needle in norm(p.get_text())];assert pages;assert all('84.944' in p.get_text() for p in pages),'Summary heading orphaned on a different page from its numeric table'
 a.store.close()
