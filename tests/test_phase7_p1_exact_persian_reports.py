"""Priority 1, Phase 7: exact Persian multi-format report contract.

All sources are synthetic; no value is injected into the production BOQ.
"""
from __future__ import annotations

import csv
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest

from core.reports import ProjectReport, ReportRequest, ReportService, build_report
from core.reports.designer import ReportLayout
from core.reports.decimal_values import decimal_value, exact_sum
from core.reports.exporters import report_date
from core.reports.formatting import money, table_html
from core.reports.production import totals as production_totals


def test_service_totals_do_not_round_decimal_cents():
    rows = [
        {"item_no": "0001", "description": "دیوار", "total": Decimal("0.10")},
        {"item_no": "0002", "description": "آرماتور", "total": "0.20"},
    ]
    actual = ReportService.build(ReportRequest("boq", "ریز متره"), rows)
    assert actual["totals"] == {"line_count": 2, "amount": Decimal("0.30")}
    assert actual["rows"] == rows
    assert actual["validation"]["valid"]


def test_service_large_and_small_values_remain_exact_without_float_overflow():
    rows = [
        {"description": "A", "total": Decimal("123456789012345678901234567890.123456789")},
        {"description": "B", "total": "0.000000000000000001"},
    ]
    actual = ReportService.build(ReportRequest("financial", "آزمون"), rows)
    assert actual["totals"]["amount"] == Decimal("123456789012345678901234567890.123456789000000001")


def test_group_subtotal_preserves_more_than_default_decimal_precision():
    rows = [
        {"group": "فولاد", "total": Decimal("123456789012345678901234567890.123456789")},
        {"group": "فولاد", "total": "0.000000000000000001"},
    ]
    out = ReportLayout(group_by="group", columns=["مبلغ"]).render_rows(rows)
    assert out[-1]["مبلغ"] == Decimal("123456789012345678901234567890.123456789000000001")


def test_shared_exact_sum_does_not_change_context_or_round_30_digit_values():
    value = exact_sum(["123456789012345678901234567890.10000000001", "-0.10000000000"])
    assert value == Decimal("123456789012345678901234567890.00000000001")


@pytest.mark.parametrize("value", [True, "", "NaN", "Infinity", "-Infinity", "bad", float("inf")])
def test_report_totals_reject_invalid_inputs(value):
    with pytest.raises(ValueError):
        ReportService.build(ReportRequest("financial", "گزارش"), [{"description": "بتن", "total": value}])


@pytest.mark.parametrize("field", ["quantity", "unit_price", "total", "مقدار", "بهای واحد", "مبلغ"])
def test_source_validation_rejects_invalid_numbers_in_both_schemas(field):
    record = {"description": "ملات", field: Decimal("-0.0001")}
    assert not build_report("گزارش", [record]).validate()["valid"]


def test_exact_validation_accepts_extremely_large_finite_decimal():
    row = {"description": "نمونه مصنوعی", "total": Decimal("1e400")}
    assert ProjectReport("گزارش", [row]).validate()["valid"]


def test_report_production_totals_are_decimal_and_immutable():
    rows = [{"مقدار": "0.1", "مبلغ": "0.2"}, {"مقدار": Decimal("0.2"), "مبلغ": Decimal("0.1")}]
    original = [dict(row) for row in rows]
    amounts = production_totals(rows)
    assert amounts == {"count": 2, "quantity": Decimal("0.3"), "amount": Decimal("0.3")}
    assert rows == original


def test_money_formatting_keeps_decimal_and_fails_closed():
    assert money(Decimal("12345678901234567890.125"), 3) == "12,345,678,901,234,567,890.125"
    assert money("0.10", 2) == "0.10"
    with pytest.raises(ValueError):
        money(float("nan"), 2)


def test_calendar_and_html_bidi_preserve_persian_content():
    assert report_date(date(2026, 10, 15)) == "۱۴۰۵/۰۷/۲۳"
    table = table_html(["شرح", "مقدار"], [["دیوار <B>", "0.10"]])
    assert 'dir="rtl"' in table
    assert "دیوار &lt;B&gt;" in table


def test_persian_multi_format_exports_share_unmodified_source(tmp_path):
    rows = [
        {"item_no": "0001", "description": "دیوار با بازشو", "quantity": Decimal("0.10"),
         "unit": "m2", "total": Decimal("0.20"), "source_id": "IFC:W-01"},
        {"item_no": "0002", "description": "ملات", "quantity": Decimal("12345678901234567890.123456789"),
         "unit": "kg", "total": Decimal("1.00"), "source_id": "MANUAL:2"},
    ]
    before = [dict(row) for row in rows]
    report = build_report("گزارش متره فارسی", rows, {"grand_total": "1.20"}, {"project_id": "P7"})
    outputs = {fmt: report.export(tmp_path / f"output.{fmt}", fmt) for fmt in ("csv", "xlsx", "docx", "pdf")}
    assert all(Path(path).is_file() and Path(path).stat().st_size for path in outputs.values())
    assert rows == before

    with open(outputs["csv"], encoding="utf-8-sig", newline="") as stream:
        csv_rows = list(csv.reader(stream))
    assert any("دیوار با بازشو" in row for row in csv_rows)
    assert any("12345678901234567890.123456789" in row for row in csv_rows)

    from openpyxl import load_workbook
    workbook = load_workbook(outputs["xlsx"])
    assert all(tab.sheet_view.rightToLeft for tab in workbook)
    cells = [cell for row in workbook["ریز متره"] for cell in row]
    assert any(cell.value == "0001" and cell.data_type == "s" for cell in cells)
    assert any(cell.value == "12345678901234567890.123456789" and cell.data_type == "s" for cell in cells)
    assert workbook["ریز متره"].oddFooter.center.text

    from docx import Document
    doc = Document(outputs["docx"])
    assert "w:bidi" in doc._element.xml
    assert "دیوار با بازشو" in doc._element.xml

    import pymupdf
    with pymupdf.open(outputs["pdf"]) as pdf:
        assert pdf.page_count >= 1
        assert all(page.rect.width > 0 and page.rect.height > 0 for page in pdf)
