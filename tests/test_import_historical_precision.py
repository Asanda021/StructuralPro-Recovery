"""Independent Decimal references and real input/output artifacts."""
from decimal import Decimal,localcontext
from pathlib import Path
import pytest
from core.pricing.catalog import PriceCatalog
from core.pricing.import_service import PricebookImportService
from core.platform.application import StructuralProApp

VALUE='1234567890123456789012345.1234567890123456789'

def csv_source(path,value=VALUE):
    path.write_text('year,group,chapter,code,description,unit,unit_price\n1405,ابنیه,01,X,مصنوعی,m2,'+value+'\n',encoding='utf-8')
    return path

@pytest.mark.parametrize('fmt',['csv','xlsx','pdf'])
def test_real_import_preserves_all_rate_digits(tmp_path,fmt):
    p=tmp_path/('input.'+fmt)
    if fmt=='csv':csv_source(p)
    elif fmt=='xlsx':
        from openpyxl import Workbook
        wb=Workbook();wb.active.append(['year','group','chapter','code','description','unit','unit_price']);wb.active.append([1405,'ابنیه','01','X','مصنوعی','m2',VALUE]);wb.save(p)
    else:
        import pymupdf
        pdf=pymupdf.open();page=pdf.new_page();page.insert_text((30,50),'X | Explicit synthetic price | m2 | '+VALUE);pdf.save(p);pdf.close()
    catalog=PriceCatalog();receipt=PricebookImportService(catalog).import_file(p,year=1405)
    assert receipt.rows==1 and Decimal(str(catalog.resolve('X',year=1405,unit='m2')['unit_price']))==Decimal(VALUE)
    exported=catalog.export_csv();copy=PriceCatalog();copy.import_csv(exported)
    assert Decimal(str(copy.get('X',1405).unit_price))==Decimal(VALUE)

@pytest.mark.parametrize('bad',[True,None,'NaN','Infinity','-1','12,34'])
def test_import_parser_fails_closed(bad):
    with pytest.raises(ValueError):PricebookImportService._number(bad)

def test_project_rate_sqlite_and_price_mismatch_are_exact(tmp_path):
    a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P')
    path=csv_source(tmp_path/'actual.csv');a.import_project_pricebook('P',path,year=1405,approved_by='explicit-synthetic-reviewer')
    a.add_takeoff('P','building','wall',source_id='W1',system='L1',length=1,height=1,count=1,openings=0,price_code='X',unit_price=VALUE)
    a.store.close();a=StructuralProApp(tmp_path/'db')
    p=a.open_project('P');assert Decimal(str(p['boq'][0]['unit_price']))==Decimal(VALUE)
    e=a.recalculate_estimate('P');assert Decimal(str(e['cost']['base']))==Decimal(VALUE)
    before=a.open_project('P');before['user_pricebook_review']['rows_by_code']['X']['unit_price']=VALUE[:-1]+'8'
    from core.reports.aec_office_package import audit_boq
    with pytest.raises(ValueError,match='قیمت'):audit_boq(before,require_pricebook=True)
    a.store.close()

def test_legacy_estimate_and_ledger_preserve_long_values():
    from core.commercial.estimation import EstimateLine,estimate,progress_claim
    from core.commercial.financial_ledger import LedgerEntry,summarize
    from core.commercial.ledger import LedgerEntry as OldEntry,ProjectLedger
    q='0.12345678901234567890123456789';factor='1.12345678901234567890123456789'
    with localcontext() as ctx:
        ctx.prec=512;amount=Decimal(q)*Decimal(VALUE)*Decimal(factor);quarter=amount/4
    line=EstimateLine('X','مصنوعی',q,VALUE,'m2',factor)
    assert Decimal(str(estimate([line])['subtotal']))==amount
    assert Decimal(str(progress_claim([line],{'X':'25'})['executed_amount']))==quarter
    assert Decimal(str(summarize([LedgerEntry('receivable',VALUE,'proof')])['totals']['receivable']))==Decimal(VALUE)
    assert Decimal(str(ProjectLedger([OldEntry('2026-10-13','receipt','explicit',VALUE)]).summary()['total']))==Decimal(VALUE)

def test_historical_contract_statement_has_no_ten_digit_rounding():
    from core.commercial.contract import Contract,ContractItem
    from core.projects.statement_engine import StatementLine
    from core.commercial.statement import build_payment_statement
    contract=Contract('C','مصنوعی',[ContractItem('X','مصنوعی','m2',1,VALUE)],retention_rate='0.1234567890123456789',tax_rate='0.09')
    result=build_payment_statement(contract,[StatementLine('X','مصنوعی','m2',1,VALUE,'0.25')])
    with localcontext() as ctx:
        ctx.prec=512;gross=Decimal(VALUE)*Decimal('.25');expected=gross*(1-Decimal(contract.retention_rate))+gross*Decimal('.09')
    assert Decimal(str(result['payable']))==expected

def test_long_exports_preserve_rate_in_all_formats(tmp_path):
    import csv,pymupdf
    from openpyxl import load_workbook
    from core.reports import exporters
    rows=[{'code':'X','unit_price':VALUE}];columns=[('code','کد'),('unit_price','بهای واحد')]
    csvpath=exporters.export_csv(rows,tmp_path/'rate.csv',columns=columns)
    assert VALUE in [cell for row in csv.reader(csvpath.open(encoding='utf-8-sig')) for cell in row]
    xlsx=exporters.export_xlsx(rows,tmp_path/'rate.xlsx',columns=columns);wb=load_workbook(xlsx)
    assert any(cell.value==VALUE for row in wb['ریز متره'] for cell in row);wb.close()
    pdf=exporters.export_pdf(rows,tmp_path/'rate.pdf',title='نرخ مصنوعی دقیق',columns=columns)
    with pymupdf.open(pdf) as doc:assert VALUE in '\n'.join(p.get_text() for p in doc)

def test_excel_numeric_long_cell_is_blocked_not_repaired(tmp_path):
    from openpyxl import Workbook
    wb=Workbook();wb.active.append(['code','description','unit','unit_price']);wb.active.append(['X','Explicit','m2',1234567890123456]);p=tmp_path/'numeric.xlsx';wb.save(p)
    with pytest.raises(ValueError,match='متنی'):PricebookImportService(PriceCatalog()).import_file(p,year=1405)

def test_bad_catalog_import_does_not_replace_existing_year():
    from core.pricing.catalog import PriceItem
    c=PriceCatalog([PriceItem(1405,'','', 'OLD','existing','m2',100)])
    with pytest.raises(ValueError):c.import_csv('year,group,chapter,code,description,unit,unit_price\n1405,,,X,valid,m2,1\n1405,,,Y,bad,,2\n',replace_year=True)
    assert c.get('OLD',1405).unit_price==100 and c.get('X',1405) is None

def test_legacy_rounded_price_is_blocked_against_real_source(tmp_path):
    a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');p=csv_source(tmp_path/'long.csv')
    a.import_project_pricebook('P',p,year=1405,approved_by='synthetic-reviewer')
    rounded=float(VALUE)
    a.add_takeoff('P','building','wall',source_id='W1',system='L1',length=1,height=1,count=1,openings=0,price_code='X',unit_price=rounded)
    old=a.open_project('P');old['user_pricebook_review']['rows_by_code']['X']['unit_price']=rounded;old['user_pricebook_review'].pop('numeric_encoding_version',None)
    from core.reports.aec_office_package import audit_boq
    with pytest.raises(ValueError,match='قیمت'):audit_boq(old,require_pricebook=True)
    a.store.close()

@pytest.mark.parametrize('field',[True,None,'NaN','Infinity','-1'])
def test_legacy_engines_reject_bad_values(field):
    from core.commercial.estimation import EstimateLine,estimate
    from core.commercial.financial_ledger import LedgerEntry,summarize
    from core.commercial.ledger import LedgerEntry as OldEntry,ProjectLedger
    with pytest.raises(ValueError):estimate([EstimateLine('X','Explicit',1,field)])
    with pytest.raises(ValueError):summarize([LedgerEntry('receivable',field,'proof')])
    with pytest.raises(ValueError):ProjectLedger([OldEntry('2026-10-13','receipt','Explicit',field)]).summary()

def test_exact_rate_and_row_total_gate_rejects_a_single_changed_digit(tmp_path):
    a=StructuralProApp(tmp_path/'db');a.create_project('مصنوعی','P');p=csv_source(tmp_path/'long.csv');a.import_project_pricebook('P',p,year=1405,approved_by='explicit-synthetic-reviewer')
    a.add_takeoff('P','building','wall',source_id='W1',system='L1',length=1,height=1,count=1,openings=0,price_code='X',unit_price=VALUE)
    bad=a.open_project('P');bad['boq'][0]['total']=VALUE[:-1]+'8'
    from core.reports.aec_office_package import audit_boq
    with pytest.raises(ValueError,match='مبلغ'):audit_boq(bad,require_pricebook=True)
    a.store.close()

def test_persian_digits_and_explicit_zero():
    assert Decimal(str(PricebookImportService._number('۱٬۲۳۴٫۵۶۷۸۹')))==Decimal('1234.56789')
    assert PricebookImportService._number('۰')==0
    assert PricebookImportService._year('۱۴۰۵')==1405
    with pytest.raises(ValueError):PricebookImportService._year('1405.5')

def test_legacy_decimal_payload_serializes_without_precision_loss():
    import json
    from core.commercial.ledger import LedgerEntry,ProjectLedger
    data=ProjectLedger([LedgerEntry('2026-10-13','receipt','Explicit',Decimal(VALUE))]).to_dict()
    restored=json.loads(json.dumps(data));assert Decimal(str(restored[0]['amount']))==Decimal(VALUE)
