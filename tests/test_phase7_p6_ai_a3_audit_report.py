"""P7 P6 A3: real persisted AI takeoff, source and four RTL output formats."""
from copy import deepcopy
from hashlib import sha256
import csv
import json

import pytest

from core.drawings.graphical_takeoff import Point
from core.drawings.takeoff_session import DrawingTakeoffSession
from core.platform.application import StructuralProApp
from core.reports.ai_a3_audit_report import (
    preview_ai_a3_audit, export_ai_a3_audit_package,
)
from drawing_source_fixture import vector_pdf
from test_ai_takeoff_persistent_approval import _proposal


def _ready(tmp_path):
    app = StructuralProApp(tmp_path / "app")
    app.create_project("آزمون سند هوش مصنوعی", "P")
    source = vector_pdf(tmp_path)
    session = DrawingTakeoffSession(source)
    session.calibrate(1, 100, 10)
    item = session.add_length([Point(0, 0), Point(100, 0)], page=1, label="خط شاهد")
    saved = app.save_drawing_takeoff_session("P", session)
    proposal = _proposal(saved["session_id"], saved["revision"], session, item)
    app.commit_ai_takeoff_proposal("P", proposal, user_confirmed=True)
    version = app.open_project("P")["_meta"]["version"]
    return app, source, version


def test_real_persisted_ai_a3_four_formats_and_manifest_without_mutation(tmp_path):
    import pymupdf
    from docx import Document
    from openpyxl import load_workbook

    app, source, version = _ready(tmp_path)
    prior = deepcopy(app.open_project("P"))
    digest = app.store.current_digest("P")
    try:
        preview = preview_ai_a3_audit(
            app, project_id="P", expected_project_version=version)
        assert preview["record_count"] == 1
        row = preview["records"][0]
        assert row["source_id"].startswith("ai-takeoff:")
        assert row["source_match"] == "EXACT_SOURCE_QUANTITY"
        assert row["unit"] == "m"
        assert row["independent_reviewer"] == "NOT_RECORDED_IN_LEGACY_RECEIPT"
        assert row["model_provenance"] == "NOT_RECORDED_IN_LEGACY_RECEIPT"
        assert preview["financial_and_engineering_approval"] is False
        assert preview["human_confirmation_scope"] == "BOOLEAN_AT_COMMIT_ONLY"
        dest = tmp_path / "A3"
        exported = export_ai_a3_audit_package(
            app, project_id="P", expected_project_version=version,
            destination=dest)
        assert exported["destination"] == str(dest)
        manifest = json.loads((dest / "manifest.json").read_text(encoding="utf-8"))
        assert manifest == exported["manifest"]
        assert manifest["record_count"] == 1
        assert manifest["source_project_digest"] == digest
        assert manifest["rejected_or_uncommitted_proposals"] == "NOT_PERSISTED_NOT_AUDITED"
        assert manifest["financial_and_engineering_approval"] is False
        assert manifest["physical_windows_office_print"] == "NOT_TESTED_PHASE8"
        assert len(manifest["files"]) == 4
        for receipt in manifest["files"]:
            path = dest / receipt["name"]
            assert receipt["sha256"] == sha256(path.read_bytes()).hexdigest()
            assert path.stat().st_size > 0
        with pymupdf.open(dest / "ai-a3-audit.pdf") as document:
            assert len(document) and any(p.get_text().strip() for p in document)
        excel = load_workbook(dest / "ai-a3-audit.xlsx")
        assert all(ws.sheet_view.rightToLeft for ws in excel)
        excel.close()
        word = Document(dest / "ai-a3-audit.docx")
        assert word.tables and "w:bidi" in word._element.xml
        with (dest / "ai-a3-audit.csv").open(
                "r", encoding="utf-8-sig", newline="") as stream:
            assert any("NOT_RECORDED_IN_LEGACY_RECEIPT" in cell
                       for line in csv.reader(stream) for cell in line)
        assert app.store.current_digest("P") == digest
        assert app.open_project("P") == prior
    finally:
        app.store.close()


def test_ai_audit_without_accepted_event_never_creates_fictitious_report(tmp_path):
    app = StructuralProApp(tmp_path / "empty")
    app.create_project("بدون AI", "P")
    version = app.open_project("P")["_meta"]["version"]
    try:
        with pytest.raises(ValueError, match="هیچ برداشت AI"):
            export_ai_a3_audit_package(
                app, project_id="P", expected_project_version=version,
                destination=tmp_path / "nothing")
        assert not (tmp_path / "nothing").exists()
    finally:
        app.store.close()


def test_changed_original_pdf_fails_closed_without_creating_package(tmp_path):
    app, source, version = _ready(tmp_path)
    before = app.store.current_digest("P")
    try:
        source = __import__("pathlib").Path(source)
        source.write_bytes(source.read_bytes() + b"tamper")
        with pytest.raises(ValueError):
            export_ai_a3_audit_package(
                app, project_id="P", expected_project_version=version,
                destination=tmp_path / "bad")
        assert app.store.current_digest("P") == before
        assert not (tmp_path / "bad").exists()
    finally:
        app.store.close()


@pytest.mark.parametrize("fault", ["quantity", "identity", "session_item", "source_id", "confidence"])
def test_tampered_ai_receipt_is_not_reported_as_verified(tmp_path, monkeypatch, fault):
    app, source, version = _ready(tmp_path)
    original = app.open_project
    def altered(pid):
        snapshot = deepcopy(original(pid))
        row = next(x for x in snapshot["takeoffs"] if x["domain"] == "drawing_ai")
        if fault == "quantity":
            row["quantities"][0]["amount"] = "999"
        elif fault == "identity":
            row["params"]["measurement_identity"] = "wrong_sha"
        elif fault == "session_item":
            row["params"]["session_item_id"] = "missing"
        elif fault == "source_id":
            row["source_id"] = "ai-takeoff:fake:1:wrong"
        elif fault == "confidence":
            row["params"]["confidence"] = 10
        return snapshot
    monkeypatch.setattr(app, "open_project", altered)
    try:
        with pytest.raises((ValueError, KeyError)):
            preview_ai_a3_audit(
                app, project_id="P", expected_project_version=version)
    finally:
        app.store.close()


def test_stale_version_and_existing_destination_rejected(tmp_path):
    app, source, version = _ready(tmp_path)
    try:
        with pytest.raises(ValueError, match="نسخه"):
            preview_ai_a3_audit(
                app, project_id="P", expected_project_version=version - 1)
        folder = tmp_path / "existing"
        folder.mkdir()
        with pytest.raises(FileExistsError):
            export_ai_a3_audit_package(
                app, project_id="P", expected_project_version=version,
                destination=folder)
    finally:
        app.store.close()
