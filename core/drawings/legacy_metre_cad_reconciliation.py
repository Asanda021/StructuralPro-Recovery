"""Read-only, fail-closed legacy Persian metre-to-CAD reconciliation.

This is NOT a pricebook importer and NEVER commits quantities, prices, or BOQ.
No textual title, shared floor area, or unverified DWG conversion can authorize
a geometric match. User-provided explicit CAD handle/metric/unit mappings are
required; every unbound row remains NOT_VERIFIED.
"""
from __future__ import annotations

from decimal import Decimal, InvalidOperation
from hashlib import sha256
from pathlib import Path
import re
from typing import Any

_MAX_FILE_SIZE = 25 * 1024 * 1024
_MAX_SHEETS = 120
_MAX_ROWS_PER_SHEET = 10000
_AMOUNT_TOLERANCE = Decimal("0.02")
_GEOMETRY_TOLERANCE = Decimal("0.001")
_SHEET_PATTERN = re.compile(r"^S_(\d+)_(\d+)$", re.IGNORECASE)
_COLUMNS = {"quantity": "I", "unit_price": "L", "amount": "M"}


def _number(value: Any) -> Decimal | None:
    if isinstance(value, bool) or value is None or value == "" or isinstance(value, str):
        return None
    try:
        n = Decimal(str(value))
    except (InvalidOperation, ValueError):
        return None
    return n if n.is_finite() else None


def _normalized_sheet(name: str) -> str:
    return str(name).strip().replace("\u200c", "").replace("\ufeff", "").replace(" ", "")


def _source(path: str | Path) -> tuple[Path, str]:
    p = Path(path)
    if p.suffix.lower() != ".xlsx" or not p.is_file():
        raise ValueError("Only an existing XLSX workbook can be audited")
    if not 0 < p.stat().st_size <= _MAX_FILE_SIZE:
        raise ValueError("Workbook exceeds audited size limits")
    return p, sha256(p.read_bytes()).hexdigest()


def _sheet_rows(ws, *, sheet: str, chapter: int, source_sha: str) -> tuple[list[dict], list[dict]]:
    rows: list[dict] = []
    issues: list[dict] = []
    if ws.max_row and ws.max_row > _MAX_ROWS_PER_SHEET:
        raise ValueError("Workbook sheet exceeds audited row limit")
    if (str(ws["B3"].value or "").strip() != "شرح عمليات" or
        str(ws["L3"].value or "").strip() != "قیمت" or
        str(ws["M3"].value or "").strip() != "قیمت کل"):
        raise ValueError("Legacy metre sheet does not have the audited source headers")
    # Read cached cell values only. No formulas or external workbooks executed.
    for index, cells in enumerate(ws.iter_rows(min_row=5, max_col=14, values_only=True), 5):
        label = cells[1] if len(cells) > 1 else None
        rate = _number(cells[11] if len(cells) > 11 else None)
        amount = _number(cells[12] if len(cells) > 12 else None)
        if not isinstance(label, str) or not label.strip() or rate is None or amount is None:
            continue
        candidate_cols = ("I", "J", "H")
        candidates = [(col, _number(cells[ord(col) - 65])) for col in candidate_cols]
        present = [(col, v) for col, v in candidates if v is not None]
        column, quantity = present[0] if present else (None, None)
        key = f"{sheet}!{index}"
        row = {
            "key": key, "sheet": sheet, "chapter": chapter, "row": index,
            "description": label.strip(), "price_code": str(cells[10] or "").strip(),
            "rate": str(rate), "amount": str(amount),
            "quantity": str(quantity) if quantity is not None else None,
            "quantity_cell": f"{column}{index}" if column else None,
            "unit": "UNDECLARED_IN_SOURCE",
            "source_sha256": source_sha,
            "cad_status": "NOT_VERIFIED_NO_EXPLICIT_GEOMETRY_BINDING",
        }
        rows.append(row)
        if quantity is None:
            issues.append({"code": "MISSING_NUMERIC_QUANTITY", "key": key})
        else:
            expected = rate * quantity
            if abs(expected - amount) > _AMOUNT_TOLERANCE:
                issues.append({
                    "code": "ROW_AMOUNT_DISAGREES_WITH_PRIMARY_QUANTITY_TIMES_RATE",
                    "key": key, "amount_cell": f"M{index}",
                    "observed_amount": str(amount),
                    "computed_amount": str(expected), "difference": str(amount - expected),
                    "quantity_cell": f"{column}{index}",
                })
    return rows, issues


def _read_summary(path: str | Path, chapter_totals: dict[int, Decimal]) -> tuple[dict, list[dict]]:
    from openpyxl import load_workbook
    p, source_sha = _source(path)
    issues: list[dict] = []
    external_links: list[dict] = []
    wb_formula = load_workbook(p, read_only=True, data_only=False, keep_links=False)
    wb_values = load_workbook(p, read_only=True, data_only=True, keep_links=False)
    try:
        if len(wb_values.worksheets) != 1:
            raise ValueError("Summary requires one explicitly recognized sheet")
        sheet_v, sheet_f = wb_values.worksheets[0], wb_formula.worksheets[0]
        # This legacy layout: row 3 = chapter 1, row 4 = chapter 2, etc.
        if sheet_v["B2"].value != "شماره فصل فهرست بها":
            raise ValueError("Summary header does not match audited legacy layout")
        for chapter, expected in sorted(chapter_totals.items()):
            idx = chapter + 2
            formula = sheet_f[f"D{idx}"].value
            observed = sheet_v[f"D{idx}"].value
            cached = _number(observed)
            if isinstance(formula, str) and "[" in formula:
                external_links.append({"cell": f"D{idx}", "chapter": chapter,
                                       "status": "UNRESOLVED_EXTERNAL_WORKBOOK_REFERENCE"})
                issues.append({"code": "EXTERNAL_LINK_NOT_INDEPENDENTLY_VERIFIED", "chapter": chapter,
                               "cell": f"D{idx}"})
            if cached is None:
                issues.append({"code": "CHAPTER_SUMMARY_NOT_NUMERIC_OR_CACHE_INVALID",
                               "chapter": chapter, "cell": f"D{idx}"})
            elif abs(cached - expected) > _AMOUNT_TOLERANCE:
                issues.append({"code": "CHAPTER_SUMMARY_DIFFERS_FROM_METRE",
                               "chapter": chapter, "cell": f"D{idx}",
                               "observed_amount": str(cached),
                               "computed_from_metre": str(expected),
                               "difference": str(cached - expected)})
            if isinstance(formula, str) and formula.startswith("=") and "!" in formula and "[" not in formula:
                issues.append({"code": "SUMMARY_FORMULA_NOT_REEVALUATED", "chapter": chapter,
                               "cell": f"D{idx}"})
        stored_total = Decimal(0)
        # G4 etc holds cached amounts after factors, NOT certified geometry or price.
        for idx in range(3, 31):
            n = _number(sheet_v[f"G{idx}"].value)
            if n is not None:
                stored_total += n
        return ({
            "source_sha256": source_sha,
            "chapter_amount_total_cached": str(stored_total),
            "external_references": external_links,
            "independent_formula_evaluation": False,
        }, issues)
    finally:
        wb_values.close()
        wb_formula.close()
        if sha256(p.read_bytes()).hexdigest() != source_sha:
            raise ValueError("Summary workbook changed while being audited")


class LegacyMetreCADReconciler:
    """Read-only inspector for explicit legacy metre row vs CAD geometry proofs."""

    def inspect(
        self, metre_path: str | Path, *,
        summary_path: str | Path | None = None,
        drawing_path: str | Path | None = None,
        bindings: list[dict] | None = None,
    ) -> dict:
        from openpyxl import load_workbook
        p, source_sha = _source(metre_path)
        wb = load_workbook(p, read_only=True, data_only=True, keep_links=False)
        rows: list[dict] = []
        issues: list[dict] = []
        chapters: dict[int, Decimal] = {}
        seen: set[str] = set()
        try:
            if len(wb.worksheets) > _MAX_SHEETS:
                raise ValueError("Too many workbook sheets")
            for ws in wb.worksheets:
                name = _normalized_sheet(ws.title)
                match = _SHEET_PATTERN.fullmatch(name)
                if match is None:
                    raise ValueError(f"Unknown legacy metre sheet: {ws.title}")
                chapter = int(match.group(1))
                if name in seen:
                    raise ValueError("Duplicate legacy metre sheet")
                seen.add(name)
                sub, findings = _sheet_rows(ws, sheet=name, chapter=chapter, source_sha=source_sha)
                rows.extend(sub)
                issues.extend(findings)
                chapters[chapter] = chapters.get(chapter, Decimal(0)) + sum(
                    (Decimal(item["amount"]) for item in sub), Decimal(0))
        finally:
            wb.close()
        if sha256(p.read_bytes()).hexdigest() != source_sha:
            raise ValueError("Metre workbook changed during audit")
        if not rows:
            raise ValueError("No priced legacy metre rows recognized")
        report = {
            "status": "READ_ONLY_AUDIT_REQUIRES_INDEPENDENT_CAD_BINDING",
            "metre_source_sha256": source_sha,
            "metre_sheets": len(seen),
            "priced_rows": len(rows),
            "chapter_totals": {str(k): str(v) for k, v in sorted(chapters.items())},
            "rows": rows, "issues": issues,
            "summary": None,
            "cad": {"status": "NOT_TESTED", "verified_rows": 0, "not_verified_rows": len(rows)},
            "boq_rows_committed": 0, "prices_approved": False,
        }
        if summary_path is not None:
            report["summary"], add = _read_summary(summary_path, chapters)
            report["issues"].extend(add)
        if drawing_path is not None:
            self._compare_cad(report, drawing_path, bindings or [])
        elif bindings:
            raise ValueError("CAD mappings cannot be evaluated without drawing file")
        return report

    @staticmethod
    def _compare_cad(report: dict, drawing_path: str | Path, bindings: list[dict]) -> None:
        from core.drawings.dwg_takeoff import DWGTakeoffEngine, cad_entity_geometry
        src = Path(drawing_path)
        if not src.is_file() or src.suffix.lower() not in {".dwg", ".dxf"}:
            raise ValueError("A real DWG or DXF file is required")
        src_digest = sha256(src.read_bytes()).hexdigest()
        cad_header = {}
        if src.suffix.lower() == ".dwg":
            from core.drawings.dwg_converter import OfflineDWGConverter
            from core.cad.external_dwg_provider_v1 import ACadSharpDWGProvider
            try:
                cad_header = OfflineDWGConverter.inspect_header(src)
            except (RuntimeError, ValueError) as exc:
                report["cad"] = {
                    "status": "BLOCKED_UNSUPPORTED_OR_CORRUPT_DWG_HEADER",
                    "source_sha256": src_digest,
                    "reason_type": type(exc).__name__,
                    "verified_rows": 0, "not_verified_rows": len(report["rows"]),
                }
                return
            from core.cad.optional_strict_python_dwg_provider import StrictPythonDWGProvider
            if (not ACadSharpDWGProvider().available and
                not OfflineDWGConverter().available and
                not StrictPythonDWGProvider().available):
                report["cad"] = {
                    "status": "BLOCKED_NATIVE_DWG_CONVERTER_UNAVAILABLE",
                    "source_sha256": src_digest, "dwg_signature": cad_header["code"],
                    "dwg_version": cad_header["version"],
                    "remediation": "Configure an approved ACadSharp/ODA converter or install pinned ezdwg[dxf]==0.17.0 prebuilt wheel; do not approve unconverted DWG",
                    "verified_rows": 0, "not_verified_rows": len(report["rows"]),
                }
                return
        # No fake DXF, auto-mapping by title, or bypass of the configured reader.
        try:
            doc = DWGTakeoffEngine.with_array_support().import_file(src)
        except (OSError, RuntimeError, ValueError) as exc:
            report["cad"] = {
                "status": "BLOCKED_NATIVE_CAD_IMPORT_NOT_VERIFIED",
                "source_sha256": src_digest,
                "dwg_signature": cad_header.get("code"),
                "dwg_version": cad_header.get("version"),
                "reason_type": type(exc).__name__,
                "verified_rows": 0,
                "not_verified_rows": len(report["rows"]),
            }
            return
        if sha256(src.read_bytes()).hexdigest() != src_digest or doc.source_sha256 != src_digest:
            raise ValueError("CAD source changed during the review")
        if doc.format.startswith("DWG") and doc.conversion.get("decoder_version") in (None, "", "unverified"):
            report["cad"] = {
                "status": "BLOCKED_DWG_DECODER_VERSION_UNVERIFIED",
                "source_sha256": src_digest, "verified_rows": 0,
                "not_verified_rows": len(report["rows"]),
            }
            return
        row_by_key = {r["key"]: r for r in report["rows"]}
        seen: set[str] = set()
        for link in bindings:
            if not isinstance(link, dict):
                raise ValueError("Explicit CAD link must be a dictionary")
            key = str(link.get("key", ""))
            if key not in row_by_key or key in seen:
                raise ValueError("Unknown or duplicate metre row key")
            seen.add(key)
            row = row_by_key[key]
            if link.get("source_sha256") != src_digest:
                raise ValueError("CAD link source SHA must be explicitly pinned")
            metric = str(link.get("metric", ""))
            unit = str(link.get("unit", ""))
            if (metric, unit) not in {("length", "m"), ("area", "m2")}:
                raise ValueError("Volume/weight/count or unknown unit cannot be inferred from 2D CAD")
            if link.get("quantity_cell") != row["quantity_cell"]:
                raise ValueError("Quantity column mismatch; explicit basis required")
            if not str(link.get("evidence_ref", "")).strip():
                raise ValueError("Independent drawing/legend evidence is mandatory")
            entity, points, measured = cad_entity_geometry(doc, link.get("handle"), metric)
            expected = Decimal(row["quantity"]) if row["quantity"] is not None else None
            actual = Decimal(str(measured))
            if expected is None:
                raise ValueError("No numeric metre quantity to compare")
            matched = abs(actual - expected) <= _GEOMETRY_TOLERANCE
            row["cad_status"] = "CANDIDATE_MATCH_AWAITING_HUMAN_APPROVAL" if matched else "CAD_QUANTITY_MISMATCH"
            row["cad_evidence"] = {
                "drawing_sha256": src_digest, "handle": entity.handle,
                "layer": entity.layer, "metric": metric, "unit": unit,
                "expected_quantity": str(expected), "measured_quantity": str(actual),
                "difference": str(actual - expected),
                "evidence_ref": str(link["evidence_ref"]),
                "independent_human_approved": False,
            }
            if not matched:
                report["issues"].append({"code": "CAD_QUANTITY_DIFFERS_FROM_METRE", "key": key})
        report["cad"] = {
            "status": "PARTIAL_EXPLICIT_GEOMETRY_CROSSCHECK_NOT_APPROVED",
            "source_sha256": src_digest,
            "document_unit": doc.units, "source_format": doc.format,
            "matched_candidates": sum(r["cad_status"] == "CANDIDATE_MATCH_AWAITING_HUMAN_APPROVAL" for r in report["rows"]),
            "mismatched_rows": sum(r["cad_status"] == "CAD_QUANTITY_MISMATCH" for r in report["rows"]),
            "unbound_rows": sum(r["cad_status"] == "NOT_VERIFIED_NO_EXPLICIT_GEOMETRY_BINDING" for r in report["rows"]),
            "verified_rows": 0,
            "not_verified_rows": len(report["rows"]),
            "human_approval_required": True,
        }
