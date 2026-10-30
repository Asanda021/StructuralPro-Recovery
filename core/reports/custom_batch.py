"""P7 E7: versioned ReportLayout projections, never a quantity/pricing engine."""
from __future__ import annotations
from copy import deepcopy
from dataclasses import asdict
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path
import re
import tempfile
from core.reports.designer import ReportLayout, _FIELD_MAP
from core.reports.decimal_values import exact_sum, decimal_value
from core.reports.project_report import ProjectReport, DEFAULT_COLUMNS
from core.reports.bundle import build_bundle_manifest

FIELDS = dict(DEFAULT_COLUMNS) | {"floor":"طبقه","member":"عضو","discipline":"رشته","zone":"زون","revision":"نسخه"}
FILTERS = {"floor","discipline","zone","revision"}

def _json(value):
    return json.dumps(value,ensure_ascii=False,sort_keys=True,separators=(",",":"))

def _digest(value):
    return sha256(_json(value).encode()).hexdigest()

def _text(value,label,limit=200):
    if not isinstance(value,str) or not value.strip() or len(value)>limit or any(ord(c)<32 for c in value):
        raise ValueError(label+" معتبر لازم است")
    return value

def _identity(value):
    if not isinstance(value,str) or not re.fullmatch(r"[A-Za-z0-9_-]{1,64}",value):
        raise ValueError("شناسه قالب معتبر نیست")
    return value

def _safe(root,*parts):
    path=root.joinpath(*parts)
    if any(p.is_symlink() for p in (path,*path.parents)) or root.resolve() not in path.resolve().parents:
        raise ValueError("مسیر خارج از پروژه یا نمادین مجاز نیست")
    return path

def _filters(value):
    if not isinstance(value,dict) or set(value)-FILTERS:raise ValueError("فیلتر ناشناخته است")
    for key,values in value.items():
        if not isinstance(values,list) or not values or not all(isinstance(v,str) for v in values) or len(set(values))!=len(values):
            raise ValueError("فهرست صریح فیلتر لازم است")
        for v in values:_text(v,"فیلتر",120)

def validate_template(payload):
    t=deepcopy(payload)
    required={"schema_version","template_id","version","report_type","layout","filters","page_format","orientation","show_source","include_totals","logo"}
    if not isinstance(t,dict) or set(t)!=required or type(t.get("schema_version")) is not int or t["schema_version"]!=2:
        raise ValueError("نسخه schema قالب پشتیبانی نمی‌شود؛ مهاجرت صریح لازم است")
    _identity(t["template_id"])
    if type(t["version"]) is not int or not 1<=t["version"]<=100:raise ValueError("نسخه قالب نامعتبر است")
    if t["report_type"] not in {"boq","takeoff","estimate"}:raise ValueError("نوع گزارش پشتیبانی نمی‌شود")
    layout=t["layout"]
    if not isinstance(layout,dict) or set(layout)!=set(asdict(ReportLayout())):raise ValueError("ساختار چیدمان ناقص است")
    _text(layout["title"],"عنوان",120)
    for key in ("header","footer"):
        if not isinstance(layout[key],str) or len(layout[key])>80 or any(ord(c)<32 for c in layout[key]):
            raise ValueError("سربرگ یا پابرگ نامعتبر است")
    if layout["rtl"] is not True or type(layout["page_size"]) is not int or not 1<=layout["page_size"]<=500:
        raise ValueError("RTL یا صفحه‌بندی نامعتبر است")
    columns=layout["columns"]
    if not isinstance(columns,list) or not columns or not all(isinstance(c,str) for c in columns):raise ValueError("ستون‌های صریح لازم است")
    layout["columns"]=[_FIELD_MAP.get(c,c) for c in columns]
    if len(set(layout["columns"]))!=len(columns) or any(c not in FIELDS for c in layout["columns"]):raise ValueError("ستون ناشناخته یا تکراری است")
    if layout["group_by"] not in ("","floor","discipline","zone","unit","chapter"):raise ValueError("گروه‌بندی نامعتبر است")
    if t["page_format"] not in {"A4","A3"} or t["orientation"] not in {"portrait","landscape"}:raise ValueError("صفحه یا جهت چاپ نامعتبر است")
    if type(t["show_source"]) is not bool or type(t["include_totals"]) is not bool:raise ValueError("نمایش منبع و جمع‌ها باید صریح باشد")
    _filters(t["filters"])
    if t["logo"] is not None and (not isinstance(t["logo"],str) or not re.fullmatch(r"[0-9a-f]{64}\.png",t["logo"])):
        raise ValueError("ارجاع لوگو داخل پروژه معتبر نیست")
    return t

def migrate_template_v1(payload):
    if not isinstance(payload,dict) or type(payload.get("schema_version")) is not int or payload["schema_version"]!=1 or set(payload)-{"schema_version","template_id","version","layout"}:
        raise ValueError("مهاجرت فقط از schema یک شناخته‌شده مجاز است")
    return validate_template(dict(payload,schema_version=2,report_type="boq",filters={},page_format="A4",orientation="landscape",show_source=True,include_totals=True,logo=None))

def _png(data):
    from PIL import Image
    if len(data)>2_000_000 or not data.startswith(b"\x89PNG\r\n\x1a\n"):raise ValueError("فقط لوگوی PNG معتبر مجاز است")
    try:
        with Image.open(BytesIO(data)) as image:
            if image.format!="PNG" or not 1<=image.width<=4096 or not 1<=image.height<=4096:raise ValueError("ابعاد لوگو نامعتبر است")
            image.verify()
    except Exception as exc:raise ValueError("لوگو قابل بازخوانی نیست") from exc

class TemplateStore:
    """Append-only template revisions outside ProjectStore/financial state."""
    def __init__(self,service,project_id):
        p=service.open_project(project_id)
        if not isinstance(p,dict) or p.get("id")!=project_id:raise ValueError("پروژه معتبر موجود نیست")
        key=sha256(project_id.encode()).hexdigest()
        self.root=Path(service.data_dir)/"report_templates"/key
        self.assets=Path(service.data_dir)/"report_assets"/key
        for path in (self.root,self.assets):
            if any(p.is_symlink() for p in (path,*path.parents)):raise ValueError("مسیر نمادین قالب مجاز نیست")
            path.mkdir(parents=True,exist_ok=True)

    def list_versions(self):
        result=[]
        for folder in sorted(self.root.iterdir()):
            _identity(folder.name)
            if folder.is_symlink() or not folder.is_dir():raise ValueError("مسیر قالب نامعتبر است")
            for file in folder.iterdir():
                if file.is_symlink() or not re.fullmatch(r"[1-9][0-9]*\.json",file.name):raise ValueError("فایل قالب نامعتبر است")
                result.append((folder.name,int(file.stem)))
        return sorted(result)

    def load(self,template_id,version):
        _identity(template_id)
        if type(version) is not int or not 1<=version<=100:raise ValueError("نسخه قالب نامعتبر است")
        path=_safe(self.root,template_id,str(version)+".json")
        try:record=json.loads(path.read_text(encoding="utf-8"))
        except (OSError,ValueError) as exc:raise ValueError("قالب قابل بازخوانی نیست") from exc
        if not isinstance(record,dict) or record.get("sha256")!=_digest({k:v for k,v in record.items() if k!="sha256"}):raise ValueError("قالب دستکاری شده است")
        t=validate_template(record["template"])
        if t!=record["template"] or t["template_id"]!=template_id or t["version"]!=version:raise ValueError("هویت قالب ناسازگار است")
        if version>1 and record["previous_sha256"]!=self.load(template_id,version-1)["sha256"]:raise ValueError("زنجیره نسخه قالب ناسازگار است")
        if version==1 and record["previous_sha256"] is not None:raise ValueError("نسخه نخست قالب نامعتبر است")
        self.logo_bytes(t)
        return record

    def save(self,template,*,actor,reason):
        t=validate_template(template);_text(actor,"ثبت‌کننده");_text(reason,"علت تغییر");self.logo_bytes(t)
        versions=[v for id,v in self.list_versions() if id==t["template_id"]]
        expected=max(versions,default=0)+1
        if t["version"]!=expected:raise ValueError("نسخه قالب هم‌زمان تغییر کرده یا تکراری است")
        previous=self.load(t["template_id"],expected-1)["sha256"] if expected>1 else None
        record=dict(template=t,actor=actor,reason=reason,created_at=datetime.now(timezone.utc).isoformat(),previous_sha256=previous)
        record["sha256"]=_digest(record)
        folder=_safe(self.root,t["template_id"]);folder.mkdir(exist_ok=True)
        target=_safe(self.root,t["template_id"],str(expected)+".json")
        with tempfile.NamedTemporaryFile(dir=folder,delete=False) as f:
            temp=Path(f.name);f.write(_json(record).encode());f.flush();os.fsync(f.fileno())
        try:os.link(temp,target)
        except FileExistsError as exc:raise ValueError("نسخه قالب هم‌زمان ثبت شد") from exc
        finally:temp.unlink(missing_ok=True)
        return record

    def logo_bytes(self,t):
        if t["logo"] is None:return None
        path=_safe(self.assets,t["logo"])
        try:data=path.read_bytes()
        except OSError as exc:raise ValueError("لوگوی پروژه موجود نیست") from exc
        if sha256(data).hexdigest()+".png"!=t["logo"]:raise ValueError("هش لوگو تغییر کرده است")
        _png(data);return data

    def import_logo(self,path):
        """Explicit selection imports bytes; persisted references stay in project assets."""
        path=Path(path)
        if not path.is_file() or path.is_symlink() or path.stat().st_size>2_000_000:raise ValueError("فایل لوگو معتبر نیست")
        data=path.read_bytes();_png(data);name=sha256(data).hexdigest()+".png";target=_safe(self.assets,name)
        try:
            with target.open("xb") as stream:stream.write(data)
        except FileExistsError:
            if target.read_bytes()!=data:raise ValueError("لوگوی موجود ناسازگار است")
        return name

def _source_rows(p,kind):
    from core.reports.aec_office_package import audit_boq,flatten_takeoffs
    from core.takeoff.manual_formula_amendments import effective_takeoffs
    from core.drawings.takeoff_bridge import session_to_boq_rows,verified_measurement_identities
    audit_boq(p,require_pricebook=kind=="estimate")
    records=effective_takeoffs(p,p.get("takeoffs") or []);record_by_source={r["source_id"]:r for r in records}
    groups={}
    for r in records:
        if r.get("domain")=="drawing":groups.setdefault(r["drawing_session_id"],[]).append(r["drawing_takeoff_id"])
    hashes={}
    for sid,ids in groups.items():
        record=p["drawing_sessions"][sid]
        measured=session_to_boq_rows(record["session"],session_id=sid,selected_item_ids=tuple(ids))
        verified_measurement_identities(record,measured);hashes[record["drawing_source"]]=record["source_sha256"]
        for row in measured:
            saved=record_by_source[row["source_id"]]
            params=saved.get("params");proof=saved.get("provenance");quantities=saved.get("quantities")
            if not isinstance(params,dict) or not isinstance(proof,dict) or not isinstance(quantities,list) or len(quantities)!=1:
                raise ValueError("رسید منبع متره ناقص است")
            for receipt in (params,proof):
                if (receipt.get("drawing_session_id")!=sid or receipt.get("source_sha256")!=record["source_sha256"]
                    or receipt.get("measurement_identity")!=row.get("measurement_identity")
                    or receipt.get("geometry_evidence")!=row.get("geometry_evidence") or receipt.get("page")!=row["page"]):
                    raise ValueError("رسید منبع متره با نشست همخوان نیست")
            if params.get("geometry")!=row["geometry"] or params.get("holes")!=row["holes"]:
                raise ValueError("هندسه یا کسورات منبع متره ناسازگار است")
            q=quantities[0]
            if (q.get("code")!=row["item_code"] or q.get("unit")!=row["unit"] or q.get("source_ref")!=row["source"]
                or decimal_value(q.get("amount"),field="stored drawing quantity")!=decimal_value(row["quantity"],field="source drawing quantity")):
                raise ValueError("کمیت، کد یا واحد متره با منبع اصلی تطبیق ندارد")
    source=flatten_takeoffs({**p,"takeoffs":records}) if kind=="takeoff" else deepcopy(p.get("boq",[]))
    result=[]
    for n,row in enumerate(source,1):
        row=deepcopy(row);ids=row.get("source_ids") or [row.get("source_id")];refs=[record_by_source[x] for x in ids]
        for field in ("floor","member","discipline","zone"):
            names={str((r.get("params",{}).get("geometry_evidence") or {}).get(field+"_id") or r.get(field+"_id") or "") for r in refs}
            row[field]=next(iter(names)) if len(names)==1 else ""
        row["revision"]=str(p["_meta"]["version"]);row["source_id"]=" | ".join(ids);row["source_ids"]=ids;row.setdefault("item_no",n)
        if kind!="estimate":
            for field in ("unit_price","total","price_code"):row.pop(field,None)
        result.append(row)
    for path,expected in hashes.items():
        if not Path(path).is_file() or sha256(Path(path).read_bytes()).hexdigest()!=expected:raise ValueError("منبع نقشه تغییر کرده است")
    return result

def preview_custom_report(service,*,project_id,expected_project_version,template,filters=None):
    t=validate_template(template);p=service.open_project(project_id)
    if not isinstance(p,dict) or type(expected_project_version) is not int or p["_meta"]["version"]!=expected_project_version:raise ValueError("نسخه پروژه تغییر کرده است")
    if filters is not None:_filters(filters)
    wanted=deepcopy(t["filters"] if filters is None else filters)
    rows=[r for r in _source_rows(p,t["report_type"]) if all(str(r.get(k,"")) in values for k,values in wanted.items())]
    if not rows:raise ValueError("هیچ داده معتبر برای فیلتر انتخاب‌شده موجود نیست")
    quantities={unit:str(exact_sum(r["quantity"] for r in rows if r["unit"]==unit)) for unit in sorted({r["unit"] for r in rows})}
    summary={"جمع مقدار ـ "+unit:value for unit,value in quantities.items()} if t["include_totals"] else {}
    if t["include_totals"] and t["report_type"]=="estimate":summary["جمع مبلغ مستند"]=str(exact_sum(r["total"] for r in rows))
    elif t["include_totals"]:summary["وضعیت مبلغ"]="فاقد مبنای قیمت در این گزارش مقداری"
    if summary:summary={"report_totals":summary,"line_count":len(rows)}
    group=t["layout"]["group_by"]
    if group:
        rows.sort(key=lambda r:str(r.get(group,"")));grouped=[]
        for key in dict.fromkeys(str(r.get(group,"")) for r in rows):
            selected=[r for r in rows if str(r.get(group,""))==key];grouped.extend(selected)
            if t["include_totals"]:
                for unit in sorted({r["unit"] for r in selected}):
                    same=[r for r in selected if r["unit"]==unit]
                    subtotal=dict(item_no=None,description="جمع "+(key or "نامشخص"),quantity=str(exact_sum(r["quantity"] for r in same)),unit=unit,source_id=" | ".join(r["source_id"] for r in same),status="خلاصه گزارش")
                    if t["report_type"]=="estimate":subtotal["total"]=str(exact_sum(r["total"] for r in same))
                    grouped.append(subtotal)
        rows=grouped
    return dict(rows=rows,summary=summary,template=t,project_version=p["_meta"]["version"],filters=wanted)

def export_custom_report_batch(service,*,project_id,expected_project_version,requests,destination):
    if not isinstance(requests,tuple) or not 1<=len(requests)<=100:raise ValueError("یک تا صد گزارش باید صریح انتخاب شود")
    target=Path(destination)
    if target.exists() or target.is_symlink():raise FileExistsError("مسیر بسته موجود است")
    if any(p.is_symlink() for p in target.parents):raise ValueError("مسیر خروجی نمادین مجاز نیست")
    store=TemplateStore(service,project_id);before=service.store.current_digest(project_id);snapshot=service.open_project(project_id);prepared=[]
    for req in deepcopy(requests):
        if not isinstance(req,dict) or set(req)!={"template_id","version","filters"}:raise ValueError("انتخاب گزارش نامعتبر است")
        record=store.load(req["template_id"],req["version"])
        preview=preview_custom_report(service,project_id=project_id,expected_project_version=expected_project_version,template=record["template"],filters=req["filters"])
        prepared.append((req,record,preview))
    if service.store.current_digest(project_id)!=before or service.open_project(project_id)!=snapshot:raise RuntimeError("پروژه حین پیش‌نمایش تغییر کرد")
    target.parent.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="structuralpro_p7_batch_",dir=target.parent) as folder:
        stage=Path(folder)/"package";stage.mkdir();files=[];reports=[]
        for n,(req,record,preview) in enumerate(prepared,1):
            t=preview["template"];columns=[(key,FIELDS[key]) for key in t["layout"]["columns"] if t["show_source"] or key!="source_id"]
            if not columns:raise ValueError("ستون قابل نمایش موجود نیست")
            report=ProjectReport(t["layout"]["title"],preview["rows"],preview["summary"],
                {"project_id":project_id,"revision":expected_project_version,"نسخه قالب":str(t["version"]),"_page_format":t["page_format"],"_orientation":t["orientation"],"_custom_header":t["layout"]["header"],"_custom_footer":t["layout"]["footer"]},selected_columns=columns)
            for fmt in ("pdf","xlsx","docx","csv"):
                path=report.export(stage/f"report_{n:03}.{fmt}",fmt);_finalize_and_check(path,fmt,t,store.logo_bytes(t));files.append(path)
            row_refs=[{k:r.get(k) for k in ("item_no","source_ids","quantity","unit","revision")} for r in preview["rows"] if r.get("item_no") is not None]
            reports.append(dict(template_id=t["template_id"],template_version=t["version"],template_sha256=record["sha256"],filters=preview["filters"],row_refs=row_refs,review_status="source_data_not_engineering_approval"))
        manifest=build_bundle_manifest(files,project_id=project_id,revision=expected_project_version)
        manifest.update(schema="structuralpro.phase7.custom-report-batch.v1",project_version=expected_project_version,project_digest=before,created_at=datetime.now(timezone.utc).isoformat(),reports=reports,financial_mutation=False)
        (stage/"manifest.json").write_text(json.dumps(manifest,ensure_ascii=False,sort_keys=True,indent=2),encoding="utf-8")
        if service.store.current_digest(project_id)!=before or service.open_project(project_id)!=snapshot:raise RuntimeError("پروژه پیش از انتشار تغییر کرد")
        for req,record,preview in prepared:
            if store.load(req["template_id"],req["version"])!=record:raise RuntimeError("قالب پیش از انتشار تغییر کرد")
            latest=preview_custom_report(service,project_id=project_id,expected_project_version=expected_project_version,template=record["template"],filters=req["filters"])
            if latest!=preview:raise RuntimeError("منبع گزارش پیش از انتشار تغییر کرد")
        _verify_source_files(snapshot,require_pricebook=any(record["template"]["report_type"]=="estimate" for _,record,_ in prepared))
        if service.store.current_digest(project_id)!=before:raise RuntimeError("پروژه در حین انتشار تغییر کرد")
        if target.exists():raise FileExistsError("مسیر بسته حین صدور ایجاد شد")
        os.rename(stage,target)
    return dict(destination=str(target),manifest=manifest)

def _finalize_and_check(path,fmt,t,logo):
    """Configure existing Office outputs, then reopen each produced format."""
    if fmt=="xlsx":
        from openpyxl import load_workbook
        book=load_workbook(path)
        for sheet in book:
            sheet.page_setup.orientation=t["orientation"];sheet.page_setup.paperSize="8" if t["page_format"]=="A3" else "9"
            sheet.oddHeader.right.text=t["layout"]["header"].replace("&","&&");sheet.oddFooter.center.text=t["layout"]["footer"].replace("&","&&")
        if logo:
            from openpyxl.drawing.image import Image
            sheet=book.create_sheet("نشان گزارش");sheet.sheet_view.rightToLeft=True;sheet["A1"]="نشان معتبر پروژه"
            image=Image(BytesIO(logo));ratio=min(320/image.width,160/image.height);image.width*=ratio;image.height*=ratio
            sheet.add_image(image,"A3");sheet.print_area="A1:H14";sheet.page_setup.orientation=t["orientation"];sheet.page_setup.paperSize="8" if t["page_format"]=="A3" else "9"
        book.save(path);load_workbook(path)
    elif fmt=="docx":
        from docx import Document
        doc=Document(path)
        if not doc.tables:raise ValueError("گزارش Word فاقد جدول است")
        if logo:
            from docx.shared import Inches
            doc.sections[0].header.paragraphs[0].add_run().add_picture(BytesIO(logo),width=Inches(.7));doc.save(path)
    elif fmt=="pdf":
        import pymupdf
        with pymupdf.open(stream=Path(path).read_bytes(),filetype="pdf") as pdf:
            if not len(pdf):raise ValueError("گزارش PDF خالی است")
            if logo:
                for page in pdf:page.insert_image(pymupdf.Rect(26,18,80,36),stream=logo,keep_proportion=True)
                result=pdf.tobytes(garbage=3,deflate=True)
        if logo:Path(path).write_bytes(result)
    else:
        import csv
        with Path(path).open(encoding="utf-8-sig",newline="") as stream:
            if len(list(csv.reader(stream,strict=True)))<2:raise ValueError("گزارش CSV خالی است")


def _verify_source_files(snapshot,*,require_pricebook):
    """Final two-pass physical source freshness, as in the accepted P2 adapter."""
    from hashlib import file_digest
    references={}
    for takeoff in snapshot.get("takeoffs",[]):
        if takeoff.get("domain")!="drawing":continue
        record=snapshot["drawing_sessions"][takeoff["drawing_session_id"]]
        path=Path(record["drawing_source"]);expected=record["source_sha256"]
        if path in references and references[path]!=expected:raise ValueError("نسخه‌های منبع مشترک ناسازگارند")
        references[path]=expected
    if require_pricebook:
        review=snapshot.get("user_pricebook_review") or {}
        path=Path(review["path"]);expected=review["sha256"].lower()
        if path in references and references[path]!=expected:raise ValueError("هش منبع فهرست‌بها ناسازگار است")
        references[path]=expected
    def signature(s):
        return (s.st_dev,s.st_ino,s.st_size,s.st_mtime_ns,s.st_ctime_ns)
    checked={}
    try:
        for path,expected in references.items():
            with path.open("rb") as stream:
                before=signature(os.fstat(stream.fileno()));actual=file_digest(stream,"sha256").hexdigest()
                after=signature(os.fstat(stream.fileno()))
            if before!=after or signature(path.stat())!=after or actual!=expected:raise ValueError("منبع گزارش هنگام کنترل نهایی تغییر کرده است")
            checked[path]=after
        if any(signature(path.stat())!=observed for path,observed in checked.items()):
            raise ValueError("منبع گزارش در حین کنترل نهایی تغییر کرده است")
    except OSError as exc:raise ValueError("منبع گزارش پیش از انتشار در دسترس نیست") from exc
