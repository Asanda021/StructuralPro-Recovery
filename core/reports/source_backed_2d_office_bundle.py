"""Phase 7 P2: source-backed 2D map in real Persian Office/PDF report bundle.

This is an additive, read-only adapter over the *existing* report exporters and
the persisted reviewed takeoff-to-BOQ crosswalk. It cannot grant an engineering
approval, infer absent geometry or change approved quantities/prices.
"""
from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path
import tempfile

from core.reports.bundle import build_bundle_manifest
from core.reports.project_report import build_report
from core.reports.source_backed_2d_overlay import link_source_backed_2d_overlay_to_boq


def _svg_png(svg: str) -> bytes:
    """Rasterize the exact source-backed SVG with existing PyMuPDF.

    Unlike QtSvgRenderer text paths, MuPDF does not require a QGuiApplication.
    PDF, Word, Excel and the PNG sidecar therefore use the *same* map bytes.
    """
    import pymupdf

    if not isinstance(svg, str) or not svg.startswith("<svg "):
        raise ValueError("نقشه برداری SVG معتبر نیست")
    try:
        with pymupdf.open(stream=svg.encode("utf-8"), filetype="svg") as doc:
            if len(doc) != 1:
                raise ValueError("نقشه SVG باید دقیقاً یک صفحه باشد")
            image = doc[0].get_pixmap(
                matrix=pymupdf.Matrix(1.25, 1.25),
                colorspace=pymupdf.csRGB,
                alpha=False,
            )
            if image.width <= 0 or image.height <= 0 or image.width > 4096 or image.height > 4096:
                raise ValueError("ابعاد تصویر SVG خارج از دامنه مجاز است")
            png = image.tobytes("png")
    except (RuntimeError, TypeError) as exc:
        raise ValueError("رندر نقشه SVG ناموفق بود") from exc
    if png[:8] != bytes((137, 80, 78, 71, 13, 10, 26, 10)):
        raise RuntimeError("خروجی PNG ناقص است")
    return png

def _append_pdf_image(path: Path, png: bytes) -> None:
    import pymupdf
    with pymupdf.open(path) as doc:
        previous_pages = len(doc)
        page = doc.new_page(width=841.89, height=595.28)  # A4 landscape
        # Dedicated bounded drawing page; the existing pages and RTL headings
        # remain untouched. The source identity appears in report + manifest.
        box = pymupdf.Rect(38, 34, 804, 560)
        page.insert_image(box, stream=png, keep_proportion=True)
        if len(doc) != previous_pages + 1:
            raise RuntimeError("صفحه نقشه به PDF گزارش اضافه نشد")
        data = doc.tobytes(garbage=3, deflate=True)
    path.write_bytes(data)


def _append_docx_image(path: Path, png: bytes, source_hash: str) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches

    doc = Document(path)
    doc.add_page_break()
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    paragraph._p.get_or_add_pPr().append(bidi)
    paragraph.add_run("نقشه دوبعدی محدوده متره مستند، نه پلان کامل ساختمان")
    image_paragraph = doc.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.add_run().add_picture(BytesIO(png), width=Inches(8.4))
    doc.add_paragraph("SHA-256: " + source_hash)
    doc.save(path)


def _append_xlsx_image(path: Path, png: bytes, links: list[dict], source_hash: str) -> None:
    from openpyxl import load_workbook
    from openpyxl.drawing.image import Image
    from core.reports.exporters import style_workbook

    book = load_workbook(path)
    if "نقشه مستند" in book.sheetnames:
        raise ValueError("برگه نقشه مستند تکراری است")
    sheet = book.create_sheet("نقشه مستند")
    sheet.sheet_view.rightToLeft = True
    sheet["A1"] = "نقشه دوبعدی محدوده متره مستند (نه پلان کامل)"
    sheet["A2"] = "SHA-256 منبع"
    sheet["B2"] = source_hash
    sheet.append(["شناسه عضو", "طبقه", "ردیف BOQ", "شناسه منبع", "مقدار مستند", "واحد"])
    for link in links:
        sheet.append([
            link["member_id"], link["floor_id"], link["boq_row_id"],
            link["source_id"], link["quantity"], link["unit"],
        ])
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 70
    sheet.column_dimensions["C"].width = 24
    sheet.column_dimensions["D"].width = 60
    sheet.column_dimensions["E"].width = 24
    sheet.column_dimensions["F"].width = 18
    picture = Image(BytesIO(png))
    picture.width = 840
    picture.height = 630
    sheet.add_image(picture, f"A{max(8, len(links) + 6)}")
    style_workbook(book, "نقشه مستند متره")
    # Existing style_workbook sets print_area to occupied CELLS only. Images
    # extend below those cells, so extend the print area to include the map.
    sheet.print_area = "A1:L50"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    book.save(path)


def _table_rows(snapshot: dict, links: list[dict]) -> list[dict]:
    """Report the actual persisted, one-source BOQ quantities, not SVG numbers."""
    boq = snapshot.get("boq", [])
    result = []
    for link in links:
        no = int(link["boq_row_id"].removeprefix("boq-item-no:"))
        matches = [row for row in boq if row.get("item_no") == no]
        if len(matches) != 1:
            raise ValueError("ردیف گزارش با BOQ ثبت‌شده یکتا نیست")
        row = matches[0]
        if row.get("source_ids") != [link["source_id"]]:
            raise ValueError("ردیف گزارش، منبع دیگری دارد")
        result.append({
            "item_no": row["item_no"],
            "item_code": row["item_code"],
            "description": row["description"],
            "quantity": row["quantity"],
            "unit": row["unit"],
            "source_id": link["source_id"],
            "notes": "عضو: " + link["member_id"] + " | طبقه: " + link["floor_id"],
        })
    return result


def export_verified_2d_office_bundle(
    service, *, project_id: str, session_id: str,
    expected_project_version: int, expected_session_revision: int,
    revision: str, page: int, selected_takeoff_ids: tuple[str, ...],
    destination: str | Path,
) -> dict:
    """Create a complete multi-format report bundle only for a *current* snapshot.

    Before atomic publication, re-check exact project digest, version, source
    file SHA and the entire original crosswalk. An externally changed project
    or drawing invalidates the whole staging directory; no partial publication.
    """
    target = Path(destination)
    if target.exists():
        raise FileExistsError("مسیر بسته خروجی از قبل موجود است")
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.is_symlink() or target.parent.is_symlink():
        raise ValueError("مسیر نمادین خروجی پذیرفته نیست")
    before_digest = service.store.current_digest(project_id)
    snapshot = service.open_project(project_id)
    if not isinstance(snapshot, dict):
        raise KeyError(project_id)
    if service.store.current_digest(project_id) != before_digest:
        raise RuntimeError("پروژه هنگام تهیه snapshot تغییر کرد")
    args = dict(
        session_id=session_id, expected_project_version=expected_project_version,
        expected_session_revision=expected_session_revision,
        revision=revision, page=page, selected_takeoff_ids=selected_takeoff_ids,
    )
    source = link_source_backed_2d_overlay_to_boq(snapshot, **args)
    linked = source["manifest"]["links"]
    rows = _table_rows(snapshot, linked)
    report = build_report(
        "گزارش نقشه و ریزمتره مستند", rows,
        metadata={
            "project_id": project_id,
            "revision": revision,
            "source_sha256": source["manifest"]["source_sha256"],
            "شناسه نشست": session_id,
            "نسخه نشست": expected_session_revision,
            "نسخه پروژه": expected_project_version,
            "توضیح": "نقشه فقط محدوده برداشت دارای هندسه معتبر را نشان می‌دهد؛ تأیید طراحی نیست",
        },
    )
    if not report.validate()["valid"]:
        raise ValueError("گزارش متره مستند با اعتبارسنجی اصلی سازگار نیست")
    png = _svg_png(source["svg"])

    with tempfile.TemporaryDirectory(prefix="structuralpro_p7_p2_", dir=target.parent) as scratch:
        stage = Path(scratch) / "package"
        stage.mkdir()
        pdf = report.export(stage / "report.pdf", "pdf")
        docx = report.export(stage / "report.docx", "docx")
        xlsx = report.export(stage / "report.xlsx", "xlsx")
        csv = report.export(stage / "report.csv", "csv")
        (stage / "overlay.svg").write_text(source["svg"], encoding="utf-8")
        (stage / "overlay.png").write_bytes(png)
        original_sha = source["manifest"]["source_sha256"]
        _append_pdf_image(pdf, png)
        _append_docx_image(docx, png, original_sha)
        _append_xlsx_image(xlsx, png, linked, original_sha)
        items = [pdf, docx, xlsx, csv, stage / "overlay.svg", stage / "overlay.png"]
        file_manifest = build_bundle_manifest(items, project_id=project_id, revision=revision)
        file_manifest.update({
            "schema": "structuralpro.phase7.source-backed-2d-office-bundle.v1",
            "snapshot_project_version": expected_project_version,
            "snapshot_session_revision": expected_session_revision,
            "snapshot_digest": before_digest,
            "overlay": deepcopy(source["manifest"]),
            "status": "source_verified_not_engineering_approval",
            "quantity_authority": False,
            "financial_mutation": False,
            "csv_image_mode": "overlay.svg and overlay.png are separate sidecars",
        })
        manifest = stage / "manifest.json"
        manifest.write_text(json.dumps(file_manifest, ensure_ascii=False, sort_keys=True, indent=2), encoding="utf-8")
        if any(not p.is_file() or p.stat().st_size == 0 for p in [*items, manifest]):
            raise RuntimeError("بسته نقشه‌دار دارای فایل خالی است")
        # Fresh retrieval and full crosswalk, not a comparison with a cached
        # build result. Verify both before and after to detect concurrent writes.
        if service.store.current_digest(project_id) != before_digest:
            raise RuntimeError("پروژه پیش از انتشار بسته تغییر کرده است")
        latest = service.open_project(project_id)
        if latest != snapshot or link_source_backed_2d_overlay_to_boq(latest, **args) != source:
            raise RuntimeError("منبع یا نسخه گزارش با داده فعلی تطبیق ندارد")
        if service.store.current_digest(project_id) != before_digest:
            raise RuntimeError("پروژه در حین انتشار تغییر کرده است")
        if target.exists():
            raise FileExistsError("مسیر خروجی حین تهیه بسته ایجاد شده است")
        os.rename(stage, target)
    return {"destination": str(target), "manifest": file_manifest}


def export_verified_2d_multifloor_bundle(
    service, *, project_id: str, expected_project_version: int,
    revision: str, selections: tuple[dict, ...], destination: str | Path,
) -> dict:
    """One read-only report over explicit, independently verified page selections.

    Geometry/BOQ validation and all four exporters are the original engines.
    Each map retains its session/page/revision identity. Repeated BOQ rows are
    rejected across pages, never silently summed or attributed twice.
    """
    if not isinstance(selections, tuple) or not 1 <= len(selections) <= 100:
        raise ValueError("انتخاب صریح یک تا صد صفحه مستند الزامی است")
    required = {"session_id", "expected_session_revision", "page", "selected_takeoff_ids"}
    selections = deepcopy(selections)
    if any(not isinstance(s, dict) or set(s) != required for s in selections):
        raise ValueError("مشخصات انتخاب صفحه نامعتبر است")
    target = Path(destination)
    if target.exists() or target.is_symlink():
        raise FileExistsError("مسیر خروجی از قبل موجود است")
    if any(p.is_symlink() for p in target.parents):
        raise ValueError("مسیر نمادین خروجی پذیرفته نیست")
    digest = service.store.current_digest(project_id)
    snapshot = service.open_project(project_id)
    if not isinstance(snapshot, dict) or snapshot.get("id") != project_id:
        raise ValueError("پروژه معتبر موجود نیست")
    if service.store.current_digest(project_id) != digest:
        raise RuntimeError("پروژه هنگام تهیه snapshot تغییر کرد")
    args = [dict(s, expected_project_version=expected_project_version, revision=revision) for s in selections]
    sources = [link_source_backed_2d_overlay_to_boq(snapshot, **s) for s in args]
    seen = set()
    rows = []
    for source in sources:
        links = source["manifest"]["links"]
        for link in links:
            if link["boq_row_id"] in seen:
                raise ValueError("ردیف BOQ در صفحات منتخب تکراری است")
            seen.add(link["boq_row_id"])
        rows.extend(_table_rows(snapshot, links))
    target.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="structuralpro_p7_multi_", dir=target.parent) as tmp:
        stage = Path(tmp) / "package"
        stage.mkdir()
        figures, pages, items = [], [], []
        for n, source in enumerate(sources, 1):
            svg, png = stage / f"map_{n:03}.svg", stage / f"map_{n:03}.png"
            svg.write_text(source["svg"], encoding="utf-8")
            png.write_bytes(_svg_png(source["svg"]))
            page = deepcopy(source["manifest"])
            page.update(svg=svg.name, png=png.name, report_page_id=f"map-{n:03}")
            pages.append(page)
            refs = " | ".join(
                f"{l['floor_id']} / {l['member_id']} / {l['boq_row_id']} / "
                + ("، ".join(l["opening_refs"]) or "بدون بازشوی ثبت‌شده")
                for l in page["links"]
            )
            caption = f"نقشه {n} ـ صفحه منبع {page['page']} ـ نسخه {revision} ـ {refs}"
            figures.append(dict(path=str(png), caption=caption))
            items.extend((svg, png))
        report = build_report("نقشه‌ها و ریزمتره چندطبقه مستند", rows, metadata={
            "project_id": project_id, "revision": revision,
            "نسخه پروژه": expected_project_version,
            "توضیح": "فقط محدوده برداشت معتبر؛ نه پلان کامل ساختمان و نه تأیید طراحی",
            "_figures": figures,
        })
        for fmt in ("pdf", "xlsx", "docx", "csv"):
            items.append(report.export(stage / f"report.{fmt}", fmt))
        # An exact, searchable map-to-BOQ index also lives in the workbook.
        from openpyxl import load_workbook
        from core.reports.exporters import style_workbook, write_excel_cell
        book = load_workbook(stage / "report.xlsx")
        index = book.create_sheet("فهرست نقشه‌ها")
        index.append(["نقشه", "نشست", "نسخه نشست", "صفحه منبع", "طبقه / عضو / ردیف BOQ", "هش منبع", "پیوند ریزمتره"])
        detail = book["ریز متره"]
        number_column = next(i for i, (key, _) in enumerate(report.columns(), 1) if key == "item_no")
        header_row = next(r for r in range(1, detail.max_row+1) if detail.cell(r, number_column).value == "ردیف")
        n = 2
        for page in pages:
            for link in page["links"]:
                number = int(link["boq_row_id"].removeprefix("boq-item-no:"))
                matches = [r for r in range(header_row+1, detail.max_row+1) if detail.cell(r, number_column).value == number]
                if len(matches) != 1:
                    raise ValueError("پیوند سطر گزارش با ردیف BOQ یکتا نیست")
                values = [page["svg"], page["session_id"], page["session_revision"], page["page"],
                          f"{link['floor_id']} / {link['member_id']} / {link['boq_row_id']}",
                          page["source_sha256"], link["boq_row_id"]]
                for col, value in enumerate(values, 1):
                    write_excel_cell(index.cell(n, col), value)
                index.cell(n, 1).hyperlink = page["svg"]
                index.cell(n, 7).hyperlink = f"#'ریز متره'!A{matches[0]}"
                n += 1
        # Preserve the existing exporters' image-aware print areas.
        areas = {s.title: s.print_area for s in book}
        style_workbook(book, report.project_name)
        for sheet in book:
            if areas[sheet.title]: sheet.print_area = areas[sheet.title]
        book.save(stage / "report.xlsx")
        # Successful return/nonzero size from an exporter does not establish a
        # readable Office file. Validate containers and selected source rows in
        # staging, before constructing an accepted manifest or publishing.
        try:
            import csv
            import pymupdf
            from docx import Document
            from zipfile import ZipFile
            # MuPDF's failed filename constructor may retain a native handle in
            # its traceback on Windows. Parse bytes so failure cannot lock the
            # staging file and prevent atomic cleanup of a rejected package.
            with pymupdf.open(stream=(stage / "report.pdf").read_bytes(), filetype="pdf") as pdf:
                if sum(bool(p.get_images()) for p in pdf) != len(pages):
                    raise ValueError("تعداد نقشه‌های PDF با انتخاب تطبیق ندارد")
            with ZipFile(stage / "report.docx") as archive:
                if archive.testzip() is not None:
                    raise ValueError("بسته Word ناقص است")
            if len(Document(stage / "report.docx").inline_shapes) != len(pages):
                raise ValueError("تعداد نقشه‌های Word با انتخاب تطبیق ندارد")
            checked_book = load_workbook(stage / "report.xlsx")
            if sum(len(s._images) for s in checked_book) != len(pages):
                raise ValueError("تعداد نقشه‌های Excel با انتخاب تطبیق ندارد")
            with (stage / "report.csv").open(encoding="utf-8-sig", newline="") as stream:
                csv_rows = list(csv.reader(stream, strict=True))
            if any(not any(row["source_id"] in cells for cells in csv_rows) for row in rows):
                raise ValueError("منابع ریزمتره در CSV بازخوانی نشد")
        except Exception as exc:
            raise ValueError("بازخوانی بسته چهارقالبی پیش از انتشار ناموفق بود") from exc
        manifest = build_bundle_manifest(items, project_id=project_id, revision=revision)
        manifest.update(schema="structuralpro.phase7.source-backed-2d-multifloor.v1",
                        snapshot_digest=digest, snapshot_project_version=expected_project_version,
                        pages=pages, quantity_authority=False, financial_mutation=False,
                        status="source_verified_not_engineering_approval",
                        csv_image_mode="separate SVG/PNG sidecars")
        (stage / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, sort_keys=True, indent=2), encoding="utf-8")
        if any(not p.is_file() or not p.stat().st_size for p in items):
            raise RuntimeError("بسته چندطبقه دارای فایل خالی است")
        if service.store.current_digest(project_id) != digest:
            raise RuntimeError("پروژه پیش از انتشار تغییر کرده است")
        latest = service.open_project(project_id)
        if latest != snapshot:
            raise RuntimeError("نسخه پروژه پیش از انتشار تغییر کرده است")
        for options, source in zip(args, sources):
            if link_source_backed_2d_overlay_to_boq(latest, **options) != source:
                raise RuntimeError("منبع یا رسید صفحه پیش از انتشار تغییر کرده است")
        # Crosswalk verification is sequential and may itself be slow. Rehash
        # every unique physical source after ALL crosswalk checks; retain and
        # compare file identity/size/timestamps after the hash pass as well.
        # This is optimistic freshness, not an OS transaction over external apps.
        from hashlib import file_digest
        def signature(stat):
            return (stat.st_dev, stat.st_ino, stat.st_size, stat.st_mtime_ns, stat.st_ctime_ns)
        checked_sources = {}
        for page in pages:
            path = Path(latest["drawing_sessions"][page["session_id"]]["drawing_source"])
            if path in checked_sources:
                continue
            try:
                with path.open("rb") as stream:
                    before_stat = os.fstat(stream.fileno())
                    actual = file_digest(stream, "sha256").hexdigest()
                    after_stat = os.fstat(stream.fileno())
                if signature(before_stat) != signature(after_stat) or signature(path.stat()) != signature(after_stat) or actual != page["source_sha256"]:
                    raise ValueError("منبع نقشه هنگام بررسی نهایی تغییر کرده است")
                checked_sources[path] = signature(after_stat)
            except OSError as exc:
                raise ValueError("منبع نقشه پیش از انتشار در دسترس نیست") from exc
        if any(signature(path.stat()) != stat for path, stat in checked_sources.items()):
            raise ValueError("منبع نقشه در حین بررسی نهایی تغییر کرده است")
        if service.store.current_digest(project_id) != digest:
            raise RuntimeError("پروژه در حین انتشار تغییر کرده است")
        if target.exists():
            raise FileExistsError("مسیر خروجی حین تهیه بسته ایجاد شده است")
        os.rename(stage, target)
    return {"destination": str(target), "manifest": manifest}
