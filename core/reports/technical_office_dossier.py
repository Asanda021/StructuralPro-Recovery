"""Phase7 P4: verified batch of existing site-office and project-control reports.

This module is a strict orchestration/acceptance layer, not a parallel schedule,
progress, materials, financial or reporting engine. It never changes a project.
Missing approved sources are blocked, not silently turned into 0% or a chart.
"""
from __future__ import annotations

from datetime import date
from hashlib import sha256
import json
import os
from pathlib import Path
import tempfile

from core.reports.bundle import build_bundle_manifest
from core.reports.custom_batch import _verify_source_files

CONTROL_KINDS = frozenset(("wbs", "activities", "schedule", "progress", "curve"))
SITE_KINDS = frozenset((
    "lookahead", "resources", "materials", "purchases", "daily",
    "meetings", "obstacles", "correspondence",
))
KINDS = CONTROL_KINDS | SITE_KINDS
FORMATS = ("pdf", "xlsx", "docx", "csv")


def _require_records(service, snapshot, project_id, kind, as_of):
    if kind == "wbs":
        data = snapshot.get("project_wbs")
    elif kind in ("activities", "schedule"):
        data = snapshot.get("project_schedule_tasks")
    elif kind in ("progress", "curve"):
        if not snapshot.get("schedule_baselines"):
            raise ValueError("مبنای زمان‌بندی تأییدشده ثبت نشده است")
        if kind == "progress":
            p = service.project_progress(project_id, as_of=as_of)
            if p.get("actual") is None or p.get("planned") is None:
                raise ValueError("پیشرفت واقعی یا برنامه مصوب قابل تعیین نیست")
            return
        data = service.project_progress_curve(project_id)
        if not data or any(r.get("actual") is None or r.get("planned") is None for r in data):
            raise ValueError("منحنی پیشرفت به داده و مبنای مصوب نیاز دارد")
        return
    elif kind == "lookahead":
        data = service.site_lookahead(project_id, start=as_of, days=7).get("activities")
    elif kind == "resources":
        data = snapshot.get("project_resources")
    elif kind in ("materials", "purchases"):
        material = service.site_materials(project_id, as_of=as_of)
        data = material.get("requirements") if kind == "materials" else material.get("purchase_drafts")
    else:
        from core.projects.site_office import snapshot as office_snapshot
        office = office_snapshot(snapshot, as_of=as_of)
        data = office[{
            "daily": "daily_reports", "meetings": "meetings",
            "obstacles": "obstacles", "correspondence": "correspondence",
        }[kind]]
    if not isinstance(data, list) or not data:
        raise ValueError("سند معتبر لازم برای گزارش " + kind + " ثبت نشده است")


def _read_back(path, fmt, kind):
    if not path.is_file() or path.stat().st_size == 0:
        raise ValueError("خروجی خالی یا گمشده است")
    if fmt == "pdf":
        import pymupdf
        with pymupdf.open(stream=path.read_bytes(), filetype="pdf") as doc:
            if not len(doc):
                raise ValueError("PDF فاقد صفحه است")
            if kind in ("schedule", "curve") and not any(p.get_images() for p in doc):
                raise ValueError("نمودار PDF مستند حذف شده است")
    elif fmt == "xlsx":
        from openpyxl import load_workbook
        with path.open("rb") as stream:
            book = load_workbook(stream)
        if not all(sheet.sheet_view.rightToLeft for sheet in book):
            raise ValueError("صفحات Excel باید RTL باشند")
        if kind in ("schedule", "curve") and not any(sheet._images for sheet in book):
            raise ValueError("نمودار Excel از خروجی حذف شده است")
        book.close()
    elif fmt == "docx":
        from docx import Document
        doc = Document(path)
        if not doc.tables or "w:bidi" not in doc._element.xml:
            raise ValueError("گزارش Word فارسی و جدولی معتبر نیست")
        if kind in ("schedule", "curve") and not doc.inline_shapes:
            raise ValueError("نمودار Word از خروجی حذف شده است")
    else:
        import csv
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream, strict=True))
        if len(rows) < 3:
            raise ValueError("CSV فاقد جدول مستند است")


def export_verified_site_controls_dossier(
    service, *, project_id: str, expected_project_version: int,
    as_of: str, selected_kinds: tuple[str, ...], destination: str | Path,
) -> dict:
    """Atomic 4-format reports from existing services; no inferred progress."""
    if (not isinstance(selected_kinds, tuple) or not selected_kinds or
            len(selected_kinds) > 12 or len(set(selected_kinds)) != len(selected_kinds) or
            any(kind not in KINDS for kind in selected_kinds)):
        raise ValueError("نوع گزارش دفتر فنی باید مجاز، صریح و یکتا باشد")
    try:
        parsed = date.fromisoformat(as_of)
    except (ValueError, TypeError) as exc:
        raise ValueError("تاریخ میلادی مبنا باید معتبر و صریح باشد") from exc
    if parsed.isoformat() != as_of:
        raise ValueError("تاریخ مبنا نرمال و بدون ابهام لازم است")
    target = Path(destination)
    if target.exists() or any(p.is_symlink() for p in (target, *target.parents)):
        raise FileExistsError("مسیر مقصد موجود یا نمادین است")
    digest = service.store.current_digest(project_id)
    snapshot = service.open_project(project_id)
    if (not isinstance(snapshot, dict) or snapshot.get("id") != project_id or
            type(expected_project_version) is not int or
            snapshot.get("_meta", {}).get("version") != expected_project_version):
        raise ValueError("شناسه یا نسخه جاری پروژه مطابقت ندارد")
    _verify_source_files(snapshot, require_pricebook=False)
    for kind in selected_kinds:
        _require_records(service, snapshot, project_id, kind, as_of)
    if service.store.current_digest(project_id) != digest:
        raise RuntimeError("پروژه حین کنترل منابع تغییر کرده است")
    target.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=target.parent, prefix="structuralpro_p7_p4_dossier_") as working:
        stage = Path(working) / "package"
        stage.mkdir()
        items, references = [], []
        for number, kind in enumerate(selected_kinds, 1):
            records = []
            for fmt in FORMATS:
                file = stage / f"{number:02}-{kind}.{fmt}"
                if kind in CONTROL_KINDS:
                    returned = service.project_controls_report(
                        project_id, kind, fmt, file, as_of=as_of)
                else:
                    returned = service.site_report(
                        project_id, kind, fmt, file, as_of=as_of)
                if Path(returned) != file:
                    raise ValueError("صادرکننده اصلی فایل دیگری برگرداند")
                _read_back(file, fmt, kind)
                records.append(file.name)
                items.append(file)
            references.append({"kind": kind, "files": records})
        # Existing chart providers write PNG figure assets next to exported
        # Office/PDF files. Include and validate those real sidecars, too.
        by_format = {p.name for p in items}
        sidecars = sorted(p for p in stage.iterdir() if p.is_file() and p.name not in by_format)
        if any(p.suffix.lower() != ".png" for p in sidecars):
            raise ValueError("پیوست ناشناخته در پوشه صدور گزارش است")
        if sidecars:
            from PIL import Image
            for png in sidecars:
                with Image.open(png) as image:
                    if image.format != "PNG" or image.width <= 0 or image.height <= 0:
                        raise ValueError("نمودار PNG قابل بازخوانی نیست")
                    image.verify()
        manifest = {
            **build_bundle_manifest([*items, *sidecars], project_id=project_id,
                                    revision=str(expected_project_version)),
            "schema": "structuralpro.phase7.source-backed-site-controls-dossier.v1",
            "as_of": as_of,
            "project_version": expected_project_version,
            "source_project_digest": digest,
            "reports": references,
            "figure_sidecars": [p.name for p in sidecars],
            "report_count": len(references),
            "source_status": "original_site_and_project_control_services",
            "physical_office_print_status": "NOT_TESTED",
            "financial_mutation": False,
        }
        (stage / "manifest.json").write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True),
            encoding="utf-8",
        )
        _verify_source_files(snapshot, require_pricebook=False)
        if (service.store.current_digest(project_id) != digest or
                service.open_project(project_id) != snapshot):
            raise RuntimeError("پروژه پیش از انتشار گزارش تغییر کرده است")
        for kind in selected_kinds:
            _require_records(service, snapshot, project_id, kind, as_of)
        if service.store.current_digest(project_id) != digest:
            raise RuntimeError("پروژه هنگام کنترل نهایی گزارش تغییر کرده است")
        if target.exists():
            raise FileExistsError("پوشه خروجی هنگام صدور ایجاد شده است")
        os.rename(stage, target)
    return {"destination": str(target), "manifest": manifest}
