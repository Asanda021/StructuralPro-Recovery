"""Reports from the versioned manual register, preserving exact textual quantities."""
import ast
from decimal import Decimal
from collections import defaultdict
from pathlib import Path
from core.takeoff.manual_formula_register import load,catalog
from core.commercial.exact_values import number,exact
from core.reports.exporters import export_pdf

UNITS={'m':'متر','m2':'متر مربع','m3':'متر مکعب','kg':'کیلوگرم','ton':'تن','عدد':'عدد','cm':'سانتی‌متر','mm':'میلی‌متر','cm2':'سانتی‌متر مربع','mm2':'میلی‌متر مربع','cm3':'سانتی‌متر مکعب','mm3':'میلی‌متر مکعب'}
STATES={'draft':'بررسی‌نشده','reviewed':'بازبینی‌شده','committed':'قطعی'}

def safe_integer(v):
    value=number(v,'عدد فرمول',True)
    if value!=value.to_integral_value() or abs(value)>=10**15:raise ValueError('حفظ دقت15رقم Excel تضمین نشده')
    return int(value)

def excel_expression(expression,variables,references):
    tree=ast.parse(expression,mode='eval')
    def visit(node):
        if isinstance(node,ast.Constant) and type(node.value) in (int,float):
            v=safe_integer(ast.get_source_segment(expression,node));return str(v),v
        if isinstance(node,ast.Name):
            cell,value=(variables|references)[node.id];return cell,safe_integer(value)
        if isinstance(node,ast.UnaryOp) and isinstance(node.op,(ast.UAdd,ast.USub)):
            code,v=visit(node.operand);return '('+('+' if isinstance(node.op,ast.UAdd) else '-')+code+')',safe_integer(v if isinstance(node.op,ast.UAdd) else -v)
        if isinstance(node,ast.BinOp) and type(node.op) in (ast.Add,ast.Sub,ast.Mult,ast.Div,ast.Pow):
            a,x=visit(node.left);b,y=visit(node.right)
            if isinstance(node.op,ast.Div):
                if not y or x%y:raise ValueError('تقسیم دقیق صحیح نیست')
                value=x//y;op='/'
            elif isinstance(node.op,ast.Pow):
                if y<0 or y>32:raise ValueError('توان برای Excel امن نیست')
                value=x**y;op='^'
            elif isinstance(node.op,ast.Mult):value=x*y;op='*'
            elif isinstance(node.op,ast.Sub):value=x-y;op='-'
            else:value=x+y;op='+'
            return '('+a+op+b+')',safe_integer(value)
        raise ValueError('گره غیرمجاز Excel')
    formula,value=visit(tree.body);return '='+formula,value

@exact
def export(service,pid,path,*,format):
    reg=load(service,pid);project=service.open_project(pid);titles={r['code']:r['title'] for r in catalog()}
    details=[];parameters=[];sources=[];deductions=[];history=[];links=[];unreviewed=[];summaries={k:defaultdict(lambda:Decimal(0)) for k in ('خلاصه متره','هر طبقه','هر آیتم','هر عضو')}
    keys=sorted(reg)
    for key in keys:
        row=reg[key];d=row['data'];q=str(row['quantity']);unit=UNITS[d['unit']]
        entry=[key,d['floor'],d['member'],titles[d['family']],d['title'],q,unit,d['expression'],d['basis'],d['source_id'],str(row['revision']),STATES[row['status']]];details.append(entry)
        if row['status']=='draft':unreviewed.append(entry)
        for name,value in d['variables'].items():parameters.append([key,name,str(value),d['basis']])
        sources.append([key,d['source_id'],d['basis'],', '.join(d['references'].values()),row.get('review',{}).get('reviewer','')])
        for item in d['deductions']:deductions.append([key,item['id'],str(item['quantity']),UNITS[item['unit']],item['basis'],item.get('source_id',item['id'])])
        matched=[b for b in project.get('boq',[]) if d['source_id'] in b.get('source_ids',[]) or b.get('source_id')==d['source_id']]
        links.append([key,d['source_id'],STATES[row['status']],str(len(matched)),q,unit])
        for name,group in [('خلاصه متره','پروژه'),('هر طبقه',d['floor']),('هر آیتم',titles[d['family']]),('هر عضو',d['member'])]:summaries[name][(group,unit,row['status'])]+=number(row['quantity'],'جمع متره')
    for item in project.get('manual_formula_history',[]):history.append([item['row_id'],str(item['revision']),item['reason'],item['at']])
    headers=['شناسه','طبقه','عضو','خانواده','شرح','مقدار دقیق','واحد','فرمول','مبنا و واحد ورودی','منبع','نسخه','وضعیت']
    datasets={'ریز متره':(headers,details),'ورودی‌ها':(['شناسه','پارامتر','مقدار دقیق','مبنا'],parameters),'کسورات':(['شناسه','منبع کسر','مقدار','واحد','مبنا','منبع هندسی'],deductions),'منابع و مدارک':(['شناسه','منبع','مبنا','وابستگی','بازبین'],sources),'تاریخچه':(['شناسه','نسخه','علت','زمان ثبت UTC'],history),'بررسی‌نشده':(headers,unreviewed),'ارتباط BOQ':(['شناسه','منبع','وضعیت','ردیف مرتبط','مقدار','واحد'],links)}
    for name,groups in summaries.items():datasets[name]=(['گروه','واحد','وضعیت','جمع دقیق'],[[a,b,STATES[c],str(v)] for (a,b,c),v in sorted(groups.items())])
    datasets['اصلاحیه قطعی']=(['شناسه اصلاحیه','ردیف','نسخه مبنا','علت','سازنده','بازبین','مقدار پیشنهادی','وضعیت'],[[a['id'],a['row_id'],str(a['base_revision']),a['reason'],a['created_by'],a.get('review',{}).get('reviewer',''),str(a['quantity']),{'draft':'بررسی‌نشده','reviewed':'بازبینی‌شده','applied':'اعمال‌شده'}[a['status']]] for a in project.get('manual_formula_amendments',{}).values()])
    from core.takeoff.manual_formula_deductions import summary as deduction_summary
    datasets['قبل و بعد کسر']=(['شناسه','پیش از کسر','خالص','کسر مؤثر','واحد'],[[key,str(v['gross']),str(v['net']),str(v['deducted']),UNITS[v['unit']]] for key in keys if reg[key]['data']['deductions'] for v in [deduction_summary(service,pid,key)]])
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    if format=='pdf':
        # Narrow key/value rows retain all fields without squeezing 12 columns.
        rows=[]
        for title,(cols,records) in datasets.items():
            for record in records:
                for label,value in zip(cols,record):rows.append({'section':title,'row':str(record[0]),'field':label,'value':str(value)})
        if not rows:raise ValueError('داده قابل گزارش موجود نیست')
        return export_pdf(rows,path,title='متره دستی مستند — '+project['name'],columns=[('section','گزارش'),('row','ردیف'),('field','عنوان'),('value','مقدار')])
    if format!='xlsx':raise ValueError('قالب گزارش معتبر نیست')
    from openpyxl import Workbook
    from openpyxl.styles import Alignment,Font
    w=Workbook();w.remove(w.active)
    def put(sheet,row,col,val,formula=False):
        cell=sheet.cell(row,col);cell.value=val if formula else str(val);cell.data_type='f' if formula else 's';cell.alignment=Alignment(horizontal='right',readingOrder=2,wrap_text=True);cell.font=Font(name='DejaVu Sans',size=10)
    for name,(cols,records) in datasets.items():
        sheet=w.create_sheet(name);sheet.sheet_view.rightToLeft=True;sheet.freeze_panes='A2';sheet.auto_filter.ref=f'A1:{__import__("openpyxl").utils.get_column_letter(len(cols))}{len(records)+1}'
        for i,values in enumerate([cols]+records,1):
            for j,val in enumerate(values,1):put(sheet,i,j,val)
        for j in range(1,len(cols)+1):sheet.column_dimensions[__import__('openpyxl').utils.get_column_letter(j)].width=24
    calc=w.create_sheet('محاسبه امن');calc.sheet_view.rightToLeft=True
    for j,v in enumerate(['شناسه','مقدار دقیق مرجع','فرمول زنده یا مقدار دقیق','هشدار'],1):put(calc,1,j,v)
    parameter_locations={};pindex={}
    for i,(key,name,value,basis) in enumerate(parameters,2):
        parameter_locations[(key,name)]=(f"'ورودی‌ها'!C{i}",value);pindex[(key,name)]=i
    cells={key:i+2 for i,key in enumerate(keys)};live={};active=set()
    def build(key):
        if key in live:return live[key]
        if key in active:raise ValueError('حلقه')
        active.add(key);row=reg[key];d=row['data'];refs={}
        try:
            for name,target in d['references'].items():
                if build(target) is None:raise ValueError('مرجع فرمول زنده دقیق نیست')
                refs[name]=(f"'محاسبه امن'!C{cells[target]}",reg[target]['quantity'])
            formula,value=excel_expression(d['expression'],{name:parameter_locations[(key,name)] for name in d['variables']},refs)
            if number(value,'Excel',True)!=number(row['quantity'],'مرجع',True):raise ValueError('خروجی ناسازگار')
            live[key]=formula
            for name,val in d['variables'].items():w['ورودی‌ها'].cell(pindex[(key,name)],3,int(number(val,'پارامتر',True)))
        except (ValueError,KeyError,OverflowError):live[key]=None
        active.remove(key);return live[key]
    for key in keys:build(key)
    for key in keys:
        i=cells[key];row=reg[key];put(calc,i,1,key);put(calc,i,2,str(row['quantity']));put(calc,i,3,live[key] or str(row['quantity']),formula=bool(live[key]));put(calc,i,4,'دقت متنی محفوظ؛ فرمول زنده به‌علت محدودیت دقت یا عملیات Excel مسدود است' if not live[key] else 'ورودی و تمام نتایج صحیح کمتر از15رقم؛ Excel واقعی باید بازحساب شود؛ تغییر فایل BOQ را تغییر نمی‌دهد')
    from openpyxl.workbook.properties import CalcProperties
    w.calculation=CalcProperties(fullCalcOnLoad=True,forceFullCalc=True,calcMode='auto');w.save(path);return path
