"""Reproducible synthetic integration; never a real-plan engineering approval."""
from pathlib import Path
from copy import deepcopy
from hashlib import sha256
from decimal import Decimal
import json,sys,csv,re,sqlite3
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from core.platform.application import StructuralProApp
from core.drawings.takeoff_session import DrawingTakeoffSession
from core.drawings.graphical_takeoff import Point
from core.drawings.roof_region import bind_source,detail_payload,source_signature
from core.takeoff.reviewed_aec import preview_system


def run(root):
 root=Path(root);root.mkdir(parents=True,exist_ok=True)
 inputs=Path(__file__).resolve().parents[1]/'tests/fixtures/phase2_priority3_closure_inputs.json'
 f=json.loads(inputs.read_text());assert f['synthetic'] and not f['engineering_approval']
 (root/'inputs.json').write_bytes(inputs.read_bytes())
 import pymupdf
 path=root/'synthetic-two-revision-pages.pdf'
 with pymupdf.open() as pdf:
  for opening in (f['initial_opening_pdf'],f['revised_opening_pdf']):
   page=pdf.new_page(width=f['pdf_page_points'][0],height=f['pdf_page_points'][1])
   for points in (f['outer_pdf'],opening):page.draw_polyline([pymupdf.Point(x*72/150,y*72/150) for x,y in points+[points[0]]],color=(0,0,0))
  pdf.save(path)
 app=StructuralProApp(root/'db');app.create_project(f['project_title'],'P');app.add_project_floor('P','طبقه صریح مصنوعی',level='L1');app.add_project_member('P','S1','L1','سقف')
 def measured(page,opening,ref):
  s=DrawingTakeoffSession(str(path));s.calibrate(page,*f['calibration']['page'+str(page)]);s.current_page=page
  item=s.add_polygon([Point(*p) for p in f['outer_pdf']],page=page,topology='exact_planar_region_v2',source='SYNTHETIC source page '+str(page),floor_id='L1',member_id='S1',quantity_owner='SYNTHETIC sole documented roof owner')
  s.add_polygon_opening(item.id,opening,ref)
  saved=app.save_drawing_takeoff_session('P',s)
  return s,item,saved,bind_source(app.load_drawing_takeoff_session('P',saved['session_id']),item.id,local_frame=f['frame'])
 s,item,first,source=measured(1,f['initial_opening_pdf'],'SYNTHETIC opening revision 1')
 assert len(s.find(item.id).geometry_evidence['components'])==2
 d=deepcopy(f['roof_data']);d['source']=source;d['source'].update(plan_ref='SYNTHETIC two-page source',detail_ref='SYNTHETIC initial detail')
 def document(data,name):
  p=root/name;p.write_text(json.dumps(detail_payload(data),ensure_ascii=False,indent=2))
  data['source'].update(detail_path=str(p),detail_sha256=sha256(p.read_bytes()).hexdigest());return p
 document(d,'initial-detail.json');app.save_aec_system_draft('P',d,expected_revision=0)
 def quantities(data,oracle):
  r=preview_system(data,app.takeoff);assert r['ready'],r['blockers'];actual={}
  for row in r['rows']:
   if row['role']=='work':actual[row['item_code']]=actual.get(row['item_code'],Decimal(0))+Decimal(str(row['quantity']))
  assert actual=={k:Decimal(v) for k,v in oracle.items()},(actual,oracle)
  return r
 quantities(d,f['initial_expected']);app.review_roof_system('P','solid_slab','L1','S1',expected_revision=1,reviewer='SYNTHETIC independent reviewer',note=f['oracle_explanation']);app.commit_aec_system('P','solid_slab','L1','S1',expected_revision=2)
 original=deepcopy(app.open_project('P'));(root/'original-project.json').write_text(json.dumps(original,ensure_ascii=False,indent=2))
 revised,j,second,new_source=measured(2,f['revised_opening_pdf'],'SYNTHETIC opening revision 2')
 proposal=deepcopy(app.get_aec_system('P','solid_slab','L1','S1')['data']);proposal.pop('roof_review',None);proposal['source']=new_source;new_source.update(plan_ref=d['source']['plan_ref'],detail_ref='SYNTHETIC revised detail')
 rd=proposal['roof_detail'];hole=[[4,0],[5,0],[5,10],[4,10]];rd['deductions'][0].update(length='1',footprint=hole,source_ref='SYNTHETIC opening revision 2')
 rd['formwork']['faces'][0].update(holes=[hole],opening_refs=['SYNTHETIC opening revision 2'])
 rd['geometry_change']={'policy':'pdf_roof_geometry_change_v1','previous_source_signature':source_signature(d['source']),'new_source_signature':source_signature(proposal['source']),'source_ref':'SYNTHETIC explicit opening correction','approved_by':'SYNTHETIC geometry reviewer'}
 proof=document(proposal,'revised-detail.json');quantities(proposal,f['revised_expected'])
 evidence={'evidence_path':str(proof),'evidence_sha256':sha256(proof.read_bytes()).hexdigest(),'evidence_reference':'SYNTHETIC reviewed revision'}
 a=app.create_roof_amendment('P','solid_slab','L1','S1',proposal,expected_revision=2,created_by='SYNTHETIC author',reason='explicit opening width changes from 2m to 1m',evidence=evidence)
 app.review_beam_amendment('P',a['id'],expected_revision=1,reviewer='SYNTHETIC independent correction reviewer',note=f['oracle_explanation']);app.apply_beam_amendment('P',a['id'],expected_revision=1)
 expected=deepcopy(app.open_project('P'));assert expected['takeoffs']==original['takeoffs'] and expected['reviewed_aec_systems']==original['reviewed_aec_systems']
 actual_boq={}
 for row in expected['boq']:actual_boq[row['item_code']]=actual_boq.get(row['item_code'],Decimal(0))+Decimal(str(row['quantity']))
 assert actual_boq=={k:Decimal(v) for k,v in f['revised_expected'].items()},actual_boq
 negatives=[]
 def rejected(label,fn):
  before=deepcopy(app.open_project('P'))
  try:fn()
  except (ValueError,RuntimeError):pass
  else:raise AssertionError(label+' was allowed')
  assert app.open_project('P')==before;negatives.append(label)
 rejected('duplicate amendment',lambda:app.apply_beam_amendment('P',a['id'],expected_revision=1))
 rejected('stale SQLite/session revision',lambda:app.save_drawing_takeoff_session('P',revised,session_id=second['session_id'],expected_revision=0))
 rejected('generic second payment for same physical roof on old page',lambda:app.commit_drawing_takeoff_to_boq('P',first['session_id'],[item.id],expected_session_revision=1))
 for fmt in ('pdf','xlsx','docx','csv'):app.export_roof_detail_report('P','solid_slab','L1','S1',root/('effective-roof.'+fmt),format=fmt)
 readback=verify_reports(root,expected['beam_amendments'][a['id']]['rows'])
 app.store.close();app=StructuralProApp(root/'db');assert app.open_project('P')==expected
 assert app.store.db.execute('pragma integrity_check').fetchone()[0]=='ok';app.store.close()
 (root/'effective-project.json').write_text(json.dumps(expected,ensure_ascii=False,indent=2))
 result={'synthetic':True,'engineering_approval':False,'initial_geometry_m2':'80','revised_geometry_m2':'90','initial_components':2,'revised_components':2,'initial_quantities':f['initial_expected'],'effective_quantities':f['revised_expected'],'original_takeoffs_unchanged':True,'original_roof_register_unchanged':True,'sqlite_integrity':'ok','reopened_snapshot_equal':True,'negative_checks':negatives,'amendment_id':a['id'],'effective_revision':3,'source_sha256':sha256(path.read_bytes()).hexdigest(),'report_readback':readback,'reports':['effective-roof.'+v for v in ('pdf','xlsx','docx','csv')]}
 (root/'receipt.json').write_text(json.dumps(result,ensure_ascii=False,indent=2));return result



def verify_reports(root,expected_rows):
 """Independent file parsers compare every effective row with persisted journal."""
 import unicodedata
 from openpyxl import load_workbook
 from docx import Document
 import pymupdf
 def text(v):return unicodedata.normalize('NFKC',str(v or '')).strip()
 def numeric(v):return Decimal(text(v).translate(str.maketrans('۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩','01234567890123456789')).replace('٫','.').replace('٬',''))
 def read_rows(matrix):
  headers=None;actual={}
  for row in matrix:
   values=[text(v) for v in row]
   if all(k in values for k in ('مقدار','واحد','منبع','فرمول')):headers={k:values.index(k) for k in ('مقدار','واحد','منبع','فرمول')};continue
   if not headers or len(row)<=max(headers.values()):continue
   source=values[headers['منبع']]
   if source in expected and values[headers['فرمول']]:
    assert source not in actual,'duplicate report row'
    actual[source]=(numeric(row[headers['مقدار']]),{'متر مکعب':'m3','متر مربع':'m2','کیلوگرم':'kg'}.get(values[headers['واحد']],values[headers['واحد']]),values[headers['فرمول']])
  return actual
 expected={r['source_id']:(Decimal(str(r['quantities'][0]['amount'])),r['quantities'][0]['unit'],r['quantities'][0]['formula']) for r in expected_rows}
 with (root/'effective-roof.csv').open(encoding='utf-8-sig',newline='') as f:csv_rows=read_rows(list(csv.reader(f)))
 assert csv_rows==expected,(csv_rows,expected)
 book=load_workbook(root/'effective-roof.xlsx',data_only=True)
 assert all(s.sheet_view.rightToLeft for s in book)
 xlsx_rows=read_rows(list(book.worksheets[0].values));assert xlsx_rows==expected
 document=Document(root/'effective-roof.docx');word_rows=read_rows([[c.text for c in r.cells] for t in document.tables for r in t.rows]);assert word_rows==expected
 assert all('bidiVisual' in t._tbl.xml for t in document.tables)
 pdf_rows={};page_count=0
 with pymupdf.open(root/'effective-roof.pdf') as pdf:
  page_count=len(pdf)
  for page in pdf:
   vertical=[];ys=set()
   for drawing in page.get_drawings():
    for item in drawing['items']:
     if item[0]!='l':continue
     a,b=item[1:3]
     if a.x==b.x:vertical.append((round(a.x,3),min(a.y,b.y),max(a.y,b.y)))
     if a.y==b.y:ys.add(round(a.y,3))
   headers=None
   levels=sorted(ys)
   for y0,y1 in zip(levels,levels[1:]):
    mid=(y0+y1)/2;xs=sorted(set(x for x,lo,hi in vertical if lo<=mid<=hi))
    if len(xs)!=8:continue
    cells=[text(page.get_text('text',clip=pymupdf.Rect(x0+.1,y0+.1,x1-.1,y1-.1))) for x0,x1 in zip(xs,xs[1:])]
    if all(k in cells for k in ('مقدار','منبع','واحد','فرمول')):headers={k:cells.index(k) for k in ('مقدار','منبع','واحد','فرمول')};continue
    if headers is None:continue
    source=''.join(cells[headers['منبع']].split())
    if source in expected and cells[headers['فرمول']]:
     units={'متر مکعب':'m3','متر مربع':'m2','کیلوگرم':'kg','عدد':'عدد'}
     unit=' '.join(cells[headers['واحد']].split());pdf_rows[source]=(numeric(cells[headers['مقدار']]),units[unit])
 assert pdf_rows=={k:v[:2] for k,v in expected.items()},(pdf_rows,expected)
 return {'csv_rows_verified':len(csv_rows),'xlsx_rows_verified':len(xlsx_rows),'docx_rows_verified':len(word_rows),'pdf_rows_verified':len(pdf_rows),'pdf_pages':page_count,'xlsx_rtl':True,'docx_table_rtl':True,'pdf_method':'rendered vector table cells, numeric/unit/source readback; visual QA separately','windows_office_acceptance':False}

if __name__=='__main__':print(json.dumps(run(sys.argv[1]),ensure_ascii=False,indent=2))
