"""Synthetic five-floor acceptance, never a real contract, rate or payment."""
from pathlib import Path
from decimal import Decimal,localcontext,ROUND_HALF_UP
import sys,json,hashlib
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from scripts.phase3_financial_probe import run as prepare
from core.platform.application import StructuralProApp


def run(root):
    # Keep one QApplication alive across both engineering and financial views.
    from PySide6.QtWidgets import QApplication
    qt=QApplication.instance() or QApplication([])
    root=Path(root).resolve();root.mkdir(parents=True,exist_ok=True)
    prepare(root/'engineering')
    a=StructuralProApp(root/'engineering'/'db')
    p=a.open_project('P');p['metadata']['financial_owner_id']='fixture-owner';a.store.save('P',p)
    proof=root/'explicit-contract.txt';proof.write_text('SYNTHETIC TEST ONLY. Sequential estimate factors. Retention .10, advance recovery .10 capped at 100 IRR, other deductions .02, tax/insurance zero. Local attributed roles; no authentication.')
    docs=[{'ref':'TEST-CONTRACT','path':str(proof),'sha256':hashlib.sha256(proof.read_bytes()).hexdigest()}]
    v=lambda:a.open_project('P')['_meta']['version']
    policy={'money':{'decimal_places':2,'rounding':'ROUND_HALF_UP','currency':'IRR'},'authorities':{'writer':['create'],'reviewer':['review'],'approver':['approve','post']}}
    a.configure_statement_amendments('P',owner_id='fixture-owner',policy=policy,source_documents=docs,expected_project_version=v())
    terms={'contract_id':'SYNTHETIC-C1','currency':'IRR','advance_paid':'100','factor_basis':'reviewed_sequential_factors','retention_rate':'.1','advance_recovery_rate':'.1','other_deduction_rate':'.02','tax_rate':'0','insurance_rate':'0'}
    a.configure_statement_workflow('P',actor='fixture-owner',terms=terms,source_documents=docs,expected_version=v())
    quantities={'concrete':Decimal(4),'formwork':Decimal(16),'rebar':Decimal('71.04'),'wall':Decimal(13)}
    prices={'concrete':Decimal(100),'formwork':Decimal(20),'rebar':Decimal(10),'wall':Decimal(50)}
    def issue(ref,fraction):
        d=a.statement_workflow('P','create',actor='writer',expected_version=v(),reference=ref,date='2026-10-23',current_quantities={k:str(q*fraction) for k,q in quantities.items()})
        for op,who in [('review','reviewer'),('approve','approver'),('issue','approver')]:d=a.statement_workflow('P',op,actor=who,expected_version=v(),id=d['id'])
        return d
    first=issue('SYNTHETIC-1',Decimal('.5'));second=issue('SYNTHETIC-2',Decimal('.25'))
    first_hash=hashlib.sha256(json.dumps(first['period'],sort_keys=True).encode()).hexdigest()
    def oracle(q,available):
        with localcontext() as ctx:
            ctx.prec=100
            gross=sum(q[k]*prices[k]*Decimal('1.1')*Decimal('1.05') for k in q)
            retention=gross*Decimal('.1');recovery=min(gross*Decimal('.1'),available);other=gross*Decimal('.02');net=gross-retention-recovery-other
            return {k:format(x.quantize(Decimal('.01'),rounding=ROUND_HALF_UP),'f') for k,x in [('gross_current',gross),('retention',retention),('advance_recovery',recovery),('other_deductions',other),('payable_current',net)]}
    one={k:q*Decimal('.5') for k,q in quantities.items()};two={k:q*Decimal('.25') for k,q in quantities.items()}
    for period,expected in [(first['period'],oracle(one,Decimal(100))),(second['period'],oracle(two,Decimal(0)))]:
        for k,value in expected.items():assert Decimal(str(period[k]))==Decimal(value),(k,period[k],value)
    source=next(row['source_ids'][0] for row in a.open_project('P')['boq'] if row['price_code']=='concrete')
    d=a.create_statement_amendment('P',1,created_by='writer',expected_project_version=v(),reason='Synthetic measured correction',source_documents=docs,lines=[{'price_code':'concrete','source_id':source,'corrected_period_quantity':'1.5','reason':'Synthetic .5 m3 correction','detail_source_ref':'TEST-CONTRACT'}])
    d=a.review_statement_amendment('P',d['amendment_id'],reviewed_by='reviewer',expected_revision=d['revision'],expected_project_version=v())
    d=a.approve_statement_amendment('P',d['amendment_id'],approved_by='approver',expected_content_digest=d['content_digest'],expected_revision=d['revision'],expected_project_version=v())
    a.post_statement_amendment('P',d['amendment_id'],posted_by='approver',expected_revision=d['revision'],expected_project_version=v())
    one['concrete']=Decimal('1.5');expected_first=oracle(one,Decimal(100))
    effective=a.statement_amendment_summary('P',1)['net']
    for k,value in expected_first.items():assert Decimal(effective[k])==Decimal(value),(k,effective[k],value)
    assert hashlib.sha256(json.dumps(a.statement_periods('P')[0],sort_keys=True).encode()).hexdigest()==first_hash
    def post(kind,reference,amount,**links):
        r=a.financial_command('P','create',actor='writer',expected_version=v(),kind=kind,reference=reference,amount=amount,currency='IRR',date='2026-10-23',counterparty='طرف حساب آزمایشی',description='سند مصنوعی؛ پرداخت واقعی نیست',source_documents=docs,**links)
        for op,who in [('review','reviewer'),('approve','approver'),('post','approver')]:r=a.financial_command('P',op,actor=who,expected_version=v(),id=r['id'])
        return r
    post('budget','BUDGET','1300');c=post('commitment','CONTRACT','1500');cost=post('cost','INVOICE','1400',commitment_id=c['id'])
    post('payment','BANK-TEST','300',cost_id=cost['id'],commitment_id=c['id']);post('receipt','RECEIPT-TEST','500',statement_no=1)
    snap=a.verified_financial_snapshot('P')
    for k,value in {'actual_cost':'1400','paid':'300','received':'500','debt':'1100','unpaid_commitments':'1200','budget_variance':'-100','cash_net':'200'}.items():assert Decimal(str(snap[k]))==Decimal(value)
    expected_net=Decimal(expected_first['payable_current'])+Decimal(oracle(two,Decimal(0))['payable_current'])
    assert Decimal(str(snap['statement_net']))==expected_net
    assert Decimal(str(snap['receivable']))==expected_net-Decimal(500)
    assert snap['over_budget'] is True
    exports=[]
    for fmt in ['pdf','xlsx','docx','csv']:
        finance=root/('financial.'+fmt);statement=root/('statement-1.'+fmt)
        a.verified_financial_report('P',fmt,finance);a.verified_statement_report('P',1,fmt,statement);exports.extend([finance,statement])
    import csv,fitz
    from openpyxl import load_workbook
    from docx import Document
    for path in exports:
        if path.suffix=='.xlsx':assert all(s.sheet_view.rightToLeft for s in load_workbook(path))
        elif path.suffix=='.docx':assert Document(path).tables
        elif path.suffix=='.csv':assert list(csv.reader(path.open(encoding='utf-8-sig')))
        else:
            with fitz.open(path) as pdf:assert len(pdf)>0 and pdf[0].get_text()
    from PySide6.QtWidgets import QApplication,QTabWidget
    from app.financial_workspace import FinancialWorkspace
    qt=QApplication.instance() or QApplication([]);w=FinancialWorkspace(a,'P');w.show();w.findChild(QTabWidget).setCurrentIndex(2);qt.processEvents();w.grab().save(str(root/'finance-qt.png'));w.close();qt.processEvents()
    before=a.open_project('P');a.store.close();a=StructuralProApp(root/'engineering'/'db')
    assert a.open_project('P')==before and a.verified_financial_snapshot('P')==snap
    assert a.store.integrity_check()['ok']
    receipt={'synthetic':True,'floors':5,'periods':2,'issued_first_preserved':True,'sqlite_integrity':'ok','reopen_equal':True,'oracle_period1_after_amendment':expected_first,'oracle_period2':oracle(two,Decimal(0)),'snapshot':snap,'exports':[x.name for x in exports]}
    (root/'receipt.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2));a.store.close();return receipt

if __name__=='__main__':print(json.dumps(run(sys.argv[1]),ensure_ascii=False,indent=2))
