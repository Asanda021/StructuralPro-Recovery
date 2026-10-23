"""Explicit synthetic five-floor financial acceptance; never official prices."""
from pathlib import Path
from decimal import Decimal,localcontext
import sys,json,hashlib,csv
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'tests'))
from test_phase2_pdf_roof import bound
from core.platform.application import StructuralProApp


def run(root):
 root=Path(root).resolve();root.mkdir(parents=True,exist_ok=True)
 app,session,item,record,data=bound(root,'solid_slab')
 for i in range(1,6):app.add_project_floor('P',f'طبقه آزمایشی {i}',level=f'L{i}')
 app.add_project_member('P','S1','L1','slab');app.add_project_member('P','W2','L2','wall')
 app.save_aec_system_draft('P',data,expected_revision=0)
 app.review_roof_system('P','solid_slab','L1','S1',expected_revision=1,reviewer='independent-fixture-reviewer',note='Explicit synthetic detail and PDF')
 app.commit_aec_system('P','solid_slab','L1','S1',expected_revision=2)
 from copy import deepcopy
 changed=deepcopy(app.get_aec_system('P','solid_slab','L1','S1')['data']);changed.pop('roof_review',None);changed['roof_detail']['geometry']['depth']='.25';changed['roof_detail']['regions'][0]['height']='.25'
 from core.drawings.roof_region import detail_payload
 evidence=root/'revision.json';evidence.write_text(json.dumps(detail_payload(changed)))
 changed['source']['detail_path']=str(evidence);changed['source']['detail_sha256']=hashlib.sha256(evidence.read_bytes()).hexdigest()
 amendment=app.create_roof_amendment('P','solid_slab','L1','S1',changed,expected_revision=2,created_by='fixture-author',reason='Explicit thickness revision from .20m to .25m',evidence={'evidence_path':str(evidence),'evidence_sha256':hashlib.sha256(evidence.read_bytes()).hexdigest(),'evidence_reference':'synthetic-R2'})
 app.review_beam_amendment('P',amendment['id'],expected_revision=1,reviewer='fixture-independent-reviewer',note='Independent synthetic check')
 app.apply_beam_amendment('P',amendment['id'],expected_revision=1)
 app.add_takeoff('P','building','wall',source_id='fixture:phase3:L2:W2',system='L2',manual_location_id='W2',length=5,height=3,openings=2,count=1)
 prices={'concrete':('m3','100'),'formwork':('m2','20'),'rebar':('kg','10'),'wall':('m2','50')}
 source=root/'synthetic-prices.csv'
 with source.open('w',newline='',encoding='utf-8') as f:
  writer=csv.writer(f);writer.writerow(['year','code','description','unit','unit_price'])
  for code,(unit,rate) in prices.items():writer.writerow([1405,code,'Synthetic fixture '+code,unit,rate])
 app.import_project_pricebook('P',source,year=1405,approved_by='fixture-reviewer',currency='IRR',source_version='synthetic-R1')
 mapping={str(i):row['item_code'] for i,row in enumerate(app.open_project('P')['boq'])}
 app.review_boq_price_mapping('P',mapping,reviewer='fixture-reviewer',expected_digest=app.store.current_digest('P'))
 doc=root/'factors.json';doc.write_text(json.dumps({'version':'synthetic-R1','year':1405,'currency':'IRR','effective_date':'2026-10-23','basis':'sequential_compound','factors':[{'name':'overhead','rate':'0.1'},{'name':'region','rate':'0.05'}]}))
 app.review_project_factors('P',doc,sha256=hashlib.sha256(doc.read_bytes()).hexdigest(),author='fixture-author',reviewer='fixture-reviewer',expected_digest=app.store.current_digest('P'))
 estimate=app.recalculate_estimate('P')
 with localcontext() as ctx:
  ctx.prec=80
  quantities={'concrete':Decimal('4'),'formwork':Decimal('16'),'rebar':Decimal('71.04'),'wall':Decimal('13')}
  base=sum(q*Decimal(prices[code][1]) for code,q in quantities.items());expected=base*Decimal('1.1')*Decimal('1.05')
 assert Decimal(str(estimate['cost']['base']))==base,(estimate,base)
 assert Decimal(str(estimate['cost']['grand_total']))==expected
 for code,q in quantities.items():assert sum(Decimal(str(r['quantity'])) for r in estimate['boq'] if r['item_code']==code)==q
 paths=[]
 for fmt in ['pdf','xlsx','docx','csv']:
  path=root/('estimate.'+fmt);app.report('P',fmt,path);assert path.stat().st_size;paths.append(path)
 from openpyxl import load_workbook
 from docx import Document
 import fitz
 assert load_workbook(paths[1],read_only=True).worksheets[0].max_row>=4
 assert len(Document(paths[2]).tables)>0
 with paths[3].open(encoding='utf-8-sig') as f:assert len(list(csv.reader(f)))>=5
 with fitz.open(paths[0]) as pdf:assert len(pdf)>0 and any(page.get_text() for page in pdf)
 snapshot=deepcopy(app.open_project('P'));digest=app.store.current_digest('P')
 try:app.review_boq_price_mapping('P',mapping,reviewer='fixture-reviewer',expected_digest='stale')
 except RuntimeError:pass
 else:raise AssertionError('stale CAS accepted')
 assert app.store.current_digest('P')==digest
 app.store.close();app=StructuralProApp(root/'db');assert app.open_project('P')==snapshot
 assert app.store.integrity_check()['ok']
 from PySide6.QtWidgets import QApplication
 from app.financial_factor_review import BOQPriceMappingDialog
 qt=QApplication.instance() or QApplication([]);dialog=BOQPriceMappingDialog(app,'P');dialog.show();qt.processEvents();dialog.grab().save(str(root/'financial-qt.png'));dialog.close();qt.processEvents()
 receipt={'synthetic':True,'floors':5,'active_floors':['L1','L2'],'oracle_quantities':{k:str(v) for k,v in quantities.items()},'base':str(base),'total':str(expected),'sqlite_integrity':'ok','reopen_equal':True,'amendment_id':amendment['id'],'exports':[p.name for p in paths]}
 (root/'receipt.json').write_text(json.dumps(receipt,indent=2));app.store.close();return receipt

if __name__=='__main__':print(json.dumps(run(sys.argv[1]),indent=2))
